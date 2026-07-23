#!/usr/bin/env python3
"""画像によるドローン検知モデルの紹介スライド（やさしい解説＋具体名つき）。

方針:
  - 専門用語はかみ砕き、各スライドに平易な導入文を添える。
  - 具体的なモデル名・ライブラリ名は明記する。
  - 冒頭に「概要」を置き、全体を再構成する。

デザイン（進捗第11回スタイル / ミニマル, 4:3）:
  白背景・黒文字（1A1A1A）・補助グレー（555555）・色なし、Yu Gothic / Courier New、
  枠はコード抜粋のみ（F7F7F7 + CCCCCC）、タイトル下に細罫線。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font)
    return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def lead(s, text, t=1.32, size=14, h=1.2):
    tf = tbox(s, 0.7, t, 8.7, h)
    run(tf.paragraphs[0], text, size, INK)
    return tf


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True)
    return p


def bullet(tf, text, size=13.5, space_before=8, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold)
    return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY)
    return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table
    tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# ── 1 表紙 ───────────────────────────────────────────
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.3, 8.2, 0.5)
run(tf.paragraphs[0], "アプリ機能の紹介", 15, GRAY)
tf = tbox(s, 0.88, 2.85, 8.4, 1.5)
run(tf.paragraphs[0], "画像によるドローン検知", 30, INK, bold=True)
hrule(s, 4.15, l=0.92, w=5.4)
tf = tbox(s, 0.92, 4.35, 8.4, 1.6)
run(tf.paragraphs[0], "スマホのカメラに写ったドローンを、AIが自動で見つけます。", 15, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "見つけた結果は、そのまま「近くにいた」という証明につながります。", 13.5, GRAY)

# ── 2 概要 ──────────────────────────────────────────
s = add_slide()
title(s, "概要")
lead(s,
     "この機能は、カメラ映像のドローンを AI「YOLO」で自動検知し、"
     "その結果をゼロ知識証明（ZK）へつなぐものです。", h=1.3)
tf = tbox(s, 0.7, 2.7, 8.7, 4.3)
subhead(tf, "ひと目でわかるポイント", 0)
bullet(tf, "AI：YOLO を TFLite 形式に変換し、スマホ内だけで推論（通信不要）",
       first=True)
bullet(tf, "主役モデル：ドローン/飛行機/鳥/ヘリを見分ける best_float32.tflite")
bullet(tf, "予備：Google ML Kit（汎用AI）へ自動で切替（止めないための保険）")
bullet(tf, "出口：「30m以内 ∧ ドローン検知」をZKで証明")
sub(tf, "アプリ本体は Flutter 製の rapidsnark_app")

# ── 3 目次 ──────────────────────────────────────────
s = add_slide()
title(s, "このスライドで分かること")
lead(s, "「カメラでドローンを見つける機能」を、具体名も交えてやさしく説明します。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
bullet(tf, "なぜカメラでの検知が必要なのか", first=True)
bullet(tf, "どういう流れでドローンを見つけているのか")
bullet(tf, "どんなAIモデルと部品（パッケージ）を使っているのか")
bullet(tf, "最近直した点と、これからの改善")

# ── 4 なぜ必要か ────────────────────────────────────
s = add_slide()
title(s, "なぜカメラでの検知が必要か")
lead(s,
     "ドローンは「電波（リモートID）」でも見つけられます。"
     "でも電波だけでは取りこぼしがあります。カメラは、その穴を別の方法で埋めます。",
     h=1.3)
tf = tbox(s, 0.7, 2.75, 8.7, 1.2)
run(tf.paragraphs[0],
    "たとえるなら、電波は「車のナンバーが自動で届く」仕組み。"
    "でも、ナンバーを出さない車や、ウソのナンバーを出す車もいます。", 13, GRAY)
table(s, 0.55, 4.05, 8.9, [
    ["こんな時", "電波だけだと", "カメラを足すと"],
    ["電波を出さない機体", "気づけない", "見た目で見つけられる"],
    ["ウソの電波（なりすまし）", "信じてしまう", "実物の有無で裏が取れる"],
    ["証拠の残し方", "数字だけ", "枠を描いた画像が残る"],
], col_w=[2.7, 2.9, 3.0], row_h=0.55, fs=11.5)

# ── 5 全体の流れ ────────────────────────────────────
s = add_slide()
title(s, "ドローンを見つけるまでの流れ")
lead(s, "カメラで撮った動画を、AIが数コマだけ取り出して1枚ずつ調べます。")
codebox(s, 0.7, 2.4, 8.6, 2.6, [
    "① 動画を撮る",
    "② 動画から数コマの静止画を取り出す（video_thumbnail）",
    "③ AI（YOLO）が「これはドローン」と判定し枠を描く",
    "④ ドローンが写っていれば『検知あり』とする",
    "⑤ そのまま近接の証明（ZK）へ進める",
], size=12.5)
tf = tbox(s, 0.7, 5.3, 8.7, 1.7)
run(tf.paragraphs[0],
    "AIの計算はすべてスマホの中だけで行います（写真を外に送りません）。",
    13.5, INK, bold=True)
sub(tf, "通信不要・プライバシーに配慮")

# ── 6 YOLOとは ─────────────────────────────────────
s = add_slide()
title(s, "使っているAI「YOLO」とは")
lead(s,
     "YOLO は「写真の中から物体を一瞬で見つけて、四角い枠で囲むAI」です。"
     "人の目のように、画像を見て『どこに・何が』あるかを答えます。", h=1.4)
tf = tbox(s, 0.7, 2.9, 8.7, 4.0)
subhead(tf, "このアプリでの使い方", 0)
bullet(tf, "ドローン・飛行機・鳥・ヘリを見分けるYOLOモデルを使用", first=True)
bullet(tf, "「TFLite」形式に変換し、tflite_flutter でスマホ内推論")
sub(tf, "TFLite = スマホで動かせるよう小さく軽くしたAIの形")
bullet(tf, "似た枠が重なったら重複を1つにまとめる（NMS という整理）")
sub(tf, "NMS = 同じ物に何個も枠が出るのを防ぐ仕組み")
bullet(tf, "AIが「ドローン」と一定以上の自信（しきい値）で答えたら検知成立")

# ── 7 モデル一覧（具体名）★ ─────────────────────────
s = add_slide()
title(s, "使うモデル（具体名）")
lead(s, "メインはドローンを見分けるYOLO。読み込めない時は汎用AIに自動で切替えます。")
table(s, 0.5, 2.35, 9.0, [
    ["モデル / ファイル", "見分けるもの", "扱い"],
    ["best_float32.tflite (new-models)", "ドローン/飛行機/鳥/ヘリ", "最優先で使用"],
    ["uav_yolo11n_uav.tflite (COCO)", "人・車など80種類", "サンプル"],
    ["uav_visdrone_finetuned.tflite", "ドローン（VisDrone学習）", "未使用（将来）"],
    ["uav_roboflow_drone.tflite", "ドローン（Roboflow）", "未使用（要配置）"],
    ["Google ML Kit", "汎用（物体/ラベル）", "予備（自動切替）"],
], col_w=[3.7, 2.9, 2.4], row_h=0.5, fs=11)
tf = tbox(s, 0.7, 5.6, 8.7, 1.3)
sub(tf, "float32 と float16 の2種を同梱（端末で軽量に動かすため）", 11.5)
sub(tf, "ML Kit は ImageLabeler + ObjectDetector（精度は控えめな保険）", 11.5)

# ── 7.5 予備：Google ML Kit のしくみ ★ ──────────────
s = add_slide()
title(s, "予備：Google ML Kit のしくみ")
lead(s,
     "メインのYOLOが使えない時に備えた『保険』。Googleが配る汎用の画像AIで、"
     "ネット不要・スマホ内だけで動きます。", h=1.3)
tf = tbox(s, 0.7, 2.65, 8.7, 2.0)
subhead(tf, "2つの機能を組み合わせて使う", 0)
bullet(tf, "ImageLabeler（画像ラベリング）：写真全体に「何が写っているか」のタグ付け",
       first=True)
sub(tf, "自信が 0.45 以上のタグだけ採用（confidenceThreshold = 0.45）")
bullet(tf, "ObjectDetector（物体検知）：写真の中の物体を四角で見つける")
sub(tf, "1枚モード・分類あり・複数物体（DetectionMode.single / classify / multiple）")
codebox(s, 0.7, 4.85, 8.6, 1.35, [
    "物体を枠で発見 → 枠を切り出して再ラベリング",
    "  → 名前に drone / uav 等が入れば「ドローン」とみなす",
], size=11.5)
tf = tbox(s, 0.7, 6.35, 8.7, 0.9)
sub(tf, "汎用モデルゆえ『ドローン』専用ラベルは出にくい＝精度は控えめな保険", 11.5)

# ── 8 パッケージ紹介 ───────────────────────────────
s = add_slide()
title(s, "使っている部品（パッケージ）")
lead(s, "アプリは、世の中の便利な部品を組み合わせて作っています。主な部品は次の通りです。")
table(s, 0.55, 2.35, 8.9, [
    ["部品（パッケージ）", "役割（やさしく）"],
    ["tflite_flutter", "AI（YOLO）をスマホの中で動かす"],
    ["google_mlkit_object_detection", "予備の物体検知AI"],
    ["google_mlkit_image_labeling", "予備の画像ラベル付けAI"],
    ["image", "画像の加工・枠の描き込み"],
    ["camera", "カメラの映像を扱う"],
    ["video_thumbnail", "動画からコマ（静止画）を取り出す"],
], col_w=[3.8, 5.1], row_h=0.46, fs=11.5)
tf = tbox(s, 0.7, 6.5, 8.7, 0.6)
sub(tf, "どれも『スマホの中だけで完結』を支える部品", 11.5)

# ── 9 検知の判定 ───────────────────────────────────
s = add_slide()
title(s, "「ドローンと判定」する条件")
lead(s, "AIが付けた名前（ラベル）に『drone』などの言葉が入っていれば、ドローンとみなします。")
codebox(s, 0.7, 2.4, 8.6, 1.7, [
    "見つけた名前に次のどれかが入っていれば「ドローン」:",
    "  drone / quadcopter / multirotor / uav",
], size=12)
tf = tbox(s, 0.7, 4.3, 8.7, 2.6)
subhead(tf, "つまり", 0)
bullet(tf, "AIが一定の自信で「Drone」と答える → 言葉が一致 → 検知成立",
       first=True)
sub(tf, "自信の度合い（しきい値）は調整可能。低すぎると誤検知が増える")

# ── 10 最近直した点（具体名）★ ─────────────────────
s = add_slide()
title(s, "最近直した点")
lead(s,
     "証明に進む経路だけ『ドローンを知らないモデル(COCO)』を使っていて、"
     "ドローンを写しても気づけませんでした。これを直しました。", h=1.4)
table(s, 0.55, 2.95, 8.9, [
    ["", "前", "後"],
    ["使うモデル", "uav_yolo11n_uav (COCO)", "best_float32 (new-models)"],
    ["結果", "写しても検知されない", "写せば検知され、証明に進める"],
], col_w=[1.7, 3.6, 3.6], row_h=0.6, fs=11.5)
tf = tbox(s, 0.7, 5.4, 8.7, 1.5)
run(tf.paragraphs[0],
    "これで「画面に映る検知」と「証明に進む検知」が同じモデルで揃いました。",
    13.5, INK, bold=True)

# ── 11 証明とのつながり ────────────────────────────
s = add_slide()
title(s, "検知から「証明」へのつながり")
lead(s,
     "見つけただけでは終わりません。"
     "『近くにいた』こと（30m以内）と『ドローンだった』ことを合わせて証明します。",
     h=1.4)
codebox(s, 0.7, 2.9, 8.6, 1.3, [
    "証明できること = 30m以内にいた  かつ  ドローンを検知した",
], size=12.5)
tf = tbox(s, 0.7, 4.5, 8.7, 2.4)
subhead(tf, "ゼロ知識証明（ZK）とは", 0)
bullet(tf, "正確な位置や写真そのものを見せずに、事実だけを証明する技術",
       first=True)
sub(tf, "「30m以内にドローンがいた」とだけ示し、中身は隠せる")
bullet(tf, "2つの条件を『かつ』で結ぶことで、ごまかしにくくなる")

# ── 12 いまの弱点 ──────────────────────────────────
s = add_slide()
title(s, "いまの弱点（正直なところ）")
lead(s, "便利ですが、まだ苦手なこともあります。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
bullet(tf, "調べるコマ数が少なく、一瞬しか映らないと見逃すことがある",
       first=True)
bullet(tf, "1枚の静止画で判断するため、鳥や凧と紛れやすい")
bullet(tf, "遠くの小さなドローンは画像がつぶれて見つけにくい")
bullet(tf, "夜や逆光は苦手（明るい昼は得意）")
bullet(tf, "AIが本物と判定しても、ニセ画像までは見破れない")
sub(tf, "→ だから電波や音と組み合わせて確かめる")

# ── 13 これから ────────────────────────────────────
s = add_slide()
title(s, "これからの改善")
lead(s, "弱点をつぶし、もっと確実にドローンを見つけられるようにします。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
bullet(tf, "調べるコマ数を増やし、複数コマの『多数決』で判断する",
       first=True)
bullet(tf, "画面を分割して小さなドローンも見つけやすくする（タイル推論）")
bullet(tf, "動きを追いかけ、鳥や凧と区別する（トラッキング）")
bullet(tf, "ドローン専用モデル（VisDrone/Roboflow）へ差し替える")
bullet(tf, "電波（リモートID）や音とも突き合わせて、より確実にする")

# ── 14 まとめ ──────────────────────────────────────
s = add_slide()
title(s, "まとめ")
tf = tbox(s, 0.7, 1.45, 8.7, 5.6)
bullet(tf, "カメラ検知は、電波だけでは取りこぼす分を補う『もう一つの目』",
       first=True, bold=True)
bullet(tf, "YOLO(best_float32.tflite)をスマホ内で動かし、ドローンを枠で検知",
       bold=True)
bullet(tf, "見つけた結果は『30m以内 ∧ 検知』の証明につながる", bold=True)
subhead(tf, "ひとことで")
bullet(tf,
       "「電波・カメラ・音」を組み合わせるほど、ごまかしにくく確かな証明になる",
       first=True)

prs.save("image_model_4x3.pptx")
print(f"saved image_model_4x3.pptx ({len(prs.slides._sldIdLst)} slides)")
