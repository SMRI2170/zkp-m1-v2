#!/usr/bin/env python3
"""Generate the Week 22 Nova research and adoption-decision deck.

Run from the repository root:
    python3 scripts/generate_week22_nova_research_slides.py

The diagrams are native PowerPoint shapes so they remain editable.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "siles-pptx" / "week22_nova_research.pptx"

PAGE_W, PAGE_H = 10.0, 7.5
FONT = "Hiragino Sans"
INK = RGBColor(0x20, 0x27, 0x35)
MUTED = RGBColor(0x65, 0x6D, 0x7A)
RULE = RGBColor(0xD9, 0xDE, 0xE6)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF5, 0xF7, 0xFA)
INDIGO = RGBColor(0x48, 0x55, 0xA5)
INDIGO_PALE = RGBColor(0xF0, 0xF1, 0xFA)
GREEN = RGBColor(0x2E, 0x75, 0x61)
GREEN_PALE = RGBColor(0xEC, 0xF5, 0xF1)
AMBER = RGBColor(0xA9, 0x70, 0x26)
AMBER_PALE = RGBColor(0xFC, 0xF4, 0xE7)
RED = RGBColor(0xA0, 0x4F, 0x49)
RED_PALE = RGBColor(0xF8, 0xEF, 0xEE)

prs = Presentation()
prs.slide_width = Inches(PAGE_W)
prs.slide_height = Inches(PAGE_H)
prs.core_properties.title = "Novaによる逐次集約 — Week 22調査"
prs.core_properties.subject = "Nova IVCの研究適合性、PoC実測、残課題と採用判断"
prs.core_properties.author = "Week 22 Nova Research"
BLANK = prs.slide_layouts[6]
slide_no = 0


def _font(run, name=FONT):
    run.font.name = name
    props = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = props.find(qn(tag))
        if el is None:
            el = props.makeelement(qn(tag), {})
            props.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, text="", size=14, color=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0,
            fit=False, font=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    if fit:
        tf.fit_text(font_family=font, max_size=size)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    _font(r, font)
    return shape


def rect(slide, x, y, w, h, fill=PANEL, line=RULE, radius=True, width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(width)
    if radius:
        sh.adjustments[0] = 0.08
    return sh


def line(slide, x1, y1, x2, y2, color=RULE, width=1.25, arrow=False):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    if arrow:
        ln = sh.line._get_or_add_ln()
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
        ln.append(tail)
    return sh


def pill(slide, x, y, w, text, fill, color, h=0.30, size=9.5):
    rect(slide, x, y, w, h, fill=fill, line=fill, radius=True, width=0.5)
    textbox(slide, x + 0.04, y + 0.015, w - 0.08, h - 0.02, text,
            size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE)


def base_slide(title, section, source=""):
    global slide_no
    slide_no += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    textbox(s, 0.64, 0.30, 7.8, 0.25, section.upper(), size=9.5, color=INDIGO, bold=True)
    textbox(s, 0.64, 0.64, 8.8, 0.48, title, size=23, color=INK, bold=True)
    line(s, 0.64, 1.22, 9.36, 1.22, color=RULE, width=0.8)
    if source:
        textbox(s, 0.66, 7.14, 8.1, 0.18, "出典: " + source, size=8, color=MUTED)
    textbox(s, 8.78, 7.10, 0.55, 0.22, f"{slide_no:02d}", size=9, color=MUTED,
            align=PP_ALIGN.RIGHT)
    return s


def card_text(slide, x, y, w, kicker, body, accent=INDIGO, fill=PANEL,
              body_size=14.5, h=1.45):
    rect(slide, x, y, w, h, fill=fill, line=RULE)
    rect(slide, x, y, 0.055, h, fill=accent, line=accent, radius=False, width=0)
    textbox(slide, x + 0.20, y + 0.15, w - 0.36, 0.22, kicker,
            size=10, color=accent, bold=True)
    textbox(slide, x + 0.20, y + 0.45, w - 0.37, h - 0.55, body,
            size=body_size, color=INK, bold=True)


def add_multiline(slide, x, y, w, h, lines, size=13, gap=5, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, (label, content, label_color) in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        p.level = 0
        r = p.add_run()
        r.text = label
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = label_color
        _font(r)
        r = p.add_run()
        r.text = content
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _font(r)
    return shape


# 1 — Cover: one concise proposition and a repeatable-state illustration.
s = prs.slides.add_slide(BLANK)
slide_no += 1
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAPER
textbox(s, 0.78, 0.58, 5.4, 0.3, "WEEK 22  /  NOVA RESEARCH", size=10,
        color=INDIGO, bold=True)
textbox(s, 0.78, 1.31, 8.3, 0.62, "Novaは、観測を", size=29,
        color=INK, bold=True)
textbox(s, 0.78, 1.94, 8.3, 0.62, "一件ずつ証明できるか", size=29,
        color=INK, bold=True)
textbox(s, 0.82, 2.82, 7.9, 0.53,
        "逐次集約への適合性と、本番化までに必要な接続", size=16, color=MUTED)

# Editable step diagram (not decorative artwork).
nodes = [(0.90, "状態 z₀", "開始"), (3.02, "状態 z₁", "観測 1"),
         (5.14, "状態 z₂", "観測 2"), (7.26, "状態 zₙ", "累積")]
for i, (x, state, cap) in enumerate(nodes):
    fill = INDIGO_PALE if i in (0, 3) else PANEL
    border = INDIGO if i in (0, 3) else RULE
    rect(s, x, 4.13, 1.42, 0.86, fill=fill, line=border)
    textbox(s, x + 0.08, 4.27, 1.26, 0.28, state, size=15, color=INK,
            bold=True, align=PP_ALIGN.CENTER)
    textbox(s, x + 0.08, 4.65, 1.26, 0.19, cap, size=9, color=MUTED,
            align=PP_ALIGN.CENTER)
    if i < len(nodes) - 1:
        line(s, x + 1.47, 4.56, x + 2.03, 4.56, color=INDIGO, width=1.5, arrow=True)
        textbox(s, x + 1.49, 4.12, 0.53, 0.22, "F(w)", size=9, color=INDIGO,
                bold=True, align=PP_ALIGN.CENTER)
textbox(s, 5.83, 5.08, 2.88, 0.25, "反復するstepを、累積proofへfold", size=10.5,
        color=GREEN, bold=True, align=PP_ALIGN.CENTER)

rect(s, 0.82, 5.65, 8.58, 0.80, fill=INDIGO_PALE, line=INDIGO_PALE)
textbox(s, 1.05, 5.83, 8.12, 0.42,
        "結論  Novaは逐次集約の有力候補。実leaf proofとEVM接続は未完了。",
        size=16, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
textbox(s, 0.82, 6.76, 7.8, 0.20,
        "調査・実装照合: 2026-10-02  /  Week 22", size=9, color=MUTED)
textbox(s, 8.78, 7.10, 0.55, 0.22, "01", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


# 2 — What Nova proves.
s = base_slide("Novaが証明するのは「同じ更新を繰り返したこと」", "01 / しくみ",
               "week22/01 §Novaとは何か、§この研究でNovaに向く部分")
textbox(s, 0.78, 1.48, 8.4, 0.43,
        "状態を一件ずつ更新し、各stepをfoldして累積proofをつくる", size=16, color=INK, bold=True)

formula = rect(s, 0.82, 2.12, 8.36, 0.76, fill=INDIGO_PALE, line=INDIGO_PALE)
textbox(s, 1.02, 2.28, 7.95, 0.40, "zᵢ₊₁  =  F(zᵢ, wᵢ)", size=22, color=INDIGO,
        bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

xs = [0.90, 3.00, 5.10, 7.20]
labels = [("前状態", "z₀"), ("step 1", "F + w₀"), ("step 2", "F + w₁"), ("最終状態", "zₙ")]
for i, (x, (lab, val)) in enumerate(zip(xs, labels)):
    fill = GREEN_PALE if i == 3 else PANEL
    border = GREEN if i == 3 else RULE
    rect(s, x, 3.42, 1.56, 0.82, fill=fill, line=border)
    textbox(s, x + 0.06, 3.53, 1.44, 0.20, lab, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    textbox(s, x + 0.06, 3.80, 1.44, 0.28, val, size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
    if i < 3:
        line(s, x + 1.62, 3.83, x + 2.00, 3.83, color=INDIGO, width=1.3, arrow=True)
textbox(s, 2.03, 4.43, 5.98, 0.24,
        "同じFを繰り返す ── 前の状態はIVCの中で引き継がれる", size=11.5,
        color=INDIGO, bold=True, align=PP_ALIGN.CENTER)

card_text(s, 0.82, 5.02, 4.07, "向いている", "端末を順番に追加し、\nstateを更新する", accent=GREEN,
          fill=GREEN_PALE, body_size=14.5, h=1.24)
card_text(s, 5.12, 5.02, 4.07, "自動では解決しない", "既存leaf proofの正しさや\n任意2集合のmerge", accent=AMBER,
          fill=AMBER_PALE, body_size=14.5, h=1.24)
textbox(s, 0.86, 6.47, 8.1, 0.34,
        "Novaのfoldingは、外部proofの検証やmembershipの代わりにはならない。",
        size=12.2, color=INK, bold=True)


# 3 — Fit to this project and state of integration.
s = base_slide("研究への適合性は高いが、Nova経路は未接続", "02 / 研究への適用",
               "week22/01 §2026-10-02時点の実装照合、week22/07 §実装ファイルの役割")
textbox(s, 0.80, 1.47, 8.35, 0.36,
        "既存BN254 leafを起点に、最終proofをEVMへ届ける経路", size=15, color=INK, bold=True)

stages = [
    (0.66, 1.28, "端末", "Circom\nBN254 leaf", "既存", GREEN),
    (2.14, 1.28, "検証", "固定VK\nBridge", "Nova未接続", AMBER),
    (3.62, 1.28, "状態", "member +\nnullifier", "Nova未実装", RED),
    (5.10, 1.28, "集約", "Nova RSSI\nstep", "3 step PoC", GREEN),
    (6.58, 1.28, "圧縮", "Decider\nproof", "未実装", RED),
    (8.06, 1.28, "検証", "EVM /\nV1 ABI", "未接続", RED),
]
for idx, (x, w, top, mid, status, c) in enumerate(stages):
    fill = GREEN_PALE if c == GREEN else (AMBER_PALE if c == AMBER else RED_PALE)
    rect(s, x, 2.20, w, 1.48, fill=fill, line=c)
    textbox(s, x + 0.08, 2.35, w - 0.16, 0.20, top, size=9, color=c, bold=True,
            align=PP_ALIGN.CENTER)
    textbox(s, x + 0.08, 2.65, w - 0.16, 0.47, mid, size=11.2, color=INK,
            bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    pill(s, x + 0.08, 3.27, w - 0.16, status, fill=fill, color=c, h=0.25, size=8.2)
    if idx < len(stages) - 1:
        line(s, x + w + 0.025, 2.93, stages[idx + 1][0] - 0.025, 2.93,
             color=RULE, width=1.05, arrow=True)

rect(s, 0.80, 4.08, 8.45, 1.80, fill=PANEL, line=RULE)
textbox(s, 1.02, 4.25, 7.95, 0.25, "現在、proofに入っている範囲", size=11, color=MUTED, bold=True)
add_multiline(s, 1.06, 4.66, 7.95, 1.00, [
    ("確認済み  ", "RSSI整合判定・件数更新・3-step IVC検証", GREEN),
    ("未接続    ", "実leaf固定VK、登録membership、重複nullifier防止", RED),
    ("未到達    ", "Decider → EVM verifier / V1 ABI", RED),
], size=12.5, gap=6)
textbox(s, 0.85, 6.28, 8.2, 0.42,
        "直接再帰側にはBridge等の部品があるが、Nova stepへはまだ接続されていない。",
        size=11.5, color=INK, bold=True)


# 4 — Evidence: desktop and phone results, kept in separate conditions.
s = base_slide("PoCは3 stepで成立し、Android実機でも検証できた", "03 / 実証結果",
               "week20/08 §5.1、week20/09 §5.3.1、week22/01 実装照合")
textbox(s, 0.82, 1.48, 8.3, 0.35,
        "正常・異常の混在を処理し、最終IVC proofの検証に成功", size=15, color=INK, bold=True)

rect(s, 0.82, 2.05, 4.03, 3.42, fill=PANEL, line=RULE)
pill(s, 1.04, 2.26, 1.52, "ANDROID / PIXEL 8", fill=GREEN_PALE, color=GREEN, h=0.30)
textbox(s, 1.08, 2.80, 3.46, 0.42, "各stepの証明生成", size=12, color=MUTED, bold=True)
textbox(s, 1.06, 3.22, 3.48, 0.44, "485 / 655 / 1,198 ms", size=17, color=INK, bold=True)
line(s, 1.07, 3.88, 4.56, 3.88, color=RULE, width=0.8)
textbox(s, 1.08, 4.04, 3.46, 0.31, "最終verify  172 ms   /   成功", size=12.5,
        color=GREEN, bold=True)
textbox(s, 1.08, 4.58, 3.40, 0.55, "count:  1 → 2 → 2\n異常:    0 → 0 → 1", size=13, color=INK,
        bold=True)

rect(s, 5.14, 2.05, 4.03, 3.42, fill=PANEL, line=RULE)
pill(s, 5.36, 2.26, 1.34, "M2 MAC", fill=INDIGO_PALE, color=INDIGO, h=0.30)
textbox(s, 5.40, 2.80, 3.46, 0.42, "native desktop / Sonobe", size=12, color=MUTED, bold=True)
textbox(s, 5.38, 3.22, 3.48, 0.44, "135 / 151 / 281 ms", size=17, color=INK, bold=True)
line(s, 5.39, 3.88, 8.88, 3.88, color=RULE, width=0.8)
textbox(s, 5.40, 4.04, 3.46, 0.31, "最終verify  28 ms   /   成功", size=12.5,
        color=GREEN, bold=True)
textbox(s, 5.40, 4.58, 3.38, 0.55, "count:  1 → 2 → 2\n異常:    0 → 0 → 1", size=13, color=INK,
        bold=True)

rect(s, 0.82, 5.78, 8.35, 0.75, fill=AMBER_PALE, line=AMBER_PALE)
textbox(s, 1.04, 5.93, 7.92, 0.41,
        "注意  3 stepの機能実証。10 step性能、leaf真正性、二重計上防止の実証ではない。",
        size=12.4, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
textbox(s, 0.84, 6.68, 8.15, 0.22,
        "DesktopとAndroidは計測条件が異なるため、数値を直接比較しない。",
        size=9.5, color=MUTED)


# 5 — Open risks grouped by trust / scale / delivery.
s = base_slide("本番判断を止めているのは「証明の意味」と接続", "04 / 残課題",
               "week22/01 §実装照合、week22/07 §完了とみなす条件")

open_items = [
    (0.82, 1.62, "01  入力の真正性", "Circom leaf proofを\n固定VKで検証していない", RED, RED_PALE),
    (5.12, 1.62, "02  一度だけ数える", "memberRoot / nullifierRoot\nがNova stateにない", RED, RED_PALE),
    (0.82, 3.24, "03  混入を拒否", "session・policy等の\n否定テストが不足", AMBER, AMBER_PALE),
    (5.12, 3.24, "04  規模と配布", "10-step性能・RAM、\niOS実機が未評価", AMBER, AMBER_PALE),
    (0.82, 4.86, "05  オンチェーン", "Decider / V1 ABI / gas\nを未接続・未測定", RED, RED_PALE),
]
for x, y, heading, body, c, bg in open_items:
    card_text(s, x, y, 4.07, heading, body, accent=c, fill=bg, body_size=13.5, h=1.35)

rect(s, 5.12, 4.86, 4.07, 1.35, fill=INDIGO_PALE, line=INDIGO_PALE)
textbox(s, 5.34, 5.02, 3.62, 0.24, "判断の軸", size=10, color=INDIGO, bold=True)
textbox(s, 5.34, 5.37, 3.58, 0.64,
        "外部の申告値ではなく、\nproofの中で根拠を束縛する。", size=13.3, color=INK, bold=True)
textbox(s, 0.86, 6.58, 8.1, 0.28,
        "高速化より先に、端末・session・重複防止を同じproofへ結び付ける。",
        size=12.3, color=INK, bold=True)


# 6 — Decision and next gates.
s = base_slide("採用候補として継続。ただし本番採用はE2E後に判断", "05 / 判断と次の一手",
               "week22/01 §採用判断を更新するための必須条件、week22/07 §完了とみなす条件")
textbox(s, 0.84, 1.49, 8.3, 0.48,
        "Novaは逐次state更新に有力。現時点では、完成したEVM検証経路ではない。",
        size=15.3, color=INK, bold=True)

gates = [
    (0.82, "GATE 1", "信頼を束縛", "leaf固定VK\nmember / nullifier", GREEN),
    (3.02, "GATE 2", "安全性を検証", "二重利用・異session\n異policyを拒否", AMBER),
    (5.22, "GATE 3", "規模を測る", "10 step + RAM\nproof size / 時間", AMBER),
    (7.42, "GATE 4", "EVMで閉じる", "Decider + V1 ABI\ngas測定", INDIGO),
]
for idx, (x, gate, head, detail, c) in enumerate(gates):
    fill = GREEN_PALE if c == GREEN else (INDIGO_PALE if c == INDIGO else AMBER_PALE)
    rect(s, x, 2.47, 1.78, 1.62, fill=fill, line=c)
    textbox(s, x + 0.10, 2.63, 1.58, 0.20, gate, size=9, color=c, bold=True,
            align=PP_ALIGN.CENTER)
    textbox(s, x + 0.09, 2.94, 1.60, 0.28, head, size=12, color=INK, bold=True,
            align=PP_ALIGN.CENTER)
    textbox(s, x + 0.08, 3.32, 1.62, 0.55, detail, size=9.8, color=MUTED,
            align=PP_ALIGN.CENTER)
    if idx < len(gates) - 1:
        line(s, x + 1.83, 3.28, x + 2.16, 3.28, color=RULE, width=1.15, arrow=True)

rect(s, 0.82, 4.55, 8.37, 1.30, fill=PANEL, line=RULE)
textbox(s, 1.06, 4.74, 1.62, 0.25, "採用ゲート", size=11, color=INDIGO, bold=True)
textbox(s, 2.68, 4.69, 6.24, 0.76,
        "有効leaf → 一度だけ計上 → 10台程度へ集約 →\nDecider proofをEVMで検証、を一本の実proofで通す。",
        size=14, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
line(s, 2.48, 4.78, 2.48, 5.60, color=INDIGO, width=1.3)

textbox(s, 0.86, 6.20, 8.18, 0.60,
        "その後、同じ安全条件・同じ計測条件で直接再帰Groth16と比較して最終選択する。",
        size=12.2, color=INK, bold=True)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Wrote {OUTPUT}")
print(f"Slides: {len(prs.slides)}")
