#!/usr/bin/env python3
"""week19/02_フレームワーク比較.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


def fw_slide(t, dekiru, merit, demerit, honban):
    """フレームワーク1枚: できること/メリット/デメリット/本案件。"""
    s = add_slide(); title(s, t)
    items = [(0, "できること = " + dekiru), (0, "メリット")]
    items += [(1, m) for m in merit]
    items += [(0, "デメリット")]
    items += [(1, d) for d in demerit]
    items += [(0, "本案件 = " + honban)]
    bullets(s, TOP, items, size0=15.5, size1=13.5)


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "フレームワーク比較", 32, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.4)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "何ができる / メリット / デメリット", 16, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 02_フレームワーク比較", 12, GRAY)

# S2 早見表
s = add_slide(); title(s, "早見表")
table(s, TOP, [
    ["フレームワーク", "エンジン", "事前準備", "モバイル証明", "一言"],
    ["circom+rapidsnark", "Groth16", "回路別", "○", "現行の定番"],
    ["gnark", "Groth16/PLONK", "回路別", "✕(重い)", "集約(サーバ)"],
    ["Noir+Barretenberg", "Honk(PLONK系)", "1回", "◎", "書きやすい・準備緩い"],
    ["halo2", "PLONK", "1回/不要", "○", "強力だが難しい"],
    ["Plonky2/3", "STARK", "不要", "○", "準備不要・再帰速い"],
    ["Nova/Sonobe", "Folding", "ほぼ不要", "◎(省メモリ)", "端末側集約向き"],
    ["RISC Zero/SP1", "zkVM(STARK)", "不要", "△(重い)", "普通のコードで書ける"],
    ["Cairo/Stwo", "STARK", "不要", "◎(実証多)", "Starknet前提"],
    ["Mopro", "(上記を包む)", "—", "◎", "証明系でなく橋渡し"],
], col_w=[2.4, 1.9, 1.2, 1.4, 1.7], fs=11, row_h=0.4)

# S3-S11 各フレームワーク
fw_slide("① circom + rapidsnark（現行）",
         "circom記述→snarkjs/rapidsnarkでGroth16。circomlibにPoseidon/EdDSA/比較、実装済",
         ["証明が最小・オンチェーン検証が最安(≈348k gas)",
          "ライブラリと実績が豊富・既に動いている",
          "rapidsnarkでarm64実機証明が実用的"],
         ["回路ごとにトラステッドセットアップの儀式(★1脆弱性)",
          "回路更新のたびに儀式のやり直し",
          "低レベルで複雑な回路は書きにくい"],
         "現状維持の基準。端末証明の土台として妥当")

fw_slide("② gnark（現行の集約器）",
         "Goで回路記述、Groth16/PLONK。再帰検証で「N個の証明を検証する証明」(M2)",
         ["プローバが高速・Goで実装しやすい",
          "再帰・Merkle・MiMC が揃い集約に向く",
          "EVM向けverifierを出力できる"],
         ["サーバ向き・端末での証明生成は重く非現実的",
          "Groth16なら回路別セットアップが必要"],
         "集約(M2)はこれで継続。端末証明器の候補ではない")

fw_slide("③ Noir + Barretenberg",
         "高水準Noirで記述、bb(UltraHonk)。Poseidon標準・EdDSA実装あり・比較可→移植で実現可能",
         ["回路別セットアップ不要(1回使い回し)→★1を根治",
          "書きやすく・クライアント証明が速い",
          "Mopro経由でiOS/Android/Flutterに載る"],
         ["オンチェーン検証が約2.4M gas(Groth16の約7倍)",
          "circom→Noirの回路移植が必要",
          "集約系を現行gnarkから組み替え"],
         "「セットアップ根治＋端末で速く」の第一候補。gasと集約再設計がトレードオフ")

fw_slide("④ halo2",
         "chip/gadgetを組んで回路化。PLONK(KZG=準備1回/IPA=不要)。Poseidon有・EdDSA自作寄り",
         ["lookup/custom gateなど高度で効率的な回路",
          "回路別セットアップ不要",
          "Moproでモバイル対応"],
         ["学習コストが高く実装が難しい",
          "Baby JubJubのEdDSA実装は手間大",
          "Groth16ほど検証は安くない"],
         "特殊機能が要らなければ費用対効果が悪い。積極採用しない")

fw_slide("⑤ Plonky2 / Plonky3",
         "Rustで記述、STARK(FRI)。PoseidonはGoldilocksネイティブ高速・再帰が速い",
         ["トラステッドセットアップ不要",
          "再帰が非常に速く集約を軽くできる",
          "証明生成が速い"],
         ["Goldilocks体で現行EdDSA(BN254系)と不一致→非ネイティブ演算で高コスト",
          "証明が大きくEVM検証はwrapが必要で高価",
          "低レベルで記述が重い"],
         "署名を含む端末回路には不向き。集約用途としては魅力")

fw_slide("⑥ Nova / Sonobe（Folding）",
         "計算を1ステップ回路に書きN個を逐次fold→最終1証明。arkworksでBN254/BabyJubJub",
         ["省メモリで端末やコーディネータ端末で集約できる",
          "Folding自体は回路別セットアップ不要",
          "「N端末を1つに」に そのまま合う"],
         ["ステップ回路への再設計が必要・事例が少ない",
          "最終圧縮の証明系でセットアップが要る場合あり"],
         "集約(M2)の代替として最有力の研究的選択肢")

fw_slide("⑦ RISC Zero / SP1（zkVM）",
         "普通のRustコードの実行の正しさをSTARKで証明。回路を意識しない",
         ["回路設計が要らず複雑なロジックを普通のコードで",
          "セットアップ不要・汎用性が非常に高い"],
         ["端末での証明生成が重い(今はサーバ/GPU前提)",
          "証明が大きくEVM検証は圧縮が必要"],
         "将来枠。今のモバイル端末証明には重い")

fw_slide("⑧ Cairo / Stwo（Starknet）",
         "Cairoで記述、STARK。Poseidonネイティブ。モバイル大規模実証あり(5秒以内)",
         ["セットアップ不要・モバイル証明の実績が豊富",
          "Poseidonが速い"],
         ["署名はStark曲線がnative・Baby JubJubは自作寄り",
          "検証がStarknet前提でEVM+gnarkとエコシステムがずれる"],
         "決済/集約レイヤをStarknetに替える話。EVM前提と不整合、見送り寄り")

fw_slide("⑨ Mopro（モバイル配布ツールキット）",
         "証明系でなく、circom/halo2/Noir等をモバイルアプリのバインディングに変換する橋渡し",
         ["現行circom回路をそのままモバイル化＝移行コスト最小",
          "Flutterパッケージがあり今のアプリと直結",
          "後からNoir/halo2へ差し替えも同じツール"],
         ["証明系の性質(セットアップ・gas)は中身次第",
          "iOSシミュレータ非対応・Androidはarm64"],
         "端末証明を今すぐ堅くする最短路")

# S12 その他・新顔
s = add_slide(); title(s, "その他・新顔（参考 — 主軸は変えない）")
bullets(s, TOP, [
    (0, "モバイルで一応検討価値がある新顔"),
    (1, "Ligero/Ligetron = 低メモリ・constrained device向け設計"),
    (1, "o1js(Mina) = TypeScriptでZK・準備不要・再帰。ただし検証はMina前提"),
    (0, "集約(層3)を攻めるなら"),
    (1, "ProtoStar/ProtoGalaxy/HyperNova = Folding改良、Spartan = 準備不要SNARK"),
    (0, "zkVM他 = Jolt / Miden / Triton / Nexus / Powdr / zkWASM"),
    (0, "プローバ/lib = Expander(GKR最速級) / Binius(二進体) / Winterfell 他"),
    (0, "言語系 = Leo(Aleo) / zkLLVM・Placeholder(=nil;)"),
    (0, "結論: 主軸(Mopro＋circom/Noir・集約gnark/folding・EVM前提)は不変"),
], size0=15, size1=13.5)

# S13 目的別のおすすめ（結論）
s = add_slide(); title(s, "目的別のおすすめ（結論）")
y = table(s, TOP, [
    ["やりたいこと", "フレームワーク"],
    ["今のまま端末証明を堅く・モバイル最適化", "Mopro（circom+rapidsnark, Flutter）"],
    ["★1トラステッドセットアップを根治", "Noir + Barretenberg（gas増は覚悟）"],
    ["多端末集約を端末側・省メモリで", "Nova / Sonobe"],
    ["集約の再帰を速く・準備不要に", "Plonky2（or 現行 gnark 継続）"],
    ["複雑ロジックを普通のコードで（将来）", "RISC Zero / SP1"],
], col_w=[4.3, 4.3], fs=12.5, row_h=0.5)
bullets(s, y + 0.1, [
    (0, "総合の現実解 = 端末はMopro、集約はgnark継続 or Novaで端末寄せ、★1消すならNoir段階導入"),
], size0=14.5)

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_02_フレームワーク比較.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
