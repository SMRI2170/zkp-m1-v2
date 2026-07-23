#!/usr/bin/env python3
"""week19/03_モバイル実行の方法.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "モバイル実行の方法", 32, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.4)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "各フレームワークを実機（Flutter）で動かすには", 16, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 03_モバイル実行の方法", 12, GRAY)

# S2 前提
s = add_slide(); title(s, "前提 — 「モバイルでZKP」とは何をすることか")
y = code(s, TOP, [
    "端末でやること : 入力 → witness生成 → proof生成   （重い＝本題）",
    "端末外でやること: verify（軽い・普通はチェーン/サーバ、端末でも可）",
])
y = bullets(s, y + 0.08, [
    (0, "モバイル化 = 証明生成器(prover)を端末で動くバイナリにして呼ぶこと"),
    (1, "ネイティブ = arm64にクロスコンパイル（速い・本命）"),
    (1, "WASM = WebAssembly（移植は楽だがネイティブ比3.5〜4倍遅い）"),
    (0, "共通で効く制約"),
    (1, "RAMが支配的（3GB以上で安定・FibRace実測）"),
    (1, "iOSシミュレータは基本不可・Androidはarm64"),
    (1, "witness生成とproof生成は別物（witnesscalc等で軽くできる）"),
])

# S3 Flutter への繋ぎ方
s = add_slide(); title(s, "Flutter への繋ぎ方（3パターン）")
y = table(s, TOP, [
    ["方法", "中身", "向き"],
    ["Mopro の Flutter パッケージ", "prover を Swift/Kotlin/Dart バインディングに自動生成", "最短・推奨"],
    ["自前 FFI（dart:ffi）", "arm64 の .so/.a を自分で作り Dart から呼ぶ", "細かく制御したい時"],
    ["MethodChannel プラグイン", "ネイティブ側に実装し channel 越しに呼ぶ", "既存ネイティブ実装がある時"],
], col_w=[2.6, 4.4, 1.6], fs=12, row_h=0.62)
y = bullets(s, y + 0.15, [
    (0, "現行アプリは rapidsnark を端末で動かす構成（③に近い）"),
    (0, "Mopro に寄せると①で標準化できる（APKビルドまで確認済 → 08）"),
])

# S4 まとめ表
s = add_slide(); title(s, "まとめ表（モバイル実行の観点）")
table(s, TOP, [
    ["FW", "端末証明", "モバイル化の方法", "Flutter連携", "一言"],
    ["circom+rapidsnark", "◎", "元からarm64ネイティブ", "Mopro / FFI", "今すぐ動く"],
    ["Noir+bb", "◎", "bbをarm64ネイティブ", "Mopro(Flutter)", "準備不要・速い"],
    ["halo2", "○", "Rust→arm64", "Mopro", "回路作りが重い"],
    ["Plonky2/3", "△", "Rust→arm64(FFI)", "自前FFI", "体不一致で不向き"],
    ["Nova/Sonobe", "◎(省メモリ)", "Rust→arm64(FFI)", "自前FFI", "端末側集約向き"],
    ["RISC Zero/SP1", "✕", "実質リモートprover", "API呼び出し", "端末単独は不可"],
    ["Cairo/Stwo", "◎", "Starknet SDK", "別SDK", "土台がStarknet"],
    ["gnark", "✕", "gomobile(重い)", "—", "集約サーバ専用"],
    ["Mopro", "—", "上記を包む", "◎公式Flutter", "配布を一本化"],
], col_w=[2.1, 1.3, 2.2, 1.6, 1.7], fs=10.5, row_h=0.4)

# S5 難易度
s = add_slide(); title(s, "モバイル実装の難易度（分岐点＝Mopro）")
y = table(s, TOP, [
    ["FW", "難易度", "理由"],
    ["circom + rapidsnark", "★☆☆ 易", "元からarm64・Mopro公式Flutter対応"],
    ["Noir + bb", "★☆☆ 易", "bbにモバイルビルド・Mopro公式Flutter対応"],
    ["Cairo / Stwo", "★★☆ 中", "実証は豊富だがSDKがStarknet系＝土台乗り換え"],
    ["halo2", "★★☆ 中", "Rust→arm64は可だが回路実装(chip)が難しい"],
    ["Plonky2 / Nova", "★★★ やや難", "Flutter成熟プラグイン無く自前FFI"],
    ["gnark / RISC Zero", "★★★ 難", "端末証明に不向き（サーバ/リモート前提）"],
], col_w=[2.4, 1.5, 4.7], fs=11.5, row_h=0.48)
y = bullets(s, y + 0.05, [
    (0, "公式対応は circom / halo2 / Noir の3つ。現行もNoirも「易」＝モバイル化リスクは低い"),
], size0=14)

# S6 結論
s = add_slide(); title(s, "あなたの Flutter アプリへの当てはめ（結論）")
bullets(s, TOP, [
    (0, "今の circom 回路をモバイルで堅くする"),
    (1, "→ Mopro の circom アダプタ＋mopro_flutter_package（移行最小）"),
    (0, "★1セットアップを消しつつ端末で速く"),
    (1, "→ Noir+bb（Mopro Flutter経由）、回路をNoirへ移植"),
    (0, "複数端末を端末側で1つに畳む実験"),
    (1, "→ Nova/Sonobe を FFI で（省メモリが効く）"),
    (0, "集約はサーバのままなら → gnark 継続（モバイル化しない）"),
    (0, "共通指針 = 端末証明はMoproで一本化・証明系は段階的に差し替え・RAM 3GB未満は要注意"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_03_モバイル実行の方法.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
