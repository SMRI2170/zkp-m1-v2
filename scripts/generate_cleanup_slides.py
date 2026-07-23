#!/usr/bin/env python3
"""不要になった機能の整理（削除）まとめスライド。

進化の過程で上位互換に置換/方針転換で孤立したものを削除。
現行フロー（多witness×ZK×チェーン＋耐改竄）に必要なものだけ残す。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、4:3。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.08)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 14, ACCENT, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.3, 8.2, 0.5)
run(tf.paragraphs[0], "リポジトリ整理", 15, GRAY)
tf = tbox(s, 0.88, 2.85, 8.4, 1.2)
run(tf.paragraphs[0], "不要になった機能の削除", 30, INK, bold=True)
hrule(s, 4.2, l=0.92, w=5.0, color=ACCENT)
tf = tbox(s, 0.92, 4.38, 8.2, 1.4)
run(tf.paragraphs[0], "上位互換に置換 / 方針転換で孤立 したものを削除し、現行フローに絞る", 14, INK)

# 2 削除方針
s = add_slide()
title(s, "削除の方針")
tf = tbox(s, 0.7, 1.5, 8.7, 3.2)
bullet(tf, "上位互換に置換済み → 旧版を削除", first=True, bold=True)
bullet(tf, "方針転換で現行フローから孤立 → 削除")
bullet(tf, "セキュリティ上のフットガン → 削除（残すと誤用リスク）")
bullet(tf, "現行フロー（多witness×ZK×チェーン＋耐改竄）に必要なものは残す")
tf = tbox(s, 0.7, 4.9, 8.7, 1.4)
subhead(tf, "判断基準", 0)
bullet(tf, "実コードからの参照有無 / 直系の後継があるか / 安全性", first=True)

# 3 削除したもの
s = add_slide()
title(s, "削除したもの")
table(s, 0.7, 1.4, 8.6, [
    ["#", "対象", "理由"],
    ["1", "remoteid_consistency（固定TOL）", "相対許容版に置換（RSSIに不適合・実参照なし）"],
    ["2", "build_witness_attestation_assets.sh", "決定論的setup＝★1脆弱性。セレモニーに置換"],
    ["3", "旧単体回路 distance/detection 4種＋build", "uav_witness_attestation に機能吸収"],
    ["4", "音響 feature_classifier ＋ scripts/acoustic", "方針転換（音響なし）で孤立"],
    ["+", "変換/比較 script・旧サンプル・設計md", "上記に付随（旧回路専用）"],
], col_w=[0.5, 4.0, 4.1], fs=10, row_h=0.52)
tf = tbox(s, 0.7, 4.9, 8.7, 1.4)
bullet(tf, "特に #2 は「残すと偽造可能な鍵を作れる」フットガンなので優先削除", size=12, first=True, bold=True)

# 4 残したもの
s = add_slide()
title(s, "残したもの（現行フロー）")
tf = tbox(s, 0.7, 1.5, 8.7, 4.4)
subhead(tf, "回路", 0)
bullet(tf, "uav_witness_attestation.circom（witness本体・public18）", first=True)
bullet(tf, "remoteid_consistency_rel.circom（RSSI整合・相対許容）")
bullet(tf, "lib_uav_safety.circom（_rel が include する土台）")
subhead(tf, "ビルド / ツール")
bullet(tf, "scripts/ceremony/（正規セレモニーの安全なビルド）", first=True)
bullet(tf, "gen_witness_sample.js / aggregator（M2）/ contracts / scripts/attestation")
subhead(tf, "確認")
bullet(tf, "現行フローが壊れないことをコンパイル・テストで検証", first=True, bold=True)

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "cleanup.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
