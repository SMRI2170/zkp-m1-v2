#!/usr/bin/env python3
"""week19/01_手法の全体像.md のスライド化。

進捗第11回スタイル: 白背景・黒文字・Yu Gothic / Courier New、4:3。
箇条書き ■=第1階層 / ・=第2階層（インデント）。体言止め。式/コード/図は薄グレー枠。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12.5, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "手法の全体像", 32, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.0)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(1.0)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "多端末の検証を1つにまとめて共有・検証する", 16, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "― 何があって、何ができるのか ―", 14, GRAY)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 01_手法の全体像", 12, GRAY)

# S2 やりたいこと（前提）
s = add_slide(); title(s, "やりたいこと（前提）")
y = code(s, TOP, [
    "各端末: RSSI検証＋距離検証＋複雑計算 を秘匿したまま証明   ← ZKP",
    "複数端末を1つにまとめる（同じドローンにK台が整合）        ← 集約の手法",
    "みんなで共有・検証（ブロックチェーンに記録）              ← 検証・共有の手法",
])
y = bullets(s, y + 0.1, [
    (0, "手法は大きく4つの層に分かれる"),
    (1, "層1 = ゼロ知識証明で何ができるか（道具）"),
    (1, "層2 = 証明系の種類（SNARK / STARK / Folding …）"),
    (1, "層3 = 複数の証明をまとめる手法（本案件の肝）"),
    (1, "層4 = 検証・共有の手法（オンチェーン / DA 層 …）"),
])

# S3 層1
s = add_slide(); title(s, "層1 — ゼロ知識証明(ZKP)で何ができるか")
y = bullets(s, TOP, [
    (0, "「中身を見せずに、計算が正しいことだけを証明する」道具"),
    (0, "秘匿 = 座標・生RSSIを出さず「整合(1)/矛盾(0)」だけを示す"),
    (0, "正しさの保証 = 「30m以内かつ相対帯で整合」をごまかせない形で証明"),
    (0, "選択的開示 = drone_id・しきい値だけを公開、他は隠す"),
    (0, "集約の土台 = 「他の証明が正しい」ことすら証明できる（→層3の再帰）"),
])
y = bullets(s, y + 0.1, [
    (0, "＝プライバシーを守ったまま、多端末の合意を検証可能にできる"),
])

# S4 層2
s = add_slide(); title(s, "層2 — 証明系の種類（性質の違うエンジン）")
y = table(s, TOP, [
    ["種類", "証明サイズ", "オンチェーン検証", "事前準備", "再帰", "ひとこと"],
    ["Groth16", "最小", "最安(≈348k)", "回路ごと", "可", "小さく安い・準備が弱点"],
    ["PLONK/Honk", "小〜中", "中〜高", "1回", "可", "準備が緩い・汎用"],
    ["STARK", "大", "高", "不要", "得意", "準備不要・証明が大"],
    ["Folding", "最終のみ", "圧縮依存", "ほぼ不要", "最得意", "N個を逐次畳む・省メモリ"],
], col_w=[1.5, 1.2, 1.7, 1.2, 0.8, 2.2], fs=11.5, row_h=0.55)
y = bullets(s, y + 0.05, [
    (0, "現行の「回路別セレモニー(★1脆弱性)」は PLONK/STARK/Folding で構造的に消える"),
], size0=15)

# S5 層3
s = add_slide(); title(s, "層3 — 複数の証明をまとめる手法（肝）")
y = table(s, TOP, [
    ["#", "手法", "まとめる場所", "オンチェーン", "現状"],
    ["1", "オンチェーン集計(K-of-N)", "契約", "台数分 O(N)", "実装済 M1"],
    ["2", "回路内バッチ(1回路にN端末)", "単一プローバ", "1回・定数", "試作"],
    ["3", "再帰的証明合成", "集約器", "1回・定数", "実装済 M2"],
    ["4", "Folding(Nova系)", "端末/コーディネータ", "最終1回", "未"],
    ["5", "専用証明集約(SnarkPack)", "集約器", "1回・対数", "未"],
    ["6", "非ZKしきい値合意(閾値署名)", "端末群", "署名1個", "未"],
], col_w=[0.4, 3.0, 2.2, 1.6, 1.4], fs=11, row_h=0.44)
y = bullets(s, y + 0.03, [
    (0, "1(集計)と3(再帰)は実装済。未検討は4(Folding)/5/6"),
], size0=14.5)

# S6 層4
s = add_slide(); title(s, "層4 — 検証・共有の手法")
y = table(s, TOP, [
    ["手法", "できること", "効く攻撃", "現状"],
    ["オンチェーンverifier", "契約が証明を検証し記録", "改ざん(暗号で不可)", "実装済"],
    ["オフチェーン検証＋コミット", "重い検証は外で結果だけ記録", "コスト削減", "任意"],
    ["DA層に提出集合コミット(網羅性)", "全員分を含めた証明", "検閲(間引き)", "設計済 #3b"],
    ["許可型チェーン(BFT合意+全複製)", "参加者全員が同結果を保持", "検閲・可用性", "設計済"],
], col_w=[3.0, 2.9, 1.7, 1.0], fs=11, row_h=0.5)

# S7 まとめ
s = add_slide(); title(s, "まとめ — 選ぶ順序")
y = code(s, TOP, [
    "① ZKPで何を秘匿し何を示すか（層1）  … 済: 座標/RSSIを隠し「整合」だけ公開",
    "② どの証明系エンジンか（層2）        … 現行Groth16 / ★1消すならPLONK系等",
    "③ 多端末をどうまとめるか（層3）      … 済:1(集計)3(再帰) / 未:4(Folding)5,6",
    "④ どう検証・共有するか（層4）        … 済: オンチェーン＋網羅性(#3b)",
])
y = bullets(s, y + 0.1, [
    (0, "フレームワークは②〜③を実装する“乗り物”にすぎない"),
    (0, "未検討で価値ありは 手法4(Folding) ＝ 端末側で省メモリに集約"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_01_手法の全体像.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
