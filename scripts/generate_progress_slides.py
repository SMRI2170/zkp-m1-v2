"""
generate_progress_slides.py
zk-UAV-Proximity-Proof 開発進捗スライド生成スクリプト
日付: 2026-05-16
"""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT_PATH = Path(__file__).parent.parent / "progress_slides_2026-05-16.pptx"

# ── サイズ ──────────────────────────────────────────────────────────────────
W  = Inches(13.33)   # 16:9
H  = Inches(7.5)
M  = Inches(0.5)
IW = W - M * 2

# ── カラーパレット ───────────────────────────────────────────────────────────
C_BG        = RGBColor(0xF8, 0xF9, 0xFC)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK      = RGBColor(0x12, 0x1A, 0x2E)   # 濃紺（タイトル背景）
C_NAVY      = RGBColor(0x1A, 0x23, 0x7E)
C_BLUE      = RGBColor(0x19, 0x76, 0xD2)
C_TEAL      = RGBColor(0x00, 0x89, 0x7B)
C_GREEN     = RGBColor(0x1B, 0x87, 0x4B)
C_ORANGE    = RGBColor(0xE6, 0x5C, 0x00)
C_RED       = RGBColor(0xC6, 0x28, 0x28)
C_PURPLE    = RGBColor(0x6A, 0x1B, 0x9A)
C_TEXT      = RGBColor(0x1A, 0x1A, 0x2E)
C_SUBTEXT   = RGBColor(0x5C, 0x67, 0x65)
C_RULE      = RGBColor(0xCC, 0xD5, 0xE8)
C_CODE_BG   = RGBColor(0xF0, 0xF4, 0xFA)
C_TAG_BG    = RGBColor(0xE3, 0xF2, 0xFD)
C_TAG_TEXT  = RGBColor(0x0D, 0x47, 0xA1)
C_OK_BG     = RGBColor(0xE8, 0xF5, 0xE9)
C_OK_TEXT   = RGBColor(0x1B, 0x5E, 0x20)
C_NG_BG     = RGBColor(0xFF, 0xEB, 0xEE)
C_NG_TEXT   = RGBColor(0xB7, 0x1C, 0x1C)
C_TBLH_BG   = RGBColor(0x1A, 0x23, 0x7E)
C_TBLH_TXT  = RGBColor(0xFF, 0xFF, 0xFF)
C_TBL_ROW1  = RGBColor(0xFF, 0xFF, 0xFF)
C_TBL_ROW2  = RGBColor(0xEE, 0xF2, 0xFF)
C_TBL_TXT   = RGBColor(0x21, 0x21, 0x21)

FONT = "Hiragino Sans"
FONT_MONO = "Courier New"

# ── ヘルパー ─────────────────────────────────────────────────────────────────
def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color=None):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color or C_BG

def rect(slide, l, t, w, h, color, border_color=None, radius=0):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if border_color:
        s.line.width = Pt(0.75)
        s.line.color.rgb = border_color
    else:
        s.line.fill.background()
    return s

def txb(slide, text, l, t, w, h,
        size=13, bold=False, color=None, align=PP_ALIGN.LEFT,
        italic=False, wrap=True, mono=False):
    color = color or C_TEXT
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT_MONO if mono else FONT
    return box

def bullets(slide, items, l, t, w, h, size=12, color=None, indent_char="  "):
    """items: list of (text, level) tuples or plain strings"""
    color = color or C_TEXT
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        prefix = "  " * level + ("• " if level > 0 else "▸ ")
        r = p.add_run()
        r.text = prefix + text
        r.font.size = Pt(size - level)
        r.font.color.rgb = color
        r.font.name = FONT
        p.space_after = Pt(3 + (2 - level) * 2)
    return box

def code_block(slide, lines, l, t, w, h, size=10):
    rect(slide, l, t, w, h, C_CODE_BG, C_RULE)
    rect(slide, l, t, Inches(0.06), h, C_BLUE)
    box = slide.shapes.add_textbox(l + Inches(0.15), t + Inches(0.1),
                                    w - Inches(0.25), h - Inches(0.15))
    tf = box.text_frame
    tf.word_wrap = False
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = FONT_MONO
        r.font.color.rgb = C_TEXT
        p.space_after = Pt(1)

def tbl(slide, headers, rows, l, t, w, col_widths=None, size=10, row_h=Inches(0.38)):
    nr = len(rows) + 1
    nc = len(headers)
    shape = slide.shapes.add_table(nr, nc, l, t, w, row_h * nr)
    tb = shape.table

    if col_widths:
        for i, cw in enumerate(col_widths):
            tb.columns[i].width = int(cw)
    else:
        for col in tb.columns:
            col.width = int(w // nc)

    def cell_fmt(cell, text, header=False, alt=False, center=False):
        cell.fill.solid()
        cell.fill.fore_color.rgb = (C_TBLH_BG if header
                                    else C_TBL_ROW2 if alt
                                    else C_TBL_ROW1)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if (header or center) else PP_ALIGN.LEFT
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = header
        r.font.color.rgb = C_TBLH_TXT if header else C_TBL_TXT
        r.font.name = FONT

    for ci, h in enumerate(headers):
        cell_fmt(tb.cell(0, ci), h, header=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell_fmt(tb.cell(ri + 1, ci), val, alt=(ri % 2 == 1))

def slide_header(slide, title, subtitle=None, date_str=None):
    bg(slide)
    rect(slide, 0, 0, W, Inches(0.06), C_NAVY)
    rect(slide, 0, Inches(0.06), Inches(0.06), H - Inches(0.06), C_BLUE)
    txb(slide, title, M, Inches(0.1), IW * 0.75, Inches(0.6),
        size=22, bold=True, color=C_NAVY)
    if subtitle:
        txb(slide, subtitle, M, Inches(0.68), IW * 0.75, Inches(0.3),
            size=11, color=C_SUBTEXT)
    if date_str:
        txb(slide, date_str, W - Inches(2.8), Inches(0.15), Inches(2.3), Inches(0.35),
            size=10, color=C_SUBTEXT, align=PP_ALIGN.RIGHT)
    rect(slide, M, Inches(1.0), IW, Pt(1.0), C_RULE)

def pg(slide, n, total):
    txb(slide, f"{n} / {total}", W - Inches(0.7), H - Inches(0.32),
        Inches(0.6), Inches(0.25), size=9, color=C_SUBTEXT, align=PP_ALIGN.RIGHT)

def section_divider(prs, label, color=C_NAVY, sub=""):
    s = blank_slide(prs)
    bg(s, color)
    txb(s, label,
        Inches(1.2), Inches(2.8), W - Inches(2.4), Inches(1.2),
        size=36, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    if sub:
        txb(s, sub,
            Inches(1.2), Inches(4.0), W - Inches(2.4), Inches(0.5),
            size=16, color=RGBColor(0xB0, 0xC8, 0xFF), align=PP_ALIGN.CENTER)
    return s

def badge(slide, text, l, t, color_bg, color_text, size=9):
    w = Inches(len(text) * 0.10 + 0.3)
    h = Inches(0.24)
    rect(slide, l, t, w, h, color_bg)
    txb(slide, text, l, t, w, h, size=size, bold=True,
        color=color_text, align=PP_ALIGN.CENTER)
    return w

# ─────────────────────────────────────────────────────────────────────────────
TOTAL = 18

def build():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    # =========================================================================
    # Slide 1: タイトル
    # =========================================================================
    s = blank_slide(prs)
    bg(s, C_DARK)
    rect(s, 0, 0, W, Inches(0.08), C_TEAL)
    rect(s, 0, H - Inches(0.08), W, Inches(0.08), C_TEAL)

    txb(s, "開発進捗レポート",
        M, Inches(1.5), IW, Inches(0.5),
        size=14, color=RGBColor(0x80, 0xCB, 0xC4))
    txb(s, "zk-UAV-Proximity-Proof",
        M, Inches(2.0), IW, Inches(1.0),
        size=44, bold=True, color=C_WHITE)
    txb(s, "UI/UX 刷新 ・ 動画→静止画リファクタ ・ 実ドローン Remote ID 完全対応",
        M, Inches(3.1), IW, Inches(0.5),
        size=16, color=RGBColor(0xB0, 0xC8, 0xFF))

    rect(s, M, Inches(3.8), Inches(2.0), Pt(2), C_TEAL)

    for i, (d, c) in enumerate([
        ("2026-05-02", C_BLUE),
        ("2026-05-15", C_TEAL),
        ("2026-05-16", C_GREEN),
    ]):
        txb(s, d, M + Inches(2.5 * i), Inches(4.2), Inches(2.2), Inches(0.3),
            size=12, bold=True, color=c)

    txb(s, "Circom / Flutter / OpenDroneID / YOLO / ZKP",
        M, H - Inches(1.0), IW, Inches(0.3),
        size=11, color=RGBColor(0x60, 0x70, 0x90))

    # =========================================================================
    # Slide 2: 変更サマリー タイムライン
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "変更サマリー", "3セッションにわたる主な改善")
    pg(s, 2, TOTAL)

    milestones = [
        ("2026-05-02", C_BLUE,   "実ドローン Remote ID 対応",
         ["Message Pack (type=0xF) 解析",
          "System メッセージ (type=4) パース",
          "DJI 旧機種 manufacturerData 対応",
          "BLE スキャンギャップ解消",
          "canProve ワークフロー修正"]),
        ("2026-05-15\n前半", C_TEAL, "動画録画 → 静止画撮影",
         ["VideoEvidence → PhotoEvidence",
          "_captureAndDetect() 新設",
          "GPS・BLE 履歴の常時蓄積化",
          "_analyzePhotoEvidence() 実装",
          "録画 UI 廃止・撮影 UI 新設"]),
        ("2026-05-15\n後半", C_GREEN, "UI/UX 直感化 & クリーンアップ",
         ["4ステップ進捗バー追加",
          "シャッターボタン一眼カメラ風",
          "ZK 証明ボタン全幅・分離",
          "詳細タブ: 検知結果・判定要約",
          "unused_field 警告 24→0件"]),
    ]

    card_w = Inches(3.7)
    for i, (date, color, title, pts) in enumerate(milestones):
        x = M + Inches(0.15) + Inches(4.1 * i)
        rect(s, x, Inches(1.2), card_w, Inches(0.5), color)
        txb(s, date, x, Inches(1.23), Inches(1.0), Inches(0.45),
            size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        txb(s, title, x + Inches(1.05), Inches(1.28), card_w - Inches(1.1), Inches(0.4),
            size=12, bold=True, color=C_WHITE)
        rect(s, x, Inches(1.7), card_w, Inches(4.8), C_WHITE, C_RULE)
        for j, pt in enumerate(pts):
            yy = Inches(1.85) + Inches(0.56 * j)
            rect(s, x + Inches(0.12), yy + Inches(0.07),
                 Inches(0.18), Inches(0.18), color)
            txb(s, pt, x + Inches(0.38), yy, card_w - Inches(0.5), Inches(0.5),
                size=11, color=C_TEXT)

    # =========================================================================
    # Section: 2026-05-02
    # =========================================================================
    section_divider(prs, "① 実ドローン Remote ID 完全対応",
                    C_NAVY, "2026-05-02")

    # =========================================================================
    # Slide 4: 検出ギャップ一覧
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "実ドローン対応：検出ギャップと修正内容",
                 "Raspberry Pi では検出できていたが実機ドローンで取れていなかった問題を全修正",
                 "2026-05-02")
    pg(s, 4, TOTAL)

    tbl(s,
        ["問題", "原因", "対応関数", "結果"],
        [
            ["Message Pack\n(type=0xF) 未対応",
             "単一メッセージしか想定していなかった",
             "extractAllMessages()",
             "BasicID+Location+System を一括展開"],
            ["System メッセージ\n(type=4) 破棄",
             "type≠0, ≠1 は null 扱いで捨てていた",
             "tryParseSystem()",
             "オペレーター位置・高度をパース"],
            ["DJI 旧機種\n(manufacturerData) 未対応",
             "serviceData のみを探索していた",
             "extractFromManufacturerData()",
             "Company ID 0x2005 を追加探索"],
            ["BLE スキャンギャップ",
             "22 秒周期タイマーで空白が発生",
             "_startContinuousBleScan()",
             "1秒ポーリング + isScanningNow で即再開"],
            ["BasicID ログなし",
             "type=0 パケットを無視していた",
             "_restartBleScan() 内",
             "UAS_ID をログに記録"],
        ],
        M, Inches(1.15), IW,
        col_widths=[Inches(2.5), Inches(3.5), Inches(3.0), Inches(3.3)],
        size=10, row_h=Inches(0.58),
    )

    # =========================================================================
    # Slide 5: BLE 検出フロー（修正後）
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "BLE 検出フロー（修正後）",
                 "serviceData → manufacturerData のフォールバック + Message Pack 展開",
                 "2026-05-02")
    pg(s, 5, TOTAL)

    code_block(s, [
        "// ① serviceData から標準 OpenDroneID (UUID 0xFFFA) を探す",
        "// ② なければ manufacturerData から DJI 旧機種 (0x2005) を探す",
        "final odidBytes =",
        "    UavTelemetryPacket.extractFromServiceData(ad.serviceData) ??",
        "    UavTelemetryPacket.extractFromManufacturerData(ad.manufacturerData);",
        "if (odidBytes == null) continue;",
        "",
        "// 単一メッセージ / Message Pack を統一処理",
        "final messages = UavTelemetryPacket.extractAllMessages(odidBytes);",
        "for (final msg in messages) {",
        "  final msgType = msg.isNotEmpty ? (msg[0] >> 4) & 0x0F : -1;",
        "  if (msgType == 0) { /* BasicID → UAS_ID をログ */ }",
        "  if (msgType == 1) { /* Location → 位置・高度・時刻を UI に反映 */ }",
        "  if (msgType == 4) { /* System  → オペレーター位置をログ */ }",
        "}",
    ], M, Inches(1.15), IW * 0.62, Inches(4.8), size=10)

    # 右側: メッセージタイプ説明
    rx = M + IW * 0.65
    rw = IW * 0.33
    for i, (mtype, label, color, desc) in enumerate([
        ("type 0", "BasicID",  C_NAVY,   "UAS_ID (ドローン識別子)"),
        ("type 1", "Location", C_BLUE,   "緯度・経度・高度・時刻"),
        ("type 4", "System",   C_PURPLE, "オペレーター位置・高度"),
        ("type F", "MessagePack", C_TEAL, "type0+1+4 を1パケットに集約"),
    ]):
        yy = Inches(1.2) + Inches(1.2 * i)
        rect(s, rx, yy, Inches(0.9), Inches(0.35), color)
        txb(s, mtype, rx, yy, Inches(0.9), Inches(0.35),
            size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        rect(s, rx + Inches(0.95), yy, rw - Inches(1.0), Inches(0.35),
             C_CODE_BG, C_RULE)
        txb(s, f"{label}  —  {desc}",
            rx + Inches(1.0), yy, rw - Inches(1.05), Inches(0.35),
            size=10, color=C_TEXT)

    txb(s, "★ 実機ドローンは多くの場合 type=F (Message Pack) で送信",
        rx, Inches(6.1), rw, Inches(0.35),
        size=10, bold=True, color=C_ORANGE)

    # =========================================================================
    # Slide 6: canProve ワークフロー修正
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "canProve ワークフロー修正",
                 "録画なしでも ZK 証明を実行できるよう条件を修正",
                 "2026-05-02")
    pg(s, 6, TOTAL)

    # Before
    txb(s, "修正前（問題）", M, Inches(1.2), IW * 0.46, Inches(0.3),
        size=13, bold=True, color=C_RED)
    code_block(s, [
        "// 録画が必須になっていた",
        "final canProve = droneReady",
        "    && locationReady",
        "    && _videoEvidence != null;  // ← 常に false",
        "",
        "// 結果: Base 回路でも証明ボタンが押せなかった",
    ], M, Inches(1.55), IW * 0.46, Inches(2.2))

    rect(s, M + IW * 0.48, Inches(2.2), Inches(0.5), Inches(0.5),
         C_CODE_BG, C_RULE)
    txb(s, "→", M + IW * 0.48, Inches(2.25), Inches(0.5), Inches(0.4),
        size=22, bold=True, color=C_BLUE, align=PP_ALIGN.CENTER)

    # After
    txb(s, "修正後（正しい挙動）", M + IW * 0.52, Inches(1.2), IW * 0.46, Inches(0.3),
        size=13, bold=True, color=C_GREEN)
    code_block(s, [
        "// BLE + GPS が揃えば証明可能",
        "final canProve = droneReady",
        "    && locationReady;",
        "",
        "// 写真は任意: あれば回路入力が豊かになる",
        "// なければデフォルト値で補完して証明",
    ], M + IW * 0.52, Inches(1.55), IW * 0.46, Inches(2.2))

    # 説明
    rect(s, M, Inches(3.9), IW, Inches(1.8), C_CODE_BG, C_RULE)
    rect(s, M, Inches(3.9), Inches(0.06), Inches(1.8), C_TEAL)
    txb(s, "設計方針",
        M + Inches(0.2), Inches(4.0), IW - Inches(0.4), Inches(0.3),
        size=12, bold=True, color=C_NAVY)
    bullets(s, [
        "BLE (Remote ID) と GPS (ユーザ位置) が揃えば距離・時刻差の ZK 証明は実行可能",
        ("写真・YOLO 検知は「証明の根拠強化」オプション — Base 回路では不要", 1),
        ("detectionCommitment 回路等では写真情報を使うが、canProve 条件には含めない", 1),
        "証明ボタンは droneReady && locationReady の 2 条件のみで有効化",
    ], M + Inches(0.2), Inches(4.35), IW - Inches(0.4), Inches(1.3), size=11)

    # =========================================================================
    # Section: 2026-05-15 前半
    # =========================================================================
    section_divider(prs, "② 動画録画 → 静止画撮影",
                    C_TEAL, "2026-05-15（前半）")

    # =========================================================================
    # Slide 8: リファクタ概要
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "動画録画 → 静止画撮影リファクタ",
                 "写真撮影・Remote ID 受信・GPS 記録を同時に行い ZK 回路へ代入",
                 "2026-05-15")
    pg(s, 8, TOTAL)

    # 設計思想
    rect(s, M, Inches(1.15), IW, Inches(0.9), C_TEAL)
    txb(s, "設計思想: 撮影ボタン 1 回で「カメラ・BLE・GPS」を同時スナップショット",
        M + Inches(0.2), Inches(1.25), IW - Inches(0.4), Inches(0.6),
        size=14, bold=True, color=C_WHITE)

    # フロー図
    steps = [
        ("takePicture()", C_NAVY, "カメラで\n写真を撮影"),
        ("_selectBestDataPair()", C_BLUE, "BLE+GPS の\n最適ペアを取得"),
        ("_analyzePhotoEvidence()", C_TEAL, "YOLO / ML Kit\nで UAV を検知"),
        ("ZK 証明へ", C_GREEN, "位置・時刻・\n検知結果を代入"),
    ]
    sw = Inches(2.5)
    for i, (func, color, desc) in enumerate(steps):
        x = M + Inches(0.3) + Inches(3.0 * i)
        rect(s, x, Inches(2.2), sw, Inches(0.5), color)
        txb(s, func, x, Inches(2.27), sw, Inches(0.4),
            size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, mono=True)
        rect(s, x, Inches(2.7), sw, Inches(1.0), C_WHITE, C_RULE)
        txb(s, desc, x, Inches(2.85), sw, Inches(0.8),
            size=11, color=C_TEXT, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            txb(s, "→", x + sw, Inches(2.6), Inches(0.5), Inches(0.4),
                size=20, bold=True, color=C_SUBTEXT, align=PP_ALIGN.CENTER)

    # API 変更表
    txb(s, "主な API 変更", M, Inches(3.85), IW, Inches(0.3),
        size=13, bold=True, color=C_NAVY)
    tbl(s,
        ["旧 (動画)", "新 (静止画)"],
        [
            ["VideoEvidence（録画パス・開始終了時刻）",
             "PhotoEvidence（写真パス・撮影時刻のみ）"],
            ["_captureEvidenceVideo()", "削除"],
            ["_analyzeVideoEvidence()", "_analyzePhotoEvidence()"],
            ["_startRecording() / _stopRecordingAndPrepare()", "削除（_captureAndDetect() に統合）"],
        ],
        M, Inches(4.2), IW,
        col_widths=[IW * 0.5, IW * 0.5],
        size=10, row_h=Inches(0.42),
    )

    # =========================================================================
    # Slide 9: 常時蓄積化
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "GPS・BLE 履歴の常時蓄積化",
                 "録画中フラグによるガードを撤廃し、アプリ起動直後から履歴を蓄積",
                 "2026-05-15")
    pg(s, 9, TOTAL)

    # Before / After コード
    txb(s, "修正前", M, Inches(1.15), IW * 0.46, Inches(0.28),
        size=12, bold=True, color=C_RED)
    code_block(s, [
        "// GPS (_startLocationStream)",
        "if (_isRecording) {",
        "  _locationHistory.add(snapshot);",
        "  if (_locationHistory.length > 20)",
        "    _locationHistory.removeAt(0);",
        "}",
        "",
        "// BLE (_applyDroneTelemetry)",
        "if (_isRecording) {",
        "  _droneHistory.add(packet);",
        "  if (_droneHistory.length > 20)",
        "    _droneHistory.removeAt(0);",
        "}",
    ], M, Inches(1.5), IW * 0.46, Inches(3.8))

    txb(s, "→", M + IW * 0.47, Inches(3.2), Inches(0.5), Inches(0.4),
        size=22, bold=True, color=C_BLUE, align=PP_ALIGN.CENTER)

    txb(s, "修正後", M + IW * 0.52, Inches(1.15), IW * 0.46, Inches(0.28),
        size=12, bold=True, color=C_GREEN)
    code_block(s, [
        "// GPS — 常時蓄積（最大20件）",
        "_locationHistory.add(snapshot);",
        "if (_locationHistory.length > 20)",
        "  _locationHistory.removeAt(0);",
        "",
        "",
        "",
        "// BLE — 常時蓄積（最大20件）",
        "_droneHistory.add(packet);",
        "if (_droneHistory.length > 20)",
        "  _droneHistory.removeAt(0);",
        "",
        "",
    ], M + IW * 0.52, Inches(1.5), IW * 0.46, Inches(3.8))

    # 効果
    rect(s, M, Inches(5.5), IW, Inches(1.0), C_OK_BG, C_RULE)
    rect(s, M, Inches(5.5), Inches(0.06), Inches(1.0), C_GREEN)
    txb(s, "効果：撮影ボタンを押した瞬間に「最良の BLE+GPS ペア」が必ず存在する。録画開始前にドローンを検知していても無駄にならない。",
        M + Inches(0.2), Inches(5.6), IW - Inches(0.4), Inches(0.8),
        size=11, color=C_OK_TEXT)

    # =========================================================================
    # Section: 2026-05-15 後半
    # =========================================================================
    section_divider(prs, "③ UI/UX 直感化 & クリーンアップ",
                    C_GREEN, "2026-05-15（後半）")

    # =========================================================================
    # Slide 11: メインタブ全面リデザイン
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "メインタブ全面リデザイン",
                 "「今何をすればいいか」が常にわかる状態駆動 UI へ",
                 "2026-05-15")
    pg(s, 11, TOTAL)

    # 左: 改善点リスト
    lw = IW * 0.50
    txb(s, "改善ポイント", M, Inches(1.15), lw, Inches(0.3),
        size=13, bold=True, color=C_NAVY)
    bullets(s, [
        "4ステップ進捗バーを上部に追加",
        ("① BLE  →  ② GPS  →  ③ 撮影  →  ④ 証明", 1),
        ("完了ステップはチェックマーク＋緑に変化", 1),
        "次のアクションヒントを常時表示",
        ("状態に応じて「BLE で Remote ID を受信してください」等を動的表示", 1),
        "シャッターボタンを一眼カメラ風の大型円形に変更",
        ("撮影前: 白丸  撮影中: オレンジスピナー", 1),
        ("撮影後: 左に「撮り直す」右に「UAV 検知バッジ」", 1),
        "ZK 証明ボタンを全幅・シャッターと明確に分離",
        ("条件が揃うと緑に変化 + verified アイコン", 1),
        ("処理中はインライン spinner＋経過時間", 1),
    ], M, Inches(1.5), lw, Inches(5.0), size=11)

    # 右: ステップバー模擬
    rx = M + lw + Inches(0.4)
    rw = IW - lw - Inches(0.4)

    txb(s, "ステップバー（上部）", rx, Inches(1.15), rw, Inches(0.3),
        size=11, bold=True, color=C_NAVY)

    step_labels = ["① BLE", "② GPS", "③ 撮影", "④ 証明"]
    step_colors = [C_GREEN, C_GREEN, C_BLUE, C_SUBTEXT]
    step_done   = [True, True, False, False]
    sw2 = (rw - Inches(0.2)) / 4
    for i, (lbl, col, done) in enumerate(zip(step_labels, step_colors, step_done)):
        xx = rx + Inches(0.1) + sw2 * i
        rect(s, xx + sw2 * 0.3, Inches(1.65), sw2 * 0.4, Inches(0.35), col)
        txb(s, "✓" if done else lbl.split()[0],
            xx + sw2 * 0.3, Inches(1.68), sw2 * 0.4, Inches(0.28),
            size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        txb(s, lbl, xx, Inches(2.05), sw2, Inches(0.25),
            size=9, color=col, align=PP_ALIGN.CENTER)
        if i < 3:
            rect(s, xx + sw2 * 0.7, Inches(1.79), sw2 * 0.3, Pt(1.5),
                 C_GREEN if done else C_RULE)

    # シャッターボタン模擬
    txb(s, "シャッターボタン（中央）", rx, Inches(2.45), rw, Inches(0.3),
        size=11, bold=True, color=C_NAVY)
    cx = rx + rw / 2
    rect(s, cx - Inches(0.55), Inches(2.8), Inches(1.1), Inches(1.1),
         C_SUBTEXT, C_WHITE)
    rect(s, cx - Inches(0.42), Inches(2.93), Inches(0.84), Inches(0.84),
         C_WHITE)
    txb(s, "撮影済", cx - Inches(0.55), Inches(4.05), Inches(1.1), Inches(0.3),
        size=9, color=C_GREEN, align=PP_ALIGN.CENTER)

    # 証明ボタン模擬
    txb(s, "証明ボタン（全幅）", rx, Inches(4.45), rw, Inches(0.3),
        size=11, bold=True, color=C_NAVY)
    rect(s, rx, Inches(4.8), rw, Inches(0.55), C_GREEN)
    txb(s, "✓ ZK 証明を実行", rx, Inches(4.88), rw, Inches(0.4),
        size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

    # =========================================================================
    # Slide 12: 詳細タブ拡充
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "詳細タブ拡充",
                 "計算されているが表示されていなかった状態を全て UI に接続",
                 "2026-05-15")
    pg(s, 12, TOTAL)

    lw = IW * 0.48

    # UAV 検知結果カード
    txb(s, "追加① UAV 検知結果カード（写真撮影後に表示）",
        M, Inches(1.15), lw, Inches(0.3), size=12, bold=True, color=C_NAVY)
    rect(s, M, Inches(1.5), lw, Inches(2.5), C_WHITE, C_RULE)
    rect(s, M, Inches(1.5), Inches(0.06), Inches(2.5), C_TEAL)
    bullets(s, [
        "_videoDetectionSummary  →  検知サマリー文",
        ("例: 「YOLO により写真内の UAV 候補を検知しました」", 1),
        "_videoDetectionLabels  →  ラベルチップ一覧",
        ("例: drone:0.87  uav:0.72  helicopter:0.61", 1),
        "_detectionBackendStatus  →  使用バックエンド",
        ("例: 「YOLO TFLite モデルを使用中」", 1),
    ], M + Inches(0.2), Inches(1.6), lw - Inches(0.3), Inches(2.3), size=10)

    # 判定要約カード
    txb(s, "追加② 判定要約・プライバシーカード（証明後に表示）",
        M, Inches(4.1), lw, Inches(0.3), size=12, bold=True, color=C_NAVY)
    rect(s, M, Inches(4.45), lw, Inches(2.6), C_WHITE, C_RULE)
    rect(s, M, Inches(4.45), Inches(0.06), Inches(2.6), C_NAVY)
    bullets(s, [
        "_verdictSummary  →  判定の根拠文",
        ("例: 「時刻は揃っていますが推定距離は 45.2m で…」", 1),
        "_privacySummary  →  プライバシー説明",
        ("例: 「Poseidon ハッシュにより座標の実値は非公開…」", 1),
        "ロック🔒アイコン付きで視覚的に分離",
    ], M + Inches(0.2), Inches(4.55), lw - Inches(0.3), Inches(2.4), size=10)

    # 右: 改善前後比較
    rx = M + lw + Inches(0.4)
    rw2 = IW - lw - Inches(0.4)
    txb(s, "改善前後", rx, Inches(1.15), rw2, Inches(0.3),
        size=12, bold=True, color=C_NAVY)

    items_before = [
        "_verdictSummary", "_privacySummary",
        "_videoDetectionSummary", "_videoDetectionLabels",
        "_detectionBackendStatus",
    ]
    txb(s, "改善前: 計算されるが表示なし", rx, Inches(1.5), rw2, Inches(0.28),
        size=10, bold=True, color=C_RED)
    for i, name in enumerate(items_before):
        rect(s, rx, Inches(1.85) + Inches(0.38 * i), rw2, Inches(0.32),
             C_NG_BG, C_RULE)
        txb(s, f"✗  {name}", rx + Inches(0.1), Inches(1.88) + Inches(0.38 * i),
            rw2 - Inches(0.15), Inches(0.28), size=9, color=C_RED, mono=True)

    txb(s, "改善後: 全て UI に接続", rx, Inches(3.9), rw2, Inches(0.28),
        size=10, bold=True, color=C_GREEN)
    for i, (name, dest) in enumerate([
        ("_verdictSummary", "→ 詳細タブ 判定カード"),
        ("_privacySummary", "→ 詳細タブ プライバシー欄"),
        ("_videoDetectionSummary", "→ 詳細タブ 検知カード"),
        ("_videoDetectionLabels", "→ 詳細タブ チップ一覧"),
        ("_detectionBackendStatus", "→ 詳細タブ バックエンド欄"),
    ]):
        rect(s, rx, Inches(4.25) + Inches(0.38 * i), rw2, Inches(0.32),
             C_OK_BG, C_RULE)
        txb(s, f"✓  {name}  {dest}",
            rx + Inches(0.1), Inches(4.28) + Inches(0.38 * i),
            rw2 - Inches(0.15), Inches(0.28), size=9, color=C_GREEN)

    # =========================================================================
    # Slide 13: コードクリーンアップ
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "コードクリーンアップ",
                 "デッドコード・重複フィールドを削除し unused_field 警告を 24 → 0 件に",
                 "2026-05-15")
    pg(s, 13, TOTAL)

    tbl(s,
        ["問題", "分類", "対応"],
        [
            ["tabs リストが build() 内で二重ビルド",
             "パフォーマンス",
             "変数削除 → 直接呼び出しに統一"],
            ["_isAnalyzingVideo が _isCapturing と完全に重複",
             "デッドフィールド",
             "フィールド・代入をすべて削除"],
            ["_currentPressureHpa が書くだけで読まれない",
             "デッドフィールド",
             "宣言・代入をすべて削除"],
            ["_lastSensorReadAt が書くだけで読まれない",
             "デッドフィールド",
             "宣言・代入をすべて削除"],
            ["_lastDroneSeenAt が書くだけで読まれない",
             "デッドフィールド",
             "宣言・代入をすべて削除"],
            ["ログに「録画終了」の記述が残存",
             "ステールコメント",
             "「撮影時点スナップショット」に修正"],
            ["initState コメントに「自動録画」が残存",
             "ステールコメント",
             "「カメラ初期化」に修正"],
        ],
        M, Inches(1.15), IW,
        col_widths=[Inches(5.5), Inches(2.3), Inches(4.5)],
        size=10, row_h=Inches(0.5),
    )

    rect(s, M, Inches(5.8), IW, Inches(0.55), C_OK_BG, C_RULE)
    rect(s, M, Inches(5.8), Inches(0.06), Inches(0.55), C_GREEN)
    txb(s, "結果:  flutter analyze の unused_field 警告  24 件 → 0 件",
        M + Inches(0.2), Inches(5.92), IW - Inches(0.4), Inches(0.35),
        size=13, bold=True, color=C_GREEN)

    # =========================================================================
    # Section: まとめ
    # =========================================================================
    section_divider(prs, "④ まとめ & 今後の課題",
                    C_DARK, "")

    # =========================================================================
    # Slide 15: 全体サマリー
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "全体サマリー", "3 セッションで達成した主な成果")
    pg(s, 15, TOTAL)

    cols = [
        (C_BLUE,   "2026-05-02\n実ドローン対応",
         ["Message Pack 解析対応",
          "System メッセージパース",
          "DJI 旧機種 対応",
          "BLE スキャンギャップ解消",
          "canProve ワークフロー修正"]),
        (C_TEAL,   "2026-05-15 前半\n動画→静止画",
         ["PhotoEvidence 新設",
          "_captureAndDetect() 1本化",
          "GPS・BLE 常時蓄積",
          "_analyzePhotoEvidence()",
          "録画 UI 廃止"]),
        (C_GREEN,  "2026-05-15 後半\nUI/UX & クリーン",
         ["4ステップ進捗バー",
          "シャッターボタン",
          "ZK 証明ボタン全幅",
          "詳細タブ拡充",
          "unused_field 0件"]),
    ]
    cw2 = Inches(3.8)
    for i, (color, title, pts) in enumerate(cols):
        x = M + Inches(0.2) + Inches(4.2 * i)
        rect(s, x, Inches(1.15), cw2, Inches(0.55), color)
        txb(s, title, x, Inches(1.2), cw2, Inches(0.5),
            size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        rect(s, x, Inches(1.7), cw2, Inches(4.3), C_WHITE, C_RULE)
        for j, pt in enumerate(pts):
            yy = Inches(1.85) + Inches(0.55 * j)
            rect(s, x + Inches(0.15), yy + Inches(0.1),
                 Inches(0.15), Inches(0.15), color)
            txb(s, pt, x + Inches(0.4), yy, cw2 - Inches(0.5), Inches(0.5),
                size=11, color=C_TEXT)

    # 数値サマリー
    kpis = [
        ("5", "検出ギャップ\n修正"),
        ("0", "unused_field\n警告"),
        ("1", "撮影ボタン\n(旧: 2)"),
        ("4", "ステップ\n可視化"),
    ]
    for i, (val, lbl) in enumerate(kpis):
        x = M + Inches(0.6) + Inches(2.9 * i)
        rect(s, x, Inches(6.1), Inches(2.2), Inches(0.95), C_CODE_BG, C_RULE)
        txb(s, val, x, Inches(6.12), Inches(1.0), Inches(0.7),
            size=36, bold=True, color=C_NAVY, align=PP_ALIGN.RIGHT)
        txb(s, lbl, x + Inches(1.05), Inches(6.2), Inches(1.1), Inches(0.65),
            size=10, color=C_SUBTEXT)

    # =========================================================================
    # Slide 16: 今後の課題
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "今後の課題・改善候補")
    pg(s, 16, TOTAL)

    items = [
        (C_BLUE,   "検出精度",
         ["YOLO モデルをドローン特化でファインチューニング（VisDrone データセット）",
          "夜間・逆光・小物体への対応強化",
          "TFLite GPU デリゲート有効化（現在 Android では OpenCL エラーで無効）"]),
        (C_ORANGE, "ZK 回路",
         ["Time Series 回路で複数フレームを一括証明（現状は静止画 1 枚）",
          "Poseidon コミットメントに写真ハッシュを組み込む",
          "証明時間の短縮（現状 ~2s、目標 <1s）"]),
        (C_PURPLE, "UI/UX",
         ["撮影した写真のプレビューをカメラ画面に重ねて表示",
          "YOLO 検知バウンディングボックスのオーバーレイ表示",
          "証明結果の QR コード / 共有機能"]),
        (C_TEAL,   "インフラ",
         ["OpenDroneID BLE 広告の完全規格準拠テスト（実機 DJI Mavic 等）",
          "Raspberry Pi GPS + BME280 の常時送信スクリプト安定化",
          "CI/CD で flutter analyze + テストを自動実行"]),
    ]

    rw3 = (IW - Inches(0.3)) / 2
    for i, (color, title, pts) in enumerate(items):
        col = i % 2
        row = i // 2
        x = M + Inches(0.15) + (rw3 + Inches(0.15)) * col
        y = Inches(1.2) + Inches(2.6) * row
        rect(s, x, y, rw3, Inches(0.38), color)
        txb(s, title, x + Inches(0.15), y + Inches(0.06), rw3 - Inches(0.3), Inches(0.28),
            size=12, bold=True, color=C_WHITE)
        rect(s, x, y + Inches(0.38), rw3, Inches(2.0), C_WHITE, C_RULE)
        for j, pt in enumerate(pts):
            yy = y + Inches(0.5) + Inches(0.55 * j)
            rect(s, x + Inches(0.12), yy + Inches(0.1),
                 Inches(0.12), Inches(0.12), color)
            txb(s, pt, x + Inches(0.32), yy, rw3 - Inches(0.45), Inches(0.5),
                size=10, color=C_TEXT)

    # =========================================================================
    # Slide 17: 技術スタック全体像
    # =========================================================================
    s = blank_slide(prs)
    slide_header(s, "技術スタック全体像", "本プロジェクトで使用している主要技術")
    pg(s, 17, TOTAL)

    layers = [
        (C_NAVY,   "ZK 証明層",
         "Circom  ·  Groth16  ·  rapidsnark  ·  circom_witnesscalc  ·  snarkjs"),
        (C_BLUE,   "Flutter アプリ層",
         "flutter_blue_plus  ·  camera  ·  geolocator  ·  tflite_flutter  ·  sensors_plus"),
        (C_TEAL,   "ML 検知層",
         "YOLO TFLite (VisDrone 事前学習)  ·  Google ML Kit Object Detection"),
        (C_GREEN,  "ドローン通信層",
         "OpenDroneID ASTM F3411-22a  ·  BLE Extended Advertising  ·  Message Pack"),
        (C_ORANGE, "UAV / センサー層",
         "Raspberry Pi  ·  GPS (gpsd)  ·  BME280 気圧センサー  ·  remote_id.py"),
    ]
    for i, (color, title, tech) in enumerate(layers):
        y = Inches(1.2) + Inches(1.08 * i)
        rect(s, M, y, Inches(2.0), Inches(0.85), color)
        txb(s, title, M, y + Inches(0.2), Inches(2.0), Inches(0.45),
            size=11, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        rect(s, M + Inches(2.05), y, IW - Inches(2.05), Inches(0.85),
             C_WHITE, C_RULE)
        rect(s, M + Inches(2.05), y, Inches(0.05), Inches(0.85), color)
        txb(s, tech,
            M + Inches(2.2), y + Inches(0.25),
            IW - Inches(2.3), Inches(0.45),
            size=11, color=C_TEXT)

    # =========================================================================
    # Slide 18: END
    # =========================================================================
    s = blank_slide(prs)
    bg(s, C_DARK)
    rect(s, 0, 0, W, Inches(0.08), C_TEAL)
    rect(s, 0, H - Inches(0.08), W, Inches(0.08), C_TEAL)
    txb(s, "zk-UAV-Proximity-Proof",
        M, Inches(2.5), IW, Inches(0.8),
        size=36, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    txb(s, "開発進捗レポート  —  2026-05-16",
        M, Inches(3.4), IW, Inches(0.4),
        size=16, color=RGBColor(0xB0, 0xC8, 0xFF), align=PP_ALIGN.CENTER)
    rect(s, W / 2 - Inches(1.5), Inches(4.0), Inches(3.0), Pt(2), C_TEAL)
    txb(s, "Circom / Flutter / OpenDroneID / YOLO / ZKP",
        M, Inches(4.3), IW, Inches(0.3),
        size=12, color=RGBColor(0x60, 0x80, 0xA0), align=PP_ALIGN.CENTER)

    prs.save(OUTPUT_PATH)
    print(f"✅ 保存: {OUTPUT_PATH}  ({TOTAL} スライド)")

if __name__ == "__main__":
    build()
