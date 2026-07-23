#!/usr/bin/env python3
"""week19/08_Mopro_PoC結果.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16, size1=14):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 2))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.07 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=12.5):
    lh = fs * 1.65 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=11.5, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(8.2), Inches(1.3)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "Mopro モバイル PoC 結果", 30, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.8)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "circom プローバを Flutter/Android に載せる（APKまで）", 15, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 08_Mopro_PoC結果", 12, GRAY)

# S2 結論
s = add_slide(); title(s, "結論 — 実機ビルドは確認済み（証明は次段）")
bullets(s, TOP, [
    (0, "circom(Groth16)プローバをMopro経由でFlutterに載せ、Android(arm64)の"),
    (0, "インストール可能なAPKまでビルドできた（＝“アプリまで作れた”段階）"),
    (1, "＝フェーズ1(Moproでモバイル標準化)が実機ビルドで裏取り。オンデバイス証明は次段"),
    (0, "生成物（実測）"),
    (1, "libmopro_uav_circom.so（arm64ネイティブ証明器・3.0MB）"),
    (1, "mopro_flutter_bindings（iOS+AndroidのDart/Kotlin/Swiftバインディング）"),
    (1, "Flutterアプリ（generateCircomProof / verifyCircomProof を呼ぶ）"),
    (1, "app-debug.apk（53MB, arm64）← 端末にインストール可能"),
    (0, "残る手動ステップ = 実機を接続して flutter run（オンデバイスで証明→検証を目視）"),
], size0=15.5, size1=13.5)

# S3 実施フロー
s = add_slide(); title(s, "実施フロー（再現手順）")
code(s, TOP, [
    "# 1. ツール",
    "rustup update stable                 # mopro依存が rustc 1.91+ → 1.90→1.97",
    "cargo install mopro-cli --locked     # mopro CLI 0.3.7",
    "rustup target add aarch64-linux-android aarch64-apple-ios ...",
    "",
    "# 2. プロジェクト生成（circomアダプタ）",
    "mopro init --adapter circom --project-name mopro_uav_circom",
    "",
    "# 3-4. Android/Flutter バインディング＋アプリ雛形",
    "mopro build --platforms android --architectures aarch64-linux-android",
    "mopro build --platforms flutter --architectures aarch64-apple-ios ...",
    "mopro create --framework flutter",
    "",
    "# 5. APKビルド（arm64のみに絞ると速い）",
    "flutter build apk --debug --target-platform android-arm64",
    "  → app-debug.apk (53MB)",
], fs=12)

# S4 つまずき
s = add_slide(); title(s, "つまずき（重要な記録）")
table(s, TOP, [
    ["事象", "原因", "回避"],
    ["nargo が tar 展開エラー", "beta.23 がバイナリ未添付", "noirup -v 1.0.0-beta.22"],
    ["cargo install mopro 失敗", "依存が rustc 1.91+ 要求", "rustup update stable(→1.97)"],
    ["Noir の Android ビルド失敗", "barretenberg が Zig の libc++要求", "Android は circom で回避"],
    ["cargo expand エラー", "flutter_rust_bridge が依存", "cargo install cargo-expand"],
    ["APK が Dart コンパイル失敗", "template の import名不一致", "mopro_example_app→実名に置換"],
    ["APK ビルドが時間切れ", "4 ABI 分の Rust を毎回ビルド", "--target-platform android-arm64"],
], col_w=[2.7, 3.1, 2.8], fs=11, row_h=0.48)

# S5 教訓＋次の一手
s = add_slide(); title(s, "教訓 ／ 次の一手")
bullets(s, TOP, [
    (0, "教訓"),
    (1, "AndroidでNoirを使うならZigが要る。circomはarkworks系でZig不要＝Androidはcircomが楽"),
    (1, "MoproのFlutterテンプレートはプロジェクト名 mopro_example_app 前提、別名init時は手直し"),
    (0, "次の一手"),
    (1, "本物の回路(uav_witness_attestation.zkey)に差し替えて同手順でビルド"),
    (1, "実機で flutter run、オンデバイスで証明生成→検証の時間を実測"),
    (1, "iOS もビルドするなら署名設定(開発者証明書)を Xcode で"),
    (0, "意義 = フェーズ1(Mopro標準化)を「設計上の推奨」→「実機ビルドで裏取り済み」に格上げ"),
], size0=16, size1=14)

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_08_Mopro_PoC結果.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
