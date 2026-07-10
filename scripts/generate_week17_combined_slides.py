#!/usr/bin/env python3
"""週17 統合デッキ。3本（全体設計 / RSSI検出 / 現状と今後）を1本に再編する。

元の3デッキは残したまま、構成の流れが一目で分かるよう章立て（セクション扉＋
冒頭の流れ図）で統合する:
  0 背景 → 1 全体設計 → 2 M1 多端末合意×ブロックチェーン
  → 3 RSSI なりすまし検出 → 4 現状と今後

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)  # 稀に使う藍色アクセント
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13.5, space_before=7, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


def section(num, text, sub_text=""):
    s = add_slide()
    tf = tbox(s, 0.9, 2.7, 8.2, 0.6)
    run(tf.paragraphs[0], num, 16, GRAY, font=MONO)
    tf = tbox(s, 0.9, 3.15, 8.2, 1.2)
    run(tf.paragraphs[0], text, 30, INK, bold=True)
    hrule(s, 4.35, l=0.92, w=4.6, color=ACCENT)
    if sub_text:
        tf2 = tbox(s, 0.92, 4.5, 8.2, 0.8)
        run(tf2.paragraphs[0], sub_text, 14, GRAY)
    return s


# ════════════════════════════════════════════════════
# 1 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.05, 8.2, 0.5)
run(tf.paragraphs[0], "進捗まとめ（2026-07 / 週17・統合版）", 15, GRAY)
tf = tbox(s, 0.88, 2.6, 8.4, 1.9)
run(tf.paragraphs[0], "マルチ端末 × ブロックチェーン", 30, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "による UAV RemoteID 検証", 30, INK, bold=True)
hrule(s, 4.5, l=0.92, w=5.6, color=ACCENT)
tf = tbox(s, 0.92, 4.68, 8.2, 1.6)
run(tf.paragraphs[0], "RemoteID の自己申告を、ドローン改造なしで裏取りする", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "背景 → 設計 → M1 実装 → RSSI 検出 → 現状と今後", 12.5, GRAY, font=MONO)

# ════════════════════════════════════════════════════
# 2 構成の流れ（アジェンダ）
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "構成の流れ")
tf = tbox(s, 0.7, 1.4, 8.7, 0.5)
run(tf.paragraphs[0], "この1本で、なぜ→どう設計→何を作った→何が残るか を辿る", 13, GRAY)
codebox(s, 0.7, 2.1, 8.6, 3.5, [
    ("0  背景        RemoteID は自己申告。1端末の証明では偽装を防げない", False),
    ("      │", True),
    ("1  全体設計    多端末合意 + なりすまし検出 / 3層アーキテクチャ", False),
    ("      │", True),
    ("2  M1 実装     多端末合意 × ブロックチェーン（実証済・テスト9/9）", False),
    ("      │", True),
    ("3  RSSI 検出   放送位置 vs 実測レンジの矛盾を ZK 検出 + アプリ実装", False),
    ("      │", True),
    ("4  現状と今後  できたこと / 残作業ロードマップ", False),
])
tf = tbox(s, 0.7, 5.85, 8.7, 1.0)
bullet(tf, "一貫した軸: 「接近の証明」ではなく「なりすまし検出の証明」", size=13, first=True, bold=True)

# ════════════════════════════════════════════════════
# 0 背景
# ════════════════════════════════════════════════════
section("0", "背景", "RemoteID は自己申告 / 1端末証明の限界")

s = add_slide()
title(s, "課題 — なぜ裏取りが要るか")
tf = tbox(s, 0.7, 1.4, 8.7, 2.2)
bullet(tf, "DJI 等は RemoteID で ID と位置を放送（アプリは受信済み）", first=True)
bullet(tf, "だが放送は詐称・リプレイ・非放送ができ、本物か裏取りできない", bold=True)
bullet(tf, "ZK は計算の正当性を示すが、入力の真正性は保証しない")
sub(tf, "自端末位置も private 入力 → 偽座標で within30=1 を捏造できる")
codebox(s, 0.7, 3.95, 8.6, 1.5, [
    "1端末の証明が無検証で信じている入力:",
    ("  drone_* 放送位置/ID   target_* 自端末位置   detection 検知結果", True),
    "→ 単独では「本当にそこにいたか」を担保できない",
])
tf = tbox(s, 0.7, 5.7, 8.7, 1.0)
bullet(tf, "アイデア: 複数端末が独立に観測し、K台一致で真偽を合意", size=13, first=True, bold=True)
sub(tf, "1台のスピーカーでは離れた複数端末に整合する観測を作れない")

# ════════════════════════════════════════════════════
# 1 全体設計
# ════════════════════════════════════════════════════
section("1", "全体設計", "3層アーキテクチャと結合方式")

s = add_slide()
title(s, "3層アーキテクチャ")
codebox(s, 0.7, 1.4, 8.6, 4.2, [
    ("[端末層]  各スマホ = 回路で Groth16 証明を生成（rapidsnark）", False),
    ("   public: within30_and_detected, 各コミットメント,", True),
    ("           drone_id_commitment(結合キー), device_pubkey, time_bucket", True),
    ("      │  同一 drone_id を共有する N 個の proof", True),
    ("      ▼", False),
    ("[集約層]  M1: 各 proof を個別にチェーンへ", False),
    ("          M2: 再帰アグリゲータ → 集約 proof 1個（新規性の本体）", True),
    ("      ▼", False),
    ("[チェーン層]  Groth16検証 + AttestationRegistry で K台合意を記録", False),
])
tf = tbox(s, 0.7, 5.8, 8.7, 0.9)
bullet(tf, "結合キー: RemoteID は共通 → drone_id_commitment が端末間で自然に一致",
       size=12.5, first=True)

s = add_slide()
title(s, "結合方式の比較 — 新規性はどこか")
table(s, 0.7, 1.5, 8.6, [
    ["方式", "内容", "新規性", "採用"],
    ["個別検証＋閾値", "各 proof を検証しレジストリで K 台数える", "低", "M1"],
    ["SnarkPack 集約", "同一 VK の proof を安く束ねる gas 最適化", "低", "×"],
    ["再帰アグリゲータ", "N proof を再帰検証＋関係を強制し1証明へ", "高", "M2"],
], col_w=[1.7, 4.3, 1.0, 1.6], fs=11.5, row_h=0.55)
tf = tbox(s, 0.7, 4.2, 8.7, 2.4)
subhead(tf, "再帰アグリゲータ（M2）が回路内で強制する関係", 0)
bullet(tf, "① 全 proof の drone_id が一致（＝同じドローン）", first=True)
bullet(tf, "② device_pubkey が相異なる（Sybil 対策）")
bullet(tf, "③ 各 within30_and_detected == 1　④（任意）幾何整合")
sub(tf, "単なる集約でなく「物理イベントへのマルチ witness 合意の証明」")

# ════════════════════════════════════════════════════
# 2 M1 実装
# ════════════════════════════════════════════════════
section("2", "M1 実装", "多端末合意 × ブロックチェーン（実証済）")

s = add_slide()
title(s, "端末回路 + コントラクト")
tf = tbox(s, 0.7, 1.35, 8.7, 1.3)
bullet(tf, "uav_witness_attestation.circom: detection_commitment を拡張", first=True)
sub(tf, "drone_id_commitment / EdDSA device_pubkey / time_bucket を追加（public14）")
codebox(s, 0.7, 2.7, 8.6, 2.9, [
    "submitAttestation(proof, input[14]):",
    "  require(input[6] == 1)                     # within30_and_detected",
    "  require(verifier.verifyProof(proof,input))  # Groth16 検証",
    "  require(registeredDevice[pubkey])          # 登録端末のみ (Sybil対策)",
    "  key = keccak(drone_id, time_bucket)",
    "  require(!counted[key][device])             # 二重計上を排除",
    "  count[key]++; if (count[key] >= K) consensusReached[key] = true",
])
tf = tbox(s, 0.7, 5.75, 8.7, 0.9)
bullet(tf, "回路 約1.2万制約（大半は EdDSA）/ 既存4回路は無改変で新規追加",
       size=12, first=True)

s = add_slide()
title(s, "E2E 実証 (K=3)")
tf = tbox(s, 0.7, 1.35, 8.7, 0.7)
bullet(tf, "実端末3台の実 Groth16 証明を、実 Verifier のオンチェーン検証に通す", first=True)
codebox(s, 0.7, 2.1, 8.6, 1.5, [
    "実証明3台(同一drone_id・distinct device_pubkey)",
    "   → 実 Groth16Verifier.sol でオンチェーン検証",
    "   → K=3 で consensusReached 発火 / count = 3",
])
tf = tbox(s, 0.7, 3.85, 8.7, 2.8)
subhead(tf, "Hardhat テスト 9/9 通過", 0)
bullet(tf, "K台合意 / 二重提出排除 / 未登録拒否 / within30チェック / 無効証明拒否", first=True)
bullet(tf, "別ドローン・別時間窓は別グループとして集計")
bullet(tf, "改ざんした public signal は実 verifier が拒否", bold=True)
sub(tf, "commit 0c64254")

# ════════════════════════════════════════════════════
# 3 RSSI 検出
# ════════════════════════════════════════════════════
section("3", "RSSI なりすまし検出", "放送位置 vs 実測レンジの矛盾を検出")

s = add_slide()
title(s, "RSSI → 距離 の較正モデル")
codebox(s, 0.7, 1.4, 8.6, 1.2, [
    "RSSI(d) = A - 10 n log10(d)   ⇔   d = 10^((A - RSSI)/(10 n))",
])
tf = tbox(s, 0.7, 2.8, 8.7, 0.6)
bullet(tf, "合成データで較正 → A/n/σ を正しく回復。距離推定の誤差帯:", first=True)
table(s, 0.7, 3.5, 8.6, [
    ["真の距離", "推定", "±2σ 帯", "measured_range_cm"],
    ["3 m", "2.9 m", "[1.6, 5.6] m", "295"],
    ["10 m", "9.9 m", "[5.2, 18.7] m", "987"],
    ["30 m", "29.8 m", "[15.7, 56.5] m", "2975"],
], col_w=[1.7, 1.5, 2.8, 2.6], fs=11.5, row_h=0.4)
tf = tbox(s, 0.7, 5.7, 8.7, 1.0)
bullet(tf, "発見: RSSI 誤差は相対%（±38%/1σ, ±90%/2σ）→ 許容は相対で設計",
       size=12.5, first=True, bold=True)
sub(tf, "粗い矛盾しか検出できない → 多witnessの冗長が必須")

s = add_slide()
title(s, "ZK 回路 + アプリ実装")
tf = tbox(s, 0.7, 1.35, 8.7, 0.6)
bullet(tf, "remoteid_consistency.circom: 放送二乗距離が実測帯に入るか証明（約540制約）", first=True)
table(s, 0.7, 2.05, 8.6, [
    ["入力", "放送位置(claimed)", "実測レンジ", "consistency_ok"],
    ["整合", "同一（≈3.4m）", "341 cm", "1"],
    ["詐称", "同一（≈3.4m）", "5000 cm", "0  ← 検出"],
], col_w=[1.3, 3.0, 2.0, 2.3], fs=11.5, row_h=0.4)
codebox(s, 0.7, 3.9, 8.6, 1.7, [
    "// アプリ: RemoteID Location 受信時（main.dart + lib/core/rssi_ranging.dart）",
    "[BLE  RSSI-check] 整合         claimed=3.4m est=2.9m band[1.6,5.6]m",
    "[WiFi RSSI-check] なりすまし疑い claimed=3.4m est=48m band[27,84]m",
])
tf = tbox(s, 0.7, 5.7, 8.7, 0.9)
bullet(tf, "ドローン改造ゼロ・端末側だけで放送の矛盾をその場判定（analyze clean）",
       size=12, first=True)

# ════════════════════════════════════════════════════
# 4 現状と今後
# ════════════════════════════════════════════════════
section("4", "現状と今後", "できたこと / 残作業ロードマップ")

s = add_slide()
title(s, "実装済み — できたこと")
table(s, 0.7, 1.45, 8.6, [
    ["領域", "成果物", "状態"],
    ["M1 端末回路", "uav_witness_attestation.circom（public14）", "✅ 実証明OK"],
    ["M1 コントラクト", "AttestationRegistry + Groth16Verifier", "✅ テスト9/9"],
    ["M1 E2E", "実3台 → 実検証器 → K=3 合意", "✅ 実証済"],
    ["RSSI 検出回路", "remoteid_consistency.circom", "✅ 整合1/詐称0"],
    ["RSSI モデル", "scripts/rssi/rssi_ranging.py", "✅ 検証済"],
    ["アプリ", "rssi_ranging.dart + main.dart 配線", "✅ analyze clean"],
], col_w=[1.9, 4.9, 1.8], fs=10.5, row_h=0.52)
tf = tbox(s, 0.7, 6.1, 8.7, 0.9)
bullet(tf, "到達点: 多端末合意の土台 + なりすまし検出の最小形が動く", size=12, first=True, bold=True)
sub(tf, "ただし「K台合意 ＝ 真正」はまだ未達（登録署名が揃っただけ / 較正は暫定）")

s = add_slide()
title(s, "これから必要なこと")
table(s, 0.7, 1.45, 8.6, [
    ["優先", "やること", "目的"],
    ["高", "回路を相対許容(±%)版へ更新", "RSSI の現実に合わせる"],
    ["高", "実機 DJI で RSSI 較正 (A/n/TOL)", "本当に検出できるかの実測検証"],
    ["高", "consistency を M1 レジストリに統合", "多端末で放送真偽を合意"],
    ["中", "M2 再帰アグリゲータ (gnark/Sonobe)", "集約proof1個・定数gas"],
    ["中", "M3 真正性: ドローン署名RemoteID / 端末Attestation", "偽装を封じる(耐偽装の核)"],
    ["低", "許可型EVM(Besu/Quorum)へデプロイ", "運用形態"],
], col_w=[0.8, 5.0, 2.8], fs=10.5, row_h=0.5)
tf = tbox(s, 0.7, 5.95, 8.7, 1.0)
bullet(tf, "次の一手: ①相対許容版へ更新 → ②実機DJI較正（本命）→ ③多端末合意へ接続",
       size=12.5, first=True, bold=True)

# ── 出力 ─────────────────────────────────────────────
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week17")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "week17_combined.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
