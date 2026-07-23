#!/usr/bin/env python3
"""week19/07_信頼モデル比較.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16, size1=14):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 2))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.07 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "信頼モデル比較", 32, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.0)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "誰がズルできて、何が止めるか", 16, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 07_信頼モデル比較", 12, GRAY)

# S2 4つの場所
s = add_slide(); title(s, "信頼が必要になる「4つの場所」")
y = code(s, TOP, [
    "① セットアップ … 鍵作りの儀式は健全か（漏れると証明を偽造できる）",
    "② プローバ     … 証明を作る人は生データを見るか（プライバシー・独占）",
    "③ 集約者       … まとめる人は結果を歪められるか（改ざん・検閲）",
    "④ 入力         … 入れたRSSI/GPS/端末は本物か（真正性・Sybil）",
])
y = bullets(s, y + 0.1, [
    (0, "①②③は暗号(ZKP)で守れる。④はZKPの外（アテステーション等）で裏取り"),
    (0, "方針 = 改ざんは暗号で不可／検閲は網羅性or合意／入力は真正性／鍵は正規セレモニー"),
])

# S3 ①セットアップ
s = add_slide(); title(s, "① セットアップの信頼（証明系ごとに違う）")
y = table(s, TOP, [
    ["証明系", "セットアップ", "信頼の中身", "綻びると", "本案件"],
    ["Groth16（現行）", "回路ごとに儀式", "1人でも正直なら安全(1-of-N)", "全員結託で偽造", "★1・実施済"],
    ["PLONK/Honk", "1回のSRS", "同上(1-of-N)・回路別不要", "同上", "★1を回路別から解放"],
    ["STARK/Folding", "不要", "信頼点なし(transparent)", "—", "★1が構造的に消える"],
], col_w=[1.7, 1.6, 2.7, 1.4, 1.2], fs=10.5, row_h=0.6)
y = bullets(s, y + 0.1, [
    (0, "現行の弱点 = 回路を変えるたびに儀式。Noirで頻度減、STARK/Foldingで儀式が消える"),
], size0=15)

# S4 ②プローバ
s = add_slide(); title(s, "② プローバの信頼（集約手法で大きく変わる）")
y = table(s, TOP, [
    ["集約手法", "誰が証明を作るか", "生データを見るのは", "壁"],
    ["オンチェーン集計(K-of-N)", "各端末が自分の分", "各端末＝自分の分のみ", "◎分散"],
    ["回路内バッチ(1回路にN端末)", "単一プローバ", "その1人が全端末分", "✕独占"],
    ["再帰的合成(recursion)", "各端末＋集約者", "集約者は証明のみ", "◎"],
    ["Folding", "各端末＋畳み込み役", "畳み込み役は証明のみ", "○"],
], col_w=[2.9, 2.0, 2.4, 1.3], fs=11, row_h=0.48)
y = bullets(s, y + 0.08, [
    (0, "重要 = 回路内バッチは「1人が全座標を持つ」→分散の趣旨と矛盾。実運用は再帰orFolding"),
], size0=14.5)

# S5 ③集約者
s = add_slide(); title(s, "③ 集約者の信頼（改ざんと検閲を分ける）")
y = table(s, TOP, [
    ["攻撃", "集約者にできるか", "何が止めるか", "本案件"],
    ["改ざん（結果を偽る）", "✕ できない", "ZKPの健全性(偽証明は落ちる)", "全証明系で担保"],
    ["検閲（有効票を間引く）", "○ できてしまう", "網羅性証明 or BFT合意＋全複製", "#3b 設計済"],
    ["なりすまし集約(空スロット)", "検証すれば不可", "集約時に各proofも検証", "M2 が担保"],
], col_w=[2.7, 1.9, 2.7, 1.3], fs=11, row_h=0.55)
y = bullets(s, y + 0.1, [
    (0, "改ざんは無理だが検閲は仕組みで別途防ぐ。証明系を替えても構図は同じ"),
], size0=15)

# S6 ④入力
s = add_slide(); title(s, "④ 入力の真正性（ZKPの外・全手法共通）")
y = table(s, TOP, [
    ["リスク", "内容", "対策", "本案件"],
    ["端末なりすまし", "偽端末が証明を投稿", "端末アテステーション(TEE)", "★2 / 設計済"],
    ["Sybil", "1人が多数端末を装う", "登録メンバーシップ＋1端末1票", "#3a 実装済"],
    ["入力捏造", "偽のRSSI/GPSを入力", "ドローン署名・freshness・相互整合", "★2 / 回路統合済"],
    ["較正すり替え", "A/nを意図的にずらす", "較正手順の固定・監査", "実機較正で対処"],
], col_w=[1.7, 2.2, 3.1, 1.6], fs=11, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "どのフレームワークを選んでも④は解決しない。ZKP外の仕組みで裏取り（概ね実装/設計済）"),
], size0=14.5)

# S7 まとめ＋結論
s = add_slide(); title(s, "まとめ — フレームワーク横断")
y = table(s, TOP, [
    ["観点", "変わるか", "効いてくるもの"],
    ["① セットアップの信頼", "証明系で変わる", "Groth16=回路別 / Noir=universal / STARK・Folding=不要"],
    ["② プローバの信頼", "集約手法で変わる", "回路内バッチは避ける / 再帰・Foldingは分散◎"],
    ["③ 集約者の検閲", "変わらない", "網羅性(#3b) or 合意という別レイヤが常に必要"],
    ["④ 入力の真正性", "変わらない", "アテステーション/登録/相互整合（ZKP外）"],
], col_w=[2.2, 1.6, 4.8], fs=10.5, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "結論 = 集約は再帰/Folding(分散)を維持、証明系は必要ならNoir/STARKで①を下げ、"),
    (0, "　　　 ③検閲・④入力は既存の #3b・アテステーションで守る"),
], size0=14.5)

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_07_信頼モデル比較.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
