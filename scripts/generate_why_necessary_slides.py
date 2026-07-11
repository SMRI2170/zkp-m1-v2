#!/usr/bin/env python3
"""「これまでの各要素がなぜ必要か」まとめスライド（必要性の連鎖）。

各層は前の層が残した穴を埋めるために存在する。1要素1枚で
「なぜ（穴）→ 何をしたか（対策）→ 状態」を示す。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 21, INK, bold=True)
    hrule(s, 1.08)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 14, ACCENT, bold=True); return p


def bullet(tf, text, size=13.5, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


def why(idx, name, gap, solution, status):
    """1要素1枚: なぜ(穴) → 何をしたか → 状態。"""
    s = add_slide()
    title(s, f"{idx}  {name}")
    tf = tbox(s, 0.7, 1.35, 8.7, 5.4)
    subhead(tf, "なぜ必要か（前の層が残した穴）", 0)
    for x in gap:
        bullet(tf, x, first=(x == gap[0]))
    subhead(tf, "何をしたか")
    for x in solution:
        bullet(tf, x, first=(x == solution[0]))
    subhead(tf, "状態")
    bullet(tf, status, first=True, bold=True)
    return s


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.15, 8.2, 0.5)
run(tf.paragraphs[0], "全体まとめ / 必要性の連鎖", 15, GRAY)
tf = tbox(s, 0.88, 2.7, 8.4, 1.4)
run(tf.paragraphs[0], "各要素はなぜ必要か", 30, INK, bold=True)
hrule(s, 4.2, l=0.92, w=5.2, color=ACCENT)
tf = tbox(s, 0.92, 4.38, 8.2, 1.6)
run(tf.paragraphs[0], "各層は「前の層が残した穴」を埋めるために存在する", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "RemoteID裏取り → 多witness → チェーン → ZK → 集約 → 真正性 → 堅牢化", 11.5, GRAY, font=MONO)

# 2 出発点
s = add_slide()
title(s, "出発点 — 何が問題か")
tf = tbox(s, 0.7, 1.4, 8.7, 2.2)
bullet(tf, "RemoteID はドローンの「自己申告」", first=True, bold=True)
bullet(tf, "放送は詐称・リプレイ・非放送ができ、本物か裏取りできない")
bullet(tf, "受信者は改ざん不能な証拠を残せず、通報で自分の位置も晒れる")
tf = tbox(s, 0.7, 3.9, 8.7, 1.2)
subhead(tf, "満たすべき3条件", 0)
bullet(tf, "真正性（本物か）/ 改ざん不能な記録 / 通報者のプライバシー", first=True, bold=True)
tf = tbox(s, 0.7, 5.3, 8.7, 1.0)
run(tf.paragraphs[0], "以降の各要素は、この3条件を段階的に満たすために積み上がる。", 12.5, GRAY)

# 3〜 各要素
why("1", "RSSI 整合検知",
    ["放送位置が本物か、放送だけ見ても分からない"],
    ["各端末が放送位置 vs RSSI実測レンジの矛盾を検出（相対許容・中央値集約）",
     "座標は隠したまま consistency_ok を出す"],
    "✅ 回路+アプリ実装、合成ROCで30m詐称をN=5で91%検知")

why("2", "多 witness 合意",
    ["単一端末・単一RSSIは偽装でき、粗い（単独64%）"],
    ["複数端末が独立に観測し、K台一致で真偽を合意",
     "1台のスピーカーでは離れた複数端末に整合する観測を作れない"],
    "✅ 8台で98%。単独より本質的に強い")

why("3", "ブロックチェーン (M1)",
    ["合意結果を改ざん不能に、第三者検証可能に残す必要"],
    ["各端末が証明を直接提出、契約がK台集計 → trusted/spoof を記録",
     "AttestationRegistry（Groth16検証+集計）"],
    "✅ Hardhat 20/20、実RPCノードでM1直接提出を実演")

why("4", "ゼロ知識証明 (rapidsnark)",
    ["座標を晒さず「30m以内で検知した」だけ示したい（プライバシー）"],
    ["端末が circom/rapidsnark で Groth16 証明を生成",
     "座標や特徴は private、条件の成否(0/1)だけ公開"],
    "✅ 端末=葉プロ―バ。回路群（距離/検知/音響/整合）")

why("5", "M2 再帰集約 (gnark)",
    ["M1は台数比例のgas。証明間の関係も契約の素朴な集計頼み"],
    ["N証明を1つに集約し、同一ドローン/端末相異/整合を回路内強制",
     "rapidsnarkは繋げないので gnark が再帰検証"],
    "✅ Spike A/B/C。検証gas 347k（N非依存）")

why("6", "M3 入力真正性",
    ["放送/位置は公開・自己申告 → 放送を知る攻撃者が偽造でき「K台合意≠真正」"],
    ["ドローン署名RemoteID（偽造対策）/ 新鮮さ（リプレイ）/",
     "端末アテステーション（Sybil・実機裏取り）"],
    "✅ 回路+契約でE2E、検証ゲート8/8、ネイティブ実装")

why("7", "トラステッドセットアップのMPC化",
    ["決定論的setupはトラップドアが既知＝誰でも偽造できた（最重要）"],
    ["真乱数の複数コントリビューション+ビーコン+監査のセレモニーに置換",
     "1-of-N 信頼（1人が乱数破棄すれば安全）"],
    "✅ 再生成でZKey Ok!、Hardhat維持")

why("8", "登録メンバーシップ (#3a)",
    ["M2集約は端末相異のみ強制し、登録済みかを未検証（Sybilの穴）"],
    ["各 device_pubkey が登録端末ルートのメンバーかを Merkle で回路内検証",
     "契約 deviceRoot と Root を照合（定数gas）"],
    "✅ 機構テスト、集約回路へ統合")

why("9", "ガバナンス強化 (#4)",
    ["owner が登録・方針を全掌握＝compromiseで即時全制御"],
    ["Ownable2Step（2段階移転）+ owner を TimelockController に",
     "方針変更は遅延＋多者承認を経る"],
    "✅ 実装+テスト3/3")

why("10", "許可型チェーン",
    ["集約者/単一主体が有効な証明を間引ける（検閲）"],
    ["Besu QBFT: BFT合意で必ず包含、全複製でDA内蔵、メンバーシップで責任追及",
     "既存 contracts を無改変で載せる"],
    "✅ デプロイ動作確認、M1直接を実演")

# 末尾 まとめ
s = add_slide()
title(s, "まとめ — 必要性の連鎖")
table(s, 0.7, 1.4, 8.6, [
    ["満たす条件", "担う要素"],
    ["真正性（本物か）", "RSSI整合 / 多witness / M3(署名・新鮮さ・端末アテスト)"],
    ["改ざん不能な記録", "ブロックチェーン(M1) / セットアップMPC化"],
    ["プライバシー", "ゼロ知識証明(rapidsnark)"],
    ["規模・効率", "M2再帰集約（定数gas）"],
    ["Sybil / 検閲 / 集権", "登録メンバーシップ / 許可型チェーン / ガバナンス"],
], col_w=[2.6, 6.0], fs=11.5, row_h=0.55)
tf = tbox(s, 0.7, 5.15, 8.7, 1.5)
bullet(tf, "各層は独立した機能でなく、前の層の穴を埋める必然として積み上がっている",
       first=True, bold=True)
sub(tf, "「改ざんは暗号で不可」＋「検閲は合意で不可」＋「入力は真正性で裏取り」で耐改竄が揃う")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "why_each_necessary.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
