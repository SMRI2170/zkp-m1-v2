#!/usr/bin/env python3
"""進捗第20回 — ZKPフレームワーク調査とモバイル実装検証（進捗報告・研究室/教授向け・拡張版）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


def section(text):
    s = add_slide()
    tf = s.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(8.2), Inches(1.0)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 26, INK, bold=True)
    hrule(s, 4.15, l=0.92, w=4.0, color=INK)
    return s


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(8.2), Inches(0.5)).text_frame
run(tf.paragraphs[0], "進捗報告 第20回", 15, GRAY)
tf = s.shapes.add_textbox(Inches(0.88), Inches(2.85), Inches(8.4), Inches(1.4)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "ZKPフレームワーク調査と\nモバイル実装検証", 27, INK, bold=True)
hrule(s, 4.55, l=0.92, w=5.6)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.72), Inches(8.2), Inches(0.6)).text_frame
run(tf.paragraphs[0], "多端末の証明を1つにまとめ、モバイルで生成する土台の選定", 14, INK)

# S2 目次
s = add_slide(); title(s, "本日の内容")
bullets(s, TOP, [
    (0, "背景・目的（多端末の証明を1つに・モバイルで）"),
    (0, "手法の全体像（4層：証明系・集約・検証共有）"),
    (0, "フレームワーク比較と決定（段階的アプローチ）"),
    (0, "実装検証（circom/Noir 実測・Mopro 実機PoC）"),
    (0, "信頼モデル・まとめ"),
])

# S3 背景・目的
s = add_slide(); title(s, "背景・目的")
bullets(s, TOP, [
    (0, "第19回 = 各端末が RSSI でなりすましを検知"),
    (0, "本回 = その証明を多端末で1つにまとめ、チェーンで共有・検証する土台を選定"),
    (0, "課題 = (a)モバイルで証明生成 (b)多端末を1つに集約 (c)★1トラステッドセットアップの根治"),
    (0, "現状スタック = circom+rapidsnark(端末) / gnark(集約M2) / EVMでGroth16検証"),
])

# S4 全体像
s = add_slide(); title(s, "やりたいことの全体像")
y = code(s, TOP, [
    "各端末: RSSI検証＋距離検証＋複雑計算 を秘匿したまま証明   ← ZKP",
    "複数端末を1つにまとめる（同じドローンにK台が整合）        ← 集約",
    "みんなで共有・検証（ブロックチェーンに記録）              ← 検証・共有",
])
y = bullets(s, y + 0.1, [
    (0, "今回 = この各段を担う手法とフレームワークを調査し、実測・実機ビルドで裏取り"),
])

# ── 手法 ──
section("手法の全体像")

# S6 4層
s = add_slide(); title(s, "手法の全体像（4つの層）")
y = code(s, TOP, [
    "層1  ZKPで何ができるか（秘匿したまま正しさを証明）",
    "層2  証明系の種類（Groth16 / PLONK / STARK / Folding）",
    "層3  複数の証明をまとめる手法（本案件の肝）",
    "層4  検証・共有（オンチェーン / DA層）",
])
y = bullets(s, y + 0.1, [
    (0, "フレームワークは層2〜3を実装する“乗り物”にすぎない"),
    (0, "先に手法を決め、その後に乗り物を選ぶのが正しい順序"),
])

# S7 層2 証明系
s = add_slide(); title(s, "層2 — 証明系の種類")
y = table(s, TOP, [
    ["種類", "証明サイズ", "オンチェーン", "事前準備", "ひとこと"],
    ["Groth16", "最小", "最安(≈348k)", "回路ごと", "小さく安い・準備が弱点"],
    ["PLONK/Honk", "小〜中", "中〜高", "1回", "準備が緩い・汎用"],
    ["STARK", "大", "高", "不要", "準備不要・証明が大"],
    ["Folding", "最終のみ", "圧縮依存", "ほぼ不要", "N個を逐次畳む・省メモリ"],
], col_w=[1.5, 1.3, 1.6, 1.2, 3.0], fs=11.5, row_h=0.5)
y = bullets(s, y + 0.05, [
    (0, "現行の回路別セレモニー(★1)は PLONK/STARK/Folding で構造的に消える"),
], size0=14.5)

# S8 層3 集約
s = add_slide(); title(s, "層3 — 複数の証明をまとめる手法")
y = table(s, TOP, [
    ["#", "手法", "オンチェーン", "現状"],
    ["1", "オンチェーン集計(K-of-N)", "台数分 O(N)", "実装済 M1"],
    ["3", "再帰的証明合成", "1回・定数", "実装済 M2"],
    ["4", "Folding(Nova系)", "最終1回", "未（端末寄せ候補）"],
    ["2/5/6", "回路内バッチ / 専用集約 / 閾値署名", "各種", "試作 / 未"],
], col_w=[1.0, 3.4, 1.9, 2.3], fs=11.5, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "1(集計)・3(再帰)は実装済。未検討で価値ありは 4(Folding)＝端末側で省メモリ集約"),
], size0=14.5)

# S9 層4 検証共有
s = add_slide(); title(s, "層4 — 検証・共有の手法")
table(s, TOP, [
    ["手法", "効く攻撃", "現状"],
    ["オンチェーン verifier", "改ざん（暗号で不可）", "実装済"],
    ["DA層に提出集合コミット（網羅性）", "検閲（間引き）", "設計済 #3b"],
    ["許可型チェーン（BFT合意＋全複製）", "検閲・可用性", "設計済"],
], col_w=[3.6, 2.7, 2.3], fs=12, row_h=0.52)

# ── フレームワーク ──
section("フレームワーク比較と決定")

# S11 早見表
s = add_slide(); title(s, "フレームワーク早見表")
table(s, TOP, [
    ["フレームワーク", "エンジン", "事前準備", "モバイル", "一言"],
    ["circom+rapidsnark", "Groth16", "回路別", "○", "現行の定番・安い検証"],
    ["Noir+Barretenberg", "Honk", "不要", "◎", "準備不要・速い・gas高"],
    ["gnark", "Groth16/PLONK", "回路別", "✕", "集約(サーバ)"],
    ["halo2", "PLONK", "1回/不要", "○", "強力だが難しい"],
    ["Plonky2/3", "STARK", "不要", "○", "準備不要・再帰速い"],
    ["Nova/Sonobe", "Folding", "ほぼ不要", "◎", "端末側集約向き"],
    ["Cairo/Stwo", "STARK", "不要", "◎", "Starknet前提"],
    ["Mopro", "(包む)", "—", "◎", "証明系でなく橋渡し"],
], col_w=[2.4, 1.9, 1.2, 1.2, 1.9], fs=11, row_h=0.42)

# S12 主要候補の要点
s = add_slide(); title(s, "主要候補の要点")
bullets(s, TOP, [
    (0, "circom+rapidsnark（現行）= 証明が最小・オンチェーン最安。★1が残る"),
    (0, "Noir+Barretenberg = 準備不要で端末証明が速い。ただしgas約7倍・回路移植要"),
    (0, "gnark = 集約(M2)のサーバ担当。端末証明には重い"),
    (0, "Nova/Sonobe = 省メモリで端末側集約向き（研究的）"),
    (0, "Mopro = 証明系でなく橋渡し。現行circomをそのままFlutterに載せられる"),
])

# S13 その他・新顔
s = add_slide(); title(s, "その他・新顔（参考 — 主軸は変えない）")
bullets(s, TOP, [
    (0, "モバイル新規に効きうる = Ligero/Ligetron（低メモリ）・o1js（TS・準備不要）"),
    (0, "集約を攻めるなら = ProtoGalaxy/HyperNova（Folding改良）・Spartan"),
    (0, "zkVM他 = Jolt / Miden / Nexus、プローバ = Expander(GKR)・Binius"),
    (0, "結論 = 主軸(Mopro＋circom/Noir・集約gnark/folding・EVM前提)は不変"),
])

# S14 決定
s = add_slide(); title(s, "決定 — 段階的アプローチ")
bullets(s, TOP, [
    (0, "評価軸 = 現行資産の再利用 / ★1根治 / gas / モバイル / 集約適性"),
    (0, "フェーズ1（すぐ）= Mopro でモバイル標準化（現行 circom のまま）"),
    (0, "フェーズ2（並行・研究）= 集約の端末寄せを Nova/Folding で検証"),
    (0, "フェーズ3（条件付き）= ★1根治のため Noir 評価（gas増を許容できる場合のみ）"),
])

# ── 実装検証 ──
section("実装検証（実測・実機ビルド）")

# S16 circom 実測
s = add_slide(); title(s, "ハンズオン — 現行 circom の実測")
y = table(s, TOP, [
    ["回路", "制約数", "証明生成", "検証"],
    ["uav_witness_attestation（フル）", "11,785", "≈1.4s(snarkjs)", "OK(≈348k gas)"],
    ["remoteid_consistency_rel（RSSI単体）", "649", "—", "—"],
], col_w=[3.8, 1.5, 1.9, 1.4], fs=12, row_h=0.5)
y = bullets(s, y + 0.12, [
    (0, "従来「~18k制約」は概算、実測は 11,785（≈12k）＝Groth16 では小規模"),
    (0, "rapidsnark(arm64) なら snarkjs より速く、端末も実用域"),
])

# S17 Noir 実測
s = add_slide(); title(s, "ハンズオン — Noir の実測（proof→verify）")
bullets(s, TOP, [(0, "環境 = nargo 1.0.0-beta.22 / bb 5.0.0-nightly（UltraHonk）")])
table(s, 1.95, [
    ["コマンド", "結果"],
    ["nargo test", "3/3 通過（正常 / 合意未達 fail / 束縛違反 fail）"],
    ["nargo info", "ACIR 92 opcodes"],
    ["bb gates", "回路サイズ 2,953 ゲート"],
    ["bb prove", "証明生成 0.11s・proof ≈14.6 KB"],
    ["bb verify", "「Proof verified successfully」・0.02s"],
], col_w=[2.2, 6.4], fs=12.5, row_h=0.52)

# S18 比較
s = add_slide(); title(s, "circom vs Noir 比較（実測）")
table(s, TOP, [
    ["項目", "circom（Groth16）", "Noir（UltraHonk）"],
    ["Poseidon ハッシュ1回", "415制約", "404ゲート ＝ ほぼ同等"],
    ["EdDSA署名検証1回", "8,086制約", "実測不可 / 文献≈24,472"],
    ["証明生成", "≈1.4s（フル回路）", "0.11s（小回路）"],
    ["証明サイズ", "定数・小（~200B）", "≈14.6 KB（大）"],
    ["セットアップ", "回路別セレモニー要", "不要（universal）"],
    ["オンチェーン検証", "≈348k gas", "≈2.4M gas（文献）"],
], col_w=[2.8, 2.9, 2.9], fs=11.5, row_h=0.45)

# S19 EdDSAとlib
s = add_slide(); title(s, "EdDSA と Noir 暗号ライブラリの成熟度")
bullets(s, TOP, [
    (0, "Poseidon は同機能で実測比較: circom 415 ≒ Noir 404（ほぼ同等）"),
    (1, "→ 移行で回路が極端に重くはならない見込み"),
    (0, "EdDSA は circom を実測（8,086制約）、Noir は実測に至らず"),
    (1, "公式 eddsa/ec の最新タグが u1（beta.22で削除）を使いコンパイル不可"),
    (0, "知見 = 署名等の重い回路は Noir だとライブラリのバージョン整合で詰まる"),
    (1, "circomlib は枯れていて即使える＝この成熟度の差自体が移行コスト"),
])

# S20 Mopro 結論
s = add_slide(); title(s, "Mopro PoC — 実機ビルドで裏取り")
bullets(s, TOP, [
    (0, "circom(Groth16)プローバをMopro経由でFlutter/Androidに載せ、APKまでビルド成功"),
    (1, "libmopro_uav_circom.so（arm64ネイティブ証明器・3.0MB）"),
    (1, "Flutterアプリ（generateCircomProof / verifyCircomProof）"),
    (1, "app-debug.apk（53MB, arm64）← 端末にインストール可能"),
    (0, "意義 = フェーズ1(Mopro標準化)を「設計上の推奨」→「実機ビルドで裏取り済み」に格上げ"),
], size0=15.5)

# S21 Mopro 手順・つまずき
s = add_slide(); title(s, "Mopro PoC — 手順とつまずき")
y = code(s, TOP, [
    "mopro init --adapter circom → mopro build (android/flutter)",
    "mopro create --framework flutter → flutter build apk --debug",
], fs=12)
y = table(s, y + 0.08, [
    ["つまずき", "回避"],
    ["Noir の Android は Zig 必須(barretenberg)", "Android は circom で回避"],
    ["mopro導入に rustc 1.91+", "rustup update(→1.97)"],
    ["templateのimport名不一致", "mopro_example_app→実名に置換"],
], col_w=[4.6, 4.0], fs=11.5, row_h=0.46)

# ── 信頼 ──
section("信頼モデル")

# S23 4つの場所
s = add_slide(); title(s, "信頼が必要になる「4つの場所」")
y = code(s, TOP, [
    "① セットアップ  … 鍵作りの儀式は健全か（証明系で変わる）",
    "② プローバ      … 生データを誰が見るか（集約手法で変わる）",
    "③ 集約者(検閲)  … 有効票を間引けるか（変わらない）",
    "④ 入力の真正性  … 入れた値は本物か（変わらない・ZKP外）",
])
y = bullets(s, y + 0.1, [
    (0, "①②は証明系/集約手法の選択で変わる。③④はフレームワークを替えても消えない"),
])

# S24 横断まとめ
s = add_slide(); title(s, "信頼モデル — 横断まとめ")
y = table(s, TOP, [
    ["観点", "変わるか", "効いてくるもの"],
    ["① セットアップ", "証明系で変わる", "Groth16=回路別 / STARK・Folding=不要"],
    ["② プローバ", "集約手法で変わる", "回路内バッチ回避 / 再帰・Foldingは分散"],
    ["③ 集約者の検閲", "変わらない", "網羅性(#3b) or 合意"],
    ["④ 入力の真正性", "変わらない", "アテステーション/登録/相互整合"],
], col_w=[2.0, 1.7, 4.9], fs=11, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "＝③④は既存の #3b・アテステーションで守る（設計/実装済）"),
], size0=14.5)

# S25 まとめ・今後
s = add_slide(); title(s, "まとめ・今後")
bullets(s, TOP, [
    (0, "まとめ = 端末は Mopro、集約は gnark継続 or Novaで端末寄せ、★1消すなら Noir 段階導入"),
    (1, "Poseidon同等・Mopro で circom を APK までビルドできることを実証"),
    (0, "今後"),
    (1, "本物の回路(uav_witness_attestation.zkey)を実機で証明時間実測"),
    (1, "Noir の EdDSA を lib vendoring して同機能比較を完成"),
    (1, "Folding(Nova)で集約の端末寄せ PoC ／ DA層の実接続(#3b)"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "siles-pptx", "進捗第20回_ZKPフレームワーク_モバイル.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
