#!/usr/bin/env python3
"""検閲耐性/包含性の解法まとめスライド。

「集約者が有効な証明を間引ける（検閲/選択バイアス）」を塞ぐ各手法を、
目的（何のため）→ 実現方法（どう）→ 最新フレームワーク で整理する。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.1)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


def solution(idx, name, purpose, how, fw):
    """1手法1枚: 目的 / 実現 / FW。"""
    s = add_slide()
    title(s, f"{idx}  {name}")
    tf = tbox(s, 0.7, 1.4, 8.7, 4.8)
    subhead(tf, "何のため（目的）", 0)
    for x in purpose:
        bullet(tf, x, first=(x == purpose[0]))
    subhead(tf, "どう実現するか")
    for x in how:
        bullet(tf, x, first=(x == how[0]))
    subhead(tf, "最新フレームワーク")
    bullet(tf, fw, first=True, bold=True)
    return s


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.15, 8.2, 0.5)
run(tf.paragraphs[0], "検閲耐性 / 包含性の解法", 15, GRAY)
tf = tbox(s, 0.88, 2.7, 8.4, 1.6)
run(tf.paragraphs[0], "「間引き」をどう防ぐか", 30, INK, bold=True)
hrule(s, 4.3, l=0.92, w=5.2, color=ACCENT)
tf = tbox(s, 0.92, 4.48, 8.2, 1.6)
run(tf.paragraphs[0], "集約で有効な証明を落とす（検閲/選択バイアス）を塞ぐ手法", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "各手法: 何のため → どう実現 → 最新フレームワーク", 12, GRAY, font=MONO)

# 2 問題の再定義
s = add_slide()
title(s, "問題の再定義")
tf = tbox(s, 0.7, 1.4, 8.7, 2.0)
bullet(tf, "改ざんは不可（Groth16の束縛＋gnark検証で弾く・実証済）", first=True, bold=True)
bullet(tf, "残る穴 = 集約者が「どの証明を含めるか」を操作できる")
sub(tf, "例: 矛盾(spoof)票だけ間引いて consistentCount を水増し")
codebox(s, 0.7, 3.6, 8.6, 1.6, [
    "中身のすり替え … ✕ 不可能（暗号）",
    "集合の間引き   … △ 可能（検閲）  ← ここを塞ぐ",
])
tf = tbox(s, 0.7, 5.4, 8.7, 1.0)
bullet(tf, "解の核: 入力集合を検証可能／可用にして、間引きを不可能 or 検知可能に",
       size=13, first=True, bold=True)

# 3-8 各手法
solution(
    "A", "データ可用性 (DA)",
    ["間引きを『検知可能』にする（落とせば露見）"],
    ["各端末が proof/コミットを DA 層に drone_id×bucket で投稿",
     "誰でもデータを取得して再集約・照合できる"],
    "Celestia / EigenDA / Avail / Ethereum blob(EIP-4844→PeerDAS)",
)
solution(
    "B", "包含性証明",
    ["間引きを『不可能』にする（全件処理を証明）"],
    ["集約回路が「提出集合の Merkle/KZG ルート」を入力に取る",
     "そのルート配下の全有効 proof を処理したことを回路内で証明"],
    "KZG/Verkle コミット / zkVM(SP1・RISC Zero・Jolt・Nexus)",
)
solution(
    "C", "暗号 mempool",
    ["内容依存の『選択的』間引きを封じる"],
    ["集約者が集合を確定するまで中身(整合/矛盾)を見られない",
     "→ spoof票だけ狙って落とせない（commit-reveal / 閾値暗号）"],
    "Shutter Network（閾値暗号）",
)
solution(
    "D", "強制包含 / 共有シーケンサ",
    ["単一主体による検閲を仕組みで排除"],
    ["提出を強制的にブロックへ含める（inclusion list）",
     "順序付けを単一主体から分離する"],
    "FOCIL(Ethereum) / based rollup / Espresso / Astria",
)
solution(
    "E", "オプティミスティック＋詐欺証明",
    ["経済的抑止で間引きを割に合わなくする"],
    ["集約者が結果を投稿 → チャレンジ窓を設ける",
     "DA データから omission を証明されたらスラッシュ"],
    "Optimistic rollup 系の fraud proof + DA",
)
solution(
    "F", "分散集約 / M1直接",
    ["単一集約者そのものを無くす／迂回する"],
    ["複数の独立集約者、またはプルーフ集約ネットワークに委託",
     "各端末が自分の proof を直接チェーンに出せる経路を常設(M1)"],
    "Nebra UPA / Aligned Layer / Succinct(SP1) / BLS集約署名",
)

# 9 まとめ
s = add_slide()
title(s, "まとめ — 何が何を解くか")
table(s, 0.7, 1.5, 8.6, [
    ["穴", "効く手法"],
    ["間引きの検知", "A データ可用性(DA)"],
    ["間引きの不可能化", "B 包含性証明（＋DAルート）"],
    ["内容依存の選択排除", "C 暗号mempool(Shutter)"],
    ["単一主体の検閲", "D 強制包含/共有シーケンサ, F 分散"],
    ["迂回（フォールバック）", "F M1直接 / BLS集約署名"],
], col_w=[3.0, 5.6], fs=11, row_h=0.5)
tf = tbox(s, 0.7, 5.2, 8.7, 1.6)
subhead(tf, "推奨の積み方（この構成）", 0)
bullet(tf, "手軽=M1直接フォールバック → 本命=DA＋包含性証明 → さらに=暗号mempool",
       first=True, bold=True)
sub(tf, "最新潮流: DA層(Celestia/EigenDA/blob) ＋ zkVM/KZG での包含性証明")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "censorship_solutions.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
