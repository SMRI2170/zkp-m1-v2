#!/usr/bin/env python3
"""許可型ブロックチェーン（Besu QBFT）による検閲耐性まとめスライド。

なぜ許可型か → アーキテクチャ(M1直接提出) → なぜ検閲に強いか →
実装(既存contracts無改変＋deploy＋Besu構成) → トレードオフ/ハイブリッド。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.1)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.15, 8.2, 0.5)
run(tf.paragraphs[0], "検閲耐性の解 / 許可型チェーン", 15, GRAY)
tf = tbox(s, 0.88, 2.7, 8.4, 1.5)
run(tf.paragraphs[0], "許可型ブロックチェーンで", 28, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "検閲・DA・包含性を『合意』で解く", 28, INK, bold=True)
hrule(s, 4.5, l=0.92, w=5.6, color=ACCENT)
tf = tbox(s, 0.92, 4.68, 8.2, 1.4)
run(tf.paragraphs[0], "Besu QBFT ／ 各端末が直接提出（M1）／ 既存 contracts を無改変", 13.5, INK)

# 2 なぜ許可型か
s = add_slide()
title(s, "なぜ許可型チェーンか")
tf = tbox(s, 0.7, 1.4, 8.7, 1.6)
bullet(tf, "残る穴 = 集約者が有効な証明を間引ける（検閲）", first=True, bold=True)
bullet(tf, "公開チェーンでは DA層＋包含性証明で塞ぐ（重い）")
bullet(tf, "許可型なら、その多くを『合意プロトコルそのもの』で解ける", bold=True)
codebox(s, 0.7, 3.35, 8.6, 1.9, [
    "公開チェーン: DA(Celestia/EigenDA/blob) + 包含性証明(zkVM/KZG)",
    "許可型チェーン: BFT合意 + 全複製 だけで 検閲/DA/包含性 がほぼ同時に片づく",
    ("  信頼を「コンソーシアムの正直クォーラム」に置く（トレードオフ）", True),
])
tf = tbox(s, 0.7, 5.5, 8.7, 1.0)
bullet(tf, "UAV/RemoteID は当局・運航事業者のコンソーシアムが自然な信頼集合＝適合",
       size=12.5, first=True)

# 3 アーキテクチャ
s = add_slide()
title(s, "アーキテクチャ — M1 直接提出")
codebox(s, 0.7, 1.4, 8.6, 3.6, [
    "各端末 = rapidsnark で Groth16 証明を生成",
    "   │  submitAttestation(proof, publicSignals[18]) を直接呼ぶ",
    "   ▼",
    "[許可型EVM: Besu QBFT 4バリデータ]",
    "   ・BFT合意で提出Txを必ずブロックへ（検閲不可）",
    "   ・全ノードが複製（DA内蔵）",
    "   ・AttestationRegistry: Groth16検証 + M3検査 + K台集計",
    "        → ConsensusReached(trusted) / SpoofAlert",
])
tf = tbox(s, 0.7, 5.5, 8.7, 1.0)
bullet(tf, "gas制約が無いので集約(gnark)は任意。必要なら AggregateVerifier も同様に載る",
       size=12, first=True)

# 4 なぜ検閲に強いか
s = add_slide()
title(s, "なぜ検閲に強いか")
table(s, 0.7, 1.5, 8.6, [
    ["検閲の穴", "許可型チェーンでの解決"],
    ["単一集約者が間引ける", "BFT合意: 正直クォーラム(≧3/4)で提出は必ず包含"],
    ["間引きの検知 (DA)", "全バリデータが複製 → omission が即バレる"],
    ["包含性", "合意の liveness ＋ 契約で inclusion 強制可"],
    ["責任追及", "メンバーシップ管理: 検閲ノードを特定・除名"],
], col_w=[2.6, 6.0], fs=11, row_h=0.55)
tf = tbox(s, 0.7, 4.7, 8.7, 1.6)
bullet(tf, "検閲するにはバリデータ・クォーラムの共謀が必要＝ガバナンスで排除可能",
       first=True, bold=True)
sub(tf, "「改ざんは暗号で不可」＋「検閲は合意で不可」→ 耐改竄が揃う")

# 5 実装
s = add_slide()
title(s, "実装 — 既存資産を無改変で載せる")
tf = tbox(s, 0.7, 1.35, 8.7, 1.3)
bullet(tf, "contracts はチェーン非依存の EVM バイトコード → Besu/Quorum に無改変", first=True, bold=True)
sub(tf, "Groth16Verifier(uint18) + AttestationRegistry（テスト14/14はそのまま有効）")
codebox(s, 0.7, 2.7, 8.6, 2.5, [
    "contracts/scripts/deploy.js      … 任意のEVMへデプロイ＋M3初期設定",
    "contracts/hardhat.config.js      … besu ネットワーク先を追加",
    "contracts/besu-qbft/             … QBFT 4バリデータ構成",
    "   qbftConfigFile.json / docker-compose.yml / README",
    "",
    "$ DEPLOYER_KEY=.. K=3 OPERATOR=.. MIN_FRESHNESS=.. \\",
    "    npx hardhat run scripts/deploy.js --network besu",
])
tf = tbox(s, 0.7, 5.5, 8.7, 0.9)
bullet(tf, "デプロイ動作を EVM ノードで確認済（Besu も同一インターフェース）", size=12, first=True)

# 6 トレードオフ / ハイブリッド
s = add_slide()
title(s, "トレードオフとハイブリッド")
tf = tbox(s, 0.7, 1.4, 8.7, 2.3)
subhead(tf, "トレードオフ（正直に）", 0)
bullet(tf, "公開チェーンの『無許可・トラストレス』な検閲耐性ではない", first=True)
bullet(tf, "信頼は「既知バリデータの正直クォーラム」に置く")
sub(tf, "ただし UAV/RemoteID は規制・コンソーシアム前提なので許容範囲")
tf = tbox(s, 0.7, 3.85, 8.7, 2.6)
subhead(tf, "ハイブリッド（公開監査性の補完）", 0)
bullet(tf, "consensus の Merkle ルートを定期的に公開チェーンへアンカー", first=True)
bullet(tf, "→ コンソーシアム外からの改ざん検知（公開監査）も得られる")
subhead(tf, "まとめ")
bullet(tf, "許可型: M1直接＋BFT＋全複製で検閲/DA/包含性を一挙に解決。実装は既存無改変",
       first=True, bold=True)

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "permissioned_chain.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
