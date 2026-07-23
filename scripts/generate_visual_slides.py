import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PATH = Path("project_highly_visual_presentation.pptx")

W, H, M = Inches(10), Inches(7.5), Inches(0.45)
IW = W - M * 2

# Colors
C_BG    = RGBColor(248, 250, 252)
C_WHITE = RGBColor(255, 255, 255)
C_NAVY  = RGBColor(15, 23, 42)
C_BLUE  = RGBColor(37, 99, 235)
C_TEAL  = RGBColor(13, 148, 136)
C_ORANGE= RGBColor(234, 88, 12)
C_GRAY  = RGBColor(100, 116, 139)
C_LIGHT = RGBColor(226, 232, 240)
C_TEXT  = RGBColor(30, 41, 59)

FONT = "Hiragino Sans"

def rect(slide, l, t, w, h, bg, border=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = bg
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    return s

def txb(slide, text, l, t, w, h, size=14, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    box.text_frame.word_wrap = True
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return box

def header(slide, title, sub, page=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C_BG
    rect(slide, 0, 0, W, Inches(0.8), C_NAVY)
    rect(slide, 0, Inches(0.8), W, Inches(0.05), C_BLUE)
    txb(slide, title, M, Inches(0.15), IW, Inches(0.5), size=24, bold=True, color=C_WHITE)
    txb(slide, sub, M, Inches(1.0), IW, Inches(0.3), size=12, color=C_GRAY)
    if page:
        txb(slide, str(page), W - Inches(0.6), H - Inches(0.4), Inches(0.4), Inches(0.3), size=10, color=C_GRAY, align=PP_ALIGN.RIGHT)

def draw_flow(slide, steps, top_y=Inches(2.5)):
    n = len(steps)
    box_w = Inches(1.8)
    box_h = Inches(1.2)
    gap = (IW - (box_w * n)) / (n - 1) if n > 1 else 0
    
    for i, step in enumerate(steps):
        x = M + (box_w + gap) * i
        
        # Box
        rect(slide, x, top_y, box_w, box_h, C_WHITE, border=C_BLUE)
        rect(slide, x, top_y, box_w, Inches(0.3), C_BLUE)
        
        txb(slide, f"Step {i+1}", x, top_y+Inches(0.02), box_w, Inches(0.25), size=11, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        txb(slide, step, x+Inches(0.1), top_y+Inches(0.35), box_w-Inches(0.2), box_h-Inches(0.35), size=11, align=PP_ALIGN.CENTER)
        
        # Arrow
        if i < n - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + box_w + Inches(0.05), top_y + box_h/2 - Inches(0.15), gap - Inches(0.1), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = C_LIGHT
            arrow.line.fill.background()

def draw_comparison(slide, left_title, left_items, right_title, right_items):
    w = (IW - Inches(0.4)) / 2
    
    # Left
    rect(slide, M, Inches(1.8), w, Inches(4.5), C_WHITE, border=C_GRAY)
    rect(slide, M, Inches(1.8), w, Inches(0.5), C_GRAY)
    txb(slide, left_title, M, Inches(1.9), w, Inches(0.4), size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    for i, item in enumerate(left_items):
        txb(slide, f"❌ {item}", M+Inches(0.2), Inches(2.5) + Inches(0.6*i), w-Inches(0.4), Inches(0.5), size=12)

    # Right
    rect(slide, M + w + Inches(0.4), Inches(1.8), w, Inches(4.5), C_WHITE, border=C_TEAL)
    rect(slide, M + w + Inches(0.4), Inches(1.8), w, Inches(0.5), C_TEAL)
    txb(slide, right_title, M + w + Inches(0.4), Inches(1.9), w, Inches(0.4), size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    for i, item in enumerate(right_items):
        txb(slide, f"✅ {item}", M + w + Inches(0.6), Inches(2.5) + Inches(0.6*i), w-Inches(0.4), Inches(0.5), size=12)

def draw_architecture(slide, layers):
    h = Inches(4.5) / len(layers)
    for i, (title, desc, color) in enumerate(layers):
        y = Inches(1.8) + (h + Inches(0.1)) * i
        rect(slide, M, y, IW, h, C_WHITE, border=color)
        rect(slide, M, y, Inches(2.5), h, color)
        txb(slide, title, M, y + h/2 - Inches(0.15), Inches(2.5), Inches(0.3), size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        txb(slide, desc, M + Inches(2.7), y + Inches(0.1), IW - Inches(2.8), h - Inches(0.2), size=12)

def draw_metrics(slide, metrics):
    w = (IW - Inches(0.4)) / len(metrics)
    for i, (val, label) in enumerate(metrics):
        x = M + (w + Inches(0.2)) * i
        rect(slide, x, Inches(2.5), w, Inches(2.0), C_WHITE, border=C_BLUE)
        txb(slide, val, x, Inches(2.8), w, Inches(0.6), size=36, bold=True, color=C_ORANGE, align=PP_ALIGN.CENTER)
        txb(slide, label, x, Inches(3.6), w, Inches(0.5), size=14, color=C_GRAY, align=PP_ALIGN.CENTER)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]
page = 1

# 1. Title
s = prs.slides.add_slide(blank)
rect(s, 0, 0, W, H, C_NAVY)
rect(s, 0, Inches(2.0), W, Inches(2.5), C_BLUE)
txb(s, "UAV Safety Proof Project", M, Inches(2.2), IW, Inches(0.8), size=44, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txb(s, "Visual Architecture & Implementation Report", M, Inches(3.2), IW, Inches(0.5), size=20, color=C_LIGHT, align=PP_ALIGN.CENTER)
txb(s, "OpenDroneID | ZKP | Flutter | YOLO", M, Inches(4.0), IW, Inches(0.5), size=14, color=C_WHITE, align=PP_ALIGN.CENTER)

# 2. Problem vs Solution (Comparison)
s = prs.slides.add_slide(blank)
header(s, "1. 背景と課題解決", "なぜゼロ知識証明(ZKP)が必要なのか", page)
page += 1
draw_comparison(s, "従来の課題 (生データ公開)", [
    "飛行ルートや現在地が常に筒抜けになる",
    "業務機密や運用者のプライバシーが侵害される",
    "通信の傍受によるスプーフィングのリスク",
    "第三者がドローンを特定・追跡しやすくなる"
], "ZKPによる解決策 (非公開証明)", [
    "生座標をスマホ内(Private Input)に留める",
    "「30m以内にいたか」のTrue/Falseだけを出力",
    "数学的な計算で事実を証明し、改ざんを防止",
    "運用者のプライバシーを守りつつ安全を担保"
])

# 3. Overall Architecture
s = prs.slides.add_slide(blank)
header(s, "2. 全体アーキテクチャ", "システムを構成する4つの主要レイヤー", page)
page += 1
draw_architecture(s, [
    ("UAV (Raspberry Pi)", "OpenDroneID (ASTM F3411-22a) 準拠のBLEブロードキャスト。GPSとBME280によるQFE相対高度。", C_NAVY),
    ("Receiver (Flutter)", "BLEスキャンによるUAV位置の取得と、自身のGPS座標・気圧の統合。動画撮影による証拠記録。", C_BLUE),
    ("AI Vision (YOLOv11)", "録画した映像からTFLiteモデルを用いてドローンを視覚的に検知。ログと映像フレームを紐付け。", C_TEAL),
    ("ZKP Engine (Circom)", "端末内でWitnessを計算し、rapidsnarkでGroth16証明を生成。オンデバイス検証。", C_ORANGE)
])

# 4. OpenDroneID Flow
s = prs.slides.add_slide(blank)
header(s, "3. OpenDroneID 通信フロー", "独自フォーマットからASTM国際標準への移行", page)
page += 1
draw_flow(s, [
    "UAV GPS/気圧\n取得 (Pi)",
    "QFE相対高度\nへの変換",
    "BLE Service Data\n(UUID 0xFFFA) 構築",
    "BasicIDとLocation\nの1秒毎ローテーション",
    "スマホでの\nBLEスキャン・パース"
])

# 5. QFE Altitude Comparison
s = prs.slides.add_slide(blank)
header(s, "4. QFE 高度方式の導入", "ハードウェア間の高度基準ズレをどう解決するか", page)
page += 1
draw_comparison(s, "GPS / 絶対気圧方式", [
    "GPSの垂直誤差は±10~30mと非常に大きい",
    "標準気圧(1013hPa)基準では天候で絶対値がズレる",
    "Piとスマホのセンサー個体差によるズレ",
    "正確な3D距離計算(30m以内判定)が不可能"
], "QFE (相対高度) 方式", [
    "両端末とも起動時の地上気圧をP₀として記録",
    "以後はP₀からの相対的な高度変化(m)のみを使用",
    "天候の変化がP₀と現在値で相殺される",
    "「地上からの高さの差」を高精度に比較可能"
])

# 6. Flutter App Pipeline
s = prs.slides.add_slide(blank)
header(s, "5. Flutter オンデバイス処理", "スマホ内で完結するZKP証明パイプライン", page)
page += 1
draw_flow(s, [
    "BLEログと自身\nのGPSを統合",
    "動画撮影と\nAI検知の実行",
    "ZKPのJSON\n入力データ生成",
    "circom_witnesscalc\nによるWitness計算",
    "rapidsnarkによる\nGroth16証明生成"
])

# 7. ZKP Optimization Metrics
s = prs.slides.add_slide(blank)
header(s, "6. ZKP回路 最適化の成果", "Circomコンパイラ --O2 オプションの効果", page)
page += 1
txb(s, "最適化により、線形制約(加算・定数乗算)のオーバーヘッドが完全に消滅。\nしかし、重いコスト要因である非線形制約(比較・ビット分解)はそのまま残る。", M, Inches(1.5), IW, Inches(0.8), size=14)
draw_metrics(s, [
    ("609 → 534", "Constraints (制約数) の削減"),
    ("0", "O2での線形制約数 (Linear Constraints)"),
    ("106 KB", "最適化後の R1CS サイズ"),
    ("~2 sec", "Pixel 8 での証明生成時間")
])

# 8. ZKP Circuit Variants Architecture
s = prs.slides.add_slide(blank)
header(s, "7. ZKP 回路のバリエーション", "ユースケースに応じた4つの回路設計", page)
page += 1
draw_architecture(s, [
    ("Detection Flag", "最もシンプル。30m以内かつ5秒以内かの「True/False」フラグのみをPublic出力。", C_NAVY),
    ("Detection Commitment", "フラグに加え、入力座標の「Poseidon Hash」を出力。ブロックチェーンでの秘匿検証に最適。", C_TEAL),
    ("Feature Classifier", "距離判定に加え、「接近速度」や「高度帯」などの危険度特徴を分類して出力。", C_BLUE),
    ("Time Series (N=20)", "録画20秒分の全データを一括処理。通信欠損やラグに強い、堅牢なバッチ証明。", C_ORANGE)
])

# 9. Future Roadmap Flow
s = prs.slides.add_slide(blank)
header(s, "8. 今後のロードマップ", "通信・セキュリティ・信頼性のさらなる拡張", page)
page += 1
draw_flow(s, [
    "通信拡張:\nBT5 Long Range\nWi-Fi NaN",
    "認証強化:\nOpenDroneID\nAuthメッセージ",
    "改ざん防止:\n飛行ログの\nMerkle Proof",
    "モデル改善:\nUAV専用\nYOLOモデル",
    "統合テスト:\n実環境での\n大規模E2E"
], top_y=Inches(3.0))

prs.save(OUTPUT_PATH)
print(f"✅ Saved: {OUTPUT_PATH}")
