#!/usr/bin/env python3
"""音響ドローン検知 × ゼロ知識証明 の詳細まとめスライド。

2つの md をできる限り詳しく1本のデッキに統合する:
  - drone_detection_design_2026-06-22.md     (設計編)
  - acoustic_detection_summary_2026-06-22.md (実装編)
さらに会話で詰めた物理的背景・なりすまし対策・実装内部も収録。

デザイン（進捗第11回スタイル / ミニマル）:
  - 背景 白 / 黒文字（1A1A1A）、補助グレー（555555）。色は使わない。
  - フォント Yu Gothic（和文・英数とも）/ コードは Courier New。
  - 文章は四角で囲まない。枠はコード/ログ/数式の抜粋のみ（F7F7F7 + CCCCCC）。
  - タイトル直下に細い CCCCCC の横罫線。
  - 4:3 (10 x 7.5 inch)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font)
    return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True)
    return p


def bullet(tf, text, size=13.5, space_before=7, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold)
    return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY)
    return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill)
        tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table
    tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            bold = header and ri == 0
            run(p, str(val), fs, INK, bold=bold)
    return tbl


def section(text, sub_text=""):
    """章扉。"""
    s = add_slide()
    tf = tbox(s, 0.9, 3.0, 8.2, 1.4)
    run(tf.paragraphs[0], text, 30, INK, bold=True)
    hrule(s, 4.25, l=0.92, w=4.6)
    if sub_text:
        tf2 = tbox(s, 0.92, 4.4, 8.2, 0.8)
        run(tf2.paragraphs[0], sub_text, 14, GRAY)
    return s


# ════════════════════════════════════════════════════
# 1 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.2, 8.2, 0.5)
run(tf.paragraphs[0], "設計 + 実装 詳細まとめ（2026-06-22）", 15, GRAY)
tf = tbox(s, 0.88, 2.75, 8.4, 1.6)
run(tf.paragraphs[0], "音響ドローン検知 ×", 32, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "ゼロ知識証明", 32, INK, bold=True)
hrule(s, 4.4, l=0.92, w=5.4)
tf = tbox(s, 0.92, 4.58, 8.2, 1.4)
run(tf.paragraphs[0], "プロペラ音の「モデル一致」を証明の中で計算する", 15, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "ブランチ acoustic-detection-zk / コミット 681eacc", 12, GRAY, font=MONO)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "実音 DroneAudioDataset → 6次元特徴 → 線形分類 → ZK回路 → 証明検証 OK", 11.5, GRAY)

# ════════════════════════════════════════════════════
# 2 目次
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "目次")
tf = tbox(s, 0.7, 1.4, 8.7, 5.6)
subhead(tf, "設計編 — なぜこの設計か", 0)
bullet(tf, "脅威モデル / ZKの本質 / 検知手法の選択 / 設計方針 / 将来構想", first=True)
subhead(tf, "音響の物理 — なぜ音で検知できるのか")
bullet(tf, "プロペラ音の署名（BPF・倍音・ドップラー）/ 本物 vs 偽物", first=True)
subhead(tf, "実装編 — 何を作りどう動いたか")
bullet(tf, "データ / 特徴抽出 / 学習結果 / 量子化 / 回路内部 / ビルド / 検証", first=True)
subhead(tf, "現在地・残作業・今後")

# ════════════════════════════════════════════════════
# 章扉: 設計編
# ════════════════════════════════════════════════════
section("設計編", "なぜこの設計にするのか — 脅威モデルから決める")

# 背景：脅威モデル
s = add_slide()
title(s, "背景：何を守る仕組みか（脅威モデル）")
tf = tbox(s, 0.7, 1.3, 8.7, 1.3)
run(tf.paragraphs[0],
    "守りたいのは「利用者が、近接の事実を改ざんせず第三者へ証明できる」こと。"
    "攻撃者は2種類いて、対策が真逆になる。", 13.5, INK)
table(s, 0.7, 2.7, 8.6, [
    ["攻撃者", "脅威", "対策"],
    ["正直な利用者", "プライバシ流出（座標/映像/音声）", "ZKコミットメントで解決"],
    ["不正な利用者", "捏造（偽音・偽座標の入力）", "ZK単独では解決できない"],
], col_w=[2.2, 3.5, 2.9], row_h=0.6)
tf = tbox(s, 0.7, 4.85, 8.7, 2.2)
subhead(tf, "設計目標の再定義", 0)
bullet(tf, "「真実の証明」ではなく…", first=True)
bullet(tf, "「捏造コストを上げ、独立した複数証拠の整合性を検証可能にする」こと", bold=True)
sub(tf, "この一点から、手法選択も回路設計もすべて決まる")

# ZKの本質（NEW 深掘り）
s = add_slide()
title(s, "なぜ「真実の証明」はできないのか")
tf = tbox(s, 0.7, 1.3, 8.7, 2.4)
subhead(tf, "ZKが証明するのは『計算の正当性』", 0)
bullet(tf, "「この入力に、この計算を正しく適用し、この結果を得た」", first=True)
bullet(tf, "入力そのものが本物か（真正性）は、ZKの保証範囲の外", bold=True)
codebox(s, 0.7, 3.7, 8.6, 1.5, [
    "偽の音声 → 特徴抽出 → 分類器を通過 = 1",
    "  ↑ この計算過程は『正しく』証明できてしまう（garbage-in）",
], size=12)
tf = tbox(s, 0.7, 5.5, 8.7, 1.6)
run(tf.paragraphs[0],
    "→ だから真正性は『ZKの外』から持ち込む必要がある："
    "署名付きRemoteID・複数センサの整合・複数端末の独立観測。", 13.5, INK)

# 検知手法の全体像
s = add_slide()
title(s, "検知手法の全体像と現実性")
table(s, 0.55, 1.3, 8.9, [
    ["手法", "スマホ単体", "この案件での価値"],
    ["RF: Remote ID (BLE/WiFi)", "実装済み", "本命（唯一の真正性アンカー候補）"],
    ["画像AI（可視光・YOLO）", "実装済み", "補強の中心"],
    ["音響（プロペラ倍音）", "可能", "補強（夜間/視界外）"],
    ["RFスペクトル探知", "要SDR", "専用HW必須 → 不要"],
    ["サーマル/IR・レーダー・LiDAR", "不可/外付け", "HW依存・距離不足 → 不要"],
    ["磁力計", "センサ有", "距離が伸びず実用外 → 不要"],
    ["マルチ端末協調", "要通信", "なりすまし耐性の決定打（将来）"],
], col_w=[3.4, 1.7, 3.8], row_h=0.45, fs=11.5)
tf = tbox(s, 0.7, 5.85, 8.7, 1.2)
run(tf.paragraphs[0],
    "スマホ単体・追加HW不要で意味があるのは「音響」と「マルチ端末協調」だけ。",
    13.5, INK, bold=True)

# 各手法を退けた理由（NEW 詳細）
s = add_slide()
title(s, "各手法を退けた理由（詳細）")
table(s, 0.55, 1.4, 8.9, [
    ["手法", "退けた理由"],
    ["RFスペクトル探知", "操縦/映像電波の直接検出は SDR が必須。スマホで生スペクトルは取れない"],
    ["サーマル/IR", "夜間に強いが外付けカメラモジュールが必要"],
    ["レーダー", "回転翼のマイクロドップラーは最も確実だが専用機。スマホ不可"],
    ["LiDAR", "一部端末のみ。測距5m程度の超近距離で実質不可"],
    ["磁力計", "モータ磁場は微弱。距離が伸びず実用外"],
], col_w=[2.2, 6.7], row_h=0.62, fs=12)
tf = tbox(s, 0.7, 5.9, 8.7, 1.1)
sub(tf, "新HWを足す方向は本案件では筋が悪い。既存のスマホ能力で完結させる", 12.5)

# ZKできること/できないこと
s = add_slide()
title(s, "ZKでできること・できないこと")
tf = tbox(s, 0.7, 1.3, 8.7, 5.8)
subhead(tf, "できること", 0)
bullet(tf, "生データ（座標/映像/音声）を出さずに条件成立を証明", first=True)
bullet(tf, "コミット済み多センサ証拠に対するポリシー検証")
bullet(tf, "安い演算：加算・乗算・比較（内積＝matched filter、しきい値）")
subhead(tf, "高コスト・回路外に出すべきもの")
bullet(tf, "FFT/メル変換（三角関数・浮動小数）", first=True)
bullet(tf, "深層CNN/Transformer 推論（数百万〜億 constraint）")
sub(tf, "→ 重い処理は回路の外で計算し、結果（特徴）だけコミットして渡す")
subhead(tf, "原則")
bullet(tf, "ニューラルネットは固定小数（整数）に量子化してから回路へ", first=True)

# 設計方針
s = add_slide()
title(s, "設計方針：必須 / 後回し / 不要")
table(s, 0.55, 1.35, 8.9, [
    ["区分", "項目"],
    ["必須", "距離+時刻のコア証明（既に動く）"],
    ["必須", "検知シグナルを1つだけ証明に束ねる（映像YOLO + frame_commitment）"],
    ["後回し", "音響検知 / RemoteID署名(EdDSA) / 多モダリティ融合(K≥2)"],
    ["将来", "マルチ端末協調（なりすまし耐性の決定打）"],
    ["不要(今)", "ドップラー強化・TEE・フル深層zkML・モデル多重保持"],
], col_w=[1.6, 7.3], row_h=0.52, fs=12.5)
tf = tbox(s, 0.7, 5.5, 8.7, 1.5)
run(tf.paragraphs[0],
    "いま着手するのは必須の2行だけ。残りは『後から足せる穴』を開けておく。",
    13.5, INK, bold=True)
sub(tf, "可変長 features[] と パラメータ K の穴を用意（YAGNI）")

# ベストプラクティス
s = add_slide()
title(s, "作る時のベストプラクティス")
tf = tbox(s, 0.7, 1.35, 8.7, 5.8)
bullet(tf, "End-to-end を先に通す（機能追加はその後）", first=True, bold=True)
bullet(tf, "一度に1機能。回路を先に分岐させない", bold=True)
sub(tf, "1本の回路をパラメータ化し、使わない機能はパラメータで殺す（YAGNI）")
bullet(tf, "重い計算は回路の外、結果だけコミット", bold=True)
bullet(tf, "各層を単体検証してから合成（回路だけ→検知だけ→1周）", bold=True)
bullet(tf, "拡張点は『穴』だけ開ける（実装しない）", bold=True)
bullet(tf, "捏造耐性は最初から完璧を狙わない", bold=True)
sub(tf, "デモ/研究は『整合性を検証可能にする土台』で十分")

# 将来構想：述語代数（NEW 詳細）
s = add_slide()
title(s, "将来構想：融合回路の述語代数")
tf = tbox(s, 0.7, 1.25, 8.7, 0.6)
run(tf.paragraphs[0], "各センサを独立した述語にし、最後に合成する（すべて安い演算）。",
    13, INK)
table(s, 0.55, 1.95, 8.9, [
    ["述語", "内容"],
    ["within30", "drone↔target 距離 ≤ 30m"],
    ["time_ok", "|drone_time − target_time| ≤ 5s"],
    ["video_ok", "uav_detected ∧ score≥thr ∧ class==UAV ∧ frame_time一致"],
    ["audio_ok", "⟨audio_features, weights⟩ ≥ thr ∧ audio_time一致"],
    ["evidence_count", "video_ok + audio_ok + remoteid_present"],
    ["final", "within30 ∧ time_ok ∧ (evidence_count ≥ K)"],
], col_w=[2.3, 6.6], row_h=0.46, fs=11.5)
tf = tbox(s, 0.7, 5.85, 8.7, 1.2)
run(tf.paragraphs[0],
    "捏造耐性は『独立≥2モダリティが同時刻窓で一致』を要求して生まれる。",
    13.5, INK, bold=True)

# 将来構想：真正性アンカー
s = add_slide()
title(s, "将来構想：真正性アンカーと役割の非対称化")
tf = tbox(s, 0.7, 1.3, 8.7, 3.0)
subhead(tf, "RemoteID = 真正性アンカー", 0)
bullet(tf, "ASTM F3411 の署名付き RemoteID を circomlib の EdDSA で回路内検証",
       first=True)
bullet(tf, "地理側だけ『主張』から『暗号的事実』へ昇格")
subhead(tf, "音響・映像 = 補強")
bullet(tf, "署名が無いのでコミットメントのみ。整合性で捏造コストを上げる役", first=True)
tf2 = tbox(s, 0.7, 4.7, 8.7, 2.3)
subhead(tf2, "マルチ端末協調（なりすまし耐性の決定打）", 0)
bullet(tf2, "複数スマホが同一機を同時刻に近接検知し、各自ZK証明", first=True)
bullet(tf2, "1台のスピーカー偽装では複数地点の整合を作れない", bold=True)

# ════════════════════════════════════════════════════
# 章扉: 音響の物理
# ════════════════════════════════════════════════════
section("音響の物理", "なぜプロペラ音で検知できるのか / 本物と偽物の見分け")

# プロペラ音の署名（NEW）
s = add_slide()
title(s, "なぜプロペラ音は検知できるのか")
tf = tbox(s, 0.7, 1.3, 8.7, 2.6)
subhead(tf, "マルチコプター音は素直な構造を持つ", 0)
bullet(tf, "基本周波数 BPF = 回転数(rps) × ブレード枚数", first=True)
sub(tf, "例: 6000rpm = 100rps × 2枚 → BPF ≈ 200Hz")
bullet(tf, "その整数倍の倍音が櫛状に並ぶ（harmonic comb：400/600/800…Hz）")
bullet(tf, "4基のモータの微妙な回転差 → うなり/ビート")
bullet(tf, "ホバ/加速で BPF が滑らかにスイープ → 時間-周波数で縞")
codebox(s, 0.7, 4.5, 8.6, 1.4, [
    "署名 = 低〜中域に等間隔の倍音列があり、それが時間でゆっくり動く",
    "  → 鳥（不規則な倍音）や風（ブロードバンド）と区別できる",
], size=12)
tf = tbox(s, 0.7, 6.1, 8.7, 0.9)
sub(tf, "弱点：風切り音・屋外騒音に弱い／距離が伸びると急速にSN低下（実用 数十m）", 12)

# 本物 vs 偽物（NEW）
s = add_slide()
title(s, "本物 vs 偽物：なりすまし対策")
tf = tbox(s, 0.7, 1.25, 8.7, 1.5)
subhead(tf, "レベルA：ドローン音 vs 環境音 → 音響分類で解ける", 0)
bullet(tf, "倍音の等間隔性・BPFの安定/スイープで鳥・風・車・人声を弾く", first=True)
subhead(tf, "レベルB：実機 vs スピーカー偽音 → 音単独では不可")
table(s, 0.55, 3.15, 8.9, [
    ["対策", "原理"],
    ["マルチモーダル整合", "音 × RemoteID × 映像 が同時刻・同方向で一致するか（最強）"],
    ["ドップラー/移動性", "実機はBPFがシフト。固定スピーカーは周波数が動かない"],
    ["空間性", "実機は空中の点音源。床置きスピーカーは到来方向が固定"],
    ["物理整合", "RemoteIDの距離と音圧/高域減衰が物理的に整合するか"],
    ["ライブネス", "再生検知（replay detection）。補助的"],
], col_w=[2.3, 6.6], row_h=0.5, fs=11.5)
tf = tbox(s, 0.7, 6.6, 8.7, 0.6)
sub(tf, "音は『もう1票』。単独で真正性を背負わせない", 12)

# ════════════════════════════════════════════════════
# 章扉: 実装編
# ════════════════════════════════════════════════════
section("実装編", "何を作り、どう動いたか — 公開データの実音で実証")

# 全体フロー
s = add_slide()
title(s, "実装：全体フロー（end-to-end 実証）")
codebox(s, 0.7, 1.45, 8.6, 3.7, [
    "DroneAudioDataset（実音 1,332件）",
    "  → features.py（6次元 解釈可能特徴）",
    "  → train_linear.py（線形分類 / ROC-AUC 0.814）",
    "      → 正/負分離の量子化重み WP / WN / ADD",
    "  → distance_30m_time_feature_classifier6.circom",
    "  → circom コンパイル → Groth16 証明生成",
    "  → snarkjs verify = OK!",
    "  → rapidsnark_app に zkey / vkey / wcd を配置",
], size=12.5)
tf = tbox(s, 0.7, 5.45, 8.7, 1.5)
run(tf.paragraphs[0],
    "「音響モデルの一致をZKの中で計算する」を、公開データの実音で一気通貫に実証。",
    13.5, INK, bold=True)

# データセット詳細（NEW）
s = add_slide()
title(s, "データセット：DroneAudioDataset")
table(s, 0.55, 1.35, 8.9, [
    ["フォルダ", "wav数", "用途"],
    ["Binary_Drone_Audio/yes_drone", "1,332", "ポジ（ドローン音）"],
    ["Binary_Drone_Audio/unknown", "10,372", "ネガ（環境音）"],
    ["Multiclass/bebop_1, membo_1", "各666", "機種別（今回未使用）"],
], col_w=[4.3, 1.5, 3.1], row_h=0.45, fs=11.5)
tf = tbox(s, 0.7, 3.5, 8.7, 1.6)
subhead(tf, "取得の工夫と教訓", 0)
bullet(tf, "フルクローン（23,409ファイル）はディスクを使い切り失敗", first=True)
bullet(tf, "sparse-checkout + blob:none filter で yes_drone / unknown のみ取得", bold=True)
bullet(tf, "ネガは unknown からランダム1,332でバランス")
tf2 = tbox(s, 0.7, 5.6, 8.7, 1.3)
sub(tf2, "本番端末マイクの自前録音を必ず混ぜる（公開データはマイク特性が異なる）", 12)
sub(tf2, "ESC-50 はネガ多様化に後で追加可（今回は自己完結を優先）", 12)

# 6次元特徴
s = add_slide()
title(s, "音響モデル：6次元の解釈可能特徴")
tf = tbox(s, 0.7, 1.25, 8.7, 0.6)
run(tf.paragraphs[0],
    "深層学習なしの最小構成（ZK親和を優先）。features.py で抽出。", 13, INK)
table(s, 0.55, 1.9, 8.9, [
    ["#", "特徴", "意味", "ドローンで"],
    ["0", "harmonic_comb_strength", "プロペラ倍音のエネルギー比", "高い"],
    ["1", "f0_norm", "推定BPF / 400Hz", "—"],
    ["2", "f0_stability", "BPFの時間安定度", "高い"],
    ["3", "tonality", "1 − スペクトル平坦度", "高い"],
    ["4", "low_mid_ratio", "100–1000Hz帯エネルギー比", "高い"],
    ["5", "doppler_slope", "BPFの時間傾き（移動音源）", "非ゼロ"],
], col_w=[0.5, 3.0, 3.7, 1.7], row_h=0.48, fs=11.5)
tf = tbox(s, 0.7, 5.6, 8.7, 1.0)
sub(tf, "解釈可能 = 後でZK回路の内積に素直に載る／挙動を説明できる", 12)

# 特徴抽出の中身（NEW）
s = add_slide()
title(s, "特徴抽出パイプラインの中身")
codebox(s, 0.7, 1.4, 8.6, 2.5, [
    "load: 16kHz mono / ピーク正規化",
    "STFT: nperseg=1024, noverlap=512, hann窓",
    "f0推定: 倍音和スペクトル（候補 60-400Hz, 倍音6本, ±1bin近傍ピーク）",
    "per-frame: f0 / harmonic_score / spectral_flatness",
    "voiced選別: harmonic_score > max(0.05, median)",
], size=11.5)
tf = tbox(s, 0.7, 4.15, 8.7, 2.9)
subhead(tf, "集約（クリップ→6次元）", 0)
bullet(tf, "harmonic_comb = mean(harmonic_score)", first=True)
bullet(tf, "f0_stability = 1 / (1 + std(f0)/20)  … 安定トーンほど1に近い")
bullet(tf, "tonality = 1 − mean(flatness)，low_mid_ratio = 100-1000Hz帯比")
bullet(tf, "doppler_slope = |f0の時間一次回帰の傾き| / 5  … 移動音源で非ゼロ")

# 学習結果
s = add_slide()
title(s, "学習結果")
codebox(s, 0.7, 1.4, 8.6, 1.6, [
    "ROC-AUC = 0.814    accuracy = 0.765   (5-fold CV)",
    "precision / recall : not_drone 0.775 / 0.745",
    "                     drone     0.755 / 0.784",
], size=12)
tf = tbox(s, 0.7, 3.3, 8.7, 3.7)
run(tf.paragraphs[0],
    "6次元・線形のベースラインとして妥当（実音でドローン/非ドローンが分離）。",
    14, INK, bold=True)
subhead(tf, "ただし実用ラインではない")
bullet(tf, "交差検証はPoC簡易版（Stratified）", first=True)
sub(tf, "本来は録音セッション単位で分割しないとリークし、数字が楽観値になる")
bullet(tf, "モデル/データ次第で改善余地が大きい")

# 限界・注意（NEW）
s = add_slide()
title(s, "限界・注意（正直なところ）")
tf = tbox(s, 0.7, 1.35, 8.7, 5.6)
subhead(tf, "重みの符号が直感と逆", 0)
bullet(tf, "harmonic/low_mid が負、tonality が強い正という学習結果", first=True)
sub(tf, "このデータ特有の癖（ネガが同じ屋内環境、ドローン録音も短く加工済み）")
bullet(tf, "別環境ではそのまま当てはまらない可能性", bold=True)
subhead(tf, "ドメインギャップ")
bullet(tf, "公開データのマイク特性と本番端末は異なる → 自前録音の混入が必須",
       first=True)
subhead(tf, "評価の正直さ")
bullet(tf, "録音セッション単位の分割で再評価すべき", first=True)

# ZK接続：負の重み
s = add_slide()
title(s, "ZK接続：負の重みをどう扱うか")
tf = tbox(s, 0.7, 1.25, 8.7, 1.2)
run(tf.paragraphs[0],
    "学習した重みは負を含む。Circomの比較は非負前提のため、"
    "正/負の重みを分離し非負オフセットで吸収する。", 13.5, INK)
codebox(s, 0.7, 2.6, 8.6, 2.0, [
    "pos = Σ WP[i]・features[i] + ADD_POS",
    "neg = Σ WN[i]・features[i] + ADD_NEG",
    "判定:  neg ≤ pos     ( ⟺  w・x + b ≥ 0 )",
], size=13)
tf = tbox(s, 0.7, 4.85, 8.7, 2.2)
subhead(tf, "量子化", 0)
bullet(tf, "features は 0..1 を ×1000、重みも ×1000 の整数", first=True)
bullet(tf, "すべて非負 → LessEqThan で安全に比較できる")
sub(tf, "回路: distance_30m_time_feature_classifier6.circom（既存4回路は無改変）")

# 量子化の導出（NEW）
s = add_slide()
title(s, "量子化の導出（worked example）")
codebox(s, 0.7, 1.35, 8.6, 2.7, [
    "決定: w・x_float + b ≥ 0      (x ∈ [0,1])",
    "x_int = round(x・S_X),  w_int = round(w・S_W)   (S_X=S_W=1000)",
    "B = round(b・S_X・S_W)",
    "WP=max(w_int,0), WN=max(-w_int,0)",
    "pos-neg+B ≥ 0  ⟺  neg + max(-B,0) ≤ pos + max(B,0)",
], size=11.5)
tf = tbox(s, 0.7, 4.3, 8.7, 0.5)
run(tf.paragraphs[0], "本ケース（b=-2.0245 → ADD_NEG=2,024,495, ADD_POS=0）：",
    12.5, INK, bold=True)
codebox(s, 0.7, 4.95, 8.6, 1.9, [
    "pos = 771・1000 + 4790・944                 = 5,292,760",
    "neg = 3771・44 + 3031・156 + 5913・60        =   993,540",
    "neg + ADD_NEG = 3,018,035  ≤  5,292,760 = pos  → model_ok = 1",
], size=11)

# 重みと検証結果
s = add_slide()
title(s, "反映した重みと証明検証")
table(s, 0.55, 1.3, 5.0, [
    ["特徴", "w", "WP", "WN"],
    ["harmonic_comb", "-3.77", "0", "3771"],
    ["f0_norm", "-3.03", "0", "3031"],
    ["f0_stability", "+0.77", "771", "0"],
    ["tonality", "+4.79", "4790", "0"],
    ["low_mid_ratio", "-5.91", "0", "5913"],
    ["doppler_slope", "-0.17", "0", "168"],
], col_w=[2.0, 1.1, 0.95, 0.95], row_h=0.4, fs=11)
codebox(s, 5.7, 1.3, 3.7, 4.1, [
    "[実ドローン入力]",
    "within30        = 1",
    "distance_ok     = 1",
    "time_ok         = 1",
    "frame_time_ok   = 1",
    "pos = 5,292,760",
    "neg = 3,018,035",
    "model_ok        = 1",
    "within30_and",
    "    _model      = 1",
], size=10.5)
tf = tbox(s, 0.7, 5.7, 8.7, 1.3)
run(tf.paragraphs[0], "snarkjs groth16 verify ... OK!", 15, INK, bold=True, font=MONO)
sub(tf, "実ドローンクリップの音響特徴で『接近 ∧ ドローン検知』が成立")

# 回路の内部構造（NEW）
s = add_slide()
title(s, "回路の内部構造")
codebox(s, 0.7, 1.35, 8.6, 2.0, [
    "private inputs: 15 (9 belong to witness)",
    "public outputs: 8",
    "wires: 764   labels: 900",
], size=12)
tf = tbox(s, 0.7, 3.55, 8.7, 3.5)
subhead(tf, "構成", 0)
bullet(tf, "Distance30mTimeCore（距離+時刻）を内包し AND 結合", first=True)
bullet(tf, "pos_acc / neg_acc で線形に内積累積（定数×signal＝非二次で安い）")
bullet(tf, "LessEqThan(32)：最大 ~3.8×10^7 < 2^32 で安全に収まる")
bullet(tf, "frame時刻差は AbsDiff + LessEqThan(64)")
subhead(tf, "公開出力（8）")
bullet(tf,
       "within30 / distance_ok / time_ok / frame_time_ok / pos / neg / "
       "model_ok / within30_and_model", first=True, size=12)

# ビルドパイプライン（NEW）
s = add_slide()
title(s, "ビルドパイプライン（ZK資産生成）")
codebox(s, 0.7, 1.4, 8.6, 3.5, [
    "circom --O2 --r1cs --wasm --sym",
    "snarkjs powersoftau new bn128 12  (pot12)",
    "snarkjs groth16 setup → zkey_0000",
    "snarkjs zkey contribute → groth16.zkey",
    "snarkjs zkey export verificationkey → vkey.json",
    "snarkjs wtns calculate → prove → public.json",
    "build-circuit (circom-witnesscalc) → .wcd   ← オンデバイス用",
], size=11.5)
tf = tbox(s, 0.7, 5.15, 8.7, 1.9)
subhead(tf, "オンデバイス（rapidsnark）", 0)
bullet(tf, ".wcd で端末内 witness 計算 → rapidsnark で証明生成・検証", first=True)
sub(tf, "build-circuit は Rust 製。cargo build に約7分（一度きり）")

# 成果物
s = add_slide()
title(s, "成果物（コミット 681eacc / 13ファイル）")
table(s, 0.55, 1.35, 8.9, [
    ["パス", "内容"],
    ["scripts/acoustic/features.py", "6次元特徴抽出"],
    ["scripts/acoustic/train_linear.py", "学習・評価・ZK量子化重み出力"],
    ["scripts/acoustic/fetch_datasets.sh", "データ取得"],
    ["circuits/...feature_classifier6.circom", "6次元ZK回路"],
    ["zk-distance/samples/...6.json", "実クリップ入力"],
    ["scripts/build_...classifier6_assets.sh", "専用ビルド"],
    ["rapidsnark_app/.../classifier6/", "zkey / vkey / wcd"],
    ["drone_detection_design_2026-06-22.md", "設計メモ"],
], col_w=[5.0, 3.9], row_h=0.42, fs=11)
tf = tbox(s, 0.7, 6.4, 8.7, 0.7)
sub(tf, "データ(data/)・学習成果物(zk_audio_params.json)は .gitignore で非追跡", 11.5)

# 再現手順
s = add_slide()
title(s, "再現手順")
codebox(s, 0.7, 1.45, 8.6, 4.5, [
    "# 1. 依存",
    "pip install -r scripts/acoustic/requirements.txt",
    "",
    "# 2. データ取得（sparse-checkout で yes_drone / unknown）",
    "#    → data/audio/{drone,not_drone} へ振り分け",
    "",
    "# 3. 学習 + ZK重み出力",
    "python scripts/acoustic/train_linear.py \\",
    "    --data-dir scripts/acoustic/data/audio",
    "",
    "# 4. WP/WN/ADD を circuit main() に反映後、ZKビルド",
    "BUILD_CIRCUIT_BIN=/path/to/build-circuit \\",
    "    ./scripts/build_distance_feature_classifier6_assets.sh",
    "",
    "# 5. 証明検証",
    "snarkjs groth16 verify <vkey> <public> <proof>",
], size=11)

# 現在地と残作業
s = add_slide()
title(s, "現在地と残作業")
table(s, 0.55, 1.35, 8.9, [
    ["項目", "状態"],
    ["音響PoC（実データ AUC 0.814）", "完了"],
    ["6次元ZK回路 + 学習済み重み", "完了"],
    ["Groth16 証明生成・検証（OK）", "完了"],
    ["アプリ assets（zkey/vkey/wcd）", "完了"],
    ["Dart 配線（オンデバイス動作）", "未着手"],
], col_w=[6.0, 2.9], row_h=0.44, fs=12.5)
tf = tbox(s, 0.7, 4.7, 8.7, 2.4)
subhead(tf, "残：Dart 配線（手順2）", 0)
bullet(tf, "マイク録音（record / iOS NSMicrophoneUsageDescription）", first=True)
bullet(tf, "features.py と同一の特徴抽出を Dart で実装（値ズレ厳禁）", bold=True)
bullet(tf, "zk_input_builder / zk_service を6次元回路に対応（features[6]）")
bullet(tf, "UI（収集タブ or 既存検知フローへ合流）")

# まとめ
s = add_slide()
title(s, "まとめ")
tf = tbox(s, 0.7, 1.35, 8.7, 5.8)
bullet(tf, "実音 → 6次元特徴 → 線形分類 → 6次元ZK回路 → 証明検証 OK を実証",
       first=True, bold=True)
bullet(tf, "負の重みは正/負分離＋非負オフセットでZKに安全に載せた", bold=True)
bullet(tf, "アプリ資産（zkey/vkey/wcd）まで生成＝オンデバイス準備完了", bold=True)
subhead(tf, "忘れてはいけない前提")
bullet(tf, "ZKは『計算の正当性』であり『入力の真正性』は保証しない", first=True)
sub(tf, "音単独ではスピーカー偽装を防げない")
bullet(tf, "偽装耐性は RemoteID / 映像との時空間整合と、将来のマルチ端末協調で担保")

# 今後
s = add_slide()
title(s, "今後")
tf = tbox(s, 0.7, 1.35, 8.7, 5.8)
bullet(tf, "Dart 配線でオンデバイス動作（録音→特徴→features[6]→証明）",
       first=True)
bullet(tf, "精度改善：録音セッション単位の分割で正直な評価")
bullet(tf, "ESC-50 でネガ多様化、ドップラー特徴の強化")
bullet(tf, "融合回路：within30 ∧ model ∧ (≥2モダリティ一致)")
bullet(tf, "真正性アンカー：RemoteID 署名の回路内検証（EdDSA）")
bullet(tf, "なりすまし耐性の本命：マルチ端末協調")

prs.save("acoustic_zk_summary_4x3.pptx")
print(f"saved acoustic_zk_summary_4x3.pptx ({len(prs.slides._sldIdLst)} slides)")
