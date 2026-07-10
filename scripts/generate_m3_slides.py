#!/usr/bin/env python3
"""M3（入力真正性）まとめスライド。

なぜM3か → ドローン署名RemoteID(偽造対策) → 新鮮さ(リプレイ対策) →
オペレータ/端末アテステーション(Sybil対策) → end-to-end結線 → 耐改竄マトリクス。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.1)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.2, 8.2, 0.5)
run(tf.paragraphs[0], "進捗まとめ / M3", 15, GRAY)
tf = tbox(s, 0.88, 2.75, 8.4, 1.4)
run(tf.paragraphs[0], "M3 — 入力の真正性", 30, INK, bold=True)
hrule(s, 4.2, l=0.92, w=5.2, color=ACCENT)
tf = tbox(s, 0.92, 4.38, 8.2, 1.6)
run(tf.paragraphs[0], "「K台合意＝真正」を成立させる最後のピース", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "ドローン署名RemoteID / 新鮮さ / オペレータ・端末アテステーション", 11.5, GRAY, font=MONO)

# 2 なぜM3か
s = add_slide()
title(s, "なぜ M3 か — 一番深い改竄の穴")
tf = tbox(s, 0.7, 1.4, 8.7, 2.1)
bullet(tf, "M1/M2 は「計算の正当性」と「束ね方」を強くしたが…", first=True)
bullet(tf, "入力の真正性には触れていなかった", bold=True)
sub(tf, "放送値も自端末位置も公開/自己申告 → 放送を知る攻撃者は偽造できた")
bullet(tf, "＝「K台合意」でも「真正」とは言えない状態だった")
codebox(s, 0.7, 3.85, 8.6, 1.7, [
    "耐改竄の5層:",
    "  1記録不変  2偽造耐性  3検閲耐性  4 Sybil  5 入力真正性 ←最深",
    "  M3 が 5(と一部4) を担い、真正性を完成させる",
])
tf = tbox(s, 0.7, 5.75, 8.7, 0.8)
bullet(tf, "方針: 外部の信用アンカーの署名を回路内検証／物理実測で裏取り", size=12, first=True, bold=True)

# 3 ドローン署名RemoteID
s = add_slide()
title(s, "① ドローン署名 RemoteID（偽造対策）")
tf = tbox(s, 0.7, 1.35, 8.7, 1.0)
bullet(tf, "放送内容をオペレータ鍵で署名し、回路内で EdDSA-Poseidon 検証", first=True)
sub(tf, "drone_auth_ok=1 のとき「登録実機由来の本物」＝放送傍受だけの偽造を封じる")
codebox(s, 0.7, 2.45, 8.6, 2.0, [
    "M_drone = Poseidon(remote_id, drone_lat, drone_lon, drone_alt, drone_time)",
    "EdDSA-Poseidon(operator_key) を回路内検証",
    "public出力: drone_auth_ok / drone_operator_commitment",
])
tf = tbox(s, 0.7, 4.65, 8.7, 1.8)
subhead(tf, "検証", 0)
bullet(tf, "正署名 → drone_auth_ok=1", first=True)
bullet(tf, "改ざん署名 → witness 生成が失敗（真正性を回路が強制）", bold=True)
bullet(tf, "soft toggle: 未署名RemoteIDは auth_ok=0 として許容（規格未普及に対応）")
sub(tf, "※ 署名付きRemoteID(DRIP/F3411 Auth)を先取り。実運用は普及に依存")

# 4 新鮮さ
s = add_slide()
title(s, "② 新鮮さビーコン（リプレイ対策）")
tf = tbox(s, 0.7, 1.35, 8.7, 1.0)
bullet(tf, "device 署名メッセージに freshness(直近ブロックハッシュ等)を混ぜる", first=True)
sub(tf, "古い proof を新しい時間窓で使い回せない")
codebox(s, 0.7, 2.45, 8.6, 1.7, [
    "M = Poseidon(drone_id, target_commit, frame_binding,",
    "            within30_and_detected, freshness)",
    "public出力: freshness_pub → チェーンが直近窓に入るか検査",
])
tf = tbox(s, 0.7, 4.35, 8.7, 1.8)
subhead(tf, "検証", 0)
bullet(tf, "freshness だけ差し替え（旧署名） → witness 失敗", first=True, bold=True)
bullet(tf, "チェーン: freshness >= minFreshness を要求（stale を revert）")

# 5 Sybil / アテステーション
s = add_slide()
title(s, "③ オペレータ・端末（Sybil対策）")
tf = tbox(s, 0.7, 1.35, 8.7, 2.3)
subhead(tf, "オペレータ登録（実装済）", 0)
bullet(tf, "drone_operator_commitment を registeredOperator で照合", first=True)
bullet(tf, "未登録オペレータの機体は revert → ドローン鍵のSybilを抑止")
subhead(tf, "端末アテステーション（骨組み・ネイティブ残）")
bullet(tf, "Android Key Attestation / iOS App Attest で device_pubkey が実機鍵と証明", first=True)
bullet(tf, "lib/core/device_attestation.dart にインターフェース＋MethodChannel骨組み")
sub(tf, "証明書チェーン検証（ネイティブ＋バックエンド/オンチェーン）が残実装")

# 6 end-to-end
s = add_slide()
title(s, "end-to-end 結線")
codebox(s, 0.7, 1.4, 8.6, 2.5, [
    "回路 uav_witness_attestation.circom (public18・約18k制約)",
    "   drone_auth_ok / drone_operator_commitment / freshness_pub を追加",
    "        │  verifier を uint[18] に再生成 (PoT 2^15)",
    "        ▼",
    "AttestationRegistry.submitAttestation:",
    "   require drone_auth_ok==1        # 偽造対策",
    "   require registeredOperator[op]  # Sybil対策",
    "   require freshness >= minFreshness # リプレイ対策",
])
tf = tbox(s, 0.7, 4.15, 8.7, 2.3)
subhead(tf, "テスト 14/14 通過", 0)
bullet(tf, "未認証ドローン / 未登録オペレータ / 古いfreshness を revert", first=True, bold=True)
bullet(tf, "実18信号 proof → 実 Groth16Verifier → 整合2台で trusted 合意（E2E）")
sub(tf, "M3 が witness回路→verifier→レジストリまで実際に通った")

# 7 耐改竄マトリクス
s = add_slide()
title(s, "耐改竄マトリクス — 現在地")
table(s, 0.7, 1.5, 8.6, [
    ["攻撃", "対策", "状態"],
    ["放送値の偽造", "ドローン署名RemoteID (drone_auth_ok)", "✅ 回路+チェーン"],
    ["リプレイ", "新鮮さ (freshness)", "✅ 回路+チェーン"],
    ["Sybil(オペレータ)", "registeredOperator", "✅ チェーン"],
    ["Sybil(端末)/端末真正", "端末アテステーション", "△ 骨組み(ネイティブ残)"],
    ["検閲(集約者の間引き)", "M1直接 / 包含性証明", "⬜ 未"],
    ["中継/ワームホール", "RSSIでは不可(要TDoA)", "✕ 原理的限界"],
], col_w=[2.5, 4.0, 2.1], fs=10.5, row_h=0.5)
tf = tbox(s, 0.7, 5.7, 8.7, 0.9)
bullet(tf, "偽造・リプレイ・オペレータSybil は閉じた。残るは端末真正性(ネイティブ)と検閲・中継",
       size=11.5, first=True, bold=True)

# 8 現状と残り
s = add_slide()
title(s, "現状と残り")
tf = tbox(s, 0.7, 1.45, 8.7, 2.4)
subhead(tf, "できたこと", 0)
bullet(tf, "M3: ドローン署名RemoteID + 新鮮さ を回路で強制、チェーンまでE2E結線", first=True)
bullet(tf, "オペレータ登録でドローン鍵のSybilを抑止。テスト14/14")
bullet(tf, "端末アテステーションのアプリ側インターフェースを用意")
tf = tbox(s, 0.7, 3.85, 8.7, 2.6)
subhead(tf, "残り", 0)
bullet(tf, "端末アテステーションのネイティブ実装＋証明書チェーン検証", first=True)
bullet(tf, "検閲耐性: M2に包含性証明 / M1フォールバック")
bullet(tf, "実機DJIでの RemoteID署名・RSSI較正（規格普及・実機依存）")
sub(tf, "到達点: 「K台の登録witnessが真正な放送を物理的に裏取りした」に前進")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "m3_input_authenticity.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
