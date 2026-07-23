#!/usr/bin/env python3
"""week19/05_決定ドキュメント_ADR.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16, size1=14):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 2))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.07 if i else 0)
    return y + h + 0.12


def quote(s, y, lines, w=W, fs=15):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.3
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(1.0); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(12); tf.margin_top = Pt(8); tf.margin_bottom = Pt(8)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 3)
        run(p, ln, fs, INK, bold=True)
    return y + bh + 0.14


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "決定ドキュメント（ADR）", 30, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.6)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "手法とフレームワークの選定 — 結論と根拠を1枚に", 15, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 05_決定ドキュメント_ADR", 12, GRAY)

# S2 Context
s = add_slide(); title(s, "決めたいこと（Context）")
bullets(s, TOP, [
    (0, "目的 = 各端末でRSSI/距離/複雑計算を秘匿証明→多端末を1つに→チェーンで共有・検証"),
    (1, "端末はFlutterアプリ。セキュリティが核（★1トラステッドセットアップ含む）"),
    (0, "現状 = circom+rapidsnark(端末) / gnark(集約M2) / EVMでGroth16検証"),
    (1, "M1(集計)・M2(再帰)・#3b(網羅性)は設計/実装済"),
    (0, "論点 = この現状をどこまで維持し何を変えるか"),
    (1, "(a)モバイル配布の標準化 (b)★1セットアップの根治 (c)集約の軽量化"),
])

# S3 評価軸
s = add_slide(); title(s, "評価軸（重み）")
table(s, TOP, [
    ["軸", "重み", "理由"],
    ["現行資産の再利用（移行コスト小）", "★★★", "動いている circom/gnark を捨てない"],
    ["セットアップの健全性（★1根治）", "★★★", "プロジェクトの核＝セキュリティ"],
    ["オンチェーン検証コスト（gas）", "★★☆", "チェーン記録が前提"],
    ["モバイル証明の現実性", "★★★", "端末で秒級・省メモリ"],
    ["集約適性（多端末を1つに）", "★★☆", "本案件の新規性"],
    ["実装難度・事例の多さ", "★★☆", "開発リソース"],
], col_w=[3.4, 1.2, 4.0], fs=12, row_h=0.46)

# S4 選択肢と評価
s = add_slide(); title(s, "選択肢と評価（要約）")
table(s, TOP, [
    ["候補", "再利用", "★1根治", "gas", "集約", "総評"],
    ["現行維持(circom+gnark)", "◎", "✕", "◎", "◎(M2)", "基準。★1が残る"],
    ["＋Moproでモバイル標準化", "◎", "✕", "◎", "◎", "低コストで底上げ"],
    ["Noir+bbへ端末を移行", "△", "◎", "✕(2.4M)", "要再設計", "★1根治だがgas増"],
    ["halo2", "△", "○", "△", "重", "費用対効果低"],
    ["Plonky2で端末回路", "✕", "◎", "△", "◎", "署名と体不一致で不向き"],
    ["Nova/Sonobeで集約", "○", "○", "圧縮次第", "◎", "端末側集約に最適"],
    ["Cairo/Stwo", "✕", "◎", "Starknet", "○", "EVM前提と不整合"],
    ["RISC Zero/SP1", "✕", "◎", "△", "○", "端末に重い(将来)"],
], col_w=[2.7, 1.0, 1.0, 1.1, 1.1, 1.7], fs=10.5, row_h=0.42)

# S5 決定（段階的）
s = add_slide(); title(s, "決定 — 段階的アプローチ")
bullets(s, TOP, [
    (0, "フェーズ1（すぐ・低リスク）= Moproでモバイル標準化"),
    (1, "現行circom+rapidsnarkのままmopro_flutter_packageに載せ替え。回路・検証・集約は変えない"),
    (0, "フェーズ2（並行・研究）= 集約の端末寄せをNova/Foldingで検証"),
    (1, "per-witness回路はそのまま、集約だけNovaでPoC。gnarkはフォールバックで残す"),
    (0, "フェーズ3（条件付き）= ★1根治のためNoir評価"),
    (1, "端末回路をNoirへ移植して比較。ただしgas増(2.4M)と集約作り直しを許容できる場合のみ"),
])

# S6 Rejected
s = add_slide(); title(s, "明示的に採らない（Rejected）")
bullets(s, TOP, [
    (0, "halo2 = 回路実装(chip/gadget)の難度が高く、必要機能に対し費用対効果が低い"),
    (0, "Plonky2を端末回路に = EdDSA(BabyJubJub)とGoldilocks体が不一致で非ネイティブ演算が重い"),
    (1, "（集約用途としては将来検討可）"),
    (0, "Cairo/Stwo = モバイル証明は優秀だが検証がStarknet前提でEVM+gnarkと不整合"),
    (0, "RISC Zero/SP1 = 端末単独の証明生成が重く、当面リモートprover前提でオンデバイスにならない"),
], size0=16.5)

# S7 Consequences
s = add_slide(); title(s, "結果として起きること（Consequences）")
bullets(s, TOP, [
    (0, "良い点"),
    (1, "現行資産を保ったままモバイルが堅く(F1)、集約の軽量化に道筋(F2)、★1根治の判断材料(F3)"),
    (1, "どの段階でも「動く状態」を保てる"),
    (0, "コスト/リスク"),
    (1, "Moproのビルド制約(iOSシミュ不可/arm64)、Novaは事例少なく自作要素、Noirはgas増と集約再設計"),
    (0, "未決定の前提"),
    (1, "実機の証明時間・メモリは未実測(→06)、信頼モデルの差(→07)"),
])

# S8 一言結論
s = add_slide(); title(s, "一言結論")
y = quote(s, TOP, [
    "現行(circom+rapidsnark+gnark)を土台に、フェーズ1でMoproによる",
    "モバイル標準化を即実施。集約はNova/Foldingを研究的に検証し、",
    "★1根治が必要になったらNoirをgasトレードオフ込みで評価する。",
])
bullets(s, y + 0.15, [
    (0, "手法（層2エンジン・層3集約）を先に固め、フレームワークはそれを実装する乗り物として選ぶ"),
], size0=15)

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_05_決定ドキュメント.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
