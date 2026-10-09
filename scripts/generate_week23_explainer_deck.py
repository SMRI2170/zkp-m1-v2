#!/usr/bin/env python3
"""Create a text-led explainer deck for the recursive UAV proof research.

Run from repository root:
    python3 scripts/generate_week23_explainer_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "siles-pptx" / "week23_解説型_複数端末ZKP集約.pptx"

W, H = 10.0, 7.5
FONT = "Hiragino Sans"
MONO = "Menlo"
INK = RGBColor(31, 38, 52)
MUTED = RGBColor(91, 100, 113)
RULE = RGBColor(210, 216, 225)
PALE = RGBColor(245, 247, 250)
INDIGO = RGBColor(63, 79, 153)
INDIGO_PALE = RGBColor(239, 242, 250)
ORANGE = RGBColor(177, 107, 35)
ORANGE_PALE = RGBColor(252, 245, 233)
GREEN = RGBColor(48, 112, 90)
GREEN_PALE = RGBColor(237, 247, 242)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "複数端末によるRemote ID位置不整合の検証"
prs.core_properties.subject = "テキスト主体の研究解説と進捗報告"
prs.core_properties.author = "zkp-m1-v2"
blank = prs.slide_layouts[6]
slide_no = 0


def set_font(run, family=FONT):
    run.font.name = family
    run._r.get_or_add_rPr().set("lang", "ja-JP")


def box(slide, x, y, w, h, text="", size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for index, line_text in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        set_font(r, font)
    return shape


def rect(slide, x, y, w, h, fill=PALE, line=RULE, radius=False, width=0.75):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def line(slide, x1, y1, x2, y2, color=RULE, width=1.0, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = "triangle"
    return shape


def bullet_text(slide, x, y, w, items, size=18, color=INK, gap=8):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(4.6))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        p.level = 0
        r = p.add_run()
        r.text = "• " + item
        r.font.size = Pt(size)
        r.font.color.rgb = color
        set_font(r)
    return shape


def base(title, section="", source=""):
    global slide_no
    slide_no += 1
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    if section:
        box(slide, 0.65, 0.30, 8.0, 0.2, section, 10, INDIGO, True)
    box(slide, 0.65, 0.62, 8.75, 0.48, title, 27, INK, True)
    line(slide, 0.65, 1.20, 9.35, 1.20, RULE, 0.7)
    if source:
        box(slide, 0.67, 7.12, 7.8, 0.18, source, 8, MUTED)
    box(slide, 8.88, 7.08, 0.45, 0.20, f"{slide_no:02d}", 9, MUTED,
        align=PP_ALIGN.RIGHT)
    return slide


def label(slide, x, y, w, text, color=INDIGO, fill=INDIGO_PALE):
    rect(slide, x, y, w, 0.30, fill, fill, radius=True, width=0.4)
    box(slide, x + 0.06, y + 0.06, w - 0.12, 0.16, text, 9, color, True,
        align=PP_ALIGN.CENTER)


def small_flow(slide, nodes, x=6.55, y=2.35, width=2.55):
    step_h = 0.55
    for i, node in enumerate(nodes):
        rect(slide, x, y + i * 0.77, width, step_h, INDIGO_PALE, INDIGO, radius=True)
        box(slide, x + 0.10, y + 0.16 + i * 0.77, width - 0.20, 0.20, node, 12,
            INK, True, align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            line(slide, x + width / 2, y + step_h + i * 0.77, x + width / 2,
                 y + 0.74 + i * 0.77, INDIGO, 1.1, True)


def simple_nodes(slide, nodes, x, y, w, h=0.55, gap=0.18, color=INDIGO):
    for i, node in enumerate(nodes):
        xi = x + i * (w + gap)
        rect(slide, xi, y, w, h, INDIGO_PALE, color, radius=True)
        box(slide, xi + 0.06, y + 0.14, w - 0.12, 0.18, node, 11, INK, True,
            align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            line(slide, xi + w, y + h / 2, xi + w + gap - 0.03, y + h / 2, color, 1.0, True)


# 1 Cover
slide = prs.slides.add_slide(blank)
slide_no += 1
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
box(slide, 0.78, 0.62, 8.3, 0.22, "WEEK 23  /  RESEARCH EXPLAINER", 10, INDIGO, True)
box(slide, 0.78, 1.40, 8.2, 0.55, "複数端末による", 32, INK, True)
box(slide, 0.78, 2.03, 8.2, 0.55, "Remote ID位置不整合の検証", 32, INK, True)
box(slide, 0.82, 2.92, 5.6, 0.72,
    "観測者の位置とRSSIを公開せずに、\n複数観測の判定を集約する。", 18, MUTED)
simple_nodes(slide, ["機体", "観測端末A/B/C", "集約proof"], 0.82, 4.45, 1.75, 0.70, 0.25)
box(slide, 0.82, 6.78, 7.2, 0.20, "2026年10月5日", 10, MUTED)
box(slide, 8.88, 7.08, 0.45, 0.20, "01", 9, MUTED, align=PP_ALIGN.RIGHT)

# 2 Current state
slide = base("研究の現在地", "01 / 現在地", "week23/00 実装進捗")
box(slide, 0.70, 1.55, 5.55, 0.70,
    "leaf、batch、Bridgeの部品検証は進んだ。\n実leafからEVMまでのproof一本化と実地評価は未完了。", 20, INK, True)
simple_nodes(slide, ["leaf / batch", "Bridge", "再帰proof", "EVM"], 0.75, 3.15, 1.75, 0.62, 0.25)
label(slide, 0.98, 4.00, 1.28, "test proof", GREEN, GREEN_PALE)
label(slide, 3.00, 4.00, 1.15, "部品テスト", ORANGE, ORANGE_PALE)
label(slide, 5.05, 4.00, 1.25, "E2E未完了", ORANGE, ORANGE_PALE)
label(slide, 7.04, 4.00, 1.25, "E2E未完了", ORANGE, ORANGE_PALE)
bullet_text(slide, 0.75, 4.70, 8.3, [
    "確認済みの範囲は、合成fixtureやtest用鍵を用いた回路・proofの検査を含む。",
    "実地BLE capture、Remote ID trust chain、最終proofのEVM検証は、まだ通していない。",
], 16)

# 3 Broadcast and actual position
slide = base("放送された位置と機体の位置", "02 / 問題")
box(slide, 0.70, 1.52, 5.35, 0.62,
    "Remote IDが伝える位置は送信者の主張である。\n受信しただけでは、機体がその位置にいるか確かめられない。", 20, INK, True)
bullet_text(slide, 0.75, 2.55, 5.55, [
    "放送位置と、受信時の物理的な状況は別に扱う。",
    "研究では、受信RSSIと放送位置の著しい矛盾を調べる。",
], 17)
rect(slide, 6.75, 2.10, 2.10, 0.72, INDIGO_PALE, INDIGO, radius=True)
box(slide, 6.90, 2.29, 1.80, 0.24, "放送された位置", 14, INK, True, align=PP_ALIGN.CENTER)
box(slide, 7.61, 3.15, 0.35, 0.32, "?", 24, ORANGE, True, align=PP_ALIGN.CENTER)
rect(slide, 6.75, 3.75, 2.10, 0.72, PALE, RULE, radius=True)
box(slide, 6.90, 3.94, 1.80, 0.24, "現地での受信状況", 14, INK, True, align=PP_ALIGN.CENTER)

# 4 RSSI
slide = base("RSSIは距離の目安", "02 / 問題")
box(slide, 0.70, 1.52, 5.60, 0.62,
    "RSSIは受信強度で、距離の大まかな手掛かりになる。\n遮蔽物や端末差があるため、正確な位置を単独で証明しない。", 20, INK, True)
bullet_text(slide, 0.75, 2.60, 5.45, [
    "近い・遠いの目安として、放送位置から計算した幾何距離と比べる。",
    "目的は絶対位置の確定ではなく、著しい不整合の検出である。",
], 17)
line(slide, 6.65, 2.35, 8.95, 2.35, INDIGO, 1.2, True)
box(slide, 6.63, 1.92, 2.40, 0.22, "RSSIが強い ←→ 弱い", 13, INDIGO, True, align=PP_ALIGN.CENTER)
rect(slide, 7.33, 2.85, 0.90, 0.58, ORANGE_PALE, ORANGE, radius=True)
box(slide, 7.43, 3.04, 0.70, 0.18, "遮蔽物", 11, ORANGE, True, align=PP_ALIGN.CENTER)
box(slide, 6.65, 3.83, 2.38, 0.52, "環境や端末による揺れを\n評価する必要がある", 13, MUTED, align=PP_ALIGN.CENTER)

# 5 Multiple observers
slide = base("複数端末で観測を照合する", "02 / 問題")
box(slide, 0.70, 1.52, 5.65, 0.62,
    "各端末が受信した観測を個別に判定する。\n集約時には、同じ対象・時間の観測であることを確かめる。", 20, INK, True)
bullet_text(slide, 0.75, 2.60, 5.55, [
    "一台だけの判断に依存せず、複数地点の観測結果を使う。",
    "端末ごとの生データを集約者へ渡さずに結果をまとめる。",
], 17)
rect(slide, 7.36, 1.82, 1.00, 0.52, INDIGO_PALE, INDIGO, radius=True)
box(slide, 7.48, 1.98, 0.76, 0.18, "放送", 12, INK, True, align=PP_ALIGN.CENTER)
for label_text, yy in [("端末A\n判定", 2.85), ("端末B\n判定", 3.78), ("端末C\n判定", 4.71)]:
    line(slide, 7.85, 2.34, 7.85, yy - 0.07, INDIGO, 0.9, True)
    rect(slide, 7.17, yy, 1.36, 0.58, PALE, RULE, radius=True)
    box(slide, 7.26, yy + 0.11, 1.18, 0.34, label_text, 11, INK, True, align=PP_ALIGN.CENTER)

# 6 Privacy
slide = base("観測値を公開すると協力者を追跡できる", "03 / プライバシー")
box(slide, 0.70, 1.52, 5.60, 0.62,
    "GPS、RSSI、時刻、端末IDは観測者の行動や参加履歴を漏らし得る。\n共有したいのは個々の値ではなく、条件を満たした件数である。", 20, INK, True)
bullet_text(slide, 0.75, 2.70, 5.60, [
    "GPSと時刻から、観測者の行動場所や移動を推測できる。",
    "RSSIや端末IDから、観測環境や繰り返し参加を推測できる。",
], 17)
small_flow(slide, ["GPS / RSSI / 時刻", "外部情報との照合", "参加者の推測"], 6.60, 2.30, 2.30)

# 7 ZK
slide = base("ZK proofは値を隠して条件を検証する", "03 / プライバシー")
box(slide, 0.70, 1.52, 5.55, 0.62,
    "ZK proofは、秘密の入力が定めた条件を満たすことを示す。\n検証者は入力の値を受け取らず、proofの正しさを確認する。", 20, INK, True)
bullet_text(slide, 0.75, 2.65, 5.55, [
    "端末の正確な位置や生RSSIを、そのまま共有しない。",
    "条件に合った観測かどうかを、検証できる形で示す。",
], 17)
simple_nodes(slide, ["秘密入力", "回路条件", "proof", "検証"], 6.35, 3.05, 0.66, 0.52, 0.12)
box(slide, 6.32, 4.00, 2.75, 0.55, "値を見せずに、\n条件が成り立つことを確認する", 13, MUTED, align=PP_ALIGN.CENTER)

# 8 Limits
slide = base("proofが保証すること・しないこと", "03 / プライバシー")
box(slide, 0.70, 1.52, 5.80, 0.60,
    "proofが保証するのは、回路に書かれた条件が入力間で成り立つこと。\n物理世界の真実やRemote IDの登録機体認証は、自動では保証されない。", 19, INK, True)
rect(slide, 0.75, 2.65, 3.95, 2.55, GREEN_PALE, GREEN, radius=True)
box(slide, 0.98, 2.91, 3.50, 0.22, "回路内で確認すること", 15, GREEN, True)
bullet_text(slide, 1.00, 3.35, 3.35, ["定めた入力関係", "proofと公開statementの整合"], 15, INK, 6)
rect(slide, 5.00, 2.65, 3.95, 2.55, ORANGE_PALE, ORANGE, radius=True)
box(slide, 5.23, 2.91, 3.50, 0.22, "現時点で未評価のこと", 15, ORANGE, True)
bullet_text(slide, 5.25, 3.35, 3.30, ["Remote ID trust chain", "物理世界での観測精度"], 15, INK, 6)
box(slide, 0.76, 5.62, 8.15, 0.34, "現在の認証状態は not_evaluated。位置判定は claimed RID に対するものとして扱う。", 14, ORANGE, True)

# 9 Target path
slide = base("目標とする証明経路", "04 / 証明経路")
box(slide, 0.70, 1.45, 8.25, 0.56,
    "端末がleaf proofを作り、Bridgeと再帰集約がproof・membership・重複防止を確認する。\n最終proofをEVMで検証し、結果を記録する。", 19, INK, True)
simple_nodes(slide, ["観測", "leaf proof", "Bridge", "再帰集約", "Final proof", "EVM"], 0.72, 2.75, 1.20, 0.57, 0.18)
for x, text, col in [(1.00, "端末", MUTED), (2.38, "test proof", GREEN), (3.76, "部品テスト", ORANGE), (5.14, "E2E未完了", ORANGE), (6.52, "E2E未完了", ORANGE), (7.90, "E2E未完了", ORANGE)]:
    box(slide, x, 3.53, 1.10, 0.18, text, 9, col, True, align=PP_ALIGN.CENTER)
bullet_text(slide, 0.75, 4.30, 8.30, [
    "leaf proofは1件の観測を扱う。Final proofは集約全体のルールが守られたことを扱う。",
    "この図は目標経路であり、全段が接続済みであることを示さない。",
], 16)

# 10 leaf proof
slide = base("1件の観測をleaf proofにする", "04 / 証明経路")
box(slide, 0.70, 1.52, 5.70, 1.10,
    "leaf proofは、1件の観測が回路条件を満たすことを示す。\nどの値が秘密で何を公開するかは、現行回路のstatementに沿って明示する。", 20, INK, True)
bullet_text(slide, 0.75, 3.00, 5.55, [
    "秘密入力には観測値や端末に関わる値を含む。",
    "公開側には、sessionを束縛する値や判定に必要なstatementを置く。",
], 16)
simple_nodes(slide, ["秘密の観測", "leaf回路", "proof / statement"], 6.42, 2.70, 0.85, 0.62, 0.18)
label(slide, 6.52, 3.67, 0.72, "非公開", INDIGO, INDIGO_PALE)
label(slide, 8.47, 3.67, 0.72, "公開", GREEN, GREEN_PALE)

# 11 membership
slide = base("登録端末かをmembershipで確かめる", "04 / 証明経路")
box(slide, 0.70, 1.52, 5.65, 0.62,
    "membershipは、端末鍵が登録一覧に含まれることを確認する。\n登録一覧全体を公開せず、rootとpathを使って確認できる。", 20, INK, True)
bullet_text(slide, 0.75, 2.65, 5.55, [
    "登録鍵の集合をMerkle rootという短い値で表す。",
    "端末は自分の鍵が含まれる経路だけをproofの入力として使う。",
], 16)
# tree
box(slide, 7.23, 2.00, 1.34, 0.22, "memberRoot", 12, INDIGO, True, align=PP_ALIGN.CENTER)
for xx in [6.80, 8.15]:
    line(slide, 7.90, 2.30, xx, 2.85, INDIGO, 0.9)
for xx in [6.55, 7.05, 7.90, 8.40]:
    line(slide, 6.80 if xx < 7.5 else 8.15, 2.93, xx, 3.50, INDIGO if xx == 6.55 else RULE, 0.9)
    rect(slide, xx - 0.25, 3.53, 0.50, 0.35, INDIGO_PALE if xx == 6.55 else PALE, INDIGO if xx == 6.55 else RULE, radius=True)
box(slide, 6.17, 4.18, 1.00, 0.36, "端末Aのpath", 11, GREEN, True, align=PP_ALIGN.CENTER)

# 12 nullifier
slide = base("nullifierで同じ端末の二重計上を防ぐ", "04 / 証明経路")
box(slide, 0.70, 1.52, 5.70, 0.62,
    "同じsessionで端末は一度だけ使える。使用前後のroot更新をproofで確認し、\n二度目は未使用条件を満たせないようにする。", 20, INK, True)
bullet_text(slide, 0.75, 2.75, 5.50, [
    "端末Aを数えると、Aのnullifier leafが未使用から使用済みに変わる。",
    "同じsessionでAを再利用すると、更新前の条件を満たせず拒否される。",
], 16)
simple_nodes(slide, ["未使用root", "端末Aを消費", "更新root", "再利用は拒否"], 6.34, 2.95, 0.65, 0.58, 0.13, ORANGE)

# 13 count updates
slide = base("判定件数を一件ずつ更新する", "04 / 証明経路", "week23/00 合成fixtureの検証例")
label(slide, 0.72, 1.48, 1.55, "合成テスト例", ORANGE, ORANGE_PALE)
box(slide, 0.70, 2.00, 5.65, 0.63,
    "合成テスト例では、整合1件と不整合2件を別々に数えた。\nこの例は回路・proofの確認用で、実地結果ではない。", 20, INK, True)
bullet_text(slide, 0.75, 3.14, 5.50, [
    "端末ごとの判定を加えながら、consistent と suspicious の一方だけを更新する。",
    "同じ鍵の重複、profile不一致、公開入力の改ざんは拒否対象にする。",
], 16)
simple_nodes(slide, ["A: 整合", "B: 不整合", "C: 不整合"], 6.45, 2.35, 0.78, 0.55, 0.12)
box(slide, 6.45, 3.48, 2.70, 0.65, "consistent = 1\nsuspicious = 2", 17, INDIGO, True, align=PP_ALIGN.CENTER)

# 14 Nova
slide = base("Novaは同じ更新を繰り返す", "05 / Nova")
box(slide, 0.70, 1.52, 5.65, 0.62,
    "Novaは、前の状態に同じ更新を順に適用する計算を累積proofにまとめる。\n複数端末を一件ずつ追加する流れに合う。", 20, INK, True)
bullet_text(slide, 0.75, 2.65, 5.55, [
    "状態 z に観測 w を加え、次の状態を作る。",
    "同じstepの履歴をfoldし、最終状態までの計算を確認する。",
], 17)
simple_nodes(slide, ["z₀", "F(w₀)", "z₁", "F(w₁)", "z₂"], 6.20, 3.00, 0.45, 0.54, 0.08)
box(slide, 6.30, 4.04, 2.60, 0.30, "zᵢ₊₁ = F(zᵢ, wᵢ)", 15, INDIGO, True, align=PP_ALIGN.CENTER, font=MONO)

# 15 Nova POC scope
slide = base("Nova PoCに入っているもの・いないもの", "05 / Nova", "week22/01 Nova調査")
box(slide, 0.70, 1.48, 8.20, 0.52,
    "現在のPoCはRSSI判定と件数更新を3 step行う。登録端末確認、nullifier、既存leaf proof検証はまだ含まない。", 19, INK, True)
rect(slide, 0.75, 2.45, 3.90, 2.80, GREEN_PALE, GREEN, radius=True)
box(slide, 1.00, 2.72, 3.40, 0.22, "現在の3-step PoC", 16, GREEN, True)
bullet_text(slide, 1.00, 3.20, 3.25, ["drone_id", "time_bucket", "consistent / suspicious", "step_count"], 16, INK, 4)
rect(slide, 5.00, 2.45, 3.90, 2.80, ORANGE_PALE, ORANGE, radius=True)
box(slide, 5.25, 2.72, 3.40, 0.22, "目標経路で追加するもの", 16, ORANGE, True)
bullet_text(slide, 5.25, 3.20, 3.25, ["memberRoot", "nullifierRoot", "policy / session", "leaf proofとの束縛"], 16, INK, 4)

# 16 bridge
slide = base("leaf proofと集約をつなぐにはBridgeが要る", "05 / Nova")
box(slide, 0.70, 1.52, 5.65, 0.62,
    "Novaへ観測値を渡すだけではleaf proofの正しさは保証されない。\nstep内でleaf proofを検証するか、leafの条件をNova回路へ移す必要がある。", 20, INK, True)
rect(slide, 0.76, 2.80, 4.90, 1.16, PALE, RULE, radius=True)
box(slide, 1.00, 3.05, 4.40, 0.53, "方法A: step内で既存leaf proofを検証\n既存回路を保てるが、stepごとの負担が増える。", 16, INK)
rect(slide, 0.76, 4.34, 4.90, 1.16, PALE, RULE, radius=True)
box(slide, 1.00, 4.59, 4.40, 0.53, "方法B: leafの条件をNova stepへ移植\n経路は単純になるが、回路の再実装が必要。", 16, INK)
simple_nodes(slide, ["既存leaf", "Bridge", "Nova step"], 6.30, 3.38, 0.78, 0.62, 0.15)

# 17 two approaches
slide = base("Novaと直接再帰PCDは別の経路", "05 / Nova", "week22/01, week23/00")
box(slide, 0.70, 1.48, 8.25, 0.54,
    "Nova PoCは逐次更新の実証である。Week 23で進めたBridgeBWやFinalBNは直接再帰PCD側の部品で、Novaへ接続済みではない。", 19, INK, True)
rect(slide, 0.76, 2.60, 3.85, 2.45, INDIGO_PALE, INDIGO, radius=True)
box(slide, 1.00, 2.88, 3.35, 0.22, "Nova IVCの検証レーン", 16, INDIGO, True)
bullet_text(slide, 1.00, 3.35, 3.25, ["RSSI判定と件数更新", "3-step PoC", "leaf等は未接続"], 15, INK, 5)
rect(slide, 5.05, 2.60, 3.85, 2.45, GREEN_PALE, GREEN, radius=True)
box(slide, 5.30, 2.88, 3.35, 0.22, "直接再帰PCDの検証レーン", 16, GREEN, True)
bullet_text(slide, 5.30, 3.35, 3.25, ["BridgeBW、NodeBLS、FinalBN", "部品・部分検証", "全段E2Eは未完了"], 15, INK, 5)
box(slide, 2.55, 5.72, 4.90, 0.28, "採用判断では、同じ安全条件と同じ計測条件で比較する。", 14, MUTED, True, align=PP_ALIGN.CENTER)

# 18 measurements
slide = base("PoCで測ったのは3-stepの動作", "06 / 実証結果", "week20/08, week20/09")
box(slide, 0.70, 1.48, 8.25, 0.52,
    "Pixel 8とM2 Macで3-step PoCのproof生成・最終verifyを確認した。環境差があるため、数値を速度順位として扱わない。", 19, INK, True)
rect(slide, 0.76, 2.45, 3.90, 2.20, INDIGO_PALE, INDIGO, radius=True)
box(slide, 1.05, 2.76, 3.30, 0.22, "Pixel 8", 18, INDIGO, True)
box(slide, 1.05, 3.20, 3.30, 0.62, "各step  485 / 655 / 1,198 ms\n最終verify  172 ms", 16, INK)
rect(slide, 5.00, 2.45, 3.90, 2.20, GREEN_PALE, GREEN, radius=True)
box(slide, 5.30, 2.76, 3.30, 0.22, "M2 Mac", 18, GREEN, True)
box(slide, 5.30, 3.20, 3.30, 0.62, "各step  135 / 151 / 281 ms\n最終verify  28 ms", 16, INK)
box(slide, 0.78, 5.32, 8.15, 0.50, "確認した結果: consistent 1 → 2 → 2、suspicious 0 → 0 → 1。count改ざんproofを拒否。\n未接続: 実leaf proof、membership、nullifier、EVM Decider。", 15, ORANGE, True)

# 19 next gates
slide = base("次の完了条件", "07 / 次の判断", "week22/07, week23/00")
box(slide, 0.70, 1.48, 8.30, 0.52,
    "leaf proof・membership・nullifierを再帰proofへ接続する。否定系を通し、最終proofをEVMで検証する。\n並行して実captureと独立真値で誤検知・見逃しを測る。", 19, INK, True)
simple_nodes(slide, ["leaf + membership", "再帰proof", "EVM検証", "実地評価"], 0.78, 3.03, 1.78, 0.64, 0.26)
for x, text in [(0.90, "固定VK・nullifier"), (2.94, "否定系を含む"), (5.00, "proof / gas"), (7.10, "校正 / holdout")]:
    box(slide, x, 3.94, 1.55, 0.30, text, 11, MUTED, align=PP_ALIGN.CENTER)
bullet_text(slide, 0.77, 4.85, 8.25, [
    "Novaの採用は、この経路を実proofで通し、直接再帰PCDと同じ条件で比較してから判断する。",
    "最終proofと実地評価が終わるまでは、部品検証の範囲で報告する。",
], 16)

# Appendix 20 terms
slide = base("補足: 用語", "APPENDIX")
terms = [
    ("leaf proof", "1件の観測が回路条件を満たすことを示すproof。"),
    ("membership", "端末鍵が登録一覧に含まれることの確認。"),
    ("nullifier", "同じ端末を同じsessionで一度だけ数えるための状態。"),
    ("IVC / Nova", "同じ状態更新を繰り返し、その履歴を累積proofにまとめる方式。"),
    ("Decider", "累積proofをEVM向けの最終検証へ接続するための仕組み。"),
    ("test-engine solve", "回路制約の充足確認。Groth16 proof生成やEVM検証を意味しない。"),
]
for i, (term, expl) in enumerate(terms):
    yy = 1.52 + i * 0.78
    box(slide, 0.75, yy, 2.05, 0.25, term, 15, INDIGO, True)
    box(slide, 2.95, yy, 5.95, 0.42, expl, 15, INK)
    line(slide, 0.75, yy + 0.53, 8.92, yy + 0.53, RULE, 0.5)

# Appendix 21 evidence
slide = base("補足: 現在の証拠の種類", "APPENDIX", "week20/08, week22/01, week23/00")
evidence = [
    ("設計", "仕様・構想として記載。動作確認はしていない。", INDIGO, INDIGO_PALE),
    ("部品テスト", "compile、unit test、test-engine solveで関係の一部を確認。", ORANGE, ORANGE_PALE),
    ("test proof", "synthetic fixtureまたはtest-only鍵のproofで確認。", GREEN, GREEN_PALE),
    ("E2E未完了", "leafから最終EVM recordまでのproofが連続していない。", ORANGE, ORANGE_PALE),
    ("実地評価未完了", "実captureと独立真値で分類性能を評価していない。", ORANGE, ORANGE_PALE),
]
for i, (name, expl, color, fill) in enumerate(evidence):
    yy = 1.56 + i * 0.94
    rect(slide, 0.75, yy, 2.15, 0.52, fill, color, radius=True)
    box(slide, 0.92, yy + 0.16, 1.80, 0.18, name, 14, color, True, align=PP_ALIGN.CENTER)
    box(slide, 3.15, yy + 0.11, 5.80, 0.34, expl, 15, INK)

# Appendix 22 comparison
slide = base("補足: 方式を比べる軸", "APPENDIX", "week22/06 再帰ZKP方式調査")
box(slide, 0.70, 1.45, 8.20, 0.50, "方式の名前だけで選ばず、同じ安全条件と同じ測定条件で比較する。", 19, INK, True)
rows = [
    ("Nova / CycleFold", "逐次state更新", "leaf Bridge、Decider、実機性能"),
    ("直接再帰Groth16", "既存leafを回路内で検証しやすい", "曲線サイクル、固定VK、setup、回路規模"),
    ("Halo2 / zkVMなど", "別の回路表現・実装経路", "回路移植、proof規模、EVM費用"),
]
box(slide, 0.78, 2.35, 2.05, 0.22, "方式", 14, INDIGO, True)
box(slide, 3.05, 2.35, 2.40, 0.22, "強み", 14, INDIGO, True)
box(slide, 5.75, 2.35, 3.05, 0.22, "主な検証課題", 14, INDIGO, True)
for i, row in enumerate(rows):
    yy = 2.82 + i * 1.00
    rect(slide, 0.75, yy, 8.15, 0.72, PALE, RULE)
    box(slide, 0.88, yy + 0.16, 1.80, 0.35, row[0], 14, INK, True)
    box(slide, 3.05, yy + 0.16, 2.35, 0.35, row[1], 14, INK)
    box(slide, 5.75, yy + 0.16, 2.88, 0.35, row[2], 14, INK)

# Appendix 23 questions
slide = base("補足: よくある質問", "APPENDIX")
qa = [
    ("RSSIで機体の本当の位置が分かるか", "絶対位置を確定しない。放送座標と受信強度の著しい矛盾を評価する。"),
    ("Remote IDは認証済みか", "いいえ。現行の認証状態は not_evaluated。claimed RIDに対する評価である。"),
    ("Nova proofをEVMで検証したか", "まだ。Decider、Solidity verifier、V1 ABIとの接続を実proofで確認する必要がある。"),
    ("Week 23のBridgeはNovaへ接続済みか", "いいえ。BridgeBWなどは直接再帰PCD側の部品であり、Nova stepへの統合は別途必要。"),
]
for i, (q, a) in enumerate(qa):
    yy = 1.48 + i * 1.22
    box(slide, 0.75, yy, 2.75, 0.35, "Q. " + q, 15, INDIGO, True)
    box(slide, 0.95, yy + 0.45, 7.85, 0.42, "A. " + a, 15, INK)
    line(slide, 0.75, yy + 1.02, 8.90, yy + 1.02, RULE, 0.5)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {OUTPUT} ({len(prs.slides)} slides)")
