from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT_PATH = Path("project_summary_4x3_enhanced.pptx")

# ── サイズ ──────────────────────────────────────────────────────────────────
W = Inches(10)
H = Inches(7.5)
M = Inches(0.45)          # 左右マージン
IW = W - M * 2            # 使用幅 = 9.1"

# ── カラーパレット（白ベース）────────────────────────────────────────────────
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_BG        = RGBColor(0xFF, 0xFF, 0xFF)
C_NAVY      = RGBColor(0x1A, 0x23, 0x7E)   # 濃紺
C_BLUE      = RGBColor(0x19, 0x76, 0xD2)   # 青
C_TEAL      = RGBColor(0x00, 0x89, 0x7B)   # 青緑
C_TEXT      = RGBColor(0x21, 0x21, 0x21)
C_SUBTEXT   = RGBColor(0x61, 0x61, 0x61)
C_RULE      = RGBColor(0xBD, 0xC8, 0xE0)
C_CODE_BG   = RGBColor(0xF3, 0xF6, 0xFB)
C_TBLH_BG   = RGBColor(0x1A, 0x23, 0x7E)
C_TBLH_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
C_TBL_ROW1  = RGBColor(0xFF, 0xFF, 0xFF)
C_TBL_ROW2  = RGBColor(0xEE, 0xF2, 0xFF)
C_TBL_TEXT  = RGBColor(0x21, 0x21, 0x21)

FONT = "Hiragino Sans"

# ── 基本ヘルパー ────────────────────────────────────────────────────────────
def rect(slide, left, top, width, height, color, border=False):
    s = slide.shapes.add_shape(1, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if border:
        s.line.width = Pt(0.5)
        s.line.color.rgb = C_RULE
    else:
        s.line.fill.background()
    return s

def set_white_bg(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG

def txb(slide, text, left, top, width, height,
        size=14, bold=False, color=C_TEXT,
        align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    return box

def bullets(slide, lines, left, top, width, height, size=13, color=C_TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
        p.space_after = Pt(4)

def table(slide, headers, rows, left, top, width, col_widths=None, font_size=11):
    nr = len(rows) + 1
    nc = len(headers)
    row_h = Inches(0.38)
    tbl = slide.shapes.add_table(nr, nc, left, top, width, row_h * nr).table
    tbl.first_row = True

    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = w
    else:
        for col in tbl.columns:
            col.width = width // nc

    def cell_style(cell, text, header=False, even=False):
        cell.fill.solid()
        if header:
            cell.fill.fore_color.rgb = C_TBLH_BG
        elif even:
            cell.fill.fore_color.rgb = C_TBL_ROW2
        else:
            cell.fill.fore_color.rgb = C_TBL_ROW1

        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = text
        r.font.size = Pt(font_size)
        r.font.bold = header
        r.font.color.rgb = C_TBLH_TEXT if header else C_TBL_TEXT
        r.font.name = FONT

    for ci, h in enumerate(headers):
        cell_style(tbl.cell(0, ci), h, header=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell_style(tbl.cell(ri + 1, ci), val, even=(ri % 2 == 1))

def header(slide, title, sub=None):
    set_white_bg(slide)
    rect(slide, 0, 0, W, Inches(0.07), C_NAVY)
    rect(slide, 0, Inches(0.07), Inches(0.07), H - Inches(0.07), C_BLUE)
    txb(slide, title, M, Inches(0.12), IW, Inches(0.55), size=22, bold=True, color=C_NAVY)
    if sub:
        txb(slide, sub, M, Inches(0.65), IW, Inches(0.32), size=11, color=C_SUBTEXT)
    rect(slide, M, Inches(0.98), IW, Pt(1.2), C_RULE)

def page_num(slide, n):
    txb(slide, str(n), W - Inches(0.55), H - Inches(0.32), Inches(0.4), Inches(0.25), size=9, color=C_SUBTEXT, align=PP_ALIGN.RIGHT)

def add_arrow(slide, left, top, width, height, text="→"):
    txb(slide, text, left, top, width, height, size=20, bold=True, color=C_SUBTEXT, align=PP_ALIGN.CENTER)

# ── スライド構築 ────────────────────────────────────────────────────────────

def build_presentation():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # --- Slide 1: Title ---
    s = prs.slides.add_slide(blank)
    set_white_bg(s)
    rect(s, 0, 0, W, Inches(1.4), C_NAVY)
    txb(s, "Project Summary", M, Inches(0.2), IW, Inches(0.35), size=10, color=RGBColor(0xB0, 0xBE, 0xFF))
    txb(s, "UAV安全距離証明アプリ", M, Inches(1.7), IW, Inches(0.8), size=30, bold=True, color=C_NAVY)
    txb(s, "実装・統合・検証の要約", M, Inches(2.45), IW, Inches(0.6), size=24, bold=True, color=C_BLUE)
    rect(s, M, Inches(3.15), Inches(1.5), Pt(2.5), C_TEAL)
    txb(s, "Circom / Flutter / OpenDroneID / YOLO によるゼロ知識証明の実機統合検証", M, Inches(3.4), IW, Inches(0.4), size=13, color=C_SUBTEXT)
    rect(s, 0, H - Inches(0.55), W, Inches(0.55), C_NAVY)
    txb(s, "ZKP-M1 Project Report | 2026", M, H - Inches(0.48), IW, Inches(0.38), size=9, color=RGBColor(0xB0, 0xBE, 0xFF))

    # --- Slide 2: 目的 ---
    s = prs.slides.add_slide(blank)
    header(s, "プロジェクトの目的と要件", "30m以内・5秒以内のドローン接近を安全に証明するシステム")
    page_num(s, 2)
    
    col_w = Inches(2.8)
    for i, (title_t, color_t, desc_t) in enumerate([
        ("課題 (Problem)", C_NAVY, "・ドローンの接近を証明したい\n・しかし、生の位置情報(GPS)を公開するとプライバシーが損なわれる"),
        ("解決策 (Solution)", C_BLUE, "・ゼロ知識証明 (ZKP) を活用\n・「30m以内にいた」という事実(フラグ)のみを暗号学的に証明\n・生座標は非公開(Private Input)"),
        ("要件 (Requirements)", C_TEAL, "・スマホ内で完結する証明・検証フロー\n・UAVと利用者のログの自動統合\n・実機(Flutter/BLE/GPS)での動作"),
    ]):
        bx = M + Inches(3.0 * i)
        rect(s, bx, Inches(2.0), col_w, Inches(0.4), color_t)
        txb(s, title_t, bx + Inches(0.1), Inches(2.05), col_w - Inches(0.2), Inches(0.3), size=12, bold=True, color=C_WHITE)
        rect(s, bx, Inches(2.4), col_w, Inches(2.5), C_CODE_BG, border=True)
        txb(s, desc_t, bx + Inches(0.1), Inches(2.5), col_w - Inches(0.2), Inches(2.3), size=12, color=C_TEXT)

    rect(s, M, Inches(5.5), IW, Inches(0.6), C_CODE_BG, border=True)
    rect(s, M, Inches(5.5), Inches(0.055), Inches(0.6), C_NAVY)
    txb(s, "目指す姿: 専用の検証装置不要で、市販スマホとドローンの標準規格(OpenDroneID)のみで成り立つ安全な証明網", 
        M + Inches(0.15), Inches(5.6), IW - Inches(0.3), Inches(0.4), size=12, bold=True, color=C_NAVY)

    # --- Slide 3: ZKP基盤実装 ---
    s = prs.slides.add_slide(blank)
    header(s, "ZKP（ゼロ知識証明）基盤の実装", "Circomによる回路設計からProof生成までの流れ")
    page_num(s, 3)

    bullets(s, [
        "■ Circom を用いた平面近似3次元距離判定回路を実装",
        "■ 緯度・経度・高度・時刻を Private Input とし、Groth16 で証明を生成",
        "■ Flutter アプリ組み込み用に .wcd / .zkey / verification key をエクスポート",
    ], M, Inches(1.3), IW, Inches(1.0), size=12)

    boxes = [
        ("1. 回路定義", "distance30m.circom\nPoseidonハッシュ\n座標の距離計算", C_NAVY),
        ("2. コンパイル", ".r1cs / .wasm 生成\nwitness長不一致解消", C_BLUE),
        ("3. セットアップ", "Groth16 Trusted Setup\n.zkey と vkey の生成", C_TEAL),
        ("4. スマホ実装", ".wcd で Witness 計算\nrapidsnark で証明生成", RGBColor(0x6A, 0x1B, 0x9A)),
    ]
    for i, (title_t, desc_t, color_t) in enumerate(boxes):
        bx = M + Inches(2.3 * i)
        rect(s, bx, Inches(3.5), Inches(2.0), Inches(0.4), color_t)
        txb(s, title_t, bx + Inches(0.1), Inches(3.55), Inches(1.8), Inches(0.3), size=11, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        rect(s, bx, Inches(3.9), Inches(2.0), Inches(1.2), C_CODE_BG, border=True)
        txb(s, desc_t, bx + Inches(0.1), Inches(4.0), Inches(1.8), Inches(1.0), size=10, color=C_TEXT, align=PP_ALIGN.CENTER)
        if i < len(boxes) - 1:
            add_arrow(s, bx + Inches(2.0), Inches(4.1), Inches(0.3), Inches(0.4))

    # --- Slide 4: アプリ統合 ---
    s = prs.slides.add_slide(blank)
    header(s, "Flutter アプリ統合", "実機でのオンデバイス証明と検証")
    page_num(s, 4)

    txb(s, "Flutter + rapidsnark の連携", M, Inches(1.3), IW, Inches(0.3), size=13, bold=True, color=C_NAVY)
    bullets(s, [
        "■ スマホ上で Witness (.wcd) を計算し、rapidsnark で Groth16 証明を生成",
        "■ 生成された証明を同じ端末内の検証モジュールで即座に検証可能",
        "■ Pixel 8 実機にて動作確認済み (証明生成: 約2秒)",
    ], M, Inches(1.7), IW, Inches(1.0), size=12)

    # 模擬UI
    rect(s, Inches(1.5), Inches(3.0), Inches(3.0), Inches(3.5), RGBColor(0xEE, 0xEE, 0xEE), border=True)
    rect(s, Inches(1.5), Inches(3.0), Inches(3.0), Inches(0.3), C_NAVY)
    txb(s, "ZKP App", Inches(1.5), Inches(3.0), Inches(3.0), Inches(0.3), size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    
    rect(s, Inches(1.7), Inches(3.5), Inches(2.6), Inches(0.4), C_WHITE, border=True)
    txb(s, "Input: [Private Data]", Inches(1.7), Inches(3.6), Inches(2.6), Inches(0.3), size=10)
    
    rect(s, Inches(1.7), Inches(4.0), Inches(2.6), Inches(0.4), C_BLUE)
    txb(s, "Generate Proof", Inches(1.7), Inches(4.05), Inches(2.6), Inches(0.3), size=11, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

    rect(s, Inches(1.7), Inches(4.5), Inches(2.6), Inches(0.7), RGBColor(220, 240, 220), border=True)
    txb(s, "✅ Verification Success\nWithin 30m: True", Inches(1.7), Inches(4.6), Inches(2.6), Inches(0.5), size=10, color=RGBColor(0, 100, 0))

    txb(s, "オンデバイス処理の流れ", Inches(5.5), Inches(3.0), Inches(3.5), Inches(0.3), size=12, bold=True, color=C_TEAL)
    bullets(s, [
        "① ユーザーが証明生成ボタンを押下",
        "② C/C++ FFI を経由して snarkjs 互換の証明器が走る",
        "③ Output と Proof を取得",
        "④ スマートコントラクト(想定)や他端末に送れる形式でエクスポート",
    ], Inches(5.5), Inches(3.4), Inches(3.5), Inches(2.0), size=11)

    # --- Slide 5: データ取得拡張 ---
    s = prs.slides.add_slide(blank)
    header(s, "データ取得の拡張 (センサー・通信)", "各種センサーを用いたリアルタイムデータの収集")
    page_num(s, 5)

    txb(s, "UAVとスマホ間のデータ連携網", M, Inches(1.3), IW, Inches(0.3), size=13, bold=True, color=C_NAVY)

    # Diagram
    # UAV
    rect(s, Inches(1.0), Inches(2.0), Inches(2.2), Inches(1.2), C_BLUE)
    txb(s, "UAV (ドローン)", Inches(1.0), Inches(2.2), Inches(2.2), Inches(0.3), size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    txb(s, "・現在位置 (GPS)\n・OpenDroneID", Inches(1.1), Inches(2.5), Inches(2.0), Inches(0.6), size=10, color=C_WHITE)

    add_arrow(s, Inches(3.3), Inches(2.4), Inches(1.0), Inches(0.4), "BLE 送信\n(毎秒)")

    # Smartphone
    rect(s, Inches(4.5), Inches(2.0), Inches(4.0), Inches(3.0), C_CODE_BG, border=True)
    rect(s, Inches(4.5), Inches(2.0), Inches(4.0), Inches(0.4), C_TEAL)
    txb(s, "利用者端末 (Flutter)", Inches(4.5), Inches(2.05), Inches(4.0), Inches(0.3), size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

    bullets(s, [
        "■ BLEスキャン: UAVからのテレメトリを自動受信",
        "■ GPS: 利用者の現在位置を1秒ごとに取得",
        "■ 気圧センサー: スマホとドローンの相対高度(QFE)を計算",
    ], Inches(4.7), Inches(2.6), Inches(3.6), Inches(1.5), size=11)

    rect(s, Inches(4.7), Inches(4.2), Inches(3.6), Inches(0.5), RGBColor(255, 240, 200), border=True)
    txb(s, "→ 回路入力パラメータ (JSON) の自動生成", Inches(4.7), Inches(4.35), Inches(3.6), Inches(0.3), size=11, bold=True, align=PP_ALIGN.CENTER)

    # --- Slide 6: 証拠・検知機能 ---
    s = prs.slides.add_slide(blank)
    header(s, "証拠・検知機能の統合 (AI / ML Kit)", "画像認識技術を用いた視覚的証拠の付与")
    page_num(s, 6)

    col_w = Inches(4.3)
    txb(s, "動画証拠の記録", M, Inches(1.3), col_w, Inches(0.3), size=12, bold=True, color=C_NAVY)
    bullets(s, [
        "・BLE 受信後、自動的にカメラによる録画を開始",
        "・録画中の各フレーム(撮影時刻)と、ZKPに入力するGPS/BLEログを厳密に対応付け",
        "・「いつ・どこで・何が映っていたか」をZKPとセットにする",
    ], M, Inches(1.7), col_w, Inches(1.5), size=11)

    txb(s, "AI物体検知 (ML Kit / YOLO)", Inches(5.1), Inches(1.3), col_w, Inches(0.3), size=12, bold=True, color=C_TEAL)
    bullets(s, [
        "・Google ML Kit を用いた軽量な動画解析",
        "・YOLO v11n モデルを使用したライブ検知画面",
        "・画面上にドローンのバウンディングボックスを表示",
        "・将来的なUAV専用カスタムモデルへの差し替えを前提としたアーキテクチャ",
    ], Inches(5.1), Inches(1.7), col_w, Inches(1.5), size=11)

    rect(s, M, Inches(3.8), IW, Inches(2.5), C_CODE_BG, border=True)
    rect(s, M, Inches(3.8), Inches(0.055), Inches(2.5), C_NAVY)
    txb(s, "システムフロー:", M + Inches(0.15), Inches(3.9), IW, Inches(0.3), size=11, bold=True)
    txb(s, "BLE受信 → 動画撮影開始 → AIによるドローン検知 → ログと映像の結合 → ZKP証明の生成", 
        M + Inches(0.3), Inches(4.3), IW - Inches(0.5), Inches(0.4), size=14, bold=True, color=C_BLUE)
    
    txb(s, "これにより、「位置情報の数学的証明」だけでなく「視覚的な事実確認」が可能になる。",
        M + Inches(0.3), Inches(5.2), IW - Inches(0.5), Inches(0.4), size=12, color=C_SUBTEXT)

    # --- Slide 7: 成果と残課題 ---
    s = prs.slides.add_slide(blank)
    header(s, "成果と残課題", "現状の達成度と今後のステップ")
    page_num(s, 7)

    table(s,
        ["ステータス", "項目", "詳細"],
        [
            ["✅ 達成", "オンデバイス証明の実現", "Flutter + rapidsnark による端末内での ZKP 証明と検証が成功"],
            ["✅ 達成", "センサーデータ統合", "BLE(OpenDroneID), GPS, 気圧センサーからの自動データ収集フロー完成"],
            ["✅ 達成", "AI検知統合", "ML Kit / YOLO によるライブ動画からのUAV検知UIの実装"],
            ["⏳ 残課題", "実機E2Eの最終確認", "実際の屋外環境での、ドローンからスマホまでのEnd-to-End疎通テスト"],
            ["⏳ 残課題", "Raspberry Pi 実運用", "ドローン側搭載用 Raspberry Pi での BLE 送信の安定性確認"],
            ["⏳ 残課題", "検知モデルの高精度化", "VisDrone データセット等でファインチューニングしたUAV専用モデルの搭載"],
        ],
        M, Inches(1.3), IW,
        col_widths=[Inches(1.2), Inches(2.5), Inches(5.4)], font_size=11)

    rect(s, M, Inches(5.0), IW, Inches(0.8), RGBColor(230, 240, 255), border=True)
    rect(s, M, Inches(5.0), Inches(0.055), Inches(0.8), C_BLUE)
    txb(s, "まとめ", M + Inches(0.15), Inches(5.1), IW, Inches(0.3), size=12, bold=True, color=C_NAVY)
    txb(s, "コアとなる「証明生成・検証のオンデバイス実行」と「通信・検知機能の統合」は完了。今後は実運用を見据えたテストとAIモデルの精度向上にシフトする。", 
        M + Inches(0.3), Inches(5.4), IW - Inches(0.5), Inches(0.4), size=11, color=C_TEXT)

    # --- 保存 ---
    prs.save(OUTPUT_PATH)
    print(f"✅ Saved: {OUTPUT_PATH}")

if __name__ == "__main__":
    build_presentation()
