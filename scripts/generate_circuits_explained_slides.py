#!/usr/bin/env python3
"""現在の Circom 回路群の説明スライド。

各回路が「何を入力に、何を計算し、何を公開するか」を積み上げ構造で説明する。
共通コア Distance30mTimeCore を土台に、検知・音響・結合・なりすまし検出へ拡張。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.1)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.3, 8.2, 0.5)
run(tf.paragraphs[0], "技術解説 / 週17", 15, GRAY)
tf = tbox(s, 0.88, 2.85, 8.4, 1.2)
run(tf.paragraphs[0], "ZK 回路の説明", 32, INK, bold=True)
hrule(s, 4.2, l=0.92, w=5.0)
tf = tbox(s, 0.92, 4.4, 8.2, 1.4)
run(tf.paragraphs[0], "共通コアを土台に、検知・音響・結合・なりすまし検出へ積み上げる", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "Circom / Groth16 / circomlib(Poseidon, EdDSA)", 12, GRAY, font=MONO)

# 2 回路ファミリー一覧
s = add_slide()
title(s, "回路ファミリー — 一覧")
tf = tbox(s, 0.7, 1.3, 8.7, 0.5)
run(tf.paragraphs[0], "すべて「座標は隠して、条件の成否だけ公開する」思想で共通コアを拡張", 12.5, GRAY)
table(s, 0.7, 1.95, 8.6, [
    ["回路", "役割", "主な公開出力"],
    ["lib_uav_safety (Core)", "距離30m + 時刻5s の判定土台", "within30"],
    ["distance_30m_time(_series)", "単一 / 時系列N点の接近判定", "within30 / any_within30"],
    ["detection_flag", "映像検知フラグを AND 結合", "within30_and_detected"],
    ["detection_commitment", "+ Poseidon で座標秘匿・フレーム束縛", "各commitment, frame_binding"],
    ["feature_classifier6", "音響6次元・線形分類を ZK 内積で判定", "model_ok, within30_and_model"],
    ["uav_witness_attestation", "結合用: drone_id + 端末署名 + 時間窓", "drone_id, device_pubkey, time_bucket"],
    ["remoteid_consistency", "放送位置 vs 実測レンジの矛盾検出", "consistency_ok"],
], col_w=[2.6, 3.7, 2.3], fs=9.8, row_h=0.5)

# 3 共通コア
s = add_slide()
title(s, "共通コア — Distance30mTimeCore")
tf = tbox(s, 0.7, 1.3, 8.7, 1.0)
bullet(tf, "全回路の土台。ドローンと自端末の位置・時刻から「30m以内かつ5秒以内」を判定", first=True)
sub(tf, "座標(lat/lon/alt)と時刻は private。出力は 0/1 のフラグのみ")
codebox(s, 0.7, 2.4, 8.6, 2.9, [
    "# 各軸の差を絶対値化し、cm スケールへ",
    "dLat = |drone_lat - target_lat| * LAT_CM_NUM   (経度も同様)",
    "# 3次元の二乗距離（平方根を使わない）",
    "dist_sq = lat^2 + lon^2 + alt^2",
    "distance_ok = ( dist_sq <= (3000cm)^2 )          # LessEqThan",
    "time_ok     = ( |drone_t - target_t| <= 5 )",
    "within30    = distance_ok * time_ok              # AND",
])
tf = tbox(s, 0.7, 5.5, 8.7, 1.0)
bullet(tf, "工夫: 平方根を避け二乗のまま比較 / 減算は絶対値化して非負に保つ", size=12, first=True)
sub(tf, "Circom の比較(LessEqThan)は非負前提のため")

# 4 検知系
s = add_slide()
title(s, "拡張① 映像検知 + コミットメント")
tf = tbox(s, 0.7, 1.3, 8.7, 1.6)
bullet(tf, "detection_flag: コアに映像検知の条件を AND 結合", first=True)
sub(tf, "score >= 閾値(bps) かつ class == UAV かつ uav_detected==1 → detection_ok")
bullet(tf, "detection_commitment: さらに Poseidon で座標を秘匿しフレームを束縛")
codebox(s, 0.7, 3.05, 8.6, 2.3, [
    "detection_ok         = uav_detected * score_ok * class_ok",
    "within30_and_detected = within30 * detection_ok",
    "",
    "drone_commitment  = Poseidon(lat, lon, alt, time)   # 座標は出さずハッシュのみ",
    "frame_binding     = Poseidon(frame_hi, frame_lo)    # 動画フレームを proof に束縛",
])
tf = tbox(s, 0.7, 5.55, 8.7, 0.9)
bullet(tf, "狙い: 「どこにいたか」を隠したまま「30m以内で検知した」ことだけ示す",
       size=12.5, first=True, bold=True)

# 5 音響系
s = add_slide()
title(s, "拡張② 音響モデル (feature_classifier6)")
tf = tbox(s, 0.7, 1.3, 8.7, 1.3)
bullet(tf, "プロペラ音の6次元特徴を、線形分類器で「ドローンらしさ」を ZK 内で判定", first=True)
sub(tf, "w・x + b >= 0 を証明。座標同様、特徴ベクトルも秘匿できる")
codebox(s, 0.7, 2.7, 8.6, 2.4, [
    "# 負の重みを扱えるよう 正/負 に分離し 非負オフセットで吸収",
    "pos = Σ WP[i]*feature[i] + ADD_POS",
    "neg = Σ WN[i]*feature[i] + ADD_NEG",
    "model_ok           = ( neg <= pos )     # ⇔ w・x + b >= 0",
    "within30_and_model = within30 * model_ok",
])
tf = tbox(s, 0.7, 5.35, 8.7, 1.0)
bullet(tf, "工夫: Circom は負数・小数が苦手 → 整数量子化(×1000)＋重みの符号分離",
       size=12, first=True)

# 6 結合用回路
s = add_slide()
title(s, "新規① uav_witness_attestation")
tf = tbox(s, 0.7, 1.3, 8.7, 1.2)
bullet(tf, "detection_commitment を、マルチ端末・オンチェーン結合用に拡張", first=True)
sub(tf, "public 出力 14 個 / 約1.2万制約（大半は EdDSA）")
codebox(s, 0.7, 2.5, 8.6, 2.5, [
    "drone_id_commitment = Poseidon(remote_id)     # 同じドローンの結合キー",
    "device_pubkey (Ax,Ay) + EdDSA 署名を回路内検証  # 端末ID / Sybil対策 / replay防止",
    "time_bucket = target_time \\ 60                # 時刻は隠し時間窓で集計",
    "",
    ("署名対象 M = Poseidon(drone_id, target_commit, frame_binding,", True),
    ("             within30_and_detected)", True),
])
tf = tbox(s, 0.7, 5.2, 8.7, 1.1)
bullet(tf, "この3つの追加で、複数端末の証明をチェーン上で「同一ドローン・別端末」に束ねられる",
       size=12, first=True, bold=True)

# 7 なりすまし検出回路
s = add_slide()
title(s, "新規② remoteid_consistency")
tf = tbox(s, 0.7, 1.3, 8.7, 1.2)
bullet(tf, "RemoteID の放送位置と、RSSI 実測レンジの整合を検証（約540制約と軽量）", first=True)
sub(tf, "ドローン改造不要。放送だけ見れば区別できない詐称を、実測で突く")
codebox(s, 0.7, 2.5, 8.6, 2.3, [
    "claimed_dist_sq = |放送位置 - 自端末位置|^2",
    "lower = max(measured - TOL, 0),  upper = measured + TOL",
    "consistency_ok = ( lower^2 <= claimed_dist_sq <= upper^2 )",
    ("→ 放送3.4m だが実測50m のような矛盾で consistency_ok=0", True),
])
tf = tbox(s, 0.7, 5.0, 8.7, 1.3)
bullet(tf, "コアと同じく二乗のまま比較。measured_range は RSSI モデルが供給", size=12, first=True)
sub(tf, "検証済: 整合→1 / 詐称→0")

# 8 積み上げ図
s = add_slide()
title(s, "まとめ — 積み上げ構造")
codebox(s, 0.7, 1.4, 8.6, 3.9, [
    "Distance30mTimeCore  (距離30m + 時刻5s)  ← 全ての土台",
    "   ├─ distance_30m_time / _series      単一 / 時系列N点",
    "   ├─ detection_flag                   + 映像検知フラグ",
    "   ├─ detection_commitment             + Poseidon 秘匿 / フレーム束縛",
    "   ├─ feature_classifier6              + 音響6次元 線形分類",
    "   └─ uav_witness_attestation          + drone_id / EdDSA / time_bucket",
    "                                         (マルチ端末結合の中核)",
    "",
    "remoteid_consistency  (放送 vs 実測レンジの矛盾検出)  ← なりすまし検出の核",
])
tf = tbox(s, 0.7, 5.5, 8.7, 1.1)
bullet(tf, "一貫した設計思想: 座標や特徴は隠し、条件の成否(0/1)だけを公開", first=True, bold=True)
sub(tf, "二乗のまま比較・整数量子化・符号分離で Circom の制約に収める")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week17")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "circuits_explained.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
