#!/usr/bin/env python3
"""week19 統合デッキ — 研究全体像から詳細への1本（進捗第11回スタイル）。

構成: SLIDES_統合デッキ_設計.md。00概要→01手法→02/04FW→03モバイル→05決定→06/08実装→07信頼→まとめ。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def callout(s, y, head, body, w=W, fs=16):
    bh = 0.95
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = INK; box.line.width = Pt(1.25); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(12); tf.margin_top = Pt(8)
    run(tf.paragraphs[0], head, 13, GRAY, bold=True)
    p = tf.add_paragraph(); p.space_before = Pt(4)
    run(p, body, fs, INK, bold=True)
    return y + bh + 0.16


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


def section(num, text):
    s = add_slide()
    tf = s.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(8.2), Inches(0.5)).text_frame
    run(tf.paragraphs[0], num, 15, GRAY, bold=True)
    tf = s.shapes.add_textbox(Inches(0.88), Inches(3.35), Inches(8.4), Inches(1.0)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 26, INK, bold=True)
    hrule(s, 4.4, l=0.92, w=4.0, color=INK)
    return s


# ════════════════════════════════════════════════════════════════════════════
# 0. 導入
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(8.3), Inches(1.5)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "zk-UAV なりすまし検知", 30, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "多端末ZK集約とモバイル実装", 22, INK, bold=True)
hrule(s, 4.5, l=0.92, w=5.6)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.66), Inches(8.2), Inches(0.6)).text_frame
run(tf.paragraphs[0], "研究全体像から、フレームワーク選定・モバイル実装検証まで", 14, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.5), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 統合", 12, GRAY)

s = add_slide(); title(s, "研究の一言と3本柱")
bullets(s, TOP, [
    (0, "ドローン無改造で、複数端末のRSSIから RemoteID 放送位置の真偽を判定し、"),
    (0, "各端末のゼロ知識証明を1つに集約してブロックチェーンに記録するシステム"),
    (0, "3本柱"),
    (1, "物理検知(RSSI) = 電波の物理から放送位置の矛盾を捕まえる"),
    (1, "ゼロ知識証明 = 座標・生RSSIを秘匿したまま「整合したか」だけを証明"),
    (1, "多端末合意＋ブロックチェーン = 改ざん不可・検閲耐性で記録"),
    (0, "「近接の証明」→「なりすまし検知の証明」へ再定義"),
], size0=15.5)

s = add_slide(); title(s, "背景・課題")
bullets(s, TOP, [
    (0, "RemoteID = ドローンが ID・位置・速度を自己申告で放送する仕組み"),
    (0, "放送に暗号署名がなく、位置フィールドの書換だけでなりすましが成立"),
    (0, "課題 = ドローン無改造・受信側だけで真偽判定"),
    (1, "かつプライバシーを守り、多端末の合意を改ざん・検閲されない形で記録したい"),
])

s = add_slide(); title(s, "検知原理の要約（詳細は第19回）")
y = bullets(s, TOP, [
    (0, "各端末は同じ放送について2通りの距離を持つ"),
    (1, "d_claim = 放送位置と自GPSの幾何距離 / d_meas = RSSIからの推定距離"),
    (1, "本物なら一致、位置詐称なら乖離"),
    (0, "誤差は乗法的 → 比(相対倍率 r=2.1)で判定、単発は中央値集約"),
    (0, "小さな詐称はノイズ以下 → 分散した複数端末の K-of-N 合意で検知（√N）"),
])

s = add_slide(); title(s, "本デッキの焦点")
y = code(s, TOP, [
    "各端末: RSSIで真偽判定 → ゼロ知識証明（座標は秘匿）",
    "  → 複数端末を1つに集約 → ブロックチェーンで trusted / spoofAlert",
])
y = bullets(s, y + 0.1, [
    (0, "焦点 = 各端末の証明を「1つに集約」し、「モバイルで生成」する土台の選定"),
    (0, "流れ = 手法 → フレームワーク → モバイル → 決定 → 実測 → 安全 → まとめ"),
])

s = add_slide(); title(s, "この資料の位置づけと結論（先に）")
bullets(s, TOP, [
    (0, "位置づけ = フレームワークの「比較」に基づく、研究方針の「決定」資料"),
    (0, "結論（研究方針）"),
    (1, "① 端末証明は Mopro で標準化（現行 circom を維持・移行コスト最小）"),
    (1, "② 集約の端末寄せを Nova/Folding で検証（省メモリ）"),
    (1, "③ 必要なら Noir で ★1(トラステッドセットアップ)を根治（gas増は覚悟）"),
    (0, "以降 = この結論に至る根拠（手法・比較・実測・安全性）"),
], size0=15.5)

s = add_slide(); title(s, "用語（補足）")
bullets(s, TOP, [
    (0, "witness = ドローン放送を受信する観測端末（周囲のスマホ）"),
    (0, "K-of-N 合意 = N 台中 K 台が「整合」と判定したら結論を確定"),
    (0, "ゼロ知識証明(ZKP) = 中身を見せずに計算の正しさだけを証明"),
    (0, "トラステッドセットアップ(★1) = 一部の証明系が要る鍵作りの儀式（漏れると偽造の恐れ）"),
    (0, "Folding = 複数の証明を1個ずつ畳み込んで1つにまとめる方式（省メモリ）"),
    (0, "gas = オンチェーンで検証する費用 ／ #3b = 検閲(間引き)対策の網羅性証明"),
], size0=14.5)

# 1. 手法の全体像
section("Part 1", "手法の全体像")

s = add_slide(); title(s, "やりたいこと＝4つの層")
y = code(s, TOP, [
    "層1  ZKPで何ができるか（秘匿したまま正しさを証明）",
    "層2  証明系の種類（Groth16 / PLONK / STARK / Folding）",
    "層3  複数の証明をまとめる手法（本案件の肝）",
    "層4  検証・共有（オンチェーン / DA層）",
])
y = bullets(s, y + 0.1, [
    (0, "フレームワークは層2〜3を実装する“乗り物”にすぎない"),
    (0, "先に手法を決め、その後に乗り物を選ぶのが正しい順序"),
])

s = add_slide(); title(s, "層2 — 証明系の種類")
y = table(s, TOP, [
    ["種類", "証明サイズ", "オンチェーン", "事前準備", "ひとこと"],
    ["Groth16", "最小", "最安(≈348k)", "回路ごと", "小さく安い・準備が弱点"],
    ["PLONK/Honk", "小〜中", "中〜高", "1回", "準備が緩い・汎用"],
    ["STARK", "大", "高", "不要", "準備不要・証明が大"],
    ["Folding", "最終のみ", "圧縮依存", "ほぼ不要", "N個を逐次畳む・省メモリ"],
], col_w=[1.5, 1.3, 1.6, 1.2, 3.0], fs=11.5, row_h=0.5)
y = bullets(s, y + 0.05, [
    (0, "現行の回路別セレモニー(★1)は PLONK/STARK/Folding で構造的に消える"),
], size0=14.5)

s = add_slide(); title(s, "層3 — 複数の証明をまとめる手法")
y = table(s, TOP, [
    ["#", "手法", "オンチェーン", "現状"],
    ["1", "オンチェーン集計(K-of-N)", "台数分 O(N)", "実装済 M1"],
    ["3", "再帰的証明合成", "1回・定数", "実装済 M2"],
    ["4", "Folding(Nova系)", "最終1回", "未（端末寄せ候補）"],
    ["2/5/6", "回路内バッチ / 専用集約 / 閾値署名", "各種", "試作 / 未"],
], col_w=[1.0, 3.4, 1.9, 2.3], fs=11.5, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "1(集計)・3(再帰)は実装済。未検討で価値ありは 4(Folding)＝端末側で省メモリ集約"),
], size0=14.5)

s = add_slide(); title(s, "層4 — 検証・共有の手法")
table(s, TOP, [
    ["手法", "効く攻撃", "現状"],
    ["オンチェーン verifier", "改ざん（暗号で不可）", "実装済"],
    ["DA層に提出集合コミット（網羅性）", "検閲（間引き）", "設計済 #3b"],
    ["許可型チェーン（BFT合意＋全複製）", "検閲・可用性", "設計済"],
], col_w=[3.6, 2.7, 2.3], fs=12, row_h=0.52)

# 2. フレームワーク
section("Part 2", "フレームワーク")

s = add_slide(); title(s, "何者か（役割別）")
bullets(s, TOP, [
    (0, "SNARK系（回路＋証明）"),
    (1, "circom+rapidsnark(iden3) / gnark(ConsenSys) / Noir(Aztec) / halo2(Zcash→PSE)"),
    (0, "STARK・再帰・Folding"),
    (1, "Plonky2(Polygon Zero) / Nova(Microsoft Research) / Cairo(StarkWare)"),
    (0, "zkVM = RISC Zero / SP1（普通のコードで証明・重い）"),
    (0, "橋渡し = Mopro(Ethereum財団PSE)＝既存証明器をモバイルアプリのバインディング化"),
], size0=15.5)

s = add_slide(); title(s, "早見表")
table(s, TOP, [
    ["フレームワーク", "エンジン", "事前準備", "モバイル", "一言"],
    ["circom+rapidsnark", "Groth16", "回路別", "○", "現行の定番・安い検証"],
    ["Noir+Barretenberg", "Honk", "不要", "◎", "準備不要・速い・gas高"],
    ["gnark", "Groth16/PLONK", "回路別", "✕", "集約(サーバ)"],
    ["halo2", "PLONK", "1回/不要", "○", "強力だが難しい"],
    ["Plonky2/3", "STARK", "不要", "○", "準備不要・再帰速い"],
    ["Nova/Sonobe", "Folding", "ほぼ不要", "◎", "端末側集約向き"],
    ["Cairo/Stwo", "STARK", "不要", "◎", "Starknet前提"],
    ["Mopro", "(包む)", "—", "◎", "証明系でなく橋渡し"],
], col_w=[2.4, 1.9, 1.2, 1.2, 1.9], fs=11, row_h=0.42)

s = add_slide(); title(s, "主要候補の要点")
bullets(s, TOP, [
    (0, "circom+rapidsnark（現行）= 証明が最小・オンチェーン最安。★1が残る"),
    (0, "Noir+Barretenberg = 準備不要で端末証明が速い。ただしgas約7倍・回路移植要"),
    (0, "gnark = 集約(M2)のサーバ担当。端末証明には重い"),
    (0, "Nova/Sonobe = 省メモリで端末側集約向き（研究的）"),
    (0, "Mopro = 証明系でなく橋渡し。現行circomをそのままFlutterに載せられる"),
])

s = add_slide(); title(s, "その他・新顔（参考 — 主軸は変えない）")
bullets(s, TOP, [
    (0, "モバイル新規に効きうる = Ligero/Ligetron（低メモリ）・o1js（TS・準備不要）"),
    (0, "集約を攻めるなら = ProtoGalaxy/HyperNova（Folding改良）・Spartan"),
    (0, "zkVM他 = Jolt / Miden / Nexus、プローバ = Expander(GKR)・Binius"),
    (0, "結論 = 主軸(Mopro＋circom/Noir・集約gnark/folding・EVM前提)は不変"),
])

# 3. モバイル
section("Part 3", "モバイルで動かすには")

s = add_slide(); title(s, "前提 — モバイルでZKPとは")
y = code(s, TOP, [
    "端末でやること : 入力 → witness生成 → proof生成   （重い＝本題）",
    "端末外でやること: verify（軽い・普通はチェーン/サーバ）",
])
y = bullets(s, y + 0.08, [
    (0, "モバイル化 = 証明生成器を端末で動くバイナリにして呼ぶこと"),
    (1, "ネイティブ = arm64にクロスコンパイル（速い・本命）"),
    (1, "WASM = 移植は楽だがネイティブ比3.5〜4倍遅い"),
    (0, "共通制約 = RAM 3GB以上で安定・iOSシミュレータ不可・Androidはarm64"),
])

s = add_slide(); title(s, "Flutter連携と難易度（分岐点＝Mopro）")
y = table(s, TOP, [
    ["FW", "難易度", "理由"],
    ["circom+rapidsnark / Noir", "★☆☆ 易", "Mopro公式Flutter対応"],
    ["halo2 / Cairo", "★★☆ 中", "回路が難しい / 土台がStarknet"],
    ["Plonky2 / Nova", "★★★ やや難", "Flutter成熟プラグイン無く自前FFI"],
    ["gnark / RISC Zero", "★★★ 難", "端末証明に不向き(サーバ/リモート)"],
], col_w=[3.0, 1.6, 4.0], fs=11.5, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "Mopro公式対応は circom / halo2 / Noir。現行もNoirも「易」＝モバイル化リスク低"),
], size0=14.5)

# 4. 決定
section("Part 4", "決定")

s = add_slide(); title(s, "評価軸（重み）")
table(s, TOP, [
    ["軸", "重み", "理由"],
    ["現行資産の再利用（移行コスト小）", "★★★", "動いている circom/gnark を捨てない"],
    ["セットアップの健全性（★1根治）", "★★★", "プロジェクトの核＝セキュリティ"],
    ["オンチェーン検証コスト（gas）", "★★☆", "チェーン記録が前提"],
    ["モバイル証明の現実性", "★★★", "端末で秒級・省メモリ"],
    ["集約適性（多端末を1つに）", "★★☆", "本案件の新規性"],
], col_w=[3.4, 1.2, 4.0], fs=12, row_h=0.46)

s = add_slide(); title(s, "決定 — 段階的アプローチ")
bullets(s, TOP, [
    (0, "フェーズ1（すぐ・低リスク）= Mopro でモバイル標準化"),
    (1, "現行 circom のまま Flutter に載せ替え。回路・検証・集約は変えない"),
    (0, "フェーズ2（並行・研究）= 集約の端末寄せを Nova/Folding で検証"),
    (1, "per-witness 回路はそのまま、集約だけ試作。gnark はフォールバック"),
    (0, "フェーズ3（条件付き）= ★1根治のため Noir 評価"),
    (1, "gas増(2.4M)と集約作り直しを許容できる場合のみ"),
])

s = add_slide(); title(s, "明示的に採らない（Rejected）")
bullets(s, TOP, [
    (0, "halo2 = 回路実装の難度が高く費用対効果が低い"),
    (0, "Plonky2を端末回路に = EdDSAとGoldilocks体が不一致で重い（集約用途は将来可）"),
    (0, "Cairo/Stwo = モバイルは優秀だが検証Starknet前提でEVMと不整合"),
    (0, "RISC Zero/SP1 = 端末単独の証明が重くリモート前提でオンデバイスにならない"),
], size0=16.5)

# 5. 実装検証
section("Part 5", "実装検証（実測・実機ビルド）")

s = add_slide(); title(s, "現行 circom の実測")
y = table(s, TOP, [
    ["回路", "制約数", "証明生成", "検証"],
    ["uav_witness_attestation（フル）", "11,785", "≈1.4s(snarkjs)", "OK(≈348k gas)"],
    ["remoteid_consistency_rel（RSSI単体）", "649", "—", "—"],
], col_w=[3.8, 1.5, 1.9, 1.4], fs=12, row_h=0.5)
y = bullets(s, y + 0.12, [
    (0, "従来「~18k制約」は概算、実測は 11,785（≈12k）＝Groth16 では小規模"),
    (0, "rapidsnark(arm64) なら snarkjs より速く端末も実用域"),
])

s = add_slide(); title(s, "Noir の実測（proof→verify）")
bullets(s, TOP, [(0, "環境 = nargo 1.0.0-beta.22 / bb 5.0.0-nightly（UltraHonk）")])
table(s, 1.95, [
    ["コマンド", "結果"],
    ["nargo test", "3/3 通過（正常 / 合意未達 fail / 束縛違反 fail）"],
    ["bb gates", "回路サイズ 2,953 ゲート"],
    ["bb prove", "証明生成 0.11s・proof ≈14.6 KB"],
    ["bb verify", "「Proof verified successfully」・0.02s"],
], col_w=[2.2, 6.4], fs=12.5, row_h=0.52)

s = add_slide(); title(s, "circom vs Noir（暫定比較・同条件でない点に注意）")
y = table(s, TOP, [
    ["項目", "circom（Groth16）", "Noir（UltraHonk）"],
    ["Poseidon 1回（同機能・公平）", "415制約", "404ゲート ＝ ほぼ同等"],
    ["セットアップ", "回路別セレモニー要", "不要（universal）"],
    ["オンチェーン検証", "≈348k gas", "≈2.4M gas（文献）"],
    ["証明生成 ※規模が違う", "≈1.4s（フル11.8k回路）", "0.11s（小PoC・署名なし）"],
    ["EdDSA署名検証", "8,086制約（実測）", "未実測（文献≈24,472）"],
], col_w=[2.9, 2.9, 2.8], fs=11.5, row_h=0.46)
y = bullets(s, y + 0.06, [
    (0, "公平に比べられるのは Poseidon（同機能）のみ ＝ ほぼ同等"),
    (0, "証明生成 0.11s vs 1.4s は回路規模が違い直接比較不可・EdDSAはNoir未実測"),
    (1, "同条件の完全比較は今後（Noirに署名を実装して計測）"),
], size0=13.5)

s = add_slide(); title(s, "Mopro PoC — 実機ビルドは確認済み（証明は次段）")
bullets(s, TOP, [
    (0, "確認済み = circom プローバを Flutter/Android の APK（53MB）までビルド"),
    (1, "arm64ネイティブ証明器(.so 3MB) / generateCircomProof を呼ぶ Flutterアプリ"),
    (0, "未達（次段）= 実機で flutter run し、オンデバイス証明の時間・メモリを実測"),
    (0, "意義 = 「モバイルで動く土台をビルドできる」ことを実機ビルドで裏取り"),
    (1, "＝“実機で証明が回った”ではなく“アプリまで作れた”段階、と正しく区別"),
], size0=15)

s = add_slide(); title(s, "Mopro PoC — つまずき（記録）")
table(s, TOP, [
    ["つまずき", "回避"],
    ["Noir の Android は Zig 必須(barretenberg)", "Android は circom で回避"],
    ["mopro導入に rustc 1.91+", "rustup update(→1.97)"],
    ["template の import名不一致", "mopro_example_app→実名に置換"],
    ["APK が4 ABIビルドで時間切れ", "--target-platform android-arm64"],
], col_w=[4.6, 4.0], fs=12, row_h=0.5)

# 6. 信頼モデル
section("Part 6", "信頼モデル")

s = add_slide(); title(s, "信頼が必要になる「4つの場所」")
y = code(s, TOP, [
    "① セットアップ  … 鍵作りの儀式は健全か（証明系で変わる）",
    "② プローバ      … 生データを誰が見るか（集約手法で変わる）",
    "③ 集約者(検閲)  … 有効票を間引けるか（変わらない）",
    "④ 入力の真正性  … 入れた値は本物か（変わらない・ZKP外）",
])
y = bullets(s, y + 0.1, [
    (0, "①②は証明系/集約手法の選択で変わる。③④はフレームワークを替えても消えない"),
])

s = add_slide(); title(s, "信頼モデル — 横断まとめ")
y = table(s, TOP, [
    ["観点", "変わるか", "効いてくるもの"],
    ["① セットアップ", "証明系で変わる", "Groth16=回路別 / STARK・Folding=不要"],
    ["② プローバ", "集約手法で変わる", "回路内バッチ回避 / 再帰・Foldingは分散"],
    ["③ 集約者の検閲", "変わらない", "網羅性(#3b) or 合意"],
    ["④ 入力の真正性", "変わらない", "アテステーション/登録/相互整合"],
], col_w=[2.0, 1.7, 4.9], fs=11, row_h=0.5)
y = bullets(s, y + 0.08, [
    (0, "＝③④は既存の #3b・アテステーションで守る（設計/実装済）"),
], size0=14.5)

# 7. まとめ
section("Part 7", "まとめ・今後")

s = add_slide(); title(s, "成果まとめ")
bullets(s, TOP, [
    (0, "検知性能（合成データ）= Δ=30m を N=5 で TPR 91%、現実ノイズでは K-of-N 合意が優位"),
    (0, "較正の重要性 = 想定 n を 0.5 ずらすと 91%→7.7%（実機較正は前提条件）"),
    (0, "回路コスト（実測）= witnessフル 11,785制約・証明 ≈1.4s"),
    (0, "モバイル実証 = Mopro で circom を Flutter/Android APK までビルド成功"),
    (0, "FW比較 = Poseidon は circom 415 ≒ Noir 404、Noir は準備不要だがgas約7倍"),
], size0=15.5)

s = add_slide(); title(s, "現状と限界（正直に）")
bullets(s, TOP, [
    (0, "中継/ワームホール攻撃は RSSI 単独では防げない（原理的限界）"),
    (0, "RSSIは環境依存 → 実機較正が必須、未較正では検知率が崩壊"),
    (0, "現状の数値は合成データによる評価（実機DJIは未実施）"),
    (0, "多機同時受信では UAS_ID ごとの窓分割が未対応"),
])

s = add_slide(); title(s, "今後")
bullets(s, TOP, [
    (0, "実機 DJI での A/n/σ 較正（最重要・検知率を分ける前提条件）"),
    (0, "多機環境対応（UAS_ID ごとの RSSI ウィンドウ分割）"),
    (0, "集約の端末寄せ（Nova/Folding で省メモリ集約の PoC）"),
    (0, "DA層の実接続（#3b の提出集合コミット）"),
    (0, "モバイル配布の標準化（Mopro）と ★1根治のための Noir 段階評価"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_統合_week19.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
