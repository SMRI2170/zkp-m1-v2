import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PATH = Path("project_simple_design_presentation.pptx")

W, H, M = Inches(10), Inches(7.5), Inches(0.45)
IW = W - M * 2

# Simple Minimalist Palette
C_WHITE = RGBColor(255, 255, 255)
C_BLACK = RGBColor(30, 30, 30)
C_GRAY_DARK = RGBColor(100, 100, 100)
C_GRAY_LIGHT = RGBColor(230, 230, 230)
C_ACCENT = RGBColor(40, 100, 200) # Soft blue accent

FONT = "Hiragino Sans"

def rect(slide, l, t, w, h, bg_color=None, border_color=None, border_width=1):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    if bg_color:
        s.fill.solid()
        s.fill.fore_color.rgb = bg_color
    else:
        s.fill.background()
    
    if border_color:
        s.line.color.rgb = border_color
        s.line.width = Pt(border_width)
    else:
        s.line.fill.background()
    return s

def txb(slide, text, l, t, w, h, size=14, bold=False, color=C_BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    box.text_frame.word_wrap = True
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return box

def header(slide, title, sub, page=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C_WHITE
    txb(slide, title, M, Inches(0.4), IW, Inches(0.5), size=24, bold=True, color=C_BLACK)
    txb(slide, sub, M, Inches(0.9), IW, Inches(0.3), size=12, color=C_GRAY_DARK)
    rect(slide, M, Inches(1.3), IW, Pt(1), bg_color=C_GRAY_LIGHT)
    if page:
        txb(slide, str(page), W - Inches(0.6), H - Inches(0.4), Inches(0.4), Inches(0.3), size=10, color=C_GRAY_DARK, align=PP_ALIGN.RIGHT)

def draw_flow_simple(slide, steps, top_y=Inches(3.0)):
    n = len(steps)
    box_w = Inches(1.6)
    box_h = Inches(0.8)
    gap = (IW - (box_w * n)) / (n - 1) if n > 1 else 0
    
    for i, step in enumerate(steps):
        x = M + (box_w + gap) * i
        
        # Simple outlined box
        rect(slide, x, top_y, box_w, box_h, bg_color=C_WHITE, border_color=C_GRAY_DARK)
        
        txb(slide, step, x, top_y + Inches(0.15), box_w, box_h - Inches(0.15), size=11, align=PP_ALIGN.CENTER)
        
        # Arrow
        if i < n - 1:
            txb(slide, "→", x + box_w, top_y + box_h/2 - Inches(0.2), gap, Inches(0.4), size=18, color=C_GRAY_DARK, align=PP_ALIGN.CENTER)

def draw_comparison_simple(slide, left_title, left_items, right_title, right_items):
    w = (IW - Inches(0.8)) / 2
    
    # Left
    txb(slide, left_title, M, Inches(1.8), w, Inches(0.4), size=16, bold=True, color=C_GRAY_DARK, align=PP_ALIGN.CENTER)
    rect(slide, M, Inches(2.2), w, Pt(1), bg_color=C_GRAY_LIGHT)
    for i, item in enumerate(left_items):
        txb(slide, f"・ {item}", M, Inches(2.5) + Inches(0.6*i), w, Inches(0.5), size=12, color=C_BLACK)

    # Right (Accent)
    txb(slide, right_title, M + w + Inches(0.8), Inches(1.8), w, Inches(0.4), size=16, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
    rect(slide, M + w + Inches(0.8), Inches(2.2), w, Pt(1), bg_color=C_GRAY_LIGHT)
    for i, item in enumerate(right_items):
        txb(slide, f"・ {item}", M + w + Inches(0.8), Inches(2.5) + Inches(0.6*i), w, Inches(0.5), size=12, color=C_BLACK)

def draw_architecture_simple(slide, layers):
    h = Inches(4.5) / len(layers)
    for i, (title, desc) in enumerate(layers):
        y = Inches(1.8) + (h + Inches(0.1)) * i
        rect(slide, M, y, IW, h, bg_color=C_WHITE, border_color=C_GRAY_LIGHT)
        txb(slide, title, M + Inches(0.2), y + h/2 - Inches(0.15), Inches(2.0), Inches(0.3), size=14, bold=True, color=C_BLACK)
        txb(slide, desc, M + Inches(2.5), y + Inches(0.15), IW - Inches(2.7), h - Inches(0.3), size=12, color=C_GRAY_DARK)

def draw_metrics_simple(slide, metrics):
    w = (IW - Inches(0.6)) / len(metrics)
    for i, (val, label) in enumerate(metrics):
        x = M + (w + Inches(0.2)) * i
        rect(slide, x, Inches(3.0), w, Inches(1.5), bg_color=C_WHITE, border_color=C_GRAY_LIGHT)
        txb(slide, val, x, Inches(3.3), w, Inches(0.6), size=28, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
        txb(slide, label, x, Inches(4.0), w, Inches(0.4), size=12, color=C_GRAY_DARK, align=PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]
page = 1

# 1. Title
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = C_WHITE
txb(s, "UAV Safety Proof Project", M, Inches(2.5), IW, Inches(0.8), size=36, bold=True, color=C_BLACK, align=PP_ALIGN.CENTER)
rect(s, M + Inches(3.0), Inches(3.5), IW - Inches(6.0), Pt(1), bg_color=C_GRAY_DARK)
txb(s, "Simple & Minimalist Architecture Report", M, Inches(3.8), IW, Inches(0.5), size=16, color=C_GRAY_DARK, align=PP_ALIGN.CENTER)
txb(s, "OpenDroneID / ZKP / Flutter / YOLO", M, Inches(4.3), IW, Inches(0.5), size=12, color=C_GRAY_DARK, align=PP_ALIGN.CENTER)

# 2. Problem vs Solution
s = prs.slides.add_slide(blank)
header(s, "背景と課題解決", "なぜゼロ知識証明(ZKP)が必要なのか", page)
page += 1
draw_comparison_simple(s, "従来の課題 (生データ公開)", [
    "飛行ルートや現在地が常に筒抜けになる",
    "業務機密や運用者のプライバシーが侵害される",
    "通信の傍受によるスプーフィングのリスク",
    "第三者がドローンを特定・追跡しやすくなる"
], "ZKPの解決策 (非公開証明)", [
    "生座標をスマホ内(Private Input)に留める",
    "「30m以内にいたか」の事実だけを出力",
    "数学的な計算で事実を証明し、改ざんを防止",
    "運用者のプライバシーを守りつつ安全を担保"
])

# 3. Overall Architecture
s = prs.slides.add_slide(blank)
header(s, "全体アーキテクチャ", "システムを構成する4つの主要レイヤー", page)
page += 1
draw_architecture_simple(s, [
    ("UAV (Raspberry Pi)", "OpenDroneID 準拠のBLEブロードキャスト。GPSとBME280によるQFE高度取得。"),
    ("Receiver (Flutter)", "BLEスキャンによるUAV位置の取得と、自身のGPS座標・気圧の統合。証拠記録。"),
    ("AI Vision (YOLOv11)", "録画した映像からTFLiteモデルを用いてドローンを視覚的に検知。ログと紐付け。"),
    ("ZKP Engine (Circom)", "端末内でWitnessを計算し、rapidsnarkでGroth16証明を生成。オンデバイス検証。")
])

# 4. OpenDroneID Flow
s = prs.slides.add_slide(blank)
header(s, "OpenDroneID 通信フロー", "独自フォーマットからASTM国際標準への移行", page)
page += 1
draw_flow_simple(s, [
    "UAV GPS/気圧\n取得 (Pi)",
    "QFE相対高度\nへの変換",
    "BLE Service Data\n構築",
    "BasicIDとLocation\n交互送信",
    "スマホ側での\nスキャン・パース"
])

# 5. QFE Altitude Comparison
s = prs.slides.add_slide(blank)
header(s, "QFE 高度方式の導入", "ハードウェア間の高度基準ズレをどう解決するか", page)
page += 1
draw_comparison_simple(s, "GPS / 絶対気圧方式", [
    "GPSの垂直誤差は±10~30mと非常に大きい",
    "標準気圧(1013hPa)基準では天候で絶対値がズレる",
    "Piとスマホのセンサー個体差によるズレ",
    "正確な3D距離計算(30m以内判定)が不可能"
], "QFE (相対高度) 方式", [
    "両端末とも起動時の地上気圧をP₀として記録",
    "以後はP₀からの相対的な高度変化(m)のみを使用",
    "天候の変化がP₀と現在値で相殺される",
    "「地上からの高さの差」を高精度に比較可能"
])

# 6. Flutter App Pipeline
s = prs.slides.add_slide(blank)
header(s, "Flutter オンデバイス処理", "スマホ内で完結するZKP証明パイプライン", page)
page += 1
draw_flow_simple(s, [
    "BLEログと自身\nのGPSを統合",
    "動画撮影と\nAI検知の実行",
    "ZKPのJSON\n入力データ生成",
    "Witness計算\n(circom)",
    "Groth16証明生成\n(rapidsnark)"
])

# 7. ZKP Optimization Metrics
s = prs.slides.add_slide(blank)
header(s, "ZKP回路 最適化の成果", "Circomコンパイラ --O2 オプションの効果", page)
page += 1
txb(s, "最適化により線形制約(加算・定数乗算)は消滅したが、非線形制約(比較・ビット分解)が主なコスト要因として残る。", M, Inches(1.8), IW, Inches(0.8), size=13, color=C_GRAY_DARK)
draw_metrics_simple(s, [
    ("609 → 534", "Constraints (制約数) の削減"),
    ("0", "O2での線形制約数"),
    ("106 KB", "最適化後の R1CS サイズ"),
    ("~2 sec", "証明生成時間 (Pixel 8)")
])

# 8. ZKP Circuit Variants Architecture
s = prs.slides.add_slide(blank)
header(s, "ZKP 回路のバリエーション", "ユースケースに応じた4つの回路設計", page)
page += 1
draw_architecture_simple(s, [
    ("Detection Flag", "最もシンプル。30m以内かつ5秒以内かの「True/False」フラグのみをPublic出力。"),
    ("Detection Commitment", "フラグに加え、入力座標の「Poseidon Hash」を出力。秘匿検証に最適。"),
    ("Feature Classifier", "距離判定に加え、「接近速度」や「高度帯」などの危険度特徴を分類して出力。"),
    ("Time Series (N=20)", "録画20秒分の全データを一括処理。通信欠損やラグに強いバッチ証明。")
])

# 9. Future Roadmap Flow
s = prs.slides.add_slide(blank)
header(s, "今後のロードマップ", "通信・セキュリティ・信頼性のさらなる拡張", page)
page += 1
draw_flow_simple(s, [
    "通信拡張\n(BT5 / Wi-Fi)",
    "認証強化\n(Authメッセージ)",
    "改ざん防止\n(Merkle Proof)",
    "モデル改善\n(UAV専用AI)",
    "統合テスト\n(実環境E2E)"
])

prs.save(OUTPUT_PATH)
print(f"✅ Saved: {OUTPUT_PATH}")
