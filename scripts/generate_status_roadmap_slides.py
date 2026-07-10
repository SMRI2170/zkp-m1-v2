#!/usr/bin/env python3
"""現状（実装済み）と今後（残作業）のまとめスライド。

マルチ端末 × ブロックチェーンによる UAV RemoteID 検証プロジェクトについて、
「いまどこまで実装できたか」「これから何が必要か」を1本に整理する。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13.5, space_before=7, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


def section(text, sub_text=""):
    s = add_slide()
    tf = tbox(s, 0.9, 3.0, 8.2, 1.4)
    run(tf.paragraphs[0], text, 30, INK, bold=True)
    hrule(s, 4.25, l=0.92, w=4.6)
    if sub_text:
        tf2 = tbox(s, 0.92, 4.4, 8.2, 0.8)
        run(tf2.paragraphs[0], sub_text, 14, GRAY)
    return s


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.2, 8.2, 0.5)
run(tf.paragraphs[0], "進捗まとめ（2026-07 / 週17）", 15, GRAY)
tf = tbox(s, 0.88, 2.75, 8.4, 1.7)
run(tf.paragraphs[0], "現状と今後", 34, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "マルチ端末 × ブロックチェーンによる UAV RemoteID 検証", 17, INK)
hrule(s, 4.5, l=0.92, w=5.6)
tf = tbox(s, 0.92, 4.68, 8.2, 1.4)
run(tf.paragraphs[0], "いまどこまで実装できたか / これから何が必要か", 14, GRAY)

# 2 全体像
s = add_slide()
title(s, "全体像 — 何を作っているか")
tf = tbox(s, 0.7, 1.4, 8.7, 2.0)
bullet(tf, "狙い: RemoteID の「自己申告」を、ドローン改造なしで裏取りする", first=True, bold=True)
bullet(tf, "複数端末が同一ドローンを独立に観測 → K台合意で真偽を判定")
bullet(tf, "座標は ZK で秘匿、合意は改ざん不能にチェーン記録")
codebox(s, 0.7, 3.5, 8.6, 2.2, [
    "各端末: ドローン検知 + RSSI実測 + ZK証明生成",
    "   │",
    "   ▼  同一 drone_id を共有する複数 proof",
    "集約/検証: 放送位置 vs 実測の整合を確認、K台合意",
    "   │",
    "   ▼",
    "チェーン: Groth16検証 + AttestationRegistry に consensus 記録",
])
tf = tbox(s, 0.7, 5.9, 8.7, 0.9)
bullet(tf, "軸: 「接近の証明」ではなく「なりすまし検出の証明」", size=12.5, first=True, bold=True)

# 3 章扉
section("現状", "いまどこまで実装できたか")

# 4 実装済み 一覧
s = add_slide()
title(s, "実装済み — できたこと")
table(s, 0.7, 1.45, 8.6, [
    ["領域", "成果物", "状態"],
    ["M1 端末回路", "uav_witness_attestation.circom（public14・約1.2万制約）", "✅ 実証明OK"],
    ["M1 コントラクト", "AttestationRegistry + Groth16Verifier(生成)", "✅ テスト9/9"],
    ["M1 E2E", "実端末3台 → 実検証器 → K=3 合意", "✅ 実証済"],
    ["RSSI 検出 回路", "remoteid_consistency.circom（約540制約）", "✅ 整合1/詐称0"],
    ["RSSI モデル", "scripts/rssi/rssi_ranging.py（較正）", "✅ 検証済"],
    ["アプリ", "lib/core/rssi_ranging.dart + main.dart 配線", "✅ analyze clean"],
], col_w=[1.9, 4.9, 1.8], fs=10.5, row_h=0.55)
tf = tbox(s, 0.7, 6.15, 8.7, 0.9)
bullet(tf, "M1 はコミット済（0c64254）。RSSI 検出機能は新規追加分（未コミット）",
       size=11.5, first=True)

# 5 現状の到達点
s = add_slide()
title(s, "現状の到達点 — 何が言えるか")
tf = tbox(s, 0.7, 1.4, 8.7, 2.6)
subhead(tf, "できている（土台）", 0)
bullet(tf, "多端末の証明を同一ドローンで束ね、K台合意する仕組みが E2E で動く", first=True)
bullet(tf, "放送位置と RSSI 実測の矛盾を ZK で検出する最小プリミティブが動く")
bullet(tf, "アプリが受信時にその場で integ/ なりすまし疑い を判定・ログ表示")
tf = tbox(s, 0.7, 4.15, 8.7, 2.6)
subhead(tf, "まだ言えない（未達）", 0)
bullet(tf, "「K台合意 ＝ 真正」はまだ成立しない（登録署名が揃っただけ）", first=True, bold=True)
bullet(tf, "RSSI 較正が暫定値（実機 DJI 未測定）→ 判定閾値が仮")
bullet(tf, "偽装耐性の核（真正性・実測裏取り）は未実装")

# 6 章扉
section("今後", "これから何が必要か")

# 7 残作業ロードマップ
s = add_slide()
title(s, "これから必要なこと")
table(s, 0.7, 1.45, 8.6, [
    ["優先", "やること", "目的"],
    ["高", "回路を相対許容(±%)版へ更新", "RSSI の現実(相対誤差)に合わせる"],
    ["高", "実機 DJI で RSSI 較正 (A/n/TOL 確定)", "本当に検出できるかの実測検証"],
    ["高", "consistency を M1 レジストリに統合", "多端末で放送真偽を合意"],
    ["中", "M2 再帰アグリゲータ (gnark/Sonobe)", "集約proof1個・定数gas・関係強制"],
    ["中", "M3 真正性: ドローン署名RemoteID / 端末Attestation", "入力の偽装を封じる(耐偽装の核)"],
    ["低", "許可型EVM(Besu/Quorum)へデプロイ / アプリ整流", "運用形態・UI 整備"],
], col_w=[0.8, 5.0, 2.8], fs=10.5, row_h=0.55)

# 8 次の一手
s = add_slide()
title(s, "次の一手（直近）")
tf = tbox(s, 0.7, 1.45, 8.7, 3.0)
subhead(tf, "1. 回路を相対許容版へ更新", 0)
bullet(tf, "固定 TOL_CM は誤り。lower=claimed*(1-r), upper=claimed*(1+r) に", first=True)
sub(tf, "整合/詐称サンプルで再検証（軽い・すぐ）")
subhead(tf, "2. 実測フィージビリティ（本命の検証）")
bullet(tf, "実機 DJI の RemoteID を複数スマホで受信し (rssi, 真距離) を収集", first=True)
bullet(tf, "fit して A/n を確定 → 30m で本当に矛盾検出できるかを見る")
sub(tf, "ここが成立して初めて、なりすまし検出に意味が出る")
tf = tbox(s, 0.7, 5.4, 8.7, 1.3)
subhead(tf, "3. 多端末合意へ接続", 0)
bullet(tf, "consistency_ok を witness 回路 / レジストリに足し、K台で放送真偽を合意",
       first=True)

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week17")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "status_and_roadmap.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
