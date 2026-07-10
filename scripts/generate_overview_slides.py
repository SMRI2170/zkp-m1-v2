#!/usr/bin/env python3
"""全体像まとめスライド（現状をわかりやすく1本に）。

何を作っているか → 処理の3層分担(回路外/回路内/チェーン) → データフロー →
RSSIなりすまし検知 → 回路の積み上げ → 現状(できたこと/数値) → 今後。

デザイン（進捗第11回スタイル / ミニマル）: 白背景・黒文字・Yu Gothic / Courier New、
枠はコード抜粋のみ、タイトル直下に横罫線、4:3 (10 x 7.5 inch)。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x33, 0x3E, 0x6E)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; return tf


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.1)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True); return p


def bullet(tf, text, size=13, space_before=6, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold); return p


def sub(tf, text, size=11, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY); return p


def codebox(s, l, t, w, h, lines, size=11):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=11, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


# 1 表紙
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.1, 8.2, 0.5)
run(tf.paragraphs[0], "全体まとめ（2026-07 / 現状）", 15, GRAY)
tf = tbox(s, 0.88, 2.65, 8.4, 1.9)
run(tf.paragraphs[0], "UAV RemoteID 検証システム", 30, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "多端末 × ZK × ブロックチェーン", 24, INK, bold=True)
hrule(s, 4.55, l=0.92, w=5.6, color=ACCENT)
tf = tbox(s, 0.92, 4.72, 8.2, 1.6)
run(tf.paragraphs[0], "RemoteID の自己申告を、ドローン改造なしで裏取りする", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "軸: 「接近の証明」ではなく「なりすまし検出の証明」", 12.5, GRAY)

# 2 何を作っているか
s = add_slide()
title(s, "何を作っているか（1枚で）")
tf = tbox(s, 0.7, 1.4, 8.7, 2.3)
bullet(tf, "問題: RemoteID は自己申告。放送位置が本物かは裏取りできない", first=True, bold=True)
bullet(tf, "解: 複数端末が独立に観測し、K 台一致で真偽を合意")
sub(tf, "1台のスピーカー偽装では、離れた複数端末に整合する観測を作れない")
bullet(tf, "秘匿: 座標は ZK で隠し、判定結果(0/1)だけを公開")
bullet(tf, "記録: 合意を改ざん不能にブロックチェーンへ")
codebox(s, 0.7, 4.4, 8.6, 1.9, [
    "複数端末 ── 各自 ZK 証明 ──▶ チェーンで K 台合意",
    "  ・within30（30m以内で検知）",
    "  ・consistency（放送位置 vs RSSI実測が整合）",
    "  → trusted（裏取り成立） / spoof（なりすまし疑い）",
])

# 3 処理の分担（回路外/回路内/チェーン）← 中心
s = add_slide()
title(s, "処理の分担 — どこで何をするか")
table(s, 0.7, 1.5, 8.6, [
    ["場所", "やること", "保証"],
    ["端末（回路外）", "RSSI受信・中央値集約・距離推定・GPS取得", "—（真正性はM3で補完）"],
    ["回路内（ZK）", "距離30m・時刻5s・映像検知・放送vs実測の整合判定", "計算の正しさ＋座標秘匿"],
    ["チェーン", "証明を検証しK台を集計（trusted/spoof）", "改ざん不能・公開監査"],
], col_w=[1.9, 4.6, 2.1], fs=10.5, row_h=0.62)
tf = tbox(s, 0.7, 4.35, 8.7, 2.1)
subhead(tf, "RSSI の扱い（よくある疑問）", 0)
bullet(tf, "RSSI→距離の「推定」は回路外（log/小数はCircomが苦手・生値は信用できない）", first=True)
bullet(tf, "「整合の判定」は回路内 → consistency_ok を公開（座標は隠したまま）", bold=True)
sub(tf, "回路は measured_range_cm を private 入力として受け取り、放送位置との整合を証明")

# 4 データフロー
s = add_slide()
title(s, "データフロー（端末 → 回路 → チェーン）")
codebox(s, 0.7, 1.45, 8.6, 4.3, [
    ("ドローン ── RemoteID 放送(位置/ID) ──▶ 複数端末", False),
    ("", False),
    ("[端末] 放送受信 + 自GPS + 映像検知 + RSSI実測(中央値)", False),
    ("   │  uav_witness_attestation.circom で Groth16 証明生成(rapidsnark)", True),
    ("   │  public: within30_and_detected, drone_id, device_pubkey,", True),
    ("   │          time_bucket, consistency_ok", True),
    ("   ▼  同一 drone_id を共有する N 個の証明", False),
    ("[チェーン] AttestationRegistry", False),
    ("   ・Groth16 検証 + 相異なる登録端末を K 台集計", True),
    ("   ・整合 K 台 → ConsensusReached(trusted)", True),
    ("   ・矛盾 K 台 → SpoofAlert(なりすまし疑い)", True),
])
tf = tbox(s, 0.7, 5.95, 8.7, 0.8)
bullet(tf, "結合キー = drone_id_commitment（RemoteID は共通なので端末間で自然に一致）",
       size=12, first=True)

# 5 RSSI なりすまし検知
s = add_slide()
title(s, "RSSI なりすまし検知の仕組み")
tf = tbox(s, 0.7, 1.35, 8.7, 1.0)
bullet(tf, "各端末が同じ放送について 2 通りの独立な距離を持つ", first=True, bold=True)
codebox(s, 0.7, 2.4, 8.6, 1.5, [
    "d_claim = |放送位置 − 自端末GPS|   （幾何）",
    "d_meas  = RSSI から推定           （電波の物理・中央値集約）",
    "整合 ⇔ d_meas*(1-r) ≤ d_claim ≤ d_meas*(1+r)   (r=±90%)",
])
tf = tbox(s, 0.7, 4.1, 8.7, 2.4)
bullet(tf, "放送だけ見れば区別できない。実測を突き合わせて矛盾を検出", first=True)
bullet(tf, "RSSI 誤差は相対%（乗法的）→ 許容は相対帯 / 単発は粗いので中央値集約")
bullet(tf, "検知できるのは粗い位置詐称（≳20–30m）。中継攻撃は原理的に不可（限界）", bold=True)
sub(tf, "だから多 witness の冗長と、真正性(M3)が必要")

# 6 回路の積み上げ
s = add_slide()
title(s, "回路の積み上げ")
codebox(s, 0.7, 1.4, 8.6, 3.4, [
    "Distance30mTimeCore  (距離30m + 時刻5s)   ← 全ての土台",
    "   ├─ detection_commitment    + 映像検知 + Poseidon秘匿",
    "   ├─ feature_classifier6      + 音響6次元 線形分類",
    "   └─ uav_witness_attestation  + drone_id / EdDSA署名 / time_bucket",
    "                                + consistency_ok（RSSI整合）",
    "",
    "remoteid_consistency(_rel)     放送 vs 実測レンジの整合（相対許容±%）",
])
tf = tbox(s, 0.7, 5.0, 8.7, 1.4)
bullet(tf, "共通思想: 座標や特徴は隠し、条件の成否(0/1)だけ公開", first=True, bold=True)
sub(tf, "平方根を避け二乗のまま比較 / 整数量子化 / 重みの符号分離で Circom に収める")

# 7 現状（できたこと + 数値）
s = add_slide()
title(s, "現状 — できたこと")
table(s, 0.7, 1.4, 8.6, [
    ["領域", "状態"],
    ["M1 マルチwitness基盤（回路+レジストリ+E2E）", "✅ Hardhat 11/11"],
    ["RSSI整合を回路・レジストリに統合（trusted/spoof）", "✅ 実証明でE2E"],
    ["相対許容回路 / 中央値ロバスト化（回路・アプリ一致）", "✅"],
    ["脅威モデル + 合成ROCで検知可否を定量化", "✅"],
    ["アプリ: RSSI整合の live判定 + CSVログ収集", "✅（未コミット）"],
], col_w=[6.0, 2.6], fs=10.5, row_h=0.5)
tf = tbox(s, 0.7, 4.5, 8.7, 2.0)
subhead(tf, "主な評価数値", 0)
bullet(tf, "位置詐称30mを N=5 で 91% 検知 / 単独64%→8台98%（多witness必須）", first=True)
bullet(tf, "較正がズレると崩壊（一致91%→n−0.5で7.7%）→ 実機較正は必須", bold=True)

# 8 今後
s = add_slide()
title(s, "今後")
tf = tbox(s, 0.7, 1.45, 8.7, 3.0)
subhead(tf, "残りの設計・実装", 0)
bullet(tf, "Phase 3 実機 DJI 較正: CSV収集 → A/n/σ fit → 閾値確定 → ROC再評価（要実機）", first=True)
bullet(tf, "M2 再帰アグリゲータ（gnark）: N証明を1つに集約、オンチェーン1回検証・定数gas")
bullet(tf, "M3 真正性: ドローン署名RemoteID / 端末アテステーション / 新鮮さ")
tf = tbox(s, 0.7, 4.3, 8.7, 2.2)
subhead(tf, "設計上の位置づけ（正直に）", 0)
bullet(tf, "現状は「K台の登録witnessが物理的矛盾を認めた」段階", first=True)
bullet(tf, "「K台合意 ＝ 真正」の完成は M3（真正性）＋実機較正に依存", bold=True)
sub(tf, "中継・Sybil・witness GPS詐称は RSSI では防げない → M3 で対処")

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week18")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "overview_current_state.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides._sldIdLst)} slides)")
