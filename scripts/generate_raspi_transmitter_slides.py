#!/usr/bin/env python3
"""ラズパイを OpenDroneID WiFi ビーコン送信機にした作業のまとめスライド。

デザイン方針（既存デッキと統一）:
  - 白背景 / 黒文字、重要部のみ青（インディゴブルー）
  - 4:3 (10 x 7.5 inch)
  - 和文 Yu Gothic / 等幅 Courier New
構成: 20 枚（表紙 + 内容 19）。図解多め。課題→対応で記述。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── 配色 ─────────────────────────────────────────────
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x1E, 0x40, 0xAF)   # インディゴブルー（強調・対応）
ISSUE = RGBColor(0xB1, 0x5A, 0x00)  # アンバー（課題）
GREEN = RGBColor(0x2E, 0x7D, 0x32)  # 成功・一致
GRAY = RGBColor(0x70, 0x70, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD9, 0xD9)
CODEBG = RGBColor(0xF2, 0xF4, 0xF8)
BLUEBG = RGBColor(0xEA, 0xEF, 0xFB)

JP_FONT = "Yu Gothic"
MONO_FONT = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

_page = [0]


def add_slide(number=True):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = slide.shapes.add_textbox(Inches(9.1), Inches(7.05), Inches(0.7), Inches(0.35)).text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(_page[0])
        r.font.size = Pt(10); r.font.color.rgb = GRAY; r.font.name = JP_FONT
    return slide


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def set_run(run, text, size, color=BLACK, bold=False, mono=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = MONO_FONT if mono else JP_FONT


def header(slide, title, kicker=None):
    if kicker:
        tf = textbox(slide, 0.6, 0.42, 9.0, 0.4)
        set_run(tf.paragraphs[0].add_run(), kicker, 13, GRAY, bold=True)
    tf = textbox(slide, 0.6, 0.74, 9.0, 0.9)
    set_run(tf.paragraphs[0].add_run(), title, 27, BLACK, bold=True)
    ln = slide.shapes.add_shape(1, Inches(0.62), Inches(1.56), Inches(2.4), Pt(3))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background(); ln.shadow.inherit = False


def bullet(tf, segments, size=15, space_before=8, level=0):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(0)
    indent = "    " * level
    marker = "・" if level == 0 else "– "
    set_run(p.add_run(), indent + marker, size, BLACK)
    for text, is_blue, is_mono in segments:
        set_run(p.add_run(), text, size, BLUE if is_blue else BLACK, bold=is_blue, mono=is_mono)


def section(tf, text, space_before=14):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    set_run(p.add_run(), text, 16, BLACK, bold=True)


def pa(tf, problem, action_segments, size=14, space_before=16):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    set_run(p.add_run(), "課題  ", size - 1, ISSUE, bold=True)
    set_run(p.add_run(), problem, size, BLACK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(2)
    set_run(p2.add_run(), "→ 対応  ", size - 1, BLUE, bold=True)
    for text, is_blue, is_mono in action_segments:
        set_run(p2.add_run(), text, size, BLUE if is_blue else BLACK, bold=is_blue, mono=is_mono)


def codebox(slide, left, top, width, height, lines, fs=11.5):
    box = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = LINE; box.line.width = Pt(1.0)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(10)
    tf.margin_top = Pt(8); tf.margin_bottom = Pt(8)
    for i, (text, is_blue) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0); p.space_after = Pt(0)
        set_run(p.add_run(), text, fs, BLUE if is_blue else BLACK, mono=True)
    return box


def rounded(slide, left, top, width, height, edge, weight=1.4, fill=WHITE):
    box = slide.shapes.add_shape(5, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.color.rgb = edge; box.line.width = Pt(weight)
    box.shadow.inherit = False
    return box


def boxlabel(slide, left, top, width, height, lines, edge=BLUE, fill=WHITE, title_color=BLUE):
    """角丸ボックス + 中央テキスト（1行目=見出し）。lines: [(text,size,color,bold)]"""
    rounded(slide, left, top, width, height, edge, fill=fill)
    tf = textbox(slide, left + 0.1, top + 0.12, width - 0.2, height - 0.2)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(3 if i else 0)
        set_run(p.add_run(), text, size, color, bold=bold)


def arrow(slide, left, top, width, text, color=BLUE):
    tf = textbox(slide, left, top, width, 0.45)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), text, 20, color, bold=True)


# ════════════════════════════════════════════════════
# 01 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
bar = s.shapes.add_shape(1, Inches(0.85), Inches(2.3), Inches(0.12), Inches(1.15))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
bar.line.fill.background(); bar.shadow.inherit = False
tf = textbox(s, 1.05, 2.42, 8.4, 1.4)
p = tf.paragraphs[0]
set_run(p.add_run(), "ラズパイから ", 32, BLACK, bold=True)
set_run(p.add_run(), "WiFi ビーコン", 32, BLUE, bold=True)
set_run(p.add_run(), "を出す", 32, BLACK, bold=True)
p2 = tf.add_paragraph()
set_run(p2.add_run(), "OpenDroneID テスト送信機を作り、スマホで受信検証", 16, GRAY)
tf = textbox(s, 1.07, 6.5, 8.2, 0.5)
set_run(tf.paragraphs[0].add_run(),
        "Raspberry Pi (raspberrypi.local) + hostapd + opendroneid-core-c", 12, GRAY, mono=True)

# ════════════════════════════════════════════════════
# 02 目的とゴール
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "目的とゴール", "Goal")
tf = textbox(s, 0.65, 1.9, 8.8, 5.0)
bullet(tf, [("スマホ側は既に ", False, False), ("WiFi Beacon 受信に対応済み", True, False),
            ("。だが実機ドローンが無いと検証できない", False, False)], 16)
bullet(tf, [("そこで ", False, False), ("ラズパイを OpenDroneID 送信機", True, False),
            (" にして、実機の代わりにテスト電波を出す", False, False)], 16)
bullet(tf, [("スマホがその電波を捕捉し、ドローン座標を表示 → ", False, False),
            ("ZK 証明フローに合流", True, False), (" するまでを確認", False, False)], 16)
section(tf, "達成基準")
bullet(tf, [("スマホのドローンカードが「未受信」→ ", False, False),
            ("緯度経度・ID 表示", True, False), (" に変わる", False, False)], 15)
bullet(tf, [("送信内容が ", False, False), ("公式 opendroneid-core-c と同一バイト", True, False),
            ("（標準準拠）であること", False, False)], 15)

# ════════════════════════════════════════════════════
# 03 背景: なぜ送信機が必要か
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "なぜ送信機が必要か", "背景")
tf = textbox(s, 0.65, 1.9, 8.8, 5.0)
bullet(tf, [("DJI 等の市販ドローンは Remote ID を ", False, False),
            ("WiFi Beacon", True, False), (" でブロードキャストする", False, False)], 16)
bullet(tf, [("受信側アプリの実装は完了。しかし ", False, False),
            ("送信側（ドローン）が無い", False, False), ("と動作確認できない", False, False)], 16)
bullet(tf, [("実機 DJI は高価・準備が大変 → ", False, False),
            ("身近なラズパイで標準準拠の電波を再現", True, False)], 16)
bullet(tf, [("受信〜証明までの ", False, False), ("エンドツーエンド検証環境", True, False),
            (" を自前で用意するのが狙い", False, False)], 16)

# ════════════════════════════════════════════════════
# 04 全体構成図
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "全体構成", "Architecture")
boxlabel(s, 0.55, 2.3, 2.7, 1.7, [
    ("ラズパイ (送信)", 14, BLUE, True),
    ("hostapd", 11, BLACK, False),
    ("vendor_elements", 10.5, BLACK, False),
    ("= OpenDroneID IE", 10.5, GRAY, False)])
arrow(s, 3.25, 2.7, 1.05, "›››")
boxlabel(s, 4.25, 2.3, 1.55, 1.7, [
    ("WiFi", 13, BLACK, True),
    ("Beacon", 12, BLACK, False),
    ("ch 9", 11, BLUE, False),
    ("2.4GHz", 10, GRAY, False)], edge=GRAY)
arrow(s, 5.8, 2.7, 1.05, "›››")
boxlabel(s, 6.85, 2.3, 2.6, 1.7, [
    ("スマホ (受信)", 14, BLUE, True),
    ("WifiOdidScanner", 10.5, BLACK, False),
    ("→ Dart パーサ", 10.5, BLACK, False),
    ("→ ZK 証明", 11, BLUE, False)])
tf = textbox(s, 0.65, 4.6, 8.8, 2.4)
bullet(tf, [("送信側＝今回構築。受信側＝既存実装をそのまま使用", False, False)], 14)
bullet(tf, [("両者をつなぐのは ", False, False),
            ("ASTM F3411 / OpenDroneID 標準フォーマット", True, False)], 14)
bullet(tf, [("検知ロジックは送信元に依存せず ", False, False),
            ("標準の識別子（OUI / type / メッセージ型）だけで判定", False, False)], 14)

# ════════════════════════════════════════════════════
# 05 ラズパイ発見
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "ラズパイを見つける", "Step 1")
tf = textbox(s, 0.65, 1.85, 8.8, 1.8)
bullet(tf, [("mDNS で探索 → ", False, False),
            ("raspberrypi.local = 192.168.11.54", True, True),
            ("（指定どおり .240 ではない別機体）", False, False)], 15)
bullet(tf, [("SSH 接続: ユーザ ", False, False), ("pi", False, True),
            (" / パスワード ", False, False), ("2170", False, True)], 15)
codebox(s, 0.65, 3.55, 8.7, 2.1, [
    ("$ dscacheutil -q host -a name raspberrypi.local", False),
    ("  ip_address: 192.168.11.54", True),
    ("$ ssh pi@192.168.11.54", False),
    ("  raspberrypi  /  Debian 12 (bookworm)", False),
    ("  Linux 6.6.51 aarch64", False)])

# ════════════════════════════════════════════════════
# 06 環境調査
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "環境を調べる", "Step 2")
tf = textbox(s, 0.65, 1.85, 8.8, 2.2)
bullet(tf, [("WiFi: ", False, False), ("wlan0 / phy0", False, True),
            ("、ドライバ ", False, False), ("brcmfmac", True, True),
            ("（内蔵 Broadcom）", False, False)], 15)
bullet(tf, [("sudo はパスワードで可、インターネット接続あり、有線 eth0 は未接続", False, False)], 15)
bullet(tf, [("SSH は ", False, False), ("wlan0 経由のみ", True, False),
            ("（AP 化すると切れる恐れ → 後述の課題②へ）", False, False)], 15)
codebox(s, 0.65, 4.35, 8.7, 1.9, [
    ("interface : wlan0 (STA, SSID Buffalo-G-6038, ch9)", False),
    ("driver    : brcmfmac", True),
    ("modes     : managed / AP / P2P  (monitor 無し)", False),
    ("eth0      : unavailable", False)])

# ════════════════════════════════════════════════════
# 07 課題① 注入非対応 → hostapd 方式
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "送信方式の選択", "課題 → 対応")
tf = textbox(s, 0.65, 1.85, 8.8, 3.0)
pa(tf, "内蔵 WiFi はモニターモード／パケット注入に非対応。aireplay 等で生フレームを撃てない。",
   [("AP モードは対応 → ", False, False),
    ("hostapd の vendor_elements", True, True),
    (" でビーコンにベンダ IE を埋め込む", False, False)])
pa(tf, "ベンダ IE を手で作るとフォーマット崩れのリスク。",
   [("送信バイトは ", False, False), ("OpenDroneID 標準の IE", True, False),
    (" をそのまま記述（次ページ）", False, False)])
codebox(s, 0.65, 5.0, 8.7, 1.3, [
    ("# hostapd は beacon に任意のベンダ IE を載せられる", False),
    ("vendor_elements=DD <len> FA0BBC 0D <counter> <message pack>", True)])

# ════════════════════════════════════════════════════
# 08 OpenDroneID WiFi Beacon フォーマット図
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "送信するIEの中身", "OpenDroneID / ASTM F3411")
parts = [
    ("DD", "element id\n221 (vendor)", 0.9),
    ("len", "長さ", 0.7),
    ("FA 0B BC", "OUI\n(ASTM 割当)", 1.5),
    ("0D", "type\nDrone RID", 1.0),
    ("cnt", "counter", 0.8),
    ("F2 19 02", "MessagePack\nヘッダ", 1.5),
]
x = 0.65
for label, desc, w in parts:
    is_key = label in ("FA 0B BC", "0D", "F2 19 02")
    boxlabel(s, x, 2.2, w, 0.95, [(label, 13, BLUE if is_key else BLACK, True)],
             edge=BLUE if is_key else GRAY, fill=BLUEBG if is_key else WHITE)
    tf = textbox(s, x, 3.2, w, 0.8)
    for j, ln in enumerate(desc.split("\n")):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), ln, 9.5, GRAY)
    x += w + 0.08
tf = textbox(s, 0.65, 4.5, 8.8, 2.5)
bullet(tf, [("OUI ", False, False), ("FA 0B BC", True, True),
            (" と type ", False, False), ("0x0D", True, True),
            (" は仕様で固定値（Remote ID 専用）", False, False)], 14)
bullet(tf, [("WiFi Beacon では複数メッセージを束ねた ", False, False),
            ("Message Pack (type 0xF)", True, False), (" で送るのが標準", False, False)], 14)
bullet(tf, [("中身: ", False, False), ("BasicID（ID 文字列）+ Location（緯度経度高度）", False, False)], 14)
bullet(tf, [("スマホの ", False, False), ("isOdidContent()", False, True),
            (" がこの OUI/type を照合して抽出する", False, False)], 14)

# ════════════════════════════════════════════════════
# 09 課題② SSH 切断回避
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "SSH を切らずに AP 化", "課題 → 対応")
tf = textbox(s, 0.65, 1.85, 8.8, 3.4)
pa(tf, "SSH は wlan0(WiFi)経由のみ。wlan0 をそのまま AP にすると接続が切れ、操作不能になる。",
   [("STA 接続を維持したまま ", False, False),
    ("仮想 AP インターフェース ap0 を追加", True, True),
    (" して、そこで hostapd を動かす", False, False)])
pa(tf, "ドライバの制約で STA と AP は同一チャンネルしか許されない。",
   [("STA が居る ", False, False), ("ch9 に AP も合わせる", True, False),
    ("（スマホは全 ch スキャンするので検知に支障なし）", False, False)])
codebox(s, 0.65, 5.3, 8.7, 1.1, [
    ("valid interface combinations:", False),
    ("  #{ managed } <= 1, #{ AP } <= 1, ... #channels <= 1", True)])

# ════════════════════════════════════════════════════
# 10 同時 STA+AP 図
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "同時 STA + AP の構成", "図解")
boxlabel(s, 3.2, 2.0, 3.6, 0.9, [("phy0 (内蔵無線 / ch9 固定)", 14, BLACK, True)], edge=GRAY)
boxlabel(s, 1.1, 3.5, 3.4, 1.6, [
    ("wlan0  — STA", 14, BLUE, True),
    ("ルータ Buffalo に接続", 11, BLACK, False),
    ("= SSH / インターネット", 11, GRAY, False),
    ("（維持する）", 11, GREEN, False)])
boxlabel(s, 5.5, 3.5, 3.4, 1.6, [
    ("ap0  — AP（新規）", 14, BLUE, True),
    ("hostapd で beacon 送出", 11, BLACK, False),
    ("SSID: ODID-RPI-TEST", 11, BLACK, False),
    ("= OpenDroneID 送信", 11, GREEN, False)])
tf = textbox(s, 0.65, 5.4, 8.8, 1.6)
bullet(tf, [("1 つの無線チップに 2 つの役割を同居（", False, False),
            ("同一 ch9 制約", True, False), ("）", False, False)], 14)
bullet(tf, [("ap0 は ", False, False), ("NetworkManager 管理外", False, True),
            (" に設定し、固定 MAC を付与", False, False)], 14)

# ════════════════════════════════════════════════════
# 11 実装手順
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "構築の手順", "Step 3")
tf = textbox(s, 0.75, 2.0, 8.6, 5.0)
steps = [
    ("hostapd を導入", "apt-get install -y hostapd（v2.10）"),
    ("仮想 AP を追加", "iw dev wlan0 interface add ap0 type __ap → ip link set ap0 up"),
    ("NM 管理外に", "nmcli dev set ap0 managed no / 固定 MAC 02:11:22:33:44:55"),
    ("IE を設定", "vendor_elements に OpenDroneID IE を記述（ch9）"),
    ("hostapd 起動", "AP-ENABLED を確認 → beacon 送出開始"),
]
for i, (t, d) in enumerate(steps, 1):
    p = tf.add_paragraph(); p.space_before = Pt(12)
    set_run(p.add_run(), f"{i}. ", 17, BLUE, bold=True)
    set_run(p.add_run(), t, 16, BLACK, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(1)
    set_run(p2.add_run(), "    " + d, 12, GRAY, mono=True)

# ════════════════════════════════════════════════════
# 12 hostapd.conf
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "hostapd 設定", "設定ファイル")
codebox(s, 0.65, 2.0, 8.7, 2.7, [
    ("interface=ap0", False),
    ("driver=nl80211", False),
    ("ssid=ODID-RPI-TEST", False),
    ("hw_mode=g", False),
    ("channel=9", True),
    ("ieee80211n=1", False),
    ("vendor_elements=dd3afa0bbc0d01f2190202125250...", True)])
tf = textbox(s, 0.65, 5.0, 8.8, 2.0)
bullet(tf, [("channel は STA と同じ ", False, False), ("9", True, False),
            ("（同一 ch 制約のため）", False, False)], 14)
bullet(tf, [("vendor_elements に ", False, False),
            ("OpenDroneID の IE（id+len 込み）", True, False), (" を 16 進で記述", False, False)], 14)

# ════════════════════════════════════════════════════
# 13 送信ペイロード生成
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "送信データを作る", "Payload")
tf = textbox(s, 0.65, 1.85, 8.8, 1.6)
bullet(tf, [("中身: ", False, False), ("BasicID = \"RPI-ODID-TEST\"", True, True),
            (" + ", False, False),
            ("Location（lat 35.6586 / lon 139.7454 / alt 80m）", True, True)], 14)
bullet(tf, [("座標は引数で変更可能（テスト中に動かせる）", False, False)], 14)
codebox(s, 0.65, 3.5, 8.7, 2.7, [
    ("# 生成された IE（先頭 DD=IE, 3A=長さ）", False),
    ("dd 3a  fa 0b bc  0d  01  f2 19 02 ...", True),
    ("        └OUI──┘ type cnt └pack┘", False),
    ("", False),
    ("lat バイト: 10124115   lon バイト: b0744b53", False),
    ("alt(geo)  : 7008       ID: 52 50 49 2d ... (RPI-..)", False)])

# ════════════════════════════════════════════════════
# 14 公式準拠へ切替（課題→対応）
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "公式実装に揃える", "課題 → 対応")
tf = textbox(s, 0.65, 1.85, 8.8, 3.6)
pa(tf, "公式の transmitter-linux はそのまま使えない（WiFi 未接続が必須 → SSH が切れる／Pi では部分動作）。",
   [("その心臓部 ", False, False),
    ("opendroneid-core-c だけを Pi にビルド", True, True),
    ("し、生成したバイトを既存 hostapd に流す", False, False)])
pa(tf, "自前エンコーダが標準と本当に一致するか不明。",
   [("公式 ", False, False), ("odid_message_build_pack()", False, True),
    (" の出力と ", False, False), ("バイト単位で突合", True, False),
    ("（次ページ）", False, False)])
tf = textbox(s, 0.65, 5.5, 8.8, 1.4)
bullet(tf, [("transmitter-linux も内部は ", False, False),
            ("結局 hostapd 方式", False, False),
            (" → 採用方式は公式と同じ仕組みだと確認できた", False, False)], 14)

# ════════════════════════════════════════════════════
# 15 自前 vs 公式 バイト比較
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "標準準拠の実証", "自前 vs 公式エンコーダ")
rows = [
    ("項目", "自前", "公式 core-c", "判定", True),
    ("Latitude バイト", "10124115", "10124115", "一致", False),
    ("Longitude バイト", "b0744b53", "b0744b53", "一致", False),
    ("AltitudeGeo", "7008", "7008", "一致", False),
    ("UASID 文字列", "RPI-ODID-TEST", "RPI-ODID-TEST", "一致", False),
    ("型/プロトコル", "同一", "同一", "一致", False),
    ("UAType/IDType", "00 (未設定)", "12 (回転翼)", "差分", False),
]
y = 2.05
for name, a, b, judge, is_head in rows:
    fill = BLUEBG if is_head else WHITE
    rounded(s, 0.65, y, 4.0, 0.62, LINE, fill=fill)
    rounded(s, 4.65, y, 1.9, 0.62, LINE, fill=fill)
    rounded(s, 6.55, y, 1.6, 0.62, LINE, fill=fill)
    rounded(s, 8.15, y, 1.2, 0.62, LINE, fill=fill)
    jcolor = BLACK if is_head else (GREEN if judge == "一致" else ISSUE)
    for left, w, text, col in [(0.65, 4.0, name, BLACK), (4.65, 1.9, a, BLACK),
                               (6.55, 1.6, b, BLACK), (8.15, 1.2, judge, jcolor)]:
        tf = textbox(s, left + 0.08, y + 0.13, w - 0.12, 0.4)
        p = tf.paragraphs[0]
        set_run(p.add_run(), text, 11.5, col, bold=is_head, mono=(not is_head and left in (4.65, 6.55)))
    y += 0.66
tf = textbox(s, 0.65, 6.65, 8.8, 0.7)
bullet(tf, [("アプリが解析する 緯度/経度/高度/ID/型 は ", False, False),
            ("完全一致", True, False),
            ("。差分は省略していた補助フィールドのみ", False, False)], 13)

# ════════════════════════════════════════════════════
# 16 結果（送信中の状態）
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "送信できた", "Result")
tf = textbox(s, 0.65, 1.9, 8.8, 2.2)
bullet(tf, [("hostapd ", False, False), ("AP-ENABLED", True, False),
            (" / SSID ", False, False), ("ODID-RPI-TEST", False, True),
            (" / ch9 で送信中", False, False)], 15)
bullet(tf, [("送信バイトは ", False, False), ("公式 core-c と同一（標準準拠）", True, False)], 15)
bullet(tf, [("SSH(STA) も維持。Pi 到達 OK（ping 0% ロス）", False, False)], 15)
codebox(s, 0.65, 4.3, 8.7, 2.0, [
    ("ap0: interface state UNINITIALIZED -> ENABLED", False),
    ("ap0: AP-ENABLED", True),
    ("wlan0: Connected to Buffalo-G-6038 (SSH alive)", False),
    ("ping 192.168.11.54 : 0.0% packet loss", False)])

# ════════════════════════════════════════════════════
# 17 運用スクリプト
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "操作用スクリプト", "Pi 上に用意")
codebox(s, 0.65, 2.05, 8.7, 2.5, [
    ("# 開始 / 座標変更 (counter lat lon alt)", False),
    ("~/odid-start.sh 1 35.6586 139.7454 80", True),
    ("", False),
    ("# 停止・仮想IF 撤去（元に戻す）", False),
    ("~/odid-stop.sh", True),
    ("", False),
    ("# 公式エンコーダで IE を生成", False),
    ("~/odid_ie 1 35.6586 139.7454 80", True)])
tf = textbox(s, 0.65, 4.85, 8.8, 1.8)
bullet(tf, [("start は ", False, False), ("ap0 を再利用", False, False),
            ("（座標変更の再実行で作り直さない）", False, False)], 14)
bullet(tf, [("hostapd 再起動時のみ STA が一瞬切れるが ", False, False),
            ("NM が数秒で自動再接続", True, False)], 14)

# ════════════════════════════════════════════════════
# 18 制約と限界（正直に）
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "制約と限界", "正直な注意点")
tf = textbox(s, 0.65, 1.9, 8.8, 5.0)
bullet(tf, [("内蔵チップは ", False, False), ("raw injection 非対応", False, False),
            (" → フレーム種別は beacon に限定", False, False)], 15)
bullet(tf, [("チャンネルは ", False, False), ("ch9", False, False),
            ("（実機 DJI は通常 ch6 / ch149）。ただしスマホは全 ch 走査するので検知は成立", False, False)], 15)
bullet(tf, [("hostapd 再起動の度に ", False, False), ("SSH が数秒瞬断", False, False),
            ("（自動復旧するが留意）", False, False)], 15)
bullet(tf, [("実機 DJI のビーコンを ", False, False),
            ("バイト比較で直接検証はしていない", False, False),
            ("（標準からの論理的帰結）", False, False)], 15)
bullet(tf, [("ただし ", False, False),
            ("スマホが解析する IE の中身は公式実装と同一", True, False),
            ("。代替テスト送信機として十分機能", False, False)], 15)

# ════════════════════════════════════════════════════
# 19 テスト手順
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "スマホでの確認手順", "Test")
tf = textbox(s, 0.75, 2.0, 8.6, 5.0)
ts = [
    ("最新 APK を入れる", "WiFi 対応版を Android 11+ 端末へ"),
    ("権限を許可", "位置情報 / 付近の WiFi デバイス"),
    ("メイン画面を見る", "ドローンカードが lat 35.6586 / lon 139.7454 / ID RPI-ODID-TEST に"),
    ("WiFi Raw タブで確認", "ODID-RPI-TEST が isOdid=true で表示されれば成功"),
    ("証明を生成", "受信データで ZK 証明が従来どおり作れることを確認"),
]
for i, (t, d) in enumerate(ts, 1):
    p = tf.add_paragraph(); p.space_before = Pt(12)
    set_run(p.add_run(), f"{i}. ", 17, BLUE, bold=True)
    set_run(p.add_run(), t, 16, BLACK, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(1)
    set_run(p2.add_run(), "    " + d, 12, GRAY)

# ════════════════════════════════════════════════════
# 20 まとめ・今後
# ════════════════════════════════════════════════════
s = add_slide()
header(s, "まとめと今後", "Summary")
tf = textbox(s, 0.65, 1.9, 8.8, 5.2)
section(tf, "やったこと", space_before=4)
bullet(tf, [("ラズパイを ", False, False),
            ("OpenDroneID WiFi ビーコン送信機", True, False), (" に構築", False, False)], 14)
bullet(tf, [("SSH を保つ ", False, False), ("STA+AP 同居", False, False),
            (" 構成、公式 core-c で ", False, False),
            ("標準準拠を実証", True, False)], 14)
section(tf, "今後")
bullet(tf, [("実機 DJI での受信テスト（世代・ファーム差の確認）", False, False)], 14)
bullet(tf, [("より忠実にするなら ", False, False),
            ("transmitter-linux（正規 ch / NAN / BT）", True, False), (" の構築", False, False)], 14)
bullet(tf, [("USB WiFi ドングルで ", False, False),
            ("monitor / injection 対応", False, False), (" にすればフレーム種別も再現可", False, False)], 14)

# ── 保存 ─────────────────────────────────────────────
out = "raspi_wifi_beacon_slides.pptx"
prs.save(out)
print(f"saved: {out}  slides: {len(prs.slides._sldIdLst)}")
