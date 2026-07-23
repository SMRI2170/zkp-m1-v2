#!/usr/bin/env python3
"""week19/06_ハンズオン結果.md のスライド化（進捗第11回スタイル）。"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16, size1=14):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 2))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.07 if i else 0)
    return y + h + 0.12


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(8.2), Inches(1.2)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "ハンズオン結果", 32, INK, bold=True)
hrule(s, 3.9, l=0.92, w=5.0)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.05), Inches(8.2), Inches(0.9)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "実測で circom と Noir を比較（proof→verifyまで）", 15, INK)
tf = s.shapes.add_textbox(Inches(0.92), Inches(6.4), Inches(8.0), Inches(0.5)).text_frame
run(tf.paragraphs[0], "week19 / 06_ハンズオン結果", 12, GRAY)

# S2 現行 circom 実測
s = add_slide(); title(s, "現行 circom 回路の実測")
y = table(s, TOP, [
    ["回路", "制約数", "証明生成", "検証"],
    ["uav_witness_attestation（フル）", "11,785", "≈1.4s(snarkjs)", "OK(≈348k gas)"],
    ["remoteid_consistency_rel（RSSI単体）", "649", "—", "—"],
], col_w=[3.8, 1.5, 1.9, 1.4], fs=12, row_h=0.5)
y = bullets(s, y + 0.15, [
    (0, "訂正 = 従来「~18k制約」は概算、実測は 11,785（≈12k）。Groth16 では小規模"),
    (0, "証明生成は snarkjs で ≈1.4s。rapidsnark(arm64) ならさらに速い＝端末も実用域"),
])

# S3 Noir 実測
s = add_slide(); title(s, "Noir 版の実測（proof→verify まで）")
bullets(s, TOP, [
    (0, "環境 = nargo 1.0.0-beta.22 / bb 5.0.0-nightly（UltraHonk）"),
])
table(s, 1.95, [
    ["コマンド", "結果"],
    ["nargo test", "3/3 通過（正常 count=2 / 合意未達 fail / 束縛違反 fail）"],
    ["nargo info", "ACIR 92 opcodes"],
    ["nargo execute", "witness 生成成功・回路出力 = 2（期待通り）"],
    ["bb gates", "回路サイズ 2,953 ゲート"],
    ["bb prove", "証明生成 0.11s・proof ≈14.6 KB"],
    ["bb verify", "「Proof verified successfully」・0.02s"],
], col_w=[2.2, 6.4], fs=12.5, row_h=0.5)

# S4 circom vs Noir 比較
s = add_slide(); title(s, "circom vs Noir（暫定比較・同条件でない）")
table(s, TOP, [
    ["項目", "circom（Groth16）", "Noir（UltraHonk）"],
    ["Poseidon 1回（同機能・公平）", "415制約", "404ゲート ＝ ほぼ同等"],
    ["EdDSA署名検証1回", "8,086制約（実測）", "未実測 / 文献≈24,472"],
    ["証明生成 ※規模が違う", "≈1.4s（フル11.8k）", "0.11s（小PoC・署名なし）"],
    ["証明サイズ", "定数・小（~200B級）", "≈14.6 KB（大）"],
    ["セットアップ", "回路別セレモニー要", "不要（universal）"],
    ["オンチェーン検証", "≈348k gas", "≈2.4M gas（文献）"],
], col_w=[2.9, 2.9, 2.8], fs=11.5, row_h=0.45)

# S5 EdDSA と Noir 暗号lib成熟度
s = add_slide(); title(s, "EdDSA と Noir 暗号ライブラリの成熟度")
bullets(s, TOP, [
    (0, "Poseidon は同機能で実測比較: circom 415 ≒ Noir 404（ほぼ同等）"),
    (1, "→ ハッシュのコストは近く、Noir移行で回路が極端に重くはならない見込み"),
    (0, "EdDSA は circom 側を実測（8,086制約）、Noir 側は実測に至らず"),
    (0, "つまずき = Noir公式 eddsa/ec の最新タグが u1（beta.22で削除）を使いコンパイル不可"),
    (1, "ec の main は修正済みだが未タグ、beta.22 は git依存に tag 必須で参照できず"),
    (0, "知見 = 署名検証のような重い回路は Noir だとライブラリのバージョン整合で詰まる"),
    (1, "circomlib は枯れていて即使える＝この成熟度の差自体が移行コスト"),
])

# S6 読み取り
s = add_slide(); title(s, "読み取り（結論）")
bullets(s, TOP, [
    (0, "現行回路は小さい（≈12k制約）＝証明が回るか自体は問題にならない"),
    (1, "デスクトップ1.4s / 端末も rapidsnark で実用域 / FibRace一般傾向 <5s"),
    (0, "差が出るのは制約数より セットアップ・オンチェーンgas・集約方式"),
    (0, "小PoCではNoirも高速。フル回路移植後に再評価（オンチェーン検証gasは約7倍）"),
    (0, "トレードオフ = Noir(準備不要・端末で速い) / Groth16(小さく安いオンチェーン検証)"),
], size0=16.5)

# S7 つまずき記録＋次の一手
s = add_slide(); title(s, "導入のつまずき ／ 次の一手")
bullets(s, TOP, [
    (0, "導入のつまずき（記録）"),
    (1, "nargo beta.23 はバイナリ未添付 → beta.22 で回避"),
    (1, "mopro 導入に rustc 1.91+ 必要 → rustup update（1.90→1.97）"),
    (1, "Noir の Android は Zig 必須（barretenberg ABI）→ Android は circom が楽"),
    (1, "Flutterテンプレートが mopro_example_app 前提 → import名を修正"),
    (0, "次の一手"),
    (1, "本物の回路(uav_witness_attestation.zkey)に差し替えて実機測定"),
    (1, "Poseidon/EdDSA を Noir版に足し、同機能での完全比較を完成（EdDSAはlib vendoring）"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "week19", "slides_06_ハンズオン結果.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
