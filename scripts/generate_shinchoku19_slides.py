#!/usr/bin/env python3
"""進捗第19回 — RSSIによる位置特定となりすまし検知（進捗報告・研究室/教授向け・拡張版）。

進捗第11回スタイル: 白背景・黒文字・Yu Gothic / Courier New、4:3。
■=第1階層 / ・=第2階層。体言止め。式/図は薄グレー枠。数値はシミュレーション実測。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]
L, W, TOP = 0.7, 8.6, 1.35


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font); return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def hrule(s, top, l=0.7, w=8.6, color=RULE):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = s.shapes.add_textbox(Inches(0.68), Inches(0.42), Inches(8.7), Inches(0.8)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 22, INK, bold=True)
    hrule(s, 1.14)


def _vlen(t):
    return int(sum(0.5 if ord(c) < 0x3000 else 1.0 for c in t))


def bullets(s, y, items, size0=16.5, size1=14.5):
    tf = s.shapes.add_textbox(Inches(L), Inches(y), Inches(W), Inches(0.4)).text_frame
    tf.word_wrap = True
    h = 0.0
    for i, (lv, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else (6 if lv == 0 else 3))
        size = size0 if lv == 0 else size1
        run(p, ("■ " if lv == 0 else "　　・ "), size, INK); run(p, text, size, INK)
        cpl = max(14, int(590 / size)); lh = size * 1.5 / 72.0
        h += max(1, -(-_vlen(text) // cpl)) * lh + (0.08 if i else 0)
    return y + h + 0.12


def code(s, y, lines, w=W, fs=13):
    lh = fs * 1.7 / 72.0
    bh = lh * len(lines) + 0.26
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 2)
        run(p, ln, fs, INK, font=MONO)
    return y + bh + 0.14


def callout(s, y, head, body, w=W, fs=16):
    bh = 0.95
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(L), Inches(y), Inches(w), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = INK; box.line.width = Pt(1.25); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(12); tf.margin_top = Pt(8)
    run(tf.paragraphs[0], head, 13, GRAY, bold=True)
    p = tf.add_paragraph(); p.space_before = Pt(4)
    run(p, body, fs, INK, bold=True)
    return y + bh + 0.16


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, y, rows, col_w, fs=12.5, row_h=0.44):
    nr, nc = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(L), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gtbl.table; tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(5); cell.margin_right = Pt(3)
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            run(cell.text_frame.paragraphs[0], str(val), fs, INK, bold=(ri == 0))
    return y + row_h * nr + 0.14


def section(text):
    s = add_slide()
    tf = s.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(8.2), Inches(1.0)).text_frame
    tf.word_wrap = True
    run(tf.paragraphs[0], text, 26, INK, bold=True)
    hrule(s, 4.15, l=0.92, w=4.0, color=INK)
    return s


# ════════════════════════════════════════════════════════════════════════════
# S1 表紙
s = add_slide(number=False)
tf = s.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(8.2), Inches(0.5)).text_frame
run(tf.paragraphs[0], "進捗報告 第19回", 15, GRAY)
tf = s.shapes.add_textbox(Inches(0.88), Inches(2.85), Inches(8.4), Inches(1.4)).text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "RSSI による位置特定と\nなりすまし検知", 28, INK, bold=True)
hrule(s, 4.55, l=0.92, w=5.4)
tf = s.shapes.add_textbox(Inches(0.92), Inches(4.72), Inches(8.2), Inches(0.6)).text_frame
run(tf.paragraphs[0], "ドローン無改造で、受信側だけで放送位置の真偽を検知", 14, INK)

# S2 目次
s = add_slide(); title(s, "本日の内容")
bullets(s, TOP, [
    (0, "背景・目的（RemoteIDのなりすまし）"),
    (0, "検知原理（2つの独立距離・RSSI距離モデル・比による判定）"),
    (0, "多端末による強化（信号対雑音・位置特定）"),
    (0, "評価（検知率・方式比較・較正誤差）＝シミュレーションで定量化"),
    (0, "実装状況・脅威モデル・今後"),
])

# S3 背景・目的
s = add_slide(); title(s, "背景・目的")
bullets(s, TOP, [
    (0, "RemoteID = ドローンが ID・位置・速度を自己申告で放送する仕組み"),
    (0, "受信側は放送を鵜呑みにするしかない構造的弱点"),
    (1, "位置フィールドの書換だけで偽装が成立"),
    (0, "目標 = ドローン無改造で、受信側だけで放送位置の真偽を判定"),
    (0, "手掛かり = 受信電波強度 RSSI と 自端末の GPS"),
])

# S4 なりすましの具体
s = add_slide(); title(s, "なりすましの具体（攻撃）")
y = code(s, TOP, [
    "実際の発信位置 A  ──(位置フィールドを書換)──▶  放送位置 B",
    "受信端末が表示するのは B（実際は A）",
])
y = bullets(s, y + 0.1, [
    (0, "放送に暗号署名がなく、内容の真偽を単体では確認できない"),
    (0, "＝「放送だけ見ても分からない」→ 物理実測(RSSI)で裏取りする"),
])

# S5 核アイデア
s = add_slide(); title(s, "核アイデア — 2つの独立な距離の突き合わせ")
y = bullets(s, TOP, [
    (0, "同じ放送に対する2通りの距離"),
    (1, "d_claim = ‖放送位置 − 自端末GPS‖ … 幾何由来"),
    (1, "d_meas = RSSI からの推定距離 … 物理由来"),
    (0, "本物なら一致、位置詐称なら乖離"),
])
y = code(s, y + 0.05, [
    "放送位置 P_claim ──(幾何)──▶ d_claim",
    "                               ⇕ 一致？",
    "受信 RSSI ───────(物理)──▶ d_meas",
])

# S6 RSSI距離モデル
s = add_slide(); title(s, "RSSI 距離モデル（log-distance path loss）")
y = code(s, TOP, [
    "RSSI(d) = A − 10·n·log10(d)",
    "d_meas  = 10^((A − RSSI) / (10·n))   ← 逆に解けば距離",
])
y = bullets(s, y + 0.05, [
    (0, "A = 1m地点のRSSI、n = 伝搬指数（自由空間2.0 / 市街地2.5〜4）"),
    (0, "A・n は実機の (距離, RSSI) を最小二乗フィットで決定"),
])
y = table(s, y + 0.05, [
    ["受信 RSSI", "−65 dBm", "−77 dBm", "−90 dBm"],
    ["推定距離", "約 10 m", "約 30 m", "約 100 m"],
], col_w=[1.8, 1.6, 1.6, 1.6])

# S7 誤差は乗法的
s = add_slide(); title(s, "誤差は乗法的 → 「比」で判定")
y = bullets(s, TOP, [
    (0, "RSSI誤差 ±4dB は距離では掛け算的に効く（約×2幅）"),
    (1, "「30mと測っても本当は約14〜63m」＝1台のRSSIは本質的に粗い"),
])
y = code(s, y + 0.05, [
    "距離の倍率 = 10^(k·σ/10n) = 10^(2·4/25) ≈ 2.1 倍（±2σ）",
    "整合 ⇔ d_meas / r ≤ d_claim ≤ d_meas × r    r = 2.1",
])
y = bullets(s, y + 0.05, [
    (0, "固定幅(±3m)は近距離で緩く遠距離で厳しく破綻 → 相対倍率で判定"),
])

# S8 中央値集約
s = add_slide(); title(s, "単発は粗い → 短窓の中央値集約")
bullets(s, TOP, [
    (0, "マルチパス（反射）で単発RSSIは大きく暴れる"),
    (0, "同一UAS_IDのRSSIを M サンプル集めて中央値へ（実装は最新10件）"),
    (1, "外れ値へのロバスト性（平均は外れ値1つで引きずられる）"),
    (1, "実効的な標準偏差の縮小"),
    (0, "集約後の値を距離 d_meas に変換"),
])

# S9 単端末の限界
s = add_slide(); title(s, "単端末の限界")
bullets(s, TOP, [
    (0, "±2σ帯は広く、単独では「粗い矛盾」しか検知できない"),
    (1, "例: 5m先と放送・実は30m → 捕捉可（粗い嘘）"),
    (1, "例: 25m を 30m と主張 → ノイズに埋もれ検知不可（細かい嘘）"),
    (0, "→ 細かい詐称の検知には多端末の冗長が必要"),
])

# S10 なぜ複数端末
s = add_slide(); title(s, "なぜ複数端末が必要か（信号対雑音）")
y = code(s, TOP, ["ΔRSSI = 10n·log10(30/25) ≈ 2 dB   （< ノイズ σ=4dB）"])
y = bullets(s, y + 0.05, [
    (0, "30m地点で5mずらしても信号は約2dB＝ノイズ以下→1台では埋もれる"),
    (0, "分散した複数端末の平均で雑音を √N 縮小（理論値）"),
])
y = table(s, y + 0.05, [
    ["端末数 N", "1 台", "5 台", "12 台"],
    ["実効ノイズ σ/√N", "4.0 dB", "1.8 dB", "1.2 dB"],
    ["2dBの嘘", "✕ 埋もれる", "○ 浮かぶ", "◎ 明確"],
], col_w=[2.4, 1.6, 1.6, 1.6], fs=12)
y = bullets(s, y + 0.02, [
    (0, "小さな位置ずれの検知には、空間的に分散した複数端末が必然"),
], size0=14.5)

# S11 位置特定
s = add_slide(); title(s, "複数端末での位置特定（多辺測量）")
y = bullets(s, TOP, [
    (0, "空間分散した N 端末の距離を連立 → 発信源の水平位置 (x,y) を特定"),
])
y = code(s, y + 0.05, [
    "d_i² = (x−x_i)² + (y−y_i)² + h²",
    "基準式との差で h² が消え、(x,y) について線形 → 2×2 最小二乗で求解",
])
y = bullets(s, y + 0.05, [
    (0, "地上(z=0)端末でも h² 消去で水平位置は解ける"),
    (0, "実装: multilaterate2d（正規方程式のクラメル解）"),
])

# S12 位置→詐称判定
s = add_slide(); title(s, "位置特定 → 詐称判定")
y = code(s, TOP, [
    "residual = ‖ P_hat(x,y) − P_claim(x,y) ‖",
    "residual > 閾値  ⇒  位置詐称",
])
y = bullets(s, y + 0.1, [
    (0, "推定発信源 P_hat と放送位置 P_claim の水平残差で判定"),
    (0, "単一偽装源では全端末の幾何を同時に満たせない構造"),
    (0, "閾値は正常データの FPR=5% を基準に決定"),
])

# ── 評価 ──
section("評価（シミュレーション）")

# S14 評価条件
s = add_slide(); title(s, "評価 — シミュレーション条件")
table(s, TOP, [
    ["項目", "条件"],
    ["真の発信源", "端末配置（半径40m）の中心付近・高度20m"],
    ["端末数 N", "1 / 3 / 5 / 8 / 12（掃引）"],
    ["RSSI ノイズ", "正規分布・σ = 2〜8 dB（暫定4dB）"],
    ["詐称量 Δ", "5 / 10 / 20 / 30 / 50 / 100 m（掃引）"],
    ["指標", "AUC・FPR=5% 時の TPR（検知率）"],
    ["試行回数", "3000 回 / 条件（集約 M=8）"],
], col_w=[2.4, 6.2], fs=12.5, row_h=0.44)

# S15 評価① Δ
s = add_slide(); title(s, "評価① — 検知率 vs 位置ずらし Δ（N=5）")
y = table(s, TOP, [
    ["Δ(m)", "5", "10", "20", "30", "50", "100"],
    ["検知率 TPR", "10%", "28%", "75%", "91%", "98%", "100%"],
], col_w=[2.4, 1.0, 1.0, 1.0, 1.0, 1.0, 1.1], fs=13, row_h=0.5)
y = bullets(s, y + 0.15, [
    (0, "位置ずらしが大きいほど検知しやすい（30mで91%）"),
    (0, "検知下限 Δ_min は運用要件（許容 FPR）から決める"),
])

# S16 評価② N
s = add_slide(); title(s, "評価② — 検知率 vs 端末数 N（Δ=30m）")
y = table(s, TOP, [
    ["N", "1", "3", "5", "8", "12"],
    ["検知率 TPR", "64%", "83%", "91%", "98%", "99%"],
], col_w=[2.4, 1.2, 1.2, 1.2, 1.2, 1.2], fs=13, row_h=0.5)
y = bullets(s, y + 0.15, [
    (0, "端末が増えるほど TPR 上昇 ＝ √N の裏付け"),
    (0, "ノイズσの影響: σ=2dB→99% / 4dB→91% / 6dB→68% / 8dB→49%"),
])

# S17 評価③ 方式比較
s = add_slide(); title(s, "評価③ — 方式比較（合意 vs 多辺測量）")
y = table(s, TOP, [
    ["Δ(m) [N=6, FPR=5%]", "10", "20", "30", "50"],
    ["中央値合意(K-of-N) TPR", "31%", "83%", "95%", "99%"],
    ["多辺測量 TPR", "13%", "60%", "96%", "100%"],
], col_w=[3.2, 1.3, 1.3, 1.3, 1.3], fs=12, row_h=0.48)
y = bullets(s, y + 0.1, [
    (0, "現実ノイズ＋少数端末では中央値合意(K-of-N)が優位"),
    (0, "→ 既定は合意、位置特定は好条件での補助（負の結果＝設計判断の裏付け）"),
])

# S18 評価④ 較正誤差
s = add_slide(); title(s, "評価④ — 較正誤差が最大の弱点")
y = callout(s, TOP, "結論", "RSSI方式の最大の弱点は、端末数よりも較正誤差")
y = table(s, y, [
    ["較正（N=5, Δ=30m）", "一致", "n:2.5→2.0", "n:2.5→3.0", "A:+6dB"],
    ["検知率 TPR", "91%", "7.7%", "72%", "8%"],
], col_w=[2.7, 1.2, 1.7, 1.5, 1.5], fs=12, row_h=0.46)
y = bullets(s, y + 0.08, [
    (0, "想定 n を 0.5 ずらすだけで 91%→7.7% に崩壊"),
    (0, "＝実機較正は前提条件（要求: A 数dB以内・n ±0.5未満）"),
])

# ── 実装 ──
section("実装状況")

# S20 実装①
s = add_slide(); title(s, "実装① — アプリ・評価基盤・ソルバ")
bullets(s, TOP, [
    (0, "アプリ = 受信〜RSSI集約〜整合判定を Pixel 8 実機で動作"),
    (1, "rssi_ranging.dart / rssi_logger.dart（オンデバイス測定・CSVログ）"),
    (0, "評価基盤 = 合成データで ROC/AUC・Δ/N/σ・較正掃引を実装"),
    (1, "eval_spoof_detection.py（依存なしで再現可能）"),
    (0, "多辺測量ソルバ = multilaterate2d（水平位置を最小二乗で推定・評価に組込み）"),
])

# S21 実装②
s = add_slide(); title(s, "実装② — ゼロ知識回路・多端末合意")
y = bullets(s, TOP, [
    (0, "ゼロ知識回路 = 相対許容の整合を秘匿証明（consistency_ok 公開・座標は非公開）"),
    (1, "r=2.1 の倍率帯を二乗比較・回路規模 約650制約"),
    (0, "witness統合 = 同一 drone_id・time_bucket の端末を束ねる"),
    (0, "多端末合意 = K-of-N で trusted / spoofAlert を記録"),
])
y = code(s, y + 0.05, [
    "受信 → 集約 → d_meas → 整合判定 → 回路 → proof",
    "  → レジストリで束ね & K 合意 → trusted / spoofAlert",
])

# S22 脅威モデル
s = add_slide(); title(s, "脅威モデルと限界（正直に）")
y = table(s, TOP, [
    ["攻撃", "RSSI での検知"],
    ["位置詐称（放送位置だけ書換）", "○ 発信源近くの端末が矛盾"],
    ["送信電力の操作（近く見せる）", "△ 全端末同時充足は困難"],
    ["中継 / ワームホール", "✕ RSSI は中継点を反映（原理的限界）"],
], col_w=[3.8, 4.8], fs=12.5, row_h=0.5)
y = bullets(s, y + 0.1, [
    (0, "主眼 = 位置詐称。中継は限界として明示（将来 TDoA / 実測裏取り）"),
    (0, "端末 GPS の信頼は端末アテステーションで別途担保"),
], size0=15)

# S23 パラメータ
s = add_slide(); title(s, "パラメータと確定方法")
table(s, TOP, [
    ["記号", "意味", "初期値", "確定"],
    ["A", "1m での RSSI", "−40 dBm(暫定)", "実機 fit"],
    ["n", "伝搬指数", "2.5", "実機 fit"],
    ["σ", "RSSI 残差", "3.5〜4 dB", "実機 fit"],
    ["M", "集約サンプル数", "10", "ノイズ/応答性"],
    ["r", "整合許容（相対倍率）", "2.1(暫定)", "評価で決定"],
    ["K", "合意しきい値", "3(暫定)", "評価で決定"],
], col_w=[1.0, 2.8, 2.4, 2.0], fs=12, row_h=0.44)

# S24 今後
s = add_slide(); title(s, "今後")
bullets(s, TOP, [
    (0, "実機 DJI での A/n/σ 較正（最重要・検知率を分ける前提条件）"),
    (1, "5/10/20/30/50m で回帰、環境別プロファイル、残差から σ 確定"),
    (0, "多機環境への対応（UAS_ID ごとに RSSI ウィンドウを分割）"),
    (0, "多端末合意 → ブロックチェーン記録の実接続"),
    (0, "次回（第20回）= 多端末証明の集約とモバイル実装（ZKPフレームワーク検証）"),
])

OUT = os.path.join(os.path.dirname(__file__), "..", "siles-pptx", "進捗第19回_RSSI位置特定.pptx")
prs.save(OUT)
print(f"saved: {os.path.abspath(OUT)}  ({len(prs.slides._sldIdLst)} slides)")
