#!/usr/bin/env python3
"""週17: マルチ端末 × ブロックチェーンによる UAV 接近証明の結合 (M1) まとめスライド。

本セッションで設計・実装・実証した内容を1本のデッキにまとめる:
  - 背景（1端末証明の限界）と多 witness 合意の狙い
  - 3層アーキテクチャ / 結合方式の比較
  - 端末回路 uav_witness_attestation.circom の追加点
  - AttestationRegistry のオンチェーン結合ロジック
  - M1 実装結果（実証明3台 → 実 Groth16Verifier → K=3 合意, テスト 9/9）
  - 正当性の評価（K台合意 ≠ 真正）/ 入力真正性の強化策(M3)
  - チェーン基盤（許可型 vs 公開）/ ロードマップ

デザイン（進捗第11回スタイル / ミニマル）:
  - 背景 白 / 黒文字（1A1A1A）、補助グレー（555555）。色は原則使わない。
  - フォント Yu Gothic（和文・英数）/ コードは Courier New。
  - 枠はコード/ログの抜粋のみ（F7F7F7 + CCCCCC）。タイトル直下に細い横罫線。
  - 4:3 (10 x 7.5 inch)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font)
    return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True)
    return p


def bullet(tf, text, size=13.5, space_before=7, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold)
    return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY)
    return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill)
        tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table
    tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            bold = header and ri == 0
            run(p, str(val), fs, INK, bold=bold)
    return tbl


def section(text, sub_text=""):
    s = add_slide()
    tf = tbox(s, 0.9, 3.0, 8.2, 1.4)
    run(tf.paragraphs[0], text, 30, INK, bold=True)
    hrule(s, 4.25, l=0.92, w=4.6)
    if sub_text:
        tf2 = tbox(s, 0.92, 4.4, 8.2, 0.8)
        run(tf2.paragraphs[0], sub_text, 14, GRAY)
    return s


# ════════════════════════════════════════════════════
# 1 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.15, 8.2, 0.5)
run(tf.paragraphs[0], "進捗まとめ（2026-07-06 / 週17）", 15, GRAY)
tf = tbox(s, 0.88, 2.68, 8.4, 1.9)
run(tf.paragraphs[0], "マルチ端末 × ブロックチェーン", 30, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "による UAV 接近証明の結合 (M1)", 30, INK, bold=True)
hrule(s, 4.6, l=0.92, w=5.6)
tf = tbox(s, 0.92, 4.78, 8.2, 1.6)
run(tf.paragraphs[0], "複数端末の ZK 証明を同一ドローンで束ね、K 台合意で「実在・接近」を検証する", 14, INK)
p = tf.add_paragraph(); p.space_before = Pt(6)
run(p, "branch acoustic-detection-zk / commit 0c64254", 12, GRAY, font=MONO)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "実端末3台の実証明 → 実 Groth16Verifier(オンチェーン) → K=3 合意 / テスト 9/9", 11.5, GRAY)


# ════════════════════════════════════════════════════
# 2 目次
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "目次")
tf = tbox(s, 0.7, 1.4, 8.7, 5.6)
subhead(tf, "背景 — なぜ複数端末で束ねるのか", 0)
bullet(tf, "1端末証明の限界（計算の正当性 ≠ 入力の真正性）と多 witness 合意", first=True)
subhead(tf, "設計 — 3層アーキテクチャと結合方式")
bullet(tf, "端末層 / 集約層 / チェーン層、再帰アグリゲータの新規性", first=True)
subhead(tf, "実装 — M1（個別検証＋閾値レジストリ）")
bullet(tf, "端末回路の追加点 / AttestationRegistry / E2E 実証結果", first=True)
subhead(tf, "評価と今後 — 正当性・入力真正性・チェーン基盤・ロードマップ")


# ════════════════════════════════════════════════════
# 章扉: 背景
# ════════════════════════════════════════════════════
section("背景", "なぜ複数端末の証明を束ねるのか")

# 3 課題
s = add_slide()
title(s, "課題 — 1端末の ZK 証明では偽装を防げない")
tf = tbox(s, 0.7, 1.4, 8.7, 2.4)
bullet(tf, "ZK は「計算の正当性」を示すが「入力の真正性」は保証しない", bold=True, first=True)
sub(tf, "回路は与えられた入力が正しく計算されたことしか言えない")
bullet(tf, "1端末では、スピーカー再生・GPS 詐称・RemoteID 放送値の流用で偽装しうる")
sub(tf, "自端末位置も private 入力 → 放送位置の近くの偽座標で within30=1 を捏造可能")
codebox(s, 0.7, 4.05, 8.6, 1.7, [
    ("# 回路が無検証で信じている入力", True),
    "drone_*  (ドローン位置/ID)  ← 放送は公開情報。傍受で一致させられる",
    "target_* (自端末位置/時刻)  ← 自己申告。任意座標を入れられる",
    "detection(映像検知の結果)   ← 実際に推論せず 1 と入れられる",
])
tf = tbox(s, 0.7, 5.95, 8.7, 1.0)
bullet(tf, "→ 単一の証明では「本当にそこにいたか」を担保できない", bold=True, first=True)

# 4 アイデア
s = add_slide()
title(s, "アイデア — 多 witness 合意による耐偽装")
tf = tbox(s, 0.7, 1.4, 8.7, 3.0)
bullet(tf, "物理的に離れた複数端末が、同一ドローンへ各自「30m以内＋検知」を独立に証明", first=True)
bullet(tf, "同じドローンについて K 台以上一致したら「実在・接近」と合意")
bullet(tf, "1台のスピーカーでは、離れた複数端末に整合する観測を同時に作れない", bold=True)
sub(tf, "＝物理的困難さを暗号学的合意で増幅する")
codebox(s, 0.7, 4.35, 8.6, 1.5, [
    "端末A ─┐",
    "端末B ─┼─ 同一 drone_id を独立に証明 ──▶ K台合意 = 信頼できる接近記録",
    "端末C ─┘",
])


# ════════════════════════════════════════════════════
# 章扉: 設計
# ════════════════════════════════════════════════════
section("設計", "3層アーキテクチャと結合方式")

# 5 アーキテクチャ
s = add_slide()
title(s, "全体設計 — 3層アーキテクチャ")
codebox(s, 0.7, 1.4, 8.6, 4.5, [
    ("[端末層]  各スマホ = uav_witness_attestation.circom で Groth16 証明", False),
    ("   private : ドローン位置/時刻, 自端末位置/時刻, 映像検知,", True),
    ("             remote_id, EdDSA署名", True),
    ("   public  : within30_and_detected, 各コミットメント,", True),
    ("             drone_id_commitment(結合キー), device_pubkey(端末ID),", True),
    ("             time_bucket(時間窓)", True),
    ("      │  N個の独立 proof（同一 drone_id_commitment を共有）", True),
    ("      ▼", False),
    ("[集約層]  M1: なし（各 proof を個別にチェーンへ）", False),
    ("          M2: 再帰アグリゲータ → 集約 proof 1個（新規性の本体）", True),
    ("      ▼", False),
    ("[チェーン層]  Groth16 検証器で検証 + AttestationRegistry で K台集計", False),
], size=11.5)
tf = tbox(s, 0.7, 6.0, 8.7, 0.9)
bullet(tf, "結合キー: 各端末が受信する RemoteID は共通 → drone_id_commitment が自然に一致", size=12, first=True)

# 6 結合方式の比較
s = add_slide()
title(s, "結合方式の比較 — 新規性はどこか")
table(s, 0.7, 1.5, 8.6, [
    ["方式", "内容", "新規性", "採用"],
    ["個別検証＋閾値", "各 proof を個別検証しレジストリで K 台数える", "低", "M1"],
    ["SnarkPack 集約", "同一 VK の proof を安く束ねる gas 最適化", "低", "×"],
    ["再帰アグリゲータ", "N proof を回路内で再帰検証＋関係を強制し1証明へ", "高", "M2"],
], col_w=[1.7, 4.3, 1.0, 1.6], fs=11.5, row_h=0.6)
tf = tbox(s, 0.7, 4.4, 8.7, 2.4)
subhead(tf, "再帰アグリゲータ（M2, 本命）が回路内で強制する関係", 0)
bullet(tf, "① 全 proof の drone_id_commitment が一致（＝同じドローン）", first=True)
bullet(tf, "② device_pubkey が相異なる（重複端末を排除 = Sybil 対策）")
bullet(tf, "③ 各 proof の within30_and_detected == 1")
bullet(tf, "④（任意）各端末位置とドローン放送位置の幾何整合（三角測量的）")
sub(tf, "単なる集約ではなく「物理イベントに対するマルチ witness 合意の証明」")


# ════════════════════════════════════════════════════
# 章扉: 実装
# ════════════════════════════════════════════════════
section("実装", "M1 — 個別検証＋閾値レジストリ")

# 7 端末回路
s = add_slide()
title(s, "端末回路 — uav_witness_attestation.circom")
tf = tbox(s, 0.7, 1.35, 8.7, 1.9)
bullet(tf, "既存 detection_commitment（距離+時刻+映像検知+コミットメント）を拡張", first=True)
bullet(tf, "オンチェーン結合のために public 出力を3つ追加", bold=True)
codebox(s, 0.7, 3.15, 8.6, 2.2, [
    "drone_id_commitment = Poseidon(remote_id)     # 同一ドローンの結合キー",
    "device_pubkey (Ax, Ay) + EdDSA 署名を回路内検証 # 端末ID / Sybil対策 / replay防止",
    "time_bucket = target_time_sec \\ 60           # 時刻は隠したまま時間窓で集計",
    "",
    ("署名対象 M = Poseidon(drone_id_commitment, target_commitment,", True),
    ("               frame_binding, within30_and_detected)", True),
])
tf = tbox(s, 0.7, 5.55, 8.7, 1.2)
bullet(tf, "コンパイル成功: public 出力 14 / 非線形制約 9,558（大半は EdDSA の楕円曲線演算）",
       size=12, first=True)
sub(tf, "circomlib primitive 上に移植（同名テンプレート衝突を回避、既存4回路は無改変）")

# 8 コントラクト
s = add_slide()
title(s, "コントラクト — AttestationRegistry")
tf = tbox(s, 0.7, 1.35, 8.7, 0.7)
bullet(tf, "同一 drone_id × 同一 time_bucket で、相異なる登録端末が K 台揃うと合意成立", first=True)
codebox(s, 0.7, 2.15, 8.6, 3.6, [
    "submitAttestation(proof, input[14]):",
    "  require(input[6] == 1)                    # within30_and_detected",
    "  require(verifier.verifyProof(proof,input)) # Groth16 検証",
    "  deviceId = keccak(input[11], input[12])   # device_pubkey",
    "  require(registeredDevice[deviceId])       # 登録端末のみ (Sybil対策)",
    "  key = keccak(input[10], input[13])        # (drone_id, time_bucket)",
    "  require(!counted[key][deviceId])          # 同一端末の二重計上を排除",
    "  counted[key][deviceId] = true; count[key]++",
    "  if (count[key] >= K) consensusReached[key] = true",
], size=11.5)
tf = tbox(s, 0.7, 5.95, 8.7, 0.9)
bullet(tf, "証明は誰が送信してもよく、合意は device_pubkey 単位で数える（リレー提出OK）",
       size=12, first=True)

# 9 E2E 結果
s = add_slide()
title(s, "実装結果 — E2E 実証 (K=3)")
tf = tbox(s, 0.7, 1.35, 8.7, 0.7)
bullet(tf, "実端末3台の実 Groth16 証明を、実 Verifier のオンチェーン検証に通し合意まで確認", first=True)
codebox(s, 0.7, 2.1, 8.6, 1.5, [
    "実証明3台(同一drone_id・distinct device_pubkey)",
    "   → 実 Groth16Verifier.sol でオンチェーン検証",
    "   → K=3 で consensusReached 発火 / count = 3",
])
tf = tbox(s, 0.7, 3.85, 8.7, 3.0)
subhead(tf, "Hardhat テスト 9/9 通過", 0)
bullet(tf, "K台合意 / 二重提出排除 / 未登録拒否 / within30チェック / 無効証明拒否", first=True)
bullet(tf, "別ドローン・別時間窓は別グループとして集計")
bullet(tf, "改ざんした public signal は実 verifier が拒否（invalid proof）", bold=True)
subhead(tf, "成果物（commit 0c64254）")
bullet(tf, "circuits/uav_witness_attestation.circom / contracts/（Hardhat）", first=True)
bullet(tf, "scripts/gen_witness_sample.js（circomlibjs で署名付き入力）")
bullet(tf, "scripts/build_witness_attestation_assets.sh（回路ビルド＋Verifier出力, PoT 2^14）")


# ════════════════════════════════════════════════════
# 章扉: 評価と今後
# ════════════════════════════════════════════════════
section("評価と今後", "正当性・入力真正性・チェーン基盤・ロードマップ")

# 10 正当性評価
s = add_slide()
title(s, "正当性の評価 — 構造は正しい、しかし穴がある")
tf = tbox(s, 0.7, 1.4, 8.7, 2.3)
subhead(tf, "妥当な点", 0)
bullet(tf, "drone_id を結合キーにする判断 / EdDSA で claim を束縛（replay 防止）", first=True)
bullet(tf, "座標は private のままコミットメントだけ公開（プライバシー維持）")
bullet(tf, "段階化（M1 個別検証 → M2 再帰集約）はリスク順として妥当")
tf = tbox(s, 0.7, 3.85, 8.7, 3.0)
subhead(tf, "最重要の穴 — 入力の真正性はまだ担保されていない", 0)
bullet(tf, "remote_id もドローン放送位置も公開情報 → 現地にいなくても一致させられる", first=True)
bullet(tf, "自端末位置も private → 偽座標で within30=1 を捏造できる")
bullet(tf, "「K台合意 ＝ 真正」ではなく「K台の登録端末の署名が揃った」に過ぎない", bold=True)
sub(tf, "真の耐偽装は次ページ(M3)の入力真正性の施策に依存する")

# 11 入力真正性(M3)
s = add_slide()
title(s, "入力の真正性を高める方法 (M3)")
tf = tbox(s, 0.7, 1.3, 8.7, 0.7)
bullet(tf, "外部の信用アンカーの署名を回路内検証、または物理実測に置き換える", first=True, bold=True)
table(s, 0.7, 2.1, 8.6, [
    ["優先", "施策", "効果"],
    ["高", "ドローン署名 RemoteID（DRIP/F3411 Auth）を回路内検証", "偽ドローンID封じ"],
    ["高", "端末アテステーション（Android/iOS）で位置・時刻を署名", "GPS注入封じ"],
    ["中", "署名 M に直近ブロックハッシュ（新鮮さ）", "リプレイ封じ"],
    ["中", "実測レンジ（RSSI/UWB）＋多端末三角測量", "遠隔偽装封じ"],
    ["低", "zkML / GNSS OSNMA / C2PA（将来）", "最強だが重い"],
], col_w=[0.8, 5.6, 2.2], fs=11.5, row_h=0.5)
tf = tbox(s, 0.7, 5.9, 8.7, 1.0)
bullet(tf, "耐偽装を主目的にするなら M3 は「任意」ではなく必須に格上げすべき", size=12.5, first=True, bold=True)

# 12 チェーン基盤
s = add_slide()
title(s, "チェーン基盤 — 許可型 vs 公開")
table(s, 0.7, 1.5, 8.6, [
    ["観点", "公開 EVM L2", "許可型 EVM（Besu/Quorum）"],
    ["Sybil対策", "ステーク等が別途必要", "参加資格で自然に担保 ◎"],
    ["プライバシー", "全 attestation が公開", "閲覧を限定 / private tx ◎"],
    ["公開検証性", "誰でも検証 ◎", "コンソーシアム信頼に依存"],
    ["コスト", "ガス変動", "トークン不要・高速 ◎"],
    ["検証器", "snarkjs Verifier をそのまま", "EVM互換なら無改変で流用可"],
], col_w=[1.7, 3.0, 3.9], fs=11, row_h=0.5)
tf = tbox(s, 0.7, 5.15, 8.7, 1.9)
bullet(tf, "本システムは合意が「登録端末」に依存 → 許可型はむしろ自然な適合", first=True, bold=True)
bullet(tf, "推奨: 許可型 EVM（Besu/Quorum, 検証器は無改変）＋ Merkle ルートを公開L2へアンカー")
sub(tf, "Sybil耐性・プライバシー・コストを取りつつ、改ざん検知の公開性も一部残すハイブリッド")

# 13 ロードマップ
s = add_slide()
title(s, "現在地とロードマップ")
tf = tbox(s, 0.7, 1.45, 8.7, 5.3)
subhead(tf, "M1 — 個別検証＋閾値レジストリ　✅ 完了", 0)
bullet(tf, "端末回路 + Groth16Verifier + AttestationRegistry、実証明で K=3 合意を実証", first=True)
sub(tf, "gas は台数比例。多端末合意の E2E が成立")
subhead(tf, "M2 — 再帰アグリゲータ（新規性の本体）　⬜")
bullet(tf, "オンチェーン集計を off-chain 再帰集約（gnark / Sonobe）へ置換", first=True)
bullet(tf, "「同一ドローン / 端末相異 / 各 within30」を回路内で強制、集約 proof 1個・定数 gas")
subhead(tf, "M3 — 耐偽装の核　⬜")
bullet(tf, "ドローン署名 RemoteID の回路内検証・端末アテステーション 等", first=True)
sub(tf, "「K台合意 ＝ 真正」を成立させる最後のピース")

# ── 出力 ─────────────────────────────────────────────
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "week17")
os.makedirs(OUT_DIR, exist_ok=True)
out_path = os.path.join(OUT_DIR, "week17_multi_witness_blockchain.pptx")
prs.save(out_path)
print(f"saved: {os.path.abspath(out_path)}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
