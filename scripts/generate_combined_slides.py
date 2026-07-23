#!/usr/bin/env python3
"""ドローン検知 × ゼロ知識証明 総合まとめスライド（やさしい解説＋具体名）。

設計・画像検知・音響検知・ZK接続・実機までを1本に統合する。
トーンは image_model_4x3 に合わせ、やさしい文章＋具体名＋各章に概要。

デザイン（進捗第11回スタイル / ミニマル, 4:3）:
  白背景・黒文字（1A1A1A）・補助グレー（555555）・色なし、Yu Gothic / Courier New、
  枠はコード抜粋のみ（F7F7F7 + CCCCCC）、タイトル下に細罫線。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xCC, 0xCC, 0xCC)
CODEBG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def _set_font(r, font=FONT):
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def run(p, text, size, color=INK, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    _set_font(r, font)
    return r


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.2), Inches(7.05), Inches(0.55), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        run(p, str(_page[0]), 10, GRAY)
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def hrule(s, top, l=0.7, w=8.6):
    ln = s.shapes.add_shape(1, Inches(l), Inches(top), Inches(w), Pt(1.0))
    ln.fill.solid(); ln.fill.fore_color.rgb = RULE
    ln.line.fill.background(); ln.shadow.inherit = False


def title(s, text):
    tf = tbox(s, 0.68, 0.42, 8.7, 0.8)
    run(tf.paragraphs[0], text, 23, INK, bold=True)
    hrule(s, 1.12)


def lead(s, text, t=1.32, size=14, h=1.2):
    tf = tbox(s, 0.7, t, 8.7, h)
    run(tf.paragraphs[0], text, size, INK)
    return tf


def subhead(tf, text, space_before=13):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, text, 15, INK, bold=True)
    return p


def bullet(tf, text, size=13.5, space_before=8, indent=False, bold=False, first=False):
    p = tf.add_paragraph(); p.space_before = Pt(2 if first else space_before)
    run(p, ("   ・ " if indent else "・ "), size, INK)
    run(p, text, size, INK, bold=bold)
    return p


def sub(tf, text, size=11.5, space_before=2):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "      " + text, size, GRAY)
    return p


def codebox(s, l, t, w, h, lines, size=11.5):
    box = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = RULE; box.line.width = Pt(0.75); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(10); tf.margin_right = Pt(8)
    tf.margin_top = Pt(7); tf.margin_bottom = Pt(7)
    for i, item in enumerate(lines):
        text, gray = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(2.5)
        run(p, text, size, GRAY if gray else INK, font=MONO)
    return box


def _cell_border(cell, color="CCCCCC", w=6350):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr); ln.append(fill); tcPr.append(ln)


def table(s, l, t, w, rows, col_w, header=True, fs=12, row_h=0.34):
    nr = len(rows); nc = len(rows[0])
    gtbl = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(row_h * nr))
    tbl = gtbl.table
    tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Pt(6); cell.margin_right = Pt(4)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_border(cell)
            p = cell.text_frame.paragraphs[0]
            run(p, str(val), fs, INK, bold=(header and ri == 0))
    return tbl


def section(no, text, sub_text=""):
    s = add_slide()
    tf = tbox(s, 0.9, 2.6, 8.2, 0.6)
    run(tf.paragraphs[0], no, 16, GRAY)
    tf = tbox(s, 0.88, 3.15, 8.4, 1.2)
    run(tf.paragraphs[0], text, 30, INK, bold=True)
    hrule(s, 4.45, l=0.92, w=4.6)
    if sub_text:
        tf2 = tbox(s, 0.92, 4.62, 8.2, 0.8)
        run(tf2.paragraphs[0], sub_text, 14, GRAY)
    return s


# ════════════════════════════════════════════════════
# 1 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
tf = tbox(s, 0.9, 2.2, 8.2, 0.5)
run(tf.paragraphs[0], "総合まとめ（2026-06）", 15, GRAY)
tf = tbox(s, 0.88, 2.75, 8.4, 1.6)
run(tf.paragraphs[0], "ドローン検知 ×", 32, INK, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "ゼロ知識証明", 32, INK, bold=True)
hrule(s, 4.4, l=0.92, w=5.4)
tf = tbox(s, 0.92, 4.58, 8.4, 1.6)
run(tf.paragraphs[0], "電波・カメラ・音でドローンを見つけ、", 15, INK)
p = tf.add_paragraph(); p.space_before = Pt(4)
run(p, "「近くにいた」ことを、中身を隠したまま証明します。", 15, INK)

# ════════════════════════════════════════════════════
# 2 概要
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "概要")
lead(s,
     "このアプリは、ドローンが自分の近くにいたことを、"
     "正確な位置や写真を見せずに証明できる仕組みです。", h=1.3)
tf = tbox(s, 0.7, 2.7, 8.7, 4.3)
subhead(tf, "ひと目でわかるポイント", 0)
bullet(tf, "見つける手段は3つ：電波（リモートID）・カメラ（画像）・音", first=True)
bullet(tf, "AIの推論も証明計算も、すべてスマホの中だけで完結（通信不要）")
bullet(tf, "証明できること：30m以内 ∧ 5秒以内 ∧ ドローン検知")
bullet(tf, "複数の手段を重ねるほど、ごまかしにくく確かになる")
sub(tf, "アプリ本体は Flutter 製の rapidsnark_app（ZKは Circom + Groth16）")

# ════════════════════════════════════════════════════
# 3 目次
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "目次")
tf = tbox(s, 0.7, 1.4, 8.7, 5.6)
bullet(tf, "1. 全体像と狙い（なぜ複数の手段か / ZKとは）", first=True)
bullet(tf, "2. カメラで見つける（画像によるドローン検知）")
bullet(tf, "3. 音で気づく（音響によるドローン検知）")
bullet(tf, "4. 証明につなげる（音の検知をZKで計算）")
bullet(tf, "5. 実機・現在地・これから")

# ════════════════════════════════════════════════════
# 第1章
# ════════════════════════════════════════════════════
section("第1章", "全体像と狙い", "なぜ複数の手段で見つけ、ZKで証明するのか")

# 3つの手段
s = add_slide()
title(s, "ドローンを見つける3つの手段")
lead(s, "1つの方法に頼らず、性質の違う3つを組み合わせます。")
table(s, 0.55, 2.35, 8.9, [
    ["手段", "中身", "得意"],
    ["電波（リモートID）", "BLE / WiFi の発信を受信", "正確な位置・時刻"],
    ["カメラ（画像）", "YOLO で見た目を検知", "電波を出さない機体・証拠画像"],
    ["音", "プロペラ音をAIで判定", "夜・視界外"],
], col_w=[2.6, 3.3, 3.0], row_h=0.55, fs=12)
tf = tbox(s, 0.7, 4.7, 8.7, 2.2)
subhead(tf, "組み合わせる理由", 0)
bullet(tf, "それぞれ穴がある（電波は偽装可、カメラは夜に弱い、音は遠距離に弱い）",
       first=True)
bullet(tf, "互いの弱点を補い合い、整合すれば確からしさが上がる", bold=True)

# なぜ複数か（脅威モデルやさしく）
s = add_slide()
title(s, "なぜ複数を重ねるのか")
lead(s,
     "守りたいのは「近くにいた事実を、改ざんせずに証明できる」こと。"
     "敵は2種類いて、対策が違います。", h=1.3)
table(s, 0.55, 2.65, 8.9, [
    ["問題", "内容", "対策"],
    ["取りこぼし", "電波を出さない機体に気づけない", "カメラ・音で補う"],
    ["なりすまし", "偽の電波・偽の音で捏造される", "複数手段の整合で裏取り"],
], col_w=[1.9, 4.0, 3.0], row_h=0.6, fs=12)
tf = tbox(s, 0.7, 4.9, 8.7, 2.0)
run(tf.paragraphs[0],
    "1つの手段だけだと、ニセモノを1つ作れば破れる。"
    "3つが同じ時刻で一致していれば、捏造はずっと難しくなる。", 13.5, INK, bold=True)

# ZKとは
s = add_slide()
title(s, "ゼロ知識証明（ZK）とは")
lead(s,
     "「正確な位置や写真そのものを見せずに、事実だけを証明する」技術です。", h=1.0)
tf = tbox(s, 0.7, 2.4, 8.7, 4.6)
subhead(tf, "このアプリでの例", 0)
bullet(tf, "見せずに証明：自分の正確な居場所やドローンの座標は隠す", first=True)
bullet(tf, "示すのは結果だけ：「30m以内にドローンがいた」という事実")
subhead(tf, "できること・できないこと")
bullet(tf, "できる：計算が正しく行われたことの証明", first=True)
bullet(tf, "できない：入力そのものが本物かの保証（＝真正性は別途必要）", bold=True)
sub(tf, "だから真正性は『電波の署名・複数手段の整合・複数端末』で外から支える")

# 設計方針
s = add_slide()
title(s, "進め方：今やること・後でやること")
lead(s, "全部を一度に作らず、動く最小形から1つずつ足します。")
table(s, 0.55, 2.35, 8.9, [
    ["区分", "内容"],
    ["今やる", "距離+時刻のコア証明 ＋ 検知を1つ束ねる"],
    ["後でやる", "音響検知 / 電波の署名検証 / 複数手段の合成"],
    ["将来", "複数端末で同じ機体を同時に確認（なりすましに最も強い）"],
], col_w=[1.7, 7.2], row_h=0.55, fs=12.5)
tf = tbox(s, 0.7, 4.6, 8.7, 2.0)
subhead(tf, "作り方の原則", 0)
bullet(tf, "まず一通り動かす → 1機能ずつ追加 → 各段階で検証", first=True)
sub(tf, "重い計算（AI推論）は証明の外で行い、結果だけ証明に渡す")

# ════════════════════════════════════════════════════
# 第2章 画像
# ════════════════════════════════════════════════════
section("第2章", "カメラで見つける", "画像によるドローン検知（YOLO）")

# 画像概要
s = add_slide()
title(s, "画像検知の概要")
lead(s,
     "カメラ映像のドローンを AI「YOLO」で自動検知し、結果を証明へつなぎます。", h=1.0)
tf = tbox(s, 0.7, 2.4, 8.7, 4.6)
subhead(tf, "ポイント", 0)
bullet(tf, "AI：YOLO を TFLite 形式に変換し、スマホ内だけで推論（通信不要）",
       first=True)
sub(tf, "TFLite = スマホで動かせるよう小さく軽くしたAIの形")
bullet(tf, "主役モデル：ドローン/飛行機/鳥/ヘリを見分ける best_float32.tflite")
bullet(tf, "予備：Google ML Kit（汎用AI）へ自動で切替")
bullet(tf, "出口：見つけたら「30m以内 ∧ 検知」の証明へ")

# 流れ
s = add_slide()
title(s, "ドローンを見つけるまでの流れ")
lead(s, "撮った動画から数コマだけ取り出し、1枚ずつAIが調べます。")
codebox(s, 0.7, 2.4, 8.6, 2.6, [
    "① 動画を撮る",
    "② 動画から数コマの静止画を取り出す（video_thumbnail）",
    "③ AI（YOLO）が「ドローン」と判定し枠を描く",
    "④ ドローンが写っていれば『検知あり』",
    "⑤ そのまま近接の証明（ZK）へ進める",
], size=12.5)
tf = tbox(s, 0.7, 5.3, 8.7, 1.5)
run(tf.paragraphs[0],
    "似た枠が重なったら1つにまとめる（NMS）。自信が一定以上で検知成立。",
    13.5, INK)

# 使うモデル
s = add_slide()
title(s, "使うモデル（具体名）")
lead(s, "メインはドローンを見分けるYOLO。読み込めない時は汎用AIに自動で切替えます。")
table(s, 0.5, 2.35, 9.0, [
    ["モデル / ファイル", "見分けるもの", "扱い"],
    ["best_float32.tflite (new-models)", "ドローン/飛行機/鳥/ヘリ", "最優先で使用"],
    ["uav_yolo11n_uav.tflite (COCO)", "人・車など80種類", "サンプル"],
    ["uav_visdrone_finetuned.tflite", "ドローン（VisDrone学習）", "未使用（将来）"],
    ["uav_roboflow_drone.tflite", "ドローン（Roboflow）", "未使用（要配置）"],
    ["Google ML Kit", "汎用（物体/ラベル）", "予備（自動切替）"],
], col_w=[3.7, 2.9, 2.4], row_h=0.5, fs=11)
tf = tbox(s, 0.7, 5.65, 8.7, 1.0)
sub(tf, "float32 / float16 の2種を同梱（端末で軽量に動かすため）", 11.5)

# ML Kit
s = add_slide()
title(s, "予備：Google ML Kit のしくみ")
lead(s,
     "メインのYOLOが使えない時の保険。Googleが配る汎用画像AIで、"
     "ネット不要・スマホ内だけで動きます。", h=1.3)
tf = tbox(s, 0.7, 2.65, 8.7, 2.0)
subhead(tf, "2つの機能を組み合わせて使う", 0)
bullet(tf, "ImageLabeler：写真全体に「何が写っているか」タグ付け（自信0.45以上）",
       first=True)
bullet(tf, "ObjectDetector：写真の中の物体を四角で見つける（1枚・分類・複数）")
codebox(s, 0.7, 4.75, 8.6, 1.35, [
    "物体を枠で発見 → 枠を切り出して再ラベリング",
    "  → 名前に drone / uav 等が入れば「ドローン」とみなす",
], size=11.5)
tf = tbox(s, 0.7, 6.25, 8.7, 0.9)
sub(tf, "汎用ゆえ『ドローン』専用ラベルは出にくい＝精度は控えめな保険", 11.5)

# パッケージ
s = add_slide()
title(s, "使っている部品（パッケージ）")
lead(s, "便利な部品を組み合わせて作っています。主な部品は次の通りです。")
table(s, 0.55, 2.25, 8.9, [
    ["部品（パッケージ）", "役割（やさしく）"],
    ["tflite_flutter", "AI（YOLO）をスマホの中で動かす"],
    ["google_mlkit_object_detection", "予備の物体検知AI"],
    ["google_mlkit_image_labeling", "予備の画像ラベル付けAI"],
    ["image", "画像の加工・枠の描き込み"],
    ["camera", "カメラの映像を扱う"],
    ["video_thumbnail", "動画からコマ（静止画）を取り出す"],
], col_w=[3.8, 5.1], row_h=0.46, fs=11.5)

# 最近直した点
s = add_slide()
title(s, "最近直した点")
lead(s,
     "証明に進む経路だけ『ドローンを知らないモデル(COCO)』を使っていて、"
     "写しても気づけませんでした。これを直しました。", h=1.4)
table(s, 0.55, 2.95, 8.9, [
    ["", "前", "後"],
    ["使うモデル", "uav_yolo11n_uav (COCO)", "best_float32 (new-models)"],
    ["結果", "写しても検知されない", "写せば検知され、証明に進める"],
], col_w=[1.7, 3.6, 3.6], row_h=0.6, fs=11.5)
tf = tbox(s, 0.7, 5.4, 8.7, 1.4)
run(tf.paragraphs[0],
    "「画面に映る検知」と「証明に進む検知」が同じモデルで揃いました。",
    13.5, INK, bold=True)

# ════════════════════════════════════════════════════
# 第3章 音
# ════════════════════════════════════════════════════
section("第3章", "音で気づく", "音響によるドローン検知")

# 音概要
s = add_slide()
title(s, "音響検知の概要")
lead(s, "ドローンのプロペラ音をスマホのマイクで聞き、AIでドローンらしさを判定します。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
subhead(tf, "なぜ音か", 0)
bullet(tf, "夜・逆光・物陰など、カメラが苦手な場面に強い", first=True)
bullet(tf, "マイクは元から付いていて、追加の機材が要らない")
subhead(tf, "考え方")
bullet(tf, "重い音の解析はアプリ側で行い、少数の『特徴』に要約する", first=True)
sub(tf, "その特徴を、後でZK証明の計算へそのまま渡せる形にする")

# なぜ音で分かるか
s = add_slide()
title(s, "なぜプロペラ音で分かるのか")
lead(s, "プロペラの音には、わかりやすい『指紋』のような特徴があります。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
bullet(tf, "回転に合わせた基本の高さ（BPF）と、その整数倍の音が並ぶ", first=True)
sub(tf, "BPF = 回転数 × ブレード枚数（例：6000rpm・2枚 → 約200Hz）")
bullet(tf, "4つのモータの差で、独特の『うなり』が出る")
bullet(tf, "飛ぶと音の高さがゆっくり動く（ドップラー）")
sub(tf, "鳥（不規則）や風（広く濁った音）とは違う形になる")

# 6次元特徴
s = add_slide()
title(s, "音から取り出す6つの特徴")
lead(s, "音を、ドローンらしさを表す6つの数字に要約します（解釈しやすい設計）。")
table(s, 0.55, 2.35, 8.9, [
    ["特徴", "意味（やさしく）"],
    ["倍音の強さ", "プロペラ特有の倍音がどれだけ出ているか"],
    ["基本の高さ(BPF)", "回転に対応する音の高さ"],
    ["高さの安定度", "音の高さがどれだけ一定か"],
    ["澄み具合", "雑音っぽくなく、音色がはっきりしているか"],
    ["低中域の比率", "低めの帯域にエネルギーが集まっているか"],
    ["高さの動き", "音の高さが時間で動くか（移動の手がかり）"],
], col_w=[2.6, 6.3], row_h=0.45, fs=11.5)

# データと結果
s = add_slide()
title(s, "データと結果")
lead(s, "公開データセットの実際の音で学習・評価しました。")
table(s, 0.55, 2.3, 8.9, [
    ["項目", "内容"],
    ["ポジ（ドローン音）", "DroneAudioDataset yes_drone 1,332件"],
    ["ネガ（環境音）", "同 unknown からランダム 1,332件"],
    ["結果", "ROC-AUC = 0.814（線形モデルのベースライン）"],
], col_w=[2.6, 6.3], row_h=0.55, fs=12)
tf = tbox(s, 0.7, 4.55, 8.7, 2.3)
subhead(tf, "正直なところ", 0)
bullet(tf, "妥当なベースラインだが、実用ラインではない（改善余地あり）",
       first=True)
sub(tf, "本番端末のマイクで録った自前データを混ぜると精度が安定する")

# 本物vs偽物
s = add_slide()
title(s, "本物 vs 偽物（なりすまし対策）")
lead(s,
     "環境音との区別はAIで解けます。難しいのは『スピーカーで流した偽の音』。"
     "これは音だけでは防げません。", h=1.4)
table(s, 0.55, 2.95, 8.9, [
    ["対策", "考え方"],
    ["他の手段との一致", "音・電波・カメラが同じ時刻に揃うか（最も有効）"],
    ["音の動き", "本物は高さが動く。固定スピーカーは動かない"],
    ["距離との整合", "電波の距離と音の大きさが物理的に合うか"],
], col_w=[2.4, 6.5], row_h=0.55, fs=12)
tf = tbox(s, 0.7, 5.4, 8.7, 1.3)
sub(tf, "音は『もう1票』。単独で本物保証は背負わせない", 12)

# ════════════════════════════════════════════════════
# 第4章 証明につなげる
# ════════════════════════════════════════════════════
section("第4章", "証明につなげる", "音の検知をゼロ知識証明の中で計算する")

# 概要
s = add_slide()
title(s, "音の判定を「証明の中」で行う")
lead(s,
     "音から取り出した6つの数字を使い、"
     "「ドローンらしさが基準を超えたか」を証明の中で計算します。", h=1.3)
codebox(s, 0.7, 2.7, 8.6, 1.5, [
    "証明できること =",
    "  30m以内にいた  かつ  音がドローンと判定された",
], size=12.5)
tf = tbox(s, 0.7, 4.5, 8.7, 2.4)
subhead(tf, "なぜ証明の中でやるのか", 0)
bullet(tf, "音そのものを見せずに、『ドローンと判定された』事実だけ示せる",
       first=True)
bullet(tf, "距離の条件と『かつ』で結べる → ごまかしにくい")

# 負の重み
s = add_slide()
title(s, "工夫：マイナスの重みをどう扱うか")
lead(s,
     "判定式にはマイナスの係数が出ます。証明の計算はマイナスが苦手なので、"
     "プラスとマイナスに分けて足し合わせます。", h=1.4)
codebox(s, 0.7, 2.9, 8.6, 1.7, [
    "プラス側の合計  と  マイナス側の合計  を別々に計算し、",
    "  「マイナス側 ≤ プラス側」かどうかで判定する",
], size=12)
tf = tbox(s, 0.7, 4.8, 8.7, 2.0)
bullet(tf, "数字はすべて整数に直してから計算（証明の計算は整数で行う）",
       first=True)
bullet(tf, "回路：distance_30m_time_feature_classifier6.circom（6つの特徴に対応）")

# 検証結果
s = add_slide()
title(s, "実際に証明が通った")
lead(s, "実際のドローン音から作った特徴で、証明の生成・検証に成功しました。")
codebox(s, 0.7, 2.4, 8.6, 2.6, [
    "[実ドローンの音から]",
    "within30           = 1   (30m以内)",
    "model_ok           = 1   (ドローンと判定)",
    "within30_and_model = 1   (接近 かつ 検知)",
    "",
    "snarkjs groth16 verify ... OK!",
], size=12)
tf = tbox(s, 0.7, 5.3, 8.7, 1.5)
run(tf.paragraphs[0],
    "「音のドローン判定」を証明の中で計算できることを、実データで確認。",
    13.5, INK, bold=True)

# ════════════════════════════════════════════════════
# 第5章 実機・まとめ
# ════════════════════════════════════════════════════
section("第5章", "実機・まとめ", "スマホで動かす / 現在地 / これから")

# 実機
s = add_slide()
title(s, "スマホ実機で動かした")
lead(s, "実際の Android スマホ（Pixel 8）に、無線でアプリを入れて動かしました。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
subhead(tf, "確認できたこと", 0)
bullet(tf, "Wi-Fi 経由でビルド・インストール・起動に成功", first=True)
bullet(tf, "カメラ・検知の画面が実機で動作")
bullet(tf, "画像検知は Drone 対応モデルへ統一済み（証明に進める）")
subhead(tf, "まだ残っていること")
bullet(tf, "音響は、録音→特徴→証明 のアプリ内配線がこれから", first=True)

# 現在地
s = add_slide()
title(s, "現在地（できたこと）")
lead(s, "土台と、音をZKに載せる部分までできています。")
table(s, 0.55, 2.3, 8.9, [
    ["項目", "状態"],
    ["距離+時刻のコア証明", "できた"],
    ["画像検知（Drone対応モデルへ統一）", "できた"],
    ["音響の検知（実データ AUC 0.814）", "できた"],
    ["音の判定をZKで計算・検証（OK）", "できた"],
    ["音響のアプリ内配線（録音→証明）", "これから"],
], col_w=[6.0, 2.9], row_h=0.46, fs=12.5)

# これから
s = add_slide()
title(s, "これからの改善")
lead(s, "弱点をつぶし、より確実でごまかしにくい証明にしていきます。")
tf = tbox(s, 0.7, 2.3, 8.7, 4.7)
bullet(tf, "画像：調べるコマ数を増やし、小さな機体や動きにも強くする",
       first=True)
bullet(tf, "音響：録音→特徴→証明 をアプリ内でつなぐ")
bullet(tf, "電波：リモートIDに署名を付け、証明の中で本物確認")
bullet(tf, "3手段（電波・画像・音）の一致を1つの証明にまとめる")
bullet(tf, "複数端末で同じ機体を確認（なりすましに最も強い）")

# まとめ
s = add_slide()
title(s, "まとめ")
tf = tbox(s, 0.7, 1.45, 8.7, 5.6)
bullet(tf, "電波・カメラ・音の3手段でドローンを見つける", first=True, bold=True)
bullet(tf, "AIも証明もスマホの中だけ。位置や写真は隠したまま証明できる",
       bold=True)
bullet(tf, "音のドローン判定を、証明の中で計算できることを実データで確認", bold=True)
subhead(tf, "ひとことで")
bullet(tf,
       "「電波・カメラ・音」を組み合わせるほど、ごまかしにくく確かな証明になる",
       first=True)

prs.save("drone_zk_overview_4x3.pptx")
print(f"saved drone_zk_overview_4x3.pptx ({len(prs.slides._sldIdLst)} slides)")
