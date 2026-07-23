#!/usr/bin/env python3
"""進捗報告スライド（文章中心 / 発表しやすい言い換え版）。

デザイン:
  - 背景 白 / シンプル
  - 上部タイトルは 青
  - 本文の重要部は 赤（赤だけ読めば要点が分かる）
  - 各スライド冒頭に「結論（概要）」を先出し
  - 専門用語はかみ砕く
  - 4:3 (10 x 7.5 inch) / 和文 Yu Gothic
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── 配色 ─────────────────────────────────────────────
BLACK = RGBColor(0x20, 0x20, 0x20)
BLUE = RGBColor(0x15, 0x5E, 0xC4)   # タイトル青
RED = RGBColor(0xC8, 0x12, 0x12)    # 本文の重要部
GRAY = RGBColor(0x6E, 0x6E, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LEADBG = RGBColor(0xF4, 0xF6, 0xFA)  # 概要ボックスの薄い背景
LINE = RGBColor(0xDD, 0xDD, 0xDD)

JP = "Yu Gothic"

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_page = [0]


def add_slide(number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    if number:
        _page[0] += 1
        tf = s.shapes.add_textbox(Inches(9.15), Inches(7.05), Inches(0.6), Inches(0.35)).text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(_page[0])
        r.font.size = Pt(10); r.font.color.rgb = GRAY; r.font.name = JP
    return s


def tbox(s, l, t, w, h):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    return tf


def run(p, text, size, color=BLACK, bold=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = JP


def title(s, text):
    """上部タイトル（青）。"""
    tf = tbox(s, 0.6, 0.5, 8.9, 1.0)
    run(tf.paragraphs[0], text, 28, BLUE, bold=True)
    ln = s.shapes.add_shape(1, Inches(0.62), Inches(1.34), Inches(8.78), Pt(2.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background(); ln.shadow.inherit = False


def lead(s, segments, top=1.62, height=1.15):
    """結論（概要）を先出しする薄い色のボックス。segments=[(text,is_red,bold)]"""
    box = s.shapes.add_shape(5, Inches(0.6), Inches(top), Inches(8.8), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = LEADBG
    box.line.color.rgb = LINE; box.line.width = Pt(1.0); box.shadow.inherit = False
    bar = s.shapes.add_shape(1, Inches(0.6), Inches(top), Inches(0.1), Inches(height))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background(); bar.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(16); tf.margin_right = Pt(12)
    tf.margin_top = Pt(8); tf.margin_bottom = Pt(8)
    p = tf.paragraphs[0]
    run(p, "結論  ", 13, BLUE, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(3)
    for text, is_red, bold in segments:
        run(p2, text, 16, RED if is_red else BLACK, bold=(bold or is_red))


def para(tf, segments, size=15.5, space_before=14):
    """本文の段落。先頭に「■」。segments=[(text,is_red,bold)]"""
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, "■ ", size, BLUE, bold=True)
    for text, is_red, bold in segments:
        run(p, text, size, RED if is_red else BLACK, bold=(bold or is_red))


def plain(tf, segments, size=14, space_before=4, indent="   "):
    p = tf.add_paragraph(); p.space_before = Pt(space_before)
    run(p, indent, size)
    for text, is_red, bold in segments:
        run(p, text, size, RED if is_red else BLACK, bold=(bold or is_red))


# ════════════════════════════════════════════════════
# 1 表紙
# ════════════════════════════════════════════════════
s = add_slide(number=False)
bar = s.shapes.add_shape(1, Inches(0.85), Inches(2.5), Inches(0.13), Inches(1.25))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
bar.line.fill.background(); bar.shadow.inherit = False
tf = tbox(s, 1.1, 2.55, 8.2, 1.6)
run(tf.paragraphs[0], "進捗報告", 20, GRAY, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(8)
run(p, "本物のドローンなしで", 30, RED, bold=True)
p2 = tf.add_paragraph()
run(p2, "WiFi 受信テストをできるようにした", 30, BLUE, bold=True)
tf = tbox(s, 1.12, 6.4, 8.0, 0.6)
run(tf.paragraphs[0], "ラズベリーパイを“疑似ドローン”にして検証環境を構築", 15, GRAY)

# ════════════════════════════════════════════════════
# 2 今回の概要（結論先出し）
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "今回やったこと（概要）")
lead(s, [("ラズパイを“疑似ドローン”に仕立て、", False, False),
         ("本物と同じ電波を発信", True, False),
         ("。スマホが", False, False),
         ("WiFi経由でドローンを検知", True, False),
         ("できるか試せる環境を作った。", False, False)], height=1.3)
tf = tbox(s, 0.62, 3.25, 8.8, 3.8)
para(tf, [("これまで受信アプリは完成していたが、", False, False),
          ("本物のドローンが無く動作確認ができなかった", True, False), ("。", False, False)])
para(tf, [("そこで身近なラズパイに、", False, False),
          ("本物そっくりのドローン電波を出させて代用", True, False),
          ("することにした。", False, False)])
para(tf, [("これで、電波の受信からアプリ内の処理まで、", False, False),
          ("実機を待たずに通して確認", True, False), ("できるようになった。", False, False)])

# ════════════════════════════════════════════════════
# 3 背景：なぜ必要だったか
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "なぜ必要だったか")
lead(s, [("多くのドローンは自分のIDや位置を", False, False),
         ("WiFiの電波で常に発信", True, False),
         ("している。これを受け取るのが本アプリ。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("ドローンには", False, False),
          ("“電子的なナンバープレート”（リモートID）", True, False),
          ("があり、ID・緯度経度・高度をWiFiで周囲にばらまいている。", False, False)])
para(tf, [("アプリはこの受信に対応済みだが、", False, False),
          ("本物のドローンが無いと動作を確かめられない", True, False), ("。", False, False)])
para(tf, [("実機（DJI 等）は高価で準備も大変。だから", False, False),
          ("ラズパイで“本物そっくりの電波”を出して代用", True, False),
          ("する方針にした。", False, False)])

# ════════════════════════════════════════════════════
# 4 何を作ったか（全体の流れ）
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "全体の流れ")
lead(s, [("ラズパイが電波を出す → スマホが受け取る → ", False, False),
         ("地図上の位置として表示 → そのまま証明まで通せる", True, False),
         ("。", False, False)])

def flowbox(s, l, head, sub, head_color=BLUE):
    box = s.shapes.add_shape(5, Inches(l), Inches(3.15), Inches(2.0), Inches(1.5))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = BLUE; box.line.width = Pt(1.4); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, head, 14, head_color, bold=True)
    for line in sub:
        pp = tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER; pp.space_before = Pt(3)
        run(pp, line, 10.5, GRAY)

flowbox(s, 0.62, "ラズパイ", ["疑似ドローン", "電波を発信"])
flowbox(s, 3.07, "スマホ", ["電波を受信", "中身を解読"])
flowbox(s, 5.52, "画面表示", ["ドローンの", "位置を表示"])
flowbox(s, 7.38, "証明", ["位置を明かさず", "近接を証明"], head_color=RED)
for x in (2.65, 5.1, 7.0):
    tf = tbox(s, x, 3.55, 0.5, 0.6)
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    run(pp, "→", 20, BLUE, bold=True)
tf = tbox(s, 0.62, 5.1, 8.8, 2.0)
para(tf, [("送信側（ラズパイ）が今回の新規。", False, False),
          ("受信側のアプリは既存のものをそのまま使用", True, False), ("した。", False, False)])
para(tf, [("最後の「証明」は、", False, False),
          ("位置情報を相手に明かさずに“近くにいた”ことだけ証明する仕組み", True, False),
          ("（ゼロ知識証明）。", False, False)])

# ════════════════════════════════════════════════════
# 5 仕組みをかみ砕く
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "ラズパイをどう“疑似ドローン”にしたか")
lead(s, [("WiFiの発信機能を使い、電波の中に", False, False),
         ("ドローンの身分証データを“国際標準の形”で埋め込んだ", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("ラズパイを、WiFiルーターのように", False, False),
          ("電波を出す状態（アクセスポイント）", True, False), ("にする。", False, False)])
para(tf, [("その電波の中に、", False, False),
          ("ドローンのID・緯度経度・高度を“決められた標準形式”で載せる", True, False),
          ("。", False, False)])
para(tf, [("スマホは電波全体ではなく、", False, False),
          ("この“ドローン情報の部分”だけを読み取って検知", True, False),
          ("する。形式が標準どおりなら本物と同じように認識される。", False, False)])

# ════════════════════════════════════════════════════
# 6 壁①
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "ぶつかった壁①：電波を“直接”は撃てない")
lead(s, [("本物と全く同じ撃ち方は内蔵WiFiが非対応 → ", False, False),
         ("発信機能の“情報枠”にデータを載せる方式で解決", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("ドローンと完全に同じ“生の電波を自由に作る方式”は、", False, False),
          ("ラズパイの内蔵WiFiでは対応していなかった", True, False), ("。", False, False)])
para(tf, [("そこで、WiFi発信機能に元々ある", False, False),
          ("“追加情報を載せる枠”を使ってドローン情報を入れる", True, False),
          ("方式に切り替えた。", False, False)])
para(tf, [("撃ち方は違っても、", False, False),
          ("スマホが読み取る中身は同じ", True, False),
          ("なので、検知の動作確認には問題ない。", False, False)])

# ════════════════════════════════════════════════════
# 7 壁②
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "ぶつかった壁②：操作の回線が切れる")
lead(s, [("操作もWiFi経由だったため発信すると切れる恐れ → ", False, False),
         ("繋いだまま同時に発信する構成で解決", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("このラズパイは有線が無く、", False, False),
          ("遠隔操作もWiFi経由", True, False), ("だった。", False, False)])
para(tf, [("単純に電波発信を始めると、", False, False),
          ("その操作回線まで切れてラズパイを触れなくなる恐れ", True, False),
          ("があった。", False, False)])
para(tf, [("そこで、", False, False),
          ("ネットに繋がったまま、発信も同時に行う構成", True, False),
          ("にして、操作を保ったまま電波を出せるようにした。", False, False)])

# ════════════════════════════════════════════════════
# 8 標準準拠の確認
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "“本物と同じ形式”だと確認した")
lead(s, [("自作データを", False, False),
         ("公式の標準ソフトの出力と照合 → 読み取る項目は完全一致", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("自分で組み立てた電波データが、", False, False),
          ("国際標準の公式ソフトが作るデータと一致するか突き合わせた", True, False),
          ("。", False, False)])
para(tf, [("結果、", False, False),
          ("スマホが実際に読む項目（ID・緯度・経度・高度）は完全に一致", True, False),
          ("した。", False, False)])
para(tf, [("つまり、", False, False),
          ("標準に準拠した“本物相当”のデータを出せている", True, False),
          ("と裏づけが取れた。", False, False)])

# ════════════════════════════════════════════════════
# 9 結果
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "結果：発信に成功")
lead(s, [("ラズパイは現在", False, False),
         ("“疑似ドローン”として電波を発信中", True, False),
         ("。操作回線も保てている。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("ラズパイは今、", False, False),
          ("ID「RPI-ODID-TEST」・位置=東京付近として電波を出し続けている", True, False),
          ("。", False, False)])
para(tf, [("遠隔操作の回線も維持できており、", False, False),
          ("いつでもスマホでの受信テストが行える状態", True, False), ("。", False, False)])
para(tf, [("位置の値（緯度経度・高度）は", False, False),
          ("コマンドひとつで自由に変えられる", True, False),
          ("ので、移動するドローンの再現もできる。", False, False)])

# ════════════════════════════════════════════════════
# 10 正直な制約
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "正直な制約（実機との違い）")
lead(s, [("電波の細部は本物と完全同一ではない。ただし", False, False),
         ("スマホが読む中身は公式と同一", True, False),
         ("で代替機として十分。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("使う電波の細かな種類やチャンネルは本物と完全同一ではない", False, False)],)
plain(tf, [("（ただしスマホは全チャンネルを探すため、検知は問題なく成立する）。", False, False)])
para(tf, [("本物のDJIの電波を捕まえて1バイト単位で突き合わせる検証はまだ未実施", False, False),
          ("。", False, False)])
para(tf, [("それでも、", False, False),
          ("スマホが解析する中身は公式実装と同一", True, False),
          ("なので、代替テスト送信機としては十分に機能する。", False, False)])

# ════════════════════════════════════════════════════
# 11 スマホでの確認手順
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "スマホでの確認手順")
lead(s, [("アプリを入れて見るだけ。", False, False),
         ("画面のドローンが位置表示に変われば成功", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.0, 8.8, 4.2)
for i, (head, sub) in enumerate([
    ("最新アプリを入れて、位置情報などの権限を許可する", ""),
    ("メイン画面のドローン表示が「未受信」から位置表示に変わる",
     "ID「RPI-ODID-TEST」・緯度経度が出れば受信成功"),
    ("「WiFi Raw」画面で、対象が“標準ドローン”として表示される", ""),
    ("そのデータで、これまでどおり証明まで作れることを確認する", ""),
], 1):
    p = tf.add_paragraph(); p.space_before = Pt(13)
    run(p, f"{i}.  ", 16, BLUE, bold=True)
    run(p, head, 15, RED if i == 2 else BLACK, bold=(i == 2))
    if sub:
        p2 = tf.add_paragraph(); p2.space_before = Pt(2)
        run(p2, "     " + sub, 12.5, GRAY)

# ════════════════════════════════════════════════════
# 12 まとめと次
# ════════════════════════════════════════════════════
s = add_slide()
title(s, "まとめと次の一手")
lead(s, [("実機なしで", False, False),
         ("“受信から証明まで”を検証できる環境ができた", True, False),
         ("。", False, False)])
tf = tbox(s, 0.62, 3.05, 8.8, 4.0)
para(tf, [("成果：", False, True),
          ("ラズパイを疑似ドローンにし、標準準拠の電波で受信〜証明を検証可能に", True, False),
          ("した。", False, False)])
para(tf, [("次①：", False, True),
          ("本物のDJIでの受信確認", True, False),
          ("（機種・ソフト差の検証）。", False, False)])
para(tf, [("次②：", False, True),
          ("より忠実な公式の送信ツール導入", False, False),
          ("や、外付けWiFiでの完全再現も検討。", False, False)])

# ── 保存 ─────────────────────────────────────────────
out = "raspi_wifi_progress.pptx"
prs.save(out)
print(f"saved: {out}  slides: {len(prs.slides._sldIdLst)}")
